"""Bounded spreadsheet import and real XLSX/DOCX result exports."""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import date, datetime, timezone

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from store import IMPORT_FIELDS, InputError


MAX_UPLOAD = 2_000_000
MAX_UNCOMPRESSED = 20_000_000


def parse_import_file(filename: str, content: bytes) -> list[dict]:
    if not filename or len(filename) > 160 or "/" in filename or "\\" in filename:
        raise InputError("文件名无效")
    if not content or len(content) > MAX_UPLOAD:
        raise InputError("文件大小须在 1 字节至 2 MB 之间")
    lower = filename.lower()
    if lower.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise InputError("CSV 须使用 UTF-8 编码") from exc
        if "\x00" in text:
            raise InputError("CSV 包含无效字符")
        reader = csv.DictReader(io.StringIO(text, newline=""))
        headers = reader.fieldnames or []
        rows = list(reader)
    elif lower.endswith(".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED:
                    raise InputError("XLSX 解压后过大")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            sheet = workbook.active
            cells = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(cells, ())]
            rows = []
            for values in cells:
                if len(rows) >= 501:
                    raise InputError("单次最多导入 500 行")
                if all(value is None or str(value).strip() == "" for value in values):
                    continue
                row = {}
                for header, value in zip(headers, values):
                    if isinstance(value, (date, datetime)):
                        value = value.strftime("%Y-%m-%d")
                    row[header] = value if value is not None else ""
                rows.append(row)
            workbook.close()
        except (ValueError, KeyError, zipfile.BadZipFile, OSError) as exc:
            if isinstance(exc, InputError):
                raise
            raise InputError("无法读取 XLSX 文件") from exc
    else:
        raise InputError("只支持 CSV 或 XLSX")
    if len(rows) > 500:
        raise InputError("单次最多导入 500 行")
    missing = [name for name in IMPORT_FIELDS if name not in headers]
    if missing:
        raise InputError("缺少表头：" + "、".join(missing))
    if len(set(headers)) != len(headers):
        raise InputError("表头不能重复")
    return rows


def import_template_csv() -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(IMPORT_FIELDS)
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def safe_cell(value: object) -> str:
    text = str(value or "")
    # Office treats these prefixes as formulas. Source data are text, never formulas.
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def export_xlsx(records: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "查询结果"
    headers = ["标题", "类型", "行业", "地区", "发布日期", "答案或状态", "摘要", "证据位置", "发布机构", "原站地址", "核验状态", "资料版本"]
    sheet.append(headers)
    for record in records:
        sheet.append([
            safe_cell(record["title"]), record["kind"], safe_cell(record["industry"]),
            safe_cell(record["region"]), record["published_at"],
            safe_cell(record["answer"] or "仅题录/暂无核验答案"), safe_cell(record["summary"]),
            safe_cell(record["evidence_locator"]), safe_cell(record["publisher"]),
            record["original_url"], record["review_status"], record["version_no"],
        ])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1B4C44")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {"A": 48, "B": 16, "C": 18, "D": 14, "E": 16, "F": 52, "G": 54, "H": 30, "I": 25, "J": 70, "K": 16, "L": 12}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    info = workbook.create_sheet("使用说明")
    info.append(["生成时间（UTC）", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    info.append(["说明", "仅包含本次选中的已发布资料；请以原站及证据位置核对。仅题录资料未由 AI 阅读全文。"])
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 100
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_docx(records: list[dict]) -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    document.add_heading("行业资料查询结果", 0)
    document.add_paragraph("生成时间（UTC）：" + datetime.now(timezone.utc).isoformat(timespec="seconds"))
    document.add_paragraph("仅包含本次选中的已发布资料。结论需结合原站、资料版本与证据位置使用。")
    for index, record in enumerate(records, 1):
        document.add_heading(f"{index}. {record['title']}", level=2)
        document.add_paragraph(
            f"{record['kind']} · {record['region'] or '地区未标注'} · "
            f"{record['industry'] or '行业未标注'} · 发布 {record['published_at'] or '日期待核'}"
        )
        document.add_paragraph("结果：" + (record["answer"] or "仅题录/暂无核验答案"))
        if record["summary"]:
            document.add_paragraph("摘要：" + record["summary"])
        document.add_paragraph("发布机构：" + record["publisher"])
        document.add_paragraph("原站：" + record["original_url"])
        document.add_paragraph("证据位置：" + (record["evidence_locator"] or "仅题录，未核验正文"))
        document.add_paragraph(f"核验状态：{record['review_status']} · 资料版本：v{record['version_no']}")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
