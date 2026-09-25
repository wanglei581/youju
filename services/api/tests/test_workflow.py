"""Behavior tests for the local source → review → query → export path."""

from __future__ import annotations

import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import formats  # noqa: E402
import store  # noqa: E402


class WorkflowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.db_path = Path(self.tmp.name) / "library.sqlite3"
        self.db = store.connection(self.db_path)
        store.init_db(self.db)

    def tearDown(self) -> None:
        self.db.close()
        self.tmp.cleanup()

    def source(self, rights: str) -> dict:
        return store.create_source(self.db, {
            "name": "测试来源", "publisher": "测试机构",
            "homepage_url": "https://example.org/column",
            "rights_level": rights,
            "rights_basis": "本地测试许可样本" if rights != "directory_only" else "",
        })

    @staticmethod
    def row(**changes: str) -> dict:
        result = {
            "title": "测试公告", "kind": "list", "industry": "医疗",
            "region": "青岛", "published_at": "2026-09-25",
            "original_url": "https://example.org/notice/1",
            "summary": "", "answer": "", "evidence_locator": "",
        }
        result.update(changes)
        return result

    def test_draft_is_not_public_and_new_version_waits_for_review(self) -> None:
        source = self.source("evidence_use")
        first = self.row(answer="测试企业入选", evidence_locator="附件第 3 行")
        batch = store.stage_import(self.db, source["id"], "test.csv", [first])
        self.assertEqual(store.commit_import(self.db, batch["batch_id"])["inserted_drafts"], 1)
        self.assertEqual(store.public_search(self.db), [])
        record_id = store.list_admin_records(self.db)[0]["id"]
        store.publish_record(self.db, record_id, "核对附件第 3 行")
        self.assertEqual(store.public_record(self.db, record_id)["answer"], "测试企业入选")

        corrected = self.row(answer="测试企业未入选，原公告已更正", evidence_locator="更正附件第 4 行")
        batch = store.stage_import(self.db, source["id"], "corrected.csv", [corrected])
        store.commit_import(self.db, batch["batch_id"])
        self.assertEqual(store.public_record(self.db, record_id)["answer"], "测试企业入选")
        store.publish_record(self.db, record_id, "核对更正附件第 4 行")
        public = store.public_record(self.db, record_id)
        self.assertEqual(public["version_no"], 2)
        self.assertIn("未入选", public["answer"])
        store.withdraw_record(self.db, record_id, "测试撤回")
        self.assertIsNone(store.public_record(self.db, record_id))
        with self.assertRaises(store.InputError):
            store.public_selected(self.db, [record_id])

    def test_directory_rights_cannot_import_answer_or_summary(self) -> None:
        source = self.source("directory_only")
        for restricted in (self.row(summary="超出题录"), self.row(answer="确定答案", evidence_locator="附件第 1 行")):
            preview = store.stage_import(self.db, source["id"], "test.csv", [restricted])
            self.assertIsNone(preview["batch_id"])
            self.assertEqual(preview["valid_count"], 0)
            self.assertEqual(len(preview["errors"]), 1)

    def test_selected_exports_contain_only_published_records(self) -> None:
        source = self.source("directory_only")
        rows = [self.row(title="已发布", original_url="https://example.org/1"),
                self.row(title="仍待审", original_url="https://example.org/2")]
        batch = store.stage_import(self.db, source["id"], "test.csv", rows)
        store.commit_import(self.db, batch["batch_id"])
        records = store.list_admin_records(self.db)
        published = next(item["id"] for item in records if item["draft_title"] == "已发布")
        pending = next(item["id"] for item in records if item["draft_title"] == "仍待审")
        store.publish_record(self.db, published, "")
        selected = store.public_selected(self.db, [published])
        xlsx = formats.export_xlsx(selected)
        docx = formats.export_docx(selected)
        for data in (xlsx, docx):
            self.assertTrue(zipfile.is_zipfile(io.BytesIO(data)))
        workbook = load_workbook(io.BytesIO(xlsx), read_only=True)
        self.assertEqual(workbook["查询结果"].max_row, 2)
        self.assertEqual(workbook["查询结果"]["A2"].value, "已发布")
        workbook.close()
        document = Document(io.BytesIO(docx))
        self.assertTrue(any("已发布" in paragraph.text for paragraph in document.paragraphs))
        self.assertFalse(any("仍待审" in paragraph.text for paragraph in document.paragraphs))
        with self.assertRaises(store.InputError):
            store.public_selected(self.db, [published, pending])

    def test_rights_restriction_masks_old_answer_and_suspension_hides_record(self) -> None:
        source = self.source("evidence_use")
        batch = store.stage_import(self.db, source["id"], "test.csv", [
            self.row(summary="获准测试摘要", answer="曾经可见的答案", evidence_locator="第 2 行")
        ])
        store.commit_import(self.db, batch["batch_id"])
        record_id = store.list_admin_records(self.db)[0]["id"]
        store.publish_record(self.db, record_id, "核验第 2 行")
        self.assertEqual(store.public_record(self.db, record_id)["answer"], "曾经可见的答案")
        store.restrict_source(self.db, source["id"], "directory_only", "仅保留题录")
        public = store.public_record(self.db, record_id)
        self.assertEqual(public["answer"], "")
        self.assertEqual(public["summary"], "")
        self.assertEqual(public["evidence_locator"], "")
        self.assertEqual(public["review_status"], "metadata")
        store.restrict_source(self.db, source["id"], "suspend", "来源要求暂停")
        self.assertIsNone(store.public_record(self.db, record_id))
        self.assertEqual(store.public_search(self.db), [])
        with self.assertRaises(store.InputError):
            store.public_selected(self.db, [record_id])

    def test_withdrawn_same_payload_can_be_reimported_for_review(self) -> None:
        source = self.source("directory_only")
        row = self.row()
        batch = store.stage_import(self.db, source["id"], "first.csv", [row])
        store.commit_import(self.db, batch["batch_id"])
        record_id = store.list_admin_records(self.db)[0]["id"]
        store.publish_record(self.db, record_id, "")
        store.withdraw_record(self.db, record_id, "原站临时撤回")
        batch = store.stage_import(self.db, source["id"], "again.csv", [row])
        result = store.commit_import(self.db, batch["batch_id"])
        self.assertEqual(result["inserted_drafts"], 1)
        self.assertEqual(store.list_admin_records(self.db)[0]["draft_version_no"], 2)
        self.assertIsNone(store.public_record(self.db, record_id))


if __name__ == "__main__":
    unittest.main()
