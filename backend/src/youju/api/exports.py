"""把公开题录与人工核对字段导出为小程序可直接打开的文件。"""

from __future__ import annotations

import hashlib
import json
import re
import threading
import time as clock
from collections import OrderedDict, deque
from dataclasses import asdict, dataclass
from datetime import datetime
from io import BytesIO
from typing import Annotated, Literal
from urllib.parse import quote, urlsplit
from zoneinfo import ZoneInfo

from docx import Document as WordDocument
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import JSONResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from youju.api.readers import NoStoreRoute
from youju.classify import applicable_region, effective_kind, list_status, region_label
from youju.config import get_settings
from youju.crawler.urls import public_url
from youju.db import get_session
from youju.dedup import canonical_ids, groups_for
from youju.export_notes import DEADLINE_HEADER, DISCLAIMER, SEARCH_NOTICE, add_excel_notes
from youju.issuer import normalize_issuer
from youju.models import Document, DocumentDeadline, DocumentDetail, DocumentField, Source
from youju.takedown import available_condition

router = APIRouter(prefix="/api/v1", tags=["小程序"])
native_router = APIRouter(route_class=NoStoreRoute)
SessionDep = Annotated[Session, Depends(get_session)]
TZ = ZoneInfo("Asia/Shanghai")
MAX_BYTES = 1000 * 1024
HEADERS = (
    "序号",
    "标题",
    "类型",
    "名单状态",
    "发文机关",
    "文号",
    "发布日期",
    "地区",
    DEADLINE_HEADER,
    "日期角色",
    "支持方式",
    "未核项",
)
NOTICE = (
    f"本文件由第三方整理，非政府平台；“{DEADLINE_HEADER}”列只列人工核对通过的记录，"
    "请以发文机关原文为准；名单状态仅按标题用词标注。"
    "“日期角色”是人工核对时确认的截止类型；“支持方式”只写申报通知里人工核对通过的；"
    "“未核项”列出这一行还没有人工核对的字段，请对照原文。"
)
EXCERPT_NOTICE = (
    f"“{DEADLINE_HEADER}”列中标“机器摘录”的日期是机器从原文摘出、未经人工核对的，请以原文为准。"
    "机器摘录的“日期角色”也标有“（机器摘录）”，未经人工核对。"
)


def _notice(excerpts: bool) -> str:
    if not excerpts:
        return NOTICE
    return (
        NOTICE.replace(
            f"“{DEADLINE_HEADER}”列只列人工核对通过的记录，",
            f"“{DEADLINE_HEADER}”列包含人工核对过的和机器摘录的，",
        ).replace(
            "“日期角色”是人工核对时确认的截止类型；",
            "“日期角色”标明截止类型，机器摘录的行另有标记；",
        )
        + EXCERPT_NOTICE
    )


# 日期角色：没有人工核对过的截止时固定写这句，不写任何日期（未核对的候选日期不进导出文件）
ROLE_SEE_ORIGINAL = "截止时间见原文"


class ExportIn(BaseModel):
    ids: list[str] = Field(min_length=1, max_length=200)
    format: Literal["xlsx", "docx"]
    # 首发版小程序（D-061）在请求体里带 excerpts=1；网址参数 ?excerpts=1 也认
    excerpts: Literal[1] | None = None
    revision: str | None = Field(default=None, pattern=r"^[a-f0-9]{64}$")


class PreviewIn(BaseModel):
    ids: list[str] = Field(min_length=1, max_length=200)
    excerpts: Literal[1] | None = None


@dataclass
class ExportItem:
    number: int
    doc_id: int
    title: str
    kind: str
    status: str
    issuer: str
    doc_number: str
    published: str
    region: str
    deadline: str
    date_role: str
    support: str
    unchecked: str
    urls: list[str]
    sites: str
    source_checked: str
    deadline_checked: str
    change_checked: str
    deadline_quote: str = ""


def _date_time(value: datetime | None) -> str:
    return value.astimezone(TZ).strftime("%Y-%m-%d") if value else ""


def _deadline_text(d: DocumentDeadline) -> str:
    day = d.approved_date.isoformat()
    clock = f" {d.approved_time.strftime('%H:%M')}" if d.approved_time else ""
    return f"{d.kind} {day}{clock}"


def _unchecked(
    kind: str,
    *,
    has_deadline: bool,
    deadline_pending: bool,
    support_approved: bool,
    support_pending: bool,
) -> list[str]:
    """这一行还没有人工核对的关键字段（固定顺序）。只看核对状态，不按正文或标题推断。

    - 截止日期（以原文为准）：没有核对过的截止；已有核对过的、
      但还有待核对的候选时写"其他截止日期（以原文为准）"（可能还有别的截止没核对）
    - 支持方式：只有申报通知有；没有核对通过的，或还有待核对（含正文变了要重核）的候选
    - 名单状态：只有名单有；只照录标题用词（D-018），没有人工核对这一步
    """
    items: list[str] = []
    if not has_deadline:
        items.append(DEADLINE_HEADER)
    elif deadline_pending:
        items.append(f"其他{DEADLINE_HEADER}")
    if kind == "申报通知" and (not support_approved or support_pending):
        items.append("支持方式")
    if kind == "名单":
        items.append("名单状态")
    return items


def _is_web_url(url: str) -> bool:
    parts = urlsplit(url)
    return parts.scheme in ("http", "https") and bool(parts.netloc)


def _items(session: Session, raw_ids: list[str], excerpts: bool = False) -> list[ExportItem]:
    valid: list[tuple[str, int]] = []
    invalid: list[str] = []
    for raw in raw_ids:
        if not re.fullmatch(r"[1-9][0-9]*", raw) or len(raw) > 19 or int(raw) > 2**63 - 1:
            invalid.append(raw)
        else:
            valid.append((raw, int(raw)))
    parsed = [doc_id for _, doc_id in valid]
    # 只导出可以公开的文件（和小程序列出的一致）；其他的当作不存在
    from youju.api import v1

    found = (
        set(
            session.scalars(
                select(Document.id)
                .join(Source, Document.source_id == Source.id)
                .where(Document.id.in_(parsed), v1.public_condition())
            )
        )
        if parsed
        else set()
    )
    invalid.extend(raw for raw, doc_id in valid if doc_id not in found)
    if invalid:
        invalid_set = set(invalid)
        raise HTTPException(
            status_code=422,
            detail={
                "不存在的文件编号": list(dict.fromkeys(i for i in raw_ids if i in invalid_set))
            },
        )

    masters = canonical_ids(session, parsed)
    ordered = list(dict.fromkeys(masters.get(doc_id, doc_id) for doc_id in parsed))
    rows = session.execute(
        select(Document, Source)
        .join(Source, Document.source_id == Source.id)
        .where(Document.id.in_(ordered))
    ).all()
    by_id = {doc.id: (doc, source) for doc, source in rows}
    groups = groups_for(session, ordered)
    member_master = {member: master for master, members in groups.items() for member in members}
    group_rows = session.execute(
        select(Document, Source)
        .join(Source, Document.source_id == Source.id)
        .where(Document.id.in_(member_master))
        .order_by(Document.id)
    ).all()
    members: dict[int, list[tuple[Document, Source]]] = {master: [] for master in ordered}
    for doc, source in group_rows:
        members[member_master[doc.id]].append((doc, source))

    # 正文不参与导出；只在数据库计算长度以沿用详情页的归并组优选规则。
    details: dict[int, list[tuple[int, str | None, str | None, int]]] = {}
    for doc_id, issuer, number, body_length in session.execute(
        select(
            DocumentDetail.document_id,
            DocumentDetail.issuer,
            DocumentDetail.doc_number,
            # 按字数算（MySQL 的 LENGTH 是字节数），和详情页用 Python len() 的口径一致
            func.char_length(DocumentDetail.body_text),
        ).where(DocumentDetail.document_id.in_(member_master))
    ):
        details.setdefault(member_master[doc_id], []).append(
            (doc_id, issuer, number, body_length or 0)
        )

    deadlines: dict[int, list[DocumentDeadline]] = {master: [] for master in ordered}
    for deadline in session.scalars(
        select(DocumentDeadline).where(
            DocumentDeadline.document_id.in_(member_master),
            DocumentDeadline.status.in_(("approved", "changed")),
        )
    ):
        deadlines[member_master[deadline.document_id]].append(deadline)
    target_ids = [
        d.changed_to_document_id
        for group in deadlines.values()
        for d in group
        if d.changed_to_document_id
    ]
    # 未核项只看有没有待核对的候选（截止、支持方式），不读候选的日期和取值
    pending_deadline = {
        member_master[doc_id]
        for doc_id in session.scalars(
            select(DocumentDeadline.document_id)
            .where(
                DocumentDeadline.document_id.in_(member_master),
                DocumentDeadline.status == "pending",
            )
            .distinct()
        )
    }
    pending_support = {
        member_master[doc_id]
        for doc_id in session.scalars(
            select(DocumentField.document_id)
            .where(
                DocumentField.document_id.in_(member_master),
                DocumentField.field == "support",
                DocumentField.status.in_(("pending", "stale")),
            )
            .distinct()
        )
    }
    # 已核对的支持方式和详情页同一口径（当前正文、一级来源、人工通过）
    approved_fields = v1._approved_fields(session, ordered)
    excerpt_views = {}
    if excerpts:
        from youju.deadline_excerpts import batch

        excerpt_views, _, _ = batch(
            session, ordered, v1._approved(session, ordered), v1.now().date()
        )
    target_masters = canonical_ids(session, target_ids)
    titles = {
        doc_id: title
        for doc_id, title in session.execute(
            select(Document.id, Document.title).where(
                Document.id.in_([target_masters.get(i, i) for i in target_ids]),
                available_condition(),
            )
        )
    }

    result = []
    for index, master in enumerate(ordered, 1):
        doc, source = by_id[master]
        best = max(
            details.get(master, []),
            key=lambda row: (bool(row[3]), row[3], row[0] == master),
            default=None,
        )
        group_deadlines = deadlines[master]
        changed = sorted(
            (d for d in group_deadlines if d.status == "changed" and d.approved_date),
            key=lambda d: (d.approved_date, d.id),
        )
        changed_keys = {(d.kind, d.approved_date, d.approved_time) for d in changed}
        approved = sorted(
            (
                d
                for d in group_deadlines
                if d.status == "approved"
                and d.approved_date
                and (d.kind, d.approved_date, d.approved_time) not in changed_keys
            ),
            key=lambda d: (d.approved_date, d.id),
        )
        notes: list[str] = []
        # 日期角色：截止日期那一格里每一条的截止类型（核对台通过时确认的类型，叫法和核对台、
        # 小程序一致），按截止日期那一格的顺序，同类只写一次
        roles: list[str] = []
        seen: set[tuple] = set()
        for d in approved:
            key = (d.kind, d.approved_date, d.approved_time)
            if key not in seen:
                notes.append(_deadline_text(d))
                roles.append(d.kind)
                seen.add(key)
        for d in changed:
            key = (d.kind, d.approved_date, d.approved_time)
            if key not in seen:
                target = target_masters.get(d.changed_to_document_id, d.changed_to_document_id)
                title = titles.get(target, "变更通知")
                notes.append(f"原定 {_deadline_text(d)} 已{d.change_reason}，见新通知《{title}》")
                roles.append(f"{d.kind}（已{d.change_reason}）")
                seen.add(key)
        kind = effective_kind(doc.title, doc.kind_override)
        excerpt_quote = "\n".join(
            dict.fromkeys(d.evidence for d in [*approved, *changed] if d.evidence)
        )
        if excerpts and kind == "申报通知" and not (approved or changed):
            from youju.deadline_excerpts import selected

            due = selected(excerpt_views.get(master, []), v1.now().date())
            if due and due.day:
                notes = [
                    due.day.isoformat()
                    + (f" {due.clock.strftime('%H:%M')}" if due.clock else "")
                    + ("（机器摘录）" if due.source == "excerpt" else "")
                ]
                roles = [due.kind + ("（机器摘录）" if due.source == "excerpt" else "")]
                excerpt_quote = due.quote or ""
            elif excerpt_views.get(master):
                notes = ["截止未核对，以原文为准"]
                roles = ["截止未核对"]
        # 支持方式：和详情页"支持方式（已核对原文）"同一来源、同样叫法，只有申报通知有；
        # 不用按标题归类的筛选标签（rule_supports）
        supports = approved_fields.get(master, {}).get("support", [])
        support = ("、".join(supports) or "见原文") if kind == "申报通知" else "—"
        unchecked = _unchecked(
            kind,
            has_deadline=bool(approved or changed),
            deadline_pending=master in pending_deadline,
            support_approved=bool(supports),
            support_pending=master in pending_support,
        )
        group = members[master]
        urls = list(dict.fromkeys(public_url(item.url) for item, _ in group))
        sites = list(
            dict.fromkeys(
                f"{src.dept_full_name}（{urlsplit(item.url).hostname or src.url}）"
                for item, src in group
            )
        )
        source_dates = sorted({src.verified_on.isoformat() for _, src in group if src.verified_on})
        review_dates = sorted({_date_time(d.reviewed_at) for d in group_deadlines if d.reviewed_at})
        change_dates = sorted({_date_time(d.changed_at) for d in changed if d.changed_at})
        result.append(
            ExportItem(
                number=index,
                doc_id=master,
                # 标题是发文机关的原标题，原样照录；标题里写着的截止信息（如"10 月 31 日前申报"）
                # 是原文的一部分，不是我们抽出的截止日期，不按"未核对不发布"处理（截止日期一栏另算）
                title=doc.title,
                kind=kind,
                status=list_status(doc.title, kind) or "",
                issuer=normalize_issuer(best[1] if best and best[1] else source.dept_full_name),
                doc_number=(best[2] if best and best[2] else ""),
                published=doc.publish_date.isoformat() if doc.publish_date else "",
                region=(
                    applicable_region(source.province, source.level, source.city, source.district)
                    if source.level in ("市级", "区县级")
                    else region_label(source.province, source.level, source.city, source.district)
                ),
                deadline="；".join(notes) if notes else "见原文",
                date_role="；".join(dict.fromkeys(roles)) if roles else ROLE_SEE_ORIGINAL,
                support=support,
                unchecked="、".join(unchecked) or "无",
                urls=urls,
                sites="；".join(sites),
                source_checked="、".join(source_dates) or "未记录",
                deadline_checked="、".join(review_dates) or "未记录",
                change_checked="、".join(change_dates) or "未记录",
                deadline_quote=excerpt_quote,
            )
        )
    return result


def _cell(sheet, row: int, column: int, value: str | int) -> None:
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"  # 政府网站题录属于外部文本，不能解释成 Excel 公式。
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _xlsx(items: list[ExportItem], stamp: str, excerpts: bool = False) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "文件清单"
    max_urls = max(len(item.urls) for item in items)
    headings = [*HEADERS, "截止原文句子"]
    headers = [*headings, "原文链接", *(f"原文链接 {n}" for n in range(2, max_urls + 1))]
    for column, title in enumerate(headers, 1):
        _cell(sheet, 1, column, title)
        sheet.cell(1, column).font = Font(bold=True)
    for row, item in enumerate(items, 2):
        values = (
            item.number,
            item.title,
            item.kind,
            item.status,
            item.issuer,
            item.doc_number,
            item.published,
            item.region,
            item.deadline,
            item.date_role,
            item.support,
            item.unchecked,
            item.deadline_quote,
            *item.urls,
        )
        for column, value in enumerate(values, 1):
            _cell(sheet, row, column, value)
        for column, url in enumerate(item.urls, len(headings) + 1):
            if _is_web_url(url):
                sheet.cell(row, column).hyperlink = url
                sheet.cell(row, column).font = Font(color="0563C1", underline="single")
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {"标题": 60, DEADLINE_HEADER: 60, "日期角色": 26, "支持方式": 22, "未核项": 36}
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = (
            60 if column > len(headings) else widths.get(headings[column - 1], 18)
        )

    note = book.create_sheet("数据说明")
    note.append(["数据说明", _notice(excerpts)])
    note.append(["导出时间（北京时间）", stamp])
    note.append(
        ["序号", "文件编号", "来源网站", "来源核对日期", "截止日期核对日期", "变更核对日期"]
    )
    for row, item in enumerate(items, 4):
        values = (
            item.number,
            str(item.doc_id),
            item.sites,
            item.source_checked,
            item.deadline_checked,
            item.change_checked,
        )
        for column, value in enumerate(values, 1):
            _cell(note, row, column, value)
    note.column_dimensions["B"].width = 65
    note.column_dimensions["C"].width = 80
    for letter in "ADEF":
        note.column_dimensions[letter].width = 22
    add_excel_notes(note)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _docx(items: list[ExportItem], stamp: str, excerpts: bool = False) -> bytes:
    book = WordDocument()
    book.add_heading("有据导出 · 数据说明", 0)
    book.add_paragraph(_notice(excerpts))
    book.add_paragraph(DISCLAIMER)
    book.add_paragraph(SEARCH_NOTICE)
    book.add_paragraph(f"导出时间（北京时间）：{stamp}")
    sources = book.add_table(rows=1, cols=6)
    for cell, title in zip(
        sources.rows[0].cells,
        ("序号", "文件编号", "来源网站", "来源核对日期", "截止日期核对日期", "变更核对日期"),
        strict=True,
    ):
        cell.text = title
    for item in items:
        values = (
            str(item.number),
            str(item.doc_id),
            item.sites,
            item.source_checked,
            item.deadline_checked,
            item.change_checked,
        )
        for cell, value in zip(sources.add_row().cells, values, strict=True):
            cell.text = value
    for item in items:
        book.add_heading(f"{item.number}. {item.title}", level=1)
        table = book.add_table(rows=0, cols=2)
        fields = (
            ("序号", str(item.number)),
            ("标题", item.title),
            ("类型", item.kind),
            ("名单状态", item.status),
            ("发文机关", item.issuer),
            ("文号", item.doc_number),
            ("发布日期", item.published),
            ("地区", item.region),
            (DEADLINE_HEADER, item.deadline),
            ("截止原文句子", item.deadline_quote),
            ("日期角色", item.date_role),
            ("支持方式", item.support),
            ("未核项", item.unchecked),
            ("来源网站", item.sites),
            ("来源核对日期", item.source_checked),
            ("截止日期核对日期", item.deadline_checked),
            ("变更核对日期", item.change_checked),
        )
        for label, value in fields:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = value
        book.add_paragraph("原文链接：")
        for url in item.urls:
            book.add_paragraph(url)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


# 导出在内存里生成文件：每个用户 10 分钟最多导出这么多次（YOUJU_EXPORTS_PER_10MIN），
# 每个实例 10 分钟最多导出这么多次、同时最多生成 4 份——资源用量有硬上限（api 服务最多 3 个实例）
_EXPORT_WINDOW = 600.0
_EXPORTS_PER_INSTANCE = 300
_EXPORT_KEYS_MAX = 10000
_EXPORT_SLOTS_WAIT = 10.0
_export_times: OrderedDict[str, deque[float]] = OrderedDict()
_instance_times: deque[float] = deque()
_export_guard = threading.Lock()
_export_slots = threading.BoundedSemaphore(4)


def _caller(request: Request) -> str:
    """限流按谁算：小程序用户按 openid；取不到（本机开发）按来源地址。"""
    from youju.api.me import request_openid

    try:
        return "u:" + request_openid(request)
    except HTTPException:
        return "ip:" + (request.client.host if request.client else "?")


def _drop_old(times: deque[float], now: float) -> None:
    while times and now - times[0] > _EXPORT_WINDOW:
        times.popleft()


def _allow_export(caller: str, now: float) -> str | None:
    """能导出返回 None，不能就返回原因：user（这个用户导出太频繁）/ busy（这个实例导出太多）。"""
    limit = get_settings().exports_per_10min
    with _export_guard:
        _drop_old(_instance_times, now)
        if len(_instance_times) >= _EXPORTS_PER_INSTANCE:
            return "busy"
        times = _export_times.get(caller) or deque()
        _drop_old(times, now)
        if len(times) >= limit:
            return "user"  # 被拒的不更新先后顺序，免得借此把别人的记录挤出去
        times.append(now)
        _export_times[caller] = times
        _export_times.move_to_end(caller)  # 放到最后：按最近成功导出的先后排
        _instance_times.append(now)
        # 记录太多时淘汰最久没导出的（不整个清空，别人没法借此重置某个用户的次数）
        while len(_export_times) > _EXPORT_KEYS_MAX:
            _export_times.popitem(last=False)
        return None


def _preview_content(items: list[ExportItem], excerpts: bool) -> dict:
    # 哈希只绑定本次目录字段和说明，不冒充正文/附件的不可变存档。
    return {"items": [asdict(item) for item in items], "notice": _notice(excerpts)}


def _revision(items: list[ExportItem], excerpts: bool) -> str:
    data = json.dumps(_preview_content(items, excerpts), ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(data.encode()).hexdigest()


@native_router.post("/exports/preview")
def preview_documents(
    payload: PreviewIn,
    session: SessionDep,
    request: Request,
    excerpts: Annotated[int | None, Query(ge=1, le=1)] = None,
) -> Response:
    # 查看目录不消耗用户保存文件的次数；实例总量仍有上限。
    refused = _allow_export("preview:" + _caller(request), clock.monotonic())
    if refused:
        raise HTTPException(status_code=429, detail="查看结果太频繁，请稍后重试")
    if not _export_slots.acquire(timeout=_EXPORT_SLOTS_WAIT):
        raise HTTPException(status_code=503, detail="正在整理结果，请稍后重试")
    try:
        excerpts = excerpts == 1 or payload.excerpts == 1
        items = _items(session, payload.ids, excerpts)
        response = JSONResponse(
            content={
                "kind": "catalogue",
                "title": "本次所选文件目录",
                **_preview_content(items, excerpts),
                "revision": _revision(items, excerpts),
                "total": len(items),
                "generatedAt": datetime.now(TZ).isoformat(),
            },
            headers={"Cache-Control": "no-store"},
        )
        if len(response.body) > 800 * 1024:
            raise HTTPException(status_code=413, detail="结果较多，请减少所选文件后重新查看")
        return response
    finally:
        _export_slots.release()


@router.post("/exports")
def export_documents(
    payload: ExportIn,
    session: SessionDep,
    request: Request,
    excerpts: Annotated[int | None, Query(ge=1, le=1)] = None,
) -> Response:
    refused = _allow_export(_caller(request), clock.monotonic())
    if refused == "user":
        raise HTTPException(status_code=429, detail="导出太频繁，请 10 分钟后再试")
    if refused == "busy":
        raise HTTPException(status_code=503, detail="现在导出的人多，请稍后再试")
    if not _export_slots.acquire(timeout=_EXPORT_SLOTS_WAIT):
        raise HTTPException(status_code=503, detail="现在导出的人多，请稍后再试")
    try:
        return _export(payload, session, excerpts == 1 or payload.excerpts == 1)
    finally:
        _export_slots.release()


def _export(payload: ExportIn, session: Session, excerpts: bool = False) -> Response:
    items = _items(session, payload.ids, True) if excerpts else _items(session, payload.ids)
    if payload.revision and payload.revision != _revision(items, excerpts):
        raise HTTPException(status_code=409, detail="目录数据已更新，请重新查看结果后再保存")
    stamp = datetime.now(TZ).strftime("%Y-%m-%d %H:%M:%S")
    if payload.format == "xlsx":
        data = _xlsx(items, stamp, True) if excerpts else _xlsx(items, stamp)
    else:
        data = _docx(items, stamp, True) if excerpts else _docx(items, stamp)
    if len(data) >= MAX_BYTES:
        raise HTTPException(status_code=413, detail="导出文件超过 1000 KiB，请减少文件数量")
    filename = f"有据导出_{stamp[:10]}.{payload.format}"
    mime = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if payload.format == "xlsx"
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    return Response(
        content=data,
        media_type=mime,
        headers={
            "Content-Disposition": (
                f"attachment; filename=youju-export_{stamp[:10]}.{payload.format}; "
                f"filename*=UTF-8''{quote(filename)}"
            )
        },
    )


router.include_router(native_router)
