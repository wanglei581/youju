"""名单比对后台引擎。

这个模块只做内部批处理：读取客户名单、按查名单的公开口径查询、生成结果和导出文件。
不保存客户上传内容，也不负责 HTTP 接口或对象存储。
"""

from __future__ import annotations

import csv
import re
import unicodedata
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from posixpath import normpath
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from youju.classify import effective_kind, list_status
from youju.export_notes import add_excel_notes
from youju.lists.lookup import ListLookupMatch, lookup_list_matches, public_condition
from youju.lists.rules import is_person_query, normalize_credit_code, normalize_name
from youju.models import Document, DocumentDetail, ListEntry, Source

MAX_INPUT_ROWS = 500
MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_XLSX_ROWS = 502
MAX_XLSX_COLUMNS = 64
MAX_XLSX_CELLS = MAX_XLSX_ROWS * MAX_XLSX_COLUMNS
MAX_XML_ELEMENTS = 100_000
MAX_ZIP_ENTRIES = 512
_NAME_HEADERS = {"企业名称", "单位名称", "公司名称"}
_CODE_HEADERS = {"统一社会信用代码", "信用代码"}
_YEAR = re.compile(r"(?<!\d)(20\d{2})(?!\d)")


class CompareInputError(ValueError):
    """客户名单不符合批处理限制。"""


@dataclass(frozen=True)
class RemovedInput:
    value: str
    reason: str


@dataclass(frozen=True)
class ListMatch:
    """一份去重后的名单通知。"""

    document_id: int
    title: str
    publish_date: date | None
    status: str | None
    issuer: str
    province: str
    doc_number: str
    url: str
    list_type: str
    year: int | None

    def as_dict(self) -> dict:
        return {
            "document_id": self.document_id,
            "title": self.title,
            "publish_date": self.publish_date.isoformat() if self.publish_date else None,
            "status": self.status,
            "issuer": self.issuer,
            "province": self.province,
            "doc_number": self.doc_number,
            "url": self.url,
            "list_type": self.list_type,
            "year": self.year,
        }


@dataclass
class CompareRow:
    input_value: str
    normalized_value: str
    input_kind: str
    status: str
    matches: list[ListMatch] = field(default_factory=list)
    note: str = ""

    def as_dict(self) -> dict:
        return {
            "input": self.input_value,
            "normalized": self.normalized_value,
            "kind": self.input_kind,
            "status": self.status,
            "matches": [m.as_dict() for m in self.matches],
            "note": self.note,
        }


@dataclass
class CompareReport:
    rows: list[CompareRow]
    removed: list[RemovedInput]
    data_version: datetime | None

    @property
    def detail_row_count(self) -> int:
        """导出明细应写出的行数（每个未命中客户也占一行）。"""
        return sum(max(1, len(row.matches)) for row in self.rows)

    def as_dict(self) -> dict:
        return {
            "rows": [row.as_dict() for row in self.rows],
            "removed": [{"value": r.value, "reason": r.reason} for r in self.removed],
            "data_version": self.data_version.isoformat() if self.data_version else None,
        }


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_key(value: object) -> str:
    return normalize_name(_text(value)).replace("：", "").replace(":", "")


def xlsx_sheet_paths(archive: ZipFile) -> list[str]:
    """按工作簿关系寻找工作表，路径不假定为 worksheets/sheet1.xml。"""
    names = set(archive.namelist())
    required = {"[Content_Types].xml", "xl/workbook.xml", "xl/_rels/workbook.xml.rels"}
    if not required <= names:
        return []
    types = ElementTree.fromstring(archive.read("[Content_Types].xml"))
    if not any(
        item.get("PartName") == "/xl/workbook.xml"
        and item.get("ContentType")
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
        for item in types
    ):
        return []
    workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    relations = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        item.get("Id"): normpath(
            item.get("Target", "").lstrip("/")
            if item.get("Target", "").startswith("/")
            else "xl/" + item.get("Target", "")
        )
        for item in relations
        if item.get("Type", "").endswith("/worksheet") and item.get("TargetMode") != "External"
    }
    paths = [
        targets.get(
            sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        )
        for sheet in workbook.findall(".//{*}sheet")
    ]
    return [path for path in paths if path in names]


def validate_xlsx(content: bytes) -> bool:
    """识别真实 XLSX；解压前检查总大小，按关系路径检查实际单元格。"""
    if not content.startswith(b"PK\x03\x04"):
        return False
    try:
        archive = ZipFile(BytesIO(content))
    except BadZipFile:
        return False
    with archive:
        if len(archive.infolist()) > MAX_ZIP_ENTRIES:
            raise CompareInputError("表格包含过多文件，请减少内容后重试")
        if sum(item.file_size for item in archive.infolist()) > MAX_FILE_BYTES:
            raise CompareInputError("表格解压后过大，请减少内容后重试")
        try:
            paths = xlsx_sheet_paths(archive)
        except ElementTree.ParseError:
            return False
        if not paths:
            return False
        cells = elements = 0
        for path in paths:
            rows = 0
            with archive.open(path) as source:
                for _, element in ElementTree.iterparse(source):
                    elements += 1
                    if elements > MAX_XML_ELEMENTS:
                        raise CompareInputError("表格内容过多，请减少内容后重试")
                    tag = element.tag.rsplit("}", 1)[-1]
                    if tag == "row":
                        rows += 1
                        row_ref = element.get("r", "0")
                        # 行号先限位数再转整数，超长数字串不进 int()
                        if (
                            rows > MAX_XLSX_ROWS
                            or not re.fullmatch(r"\d{1,7}", row_ref)
                            or int(row_ref) > MAX_XLSX_ROWS
                        ):
                            raise CompareInputError("表格行号过大，请移除多余空行后重试")
                    if tag == "c":
                        cells += 1
                        reference = element.get("r", "")
                        # Excel 最多到 XFD 列（3 个字母）、行号不超过 7 位，先限长度再计算
                        match = re.fullmatch(r"([A-Z]{1,3})([1-9]\d{0,6})", reference)
                        column = 0
                        if match:
                            for ch in match[1]:
                                column = column * 26 + ord(ch) - 64
                        if not match or column > MAX_XLSX_COLUMNS:
                            raise CompareInputError(
                                "表格列数过多或单元格坐标无效，请只保留客户字段"
                            )
                        if int(match[2]) > MAX_XLSX_ROWS:
                            raise CompareInputError("表格行号过大，请移除多余空行后重试")
                        if cells > MAX_XLSX_CELLS:
                            raise CompareInputError("表格单元格过多，请减少内容后重试")
                    element.clear()
        return True


def _file_rows(path: Path, *, content: bytes | None = None) -> list[list[str]]:
    if content is None and not path.is_file():
        raise CompareInputError(f"文件不存在：{path}")
    if (len(content) if content is not None else path.stat().st_size) > MAX_FILE_BYTES:
        raise CompareInputError("客户名单文件不能超过 2MB")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        raw = content if content is not None else path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("gbk")
            except UnicodeDecodeError as exc:
                raise CompareInputError("CSV 不是 UTF-8 或 GBK 编码") from exc
        try:
            return [
                [_text(cell) for cell in row]
                for row in csv.reader(StringIO(text, newline=""), strict=True)
            ]
        except csv.Error as exc:
            raise CompareInputError("CSV 格式有误，请检查引号是否闭合及分隔符") from exc
    if suffix != ".xlsx":
        raise CompareInputError("只支持 .xlsx 和 .csv 文件")
    try:
        raw = content if content is not None else path.read_bytes()
        if not validate_xlsx(raw):
            raise CompareInputError("无法读取 XLSX 文件")
        book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except CompareInputError:
        raise
    except Exception as exc:  # openpyxl 的异常类型随版本变化
        raise CompareInputError("无法读取 XLSX 文件") from exc
    try:
        sheet = book.worksheets[0]
        # dimension 不可信；同时给解析器明确列上限和迭代上限。
        sheet.reset_dimensions()
        rows = []
        for index, row in enumerate(sheet.iter_rows(values_only=True, max_col=MAX_XLSX_COLUMNS)):
            if index >= MAX_XLSX_ROWS:
                raise CompareInputError("表格行数过多，请减少内容后重试")
            rows.append([_text(cell) for cell in row])
        return rows
    finally:
        book.close()


def read_customer_file(path: str | Path, *, content: bytes | None = None) -> list[str]:
    """按表头读取客户名称或信用代码列；没有表头时取第一列。"""
    rows = _file_rows(Path(path), content=content)
    if not rows:
        return []
    header = rows[0]
    name_col = next(
        (i for i, value in enumerate(header) if _header_key(value) in _NAME_HEADERS), None
    )
    code_col = next(
        (i for i, value in enumerate(header) if _header_key(value) in _CODE_HEADERS), None
    )
    column = name_col if name_col is not None else code_col
    data = rows[1:] if column is not None else rows
    if column is None:
        column = 0
    values = [_text(row[column]) if column < len(row) else "" for row in data]
    if len(values) > MAX_INPUT_ROWS:
        raise CompareInputError("客户名单最多 500 行")
    return values


def _looks_like_code(value: str) -> bool:
    normalized = re.sub(r"\s+", "", unicodedata.normalize("NFKC", value)).upper()
    return len(normalized) == 18 and bool(re.fullmatch(r"[0-9A-Z]+", normalized))


def _input_rows(values: Iterable[str]) -> tuple[list[tuple[str, str, str]], list[RemovedInput]]:
    rows: list[tuple[str, str, str]] = []
    removed: list[RemovedInput] = []
    seen: set[tuple[str, str]] = set()
    all_values = list(values)
    if len(all_values) > MAX_INPUT_ROWS:
        raise CompareInputError("客户名单最多 500 行")
    for raw in all_values:
        original = _text(raw)
        if not original:
            removed.append(RemovedInput(original, "空行"))
            continue
        code = normalize_credit_code(original)
        if code:
            key = ("code", code)
            normalized, kind = code, "信用代码"
        elif _looks_like_code(original):
            key = ("invalid_code", unicodedata.normalize("NFKC", original).upper())
            normalized, kind = key[1], "信用代码"
        else:
            normalized, kind = normalize_name(original), "企业名称"
            key = ("name", normalized)
        if key in seen:
            removed.append(RemovedInput(original, "重复行"))
            continue
        seen.add(key)
        rows.append((original, normalized, kind))
    return rows, removed


def _make_match(
    entry: ListLookupMatch,
    document: tuple[int, str, date | None, str, str | None, str, str, str | None, str | None],
) -> ListMatch:
    (
        document_id,
        title,
        publish_date,
        url,
        issuer,
        doc_number,
        province,
        dept_full_name,
        kind_override,
    ) = document
    kind = effective_kind(title, kind_override)
    issuer = issuer or dept_full_name or ""
    return ListMatch(
        document_id=document_id,
        title=title,
        publish_date=publish_date,
        status=list_status(title, kind),
        issuer=issuer,
        province=province,
        doc_number=doc_number or "",
        url=url,
        list_type=entry.category or "其他",
        year=int(_YEAR.search(title).group(1)) if _YEAR.search(title) else None,
    )


def _canonical_metadata(session: Session, ids: set[int]):
    """取归并主文件的必要题录列；展示链接与查名单接口保持一致。"""
    if not ids:
        return {}
    rows = session.execute(
        select(
            Document.id,
            Document.title,
            Document.publish_date,
            Document.url,
            DocumentDetail.issuer,
            DocumentDetail.doc_number,
            Source.province,
            Source.dept_full_name,
            Document.kind_override,
        )
        .join(Source, Source.id == Document.source_id)
        .outerjoin(DocumentDetail, DocumentDetail.document_id == Document.id)
        .where(Document.id.in_(ids))
    ).all()
    return {row.id: tuple(row) for row in rows}


def _beijing_minute(value: datetime) -> str:
    """导出表头用北京时间、精确到分钟；数据库取回的无时区时间按 UTC 处理。"""
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M")


def _data_version(session: Session) -> datetime | None:
    return session.scalar(
        select(func.max(Document.last_seen_at))
        .join(Source, Source.id == Document.source_id)
        .where(
            public_condition(),
            or_(Document.kind_override.is_(None), Document.kind_override == "名单"),
            Document.id.in_(select(ListEntry.document_id)),
        )
    )


def compare_lists(session: Session, values: Iterable[str]) -> CompareReport:
    """比较客户名单，返回稳定的逐客户报告。"""
    inputs, removed = _input_rows(values)
    result: list[CompareRow] = []
    for original, normalized, kind in inputs:
        if kind == "企业名称" and is_person_query(normalized):
            result.append(
                CompareRow(
                    original, normalized, kind, "像个人姓名，未查询", note="像个人姓名，未查询"
                )
            )
            continue
        if kind == "信用代码" and normalize_credit_code(normalized) is None:
            result.append(
                CompareRow(
                    original, normalized, kind, "代码格式不对", note="统一社会信用代码校验位不正确"
                )
            )
            continue
        lookup = lookup_list_matches(session, normalized, scan_limit=None, result_limit=None)
        metadata = _canonical_metadata(session, set(lookup.master_ids))
        matches = [
            _make_match(entry, metadata[entry.master_id])
            for entry in lookup.matches
            if entry.master_id in metadata
        ]
        matches.sort(
            key=lambda match: (
                match.publish_date is not None,
                match.publish_date or date.min,
                match.document_id,
            ),
            reverse=True,
        )
        status = "多个同名" if lookup.ambiguous else "命中" if matches else "未命中"
        note = "名称对应多个信用代码，请核对" if lookup.ambiguous else ""
        result.append(CompareRow(original, normalized, kind, status, matches, note))
    return CompareReport(result, removed, _data_version(session))


def _safe_cell(value: object) -> object:
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _date_text(value: date | None) -> str:
    return value.isoformat() if value else ""


def export_compare_excel(report: CompareReport, output: str | Path) -> int:
    """导出明细和汇总，返回明细数据行数。"""
    book = Workbook()
    detail = book.active
    detail.title = "明细"
    version = _beijing_minute(report.data_version) if report.data_version else "暂无"
    note = f"第三方整理 · 非政府平台；名单状态照录标题用词，以原文为准；数据截至 {version}"
    if report.removed:
        removed_counts = defaultdict(int)
        for item in report.removed:
            removed_counts[item.reason] += 1
        note += "；已去掉" + "、".join(
            f"{reason}{count}行" for reason, count in removed_counts.items()
        )
    detail.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
    detail.cell(1, 1, _safe_cell(note))
    headers = [
        "客户输入",
        "比对用的名称",
        "结果",
        "名单标题",
        "发布日期",
        "名单状态",
        "发文机关",
        "省份",
        "文件编号",
        "原文链接",
        "名单类型",
        "年度",
        "说明",
    ]
    for col, header in enumerate(headers, 1):
        detail.cell(2, col, header)
    row_count = 0
    summary_counts: dict[tuple[str, int | None], set[str]] = defaultdict(set)
    for index, row in enumerate(report.rows):
        matches = row.matches or [None]
        for match in matches:
            values = [
                row.input_value,
                row.normalized_value,
                row.status,
                match.title if match else "",
                _date_text(match.publish_date) if match else "",
                match.status if match else "",
                match.issuer if match else "",
                match.province if match else "",
                match.doc_number if match else "",
                match.url if match else "",
                match.list_type if match else "",
                match.year if match else "",
                row.note,
            ]
            for col, value in enumerate(values, 1):
                detail.cell(row_count + 3, col, _safe_cell(value))
            if match and row.status in ("命中", "多个同名"):
                summary_counts[(match.list_type, match.year)].add(str(index))
            row_count += 1
    summary = book.create_sheet("汇总")
    summary.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    summary.cell(1, 1, _safe_cell(note))
    for col, header in enumerate(("名单类型", "年度", "命中家数"), 1):
        summary.cell(2, col, header)
    for row_number, ((list_type, year), customers) in enumerate(
        sorted(
            summary_counts.items(),
            key=lambda item: (item[0][0], item[0][1] is None, item[0][1] or 0),
        ),
        3,
    ):
        summary.cell(row_number, 1, _safe_cell(list_type))
        summary.cell(row_number, 2, year or "")
        summary.cell(row_number, 3, len(customers))
    add_excel_notes(book.create_sheet("数据说明"))
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    book.save(output_path)
    return row_count


# 给调用方和测试保留更直观的短名称。
compare = compare_lists
compare_customer_list = compare_lists
read_customer_list = read_customer_file
load_customer_file = read_customer_file
export_excel = export_compare_excel
write_excel = export_compare_excel
