"""R1.5 工具任务队列、额度和结果文件处理。"""

from __future__ import annotations

import copy
import json
import logging
import math
import re
import tempfile
import time as monotonic_clock
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from threading import Lock
from typing import Any
from zipfile import ZipFile
from zoneinfo import ZoneInfo

import httpx
from openpyxl import load_workbook
from sqlalchemy import and_, delete, func, or_, select, text, update
from sqlalchemy.dialects.mysql import insert
from sqlalchemy.orm import Session, sessionmaker

from youju import access_log, safety_words, wx_security
from youju.ai import advisor_endpoint, clause, heartbeat, interpret, result_safety
from youju.ai.client import AIError, ChatClient
from youju.ai.skills import AI_KINDS, get_skill
from youju.ai.usage import record_usage
from youju.backup import BackupError, _open_api, upload
from youju.compare import CompareInputError, compare_lists, export_compare_excel
from youju.config import Settings, get_settings
from youju.models import (
    AdvisorFileTrash,
    AdvisorTask,
    AdvisorUsage,
    AISafetyHit,
    AIUsage,
    Document,
    User,
    WorkbenchPass,
)
from youju.track import export_track_excel, track_clients

log = logging.getLogger("youju.advisor")
BEIJING = ZoneInfo("Asia/Shanghai")
MAX_ROWS = 50
MAX_BODY_BYTES = 100 * 1024
MAX_UPLOAD_BYTES = 60 * 1024
ADVISOR_AGREEMENT_VERSION = "advisor-2026-11-v2"
TASK_TIMEOUT = timedelta(minutes=5)
STALE_RUNNING = timedelta(minutes=10)
# advisor-run 每分钟一轮（crontab）。队里有 AI 任务就一直领（和原来一样领到空）；
# 队里空了，在开始后的 POLL_WINDOW 秒里每隔 POLL_INTERVAL 秒再看一次，用户点了
# 几秒内就开始，不用等下一分钟（10/27 实测排队 48.6 秒）。过了这 50 秒、别的线程
# 还在做任务时也接着看（反正要等它做完这一轮才结束）；全都空闲就退出。50 秒给
# 启动、退出留余量，空闲时不和下一分钟那一轮撞上。开始满 CLAIM_LIMIT 秒不再领
# 新的，手上的做完就退出，一轮最长约 4 分钟加一个任务（任务最长 5 分钟），
# 不到 job.sh 当作"卡住了"的 10 分钟
POLL_WINDOW = 50.0
POLL_INTERVAL = 3.0
CLAIM_LIMIT = 240.0
# 处理中超过这么久的一定已经没人在跑（清理过期任务时才用，见 cleanup_advisor）
ABANDONED = timedelta(days=1)
INPUT_RETENTION = timedelta(days=7)
FILE_RETENTION = timedelta(days=30)
REFUSAL_CACHE = timedelta(days=1)
AI_CAP_LOCK = "youju_ai_advisor_daily_cap"
AI_CAP_LOCAL_LOCK = Lock()
AI_CAP_MESSAGE = "今天的 AI 使用次数已满，请明天再试"
# 任务排队时失去授权（AI 关了、移出白名单）退回次数时记的原因；接口据此给错误代码 AI_OFF
AI_OFF_MESSAGE = "功能未开放"
# 任务结果里的错误代码（D-075）：小程序按代码写原因，不靠比对文字
ERROR_CODES = {
    AI_CAP_MESSAGE: "AI_DAILY_CAP",
    AI_OFF_MESSAGE: "AI_OFF",
    heartbeat.UNAVAILABLE: "AI_UNAVAILABLE",
    # 内容安全（门槛第 4 道）：结果没通过、检测接口出错。小程序没有这两个代码的固定说法时照写 error
    result_safety.UNSAFE_MESSAGE: "AI_UNSAFE",
    result_safety.REMOVED_MESSAGE: "AI_UNSAFE",
    result_safety.CHECK_DOWN_MESSAGE: "AI_CHECK_DOWN",
}


class AdvisorBusy(Exception):
    """全站额度锁未及时获得。"""


class AICapReached(Exception):
    """实际模型调用已达到当天上限。"""


@contextmanager
def ai_cap_lock(session: Session):
    """锁在独立连接上；额度事务提交后才释放命名锁。"""
    with AI_CAP_LOCAL_LOCK, session.get_bind().connect() as conn:
        if conn.scalar(text("SELECT GET_LOCK(:name, 2)"), {"name": AI_CAP_LOCK}) != 1:
            raise AdvisorBusy
        try:
            yield
        finally:
            conn.execute(text("SELECT RELEASE_LOCK(:name)"), {"name": AI_CAP_LOCK})


def global_ai_calls(session: Session, day: date) -> int:
    return int(
        session.scalar(
            select(func.coalesce(func.sum(AIUsage.calls), 0)).where(
                AIUsage.day == day, AIUsage.task.in_(AI_KINDS)
            )
        )
        or 0
    )


def reserve_ai_call(session: Session, kind: str, model: str, stamp: datetime, cap: int) -> None:
    """调用前持久化一次不可退的调用；进程中断时宁可少调用。"""
    day = usage_day(stamp)
    with ai_cap_lock(session):
        if global_ai_calls(session, day) >= cap:
            raise AICapReached
        table = AIUsage.__table__
        stmt = insert(table).values(
            day=day, task=kind, model=model, calls=1, errors=0, prompt_tokens=0, completion_tokens=0
        )
        session.execute(stmt.on_duplicate_key_update(calls=table.c.calls + 1))
        session.commit()


def queue_info(session: Session, task: AdvisorTask, concurrency: int) -> dict[str, int]:
    """只统计排在本任务前的 AI 任务；ETA 包含本任务完成所需时间。"""
    ahead = int(
        session.scalar(
            select(func.count(AdvisorTask.id)).where(
                AdvisorTask.kind.in_(AI_KINDS),
                AdvisorTask.status == "queued",
                AdvisorTask.id < task.id,
            )
        )
        or 0
    )
    durations = session.execute(
        select(AdvisorTask.started_at, AdvisorTask.finished_at)
        .where(
            AdvisorTask.kind.in_(AI_KINDS),
            AdvisorTask.status == "done",
            AdvisorTask.started_at.is_not(None),
            AdvisorTask.finished_at.is_not(None),
        )
        .order_by(AdvisorTask.finished_at.desc())
        .limit(20)
    ).all()
    average = (
        (
            sum(max(0, (_utc(end) - _utc(start)).total_seconds()) for start, end in durations)
            / len(durations)
        )
        if durations
        else 30
    )
    return {"queuePosition": ahead, "etaSeconds": math.ceil((ahead // concurrency + 1) * average)}


def business_now(value: datetime | None = None) -> datetime:
    value = value or datetime.now(UTC)
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(BEIJING)


def _utc(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value.replace(tzinfo=UTC)
    return value.astimezone(UTC)


def _friendly_error(exc: Exception) -> str:
    if isinstance(exc, CompareInputError):
        return str(exc)
    if isinstance(exc, BackupError):
        return "结果文件暂时无法保存，请稍后重试"
    if isinstance(exc, TimeoutError):
        return "任务处理超时，请稍后重试"
    return "任务处理失败，请稍后重试"


def result_path(user: User, task_id: int, *, engine=None) -> str:
    """对象存储路径只使用程序生成的去标识用户编号。"""
    return f"advisor/out/{access_log.log_uid(user.openid, engine)}/{task_id}.xlsx"


def retained_input(task: AdvisorTask, stamp: datetime | None = None) -> dict | None:
    """清理任务尚未执行时也不展示或导出超过七天的企业输入。"""
    stamp = _utc(stamp or datetime.now(UTC))
    if task.created_at is not None and _utc(task.created_at) <= stamp - INPUT_RETENTION:
        return None
    return task.input_json


def task_dict(
    task: AdvisorTask, *, include_result: bool = True, stamp: datetime | None = None
) -> dict[str, Any]:
    """API 公共字段；不返回 input_rows，避免名单内容进入查询响应。"""
    result = {
        "id": task.id,
        "kind": task.kind,
        "status": task.status,
        "rowCount": task.row_count,
        "attempts": task.attempts,
        "createdAt": task.created_at,
        "startedAt": task.started_at,
        "finishedAt": task.finished_at,
        "expiresAt": task.expires_at,
        "error": task.error,
        # D-075：AI 任务失败原因的代码（全站上限、AI 关了），小程序按代码写原因，不靠比对文字
        "errorCode": ERROR_CODES.get(task.error) if task.kind in AI_KINDS and task.error else None,
        "usageSource": task.usage_source,
    }
    if include_result:
        result.update(
            {
                "resultRows": task.result_rows,
                "hasFile": bool(task.result_file_id),
            }
        )
    if task.kind in AI_KINDS and include_result:
        stored = (
            clause.with_company(task.result_json, retained_input(task, stamp))
            if task.kind == "clause" and task.result_json is not None
            else task.result_json
        )
        result["result"] = (
            {k: v for k, v in stored.items() if k != "dropped_reasons"}
            if task.result_json is not None
            else None
        )
    return result


def usage_day(value: datetime) -> date:
    """把任务时间转换为额度账本使用的北京时间日期。"""
    return _utc(value).astimezone(BEIJING).date()


def usage_limit(kind: str) -> int:
    return 5 if kind in AI_KINDS else 3


# 加量体验期内的提交（两种数据工具合计）记在这个账本上，按北京日期计，不进每月免费汇总（boost.py）
TRIAL_USAGE = "trial"


def ledger_kind(task: AdvisorTask) -> str:
    """任务占用的是哪本账：体验期的提交记在体验账本，其余按任务类型记。"""
    return TRIAL_USAGE if task.usage_source == "trial" else task.kind


def global_ai_usage(session: Session, day: date) -> int:
    """当天已占用的解读和条款对照次数；失败退款已从账本扣除。"""
    return session.scalar(
        select(func.coalesce(func.sum(AdvisorUsage.used_count), 0)).where(
            AdvisorUsage.day == day, AdvisorUsage.kind.in_(AI_KINDS)
        )
    )


def usage_counts(
    session: Session, user_id: int, day: date, *, lock: bool = False
) -> dict[str, int]:
    """日账本保持不变；数据工具汇总北京时间自然月，解读工具仍按日。"""
    month = day.replace(day=1)
    next_month = (month + timedelta(days=32)).replace(day=1)
    stmt = select(AdvisorUsage.kind, AdvisorUsage.used_count).where(
        AdvisorUsage.user_id == user_id,
        or_(
            and_(
                AdvisorUsage.kind.in_(("compare", "track")),
                AdvisorUsage.day >= month,
                AdvisorUsage.day < next_month,
            ),
            and_(AdvisorUsage.kind.not_in(("compare", "track")), AdvisorUsage.day == day),
        ),
    )
    # 占用额度时必须读最新值；普通快照可能早于等待用户行锁的并发请求。
    if lock:
        stmt = stmt.with_for_update()
    counts = {}
    for kind, used in session.execute(stmt):
        counts[kind] = counts.get(kind, 0) + used
    return counts


def reserve_usage(
    session: Session, user_id: int, kind: str, day: date, *, limit: int | None = None
) -> bool:
    """在调用方已锁定用户行后占用一次额度。kind 是账本类型（体验账本是 TRIAL_USAGE）；
    limit 不传按 usage_limit（申报季加量、加量体验由调用方传入当时的上限）。"""
    if limit is None:
        limit = usage_limit(kind)
    if usage_counts(session, user_id, day, lock=True).get(kind, 0) >= limit:
        return False
    row = session.scalar(
        select(AdvisorUsage)
        .where(
            AdvisorUsage.user_id == user_id,
            AdvisorUsage.kind == kind,
            AdvisorUsage.day == day,
        )
        .with_for_update()
    )
    if row is None:
        row = AdvisorUsage(user_id=user_id, kind=kind, day=day, used_count=0)
        session.add(row)
        session.flush()
    row.used_count += 1
    return True


def release_usage(session: Session, task: AdvisorTask) -> None:
    """释放一次已占用额度；重复释放安全。"""
    if task.usage_source == "pass" or task.usage_released:
        return
    # 原子标记保证两个释放路径至多退回一次，账本仍按原提交日期退款。
    changed = session.execute(
        update(AdvisorTask)
        .where(
            AdvisorTask.id == task.id,
            AdvisorTask.usage_released == False,  # noqa: E712
        )
        .values(usage_released=True)
    ).rowcount
    if not changed:
        return
    task.usage_released = True
    row = session.scalar(
        select(AdvisorUsage)
        .where(
            AdvisorUsage.user_id == task.user_id,
            AdvisorUsage.kind == ledger_kind(task),
            AdvisorUsage.day == usage_day(task.created_at),
        )
        .with_for_update()
    )
    if row is not None and row.used_count > 0:
        row.used_count -= 1


def _successful_ids(result: object, file_ids: list[str]) -> set[str]:
    """兼容旧测试替身返回数量；正式实现返回成功的文件 ID。"""
    if isinstance(result, (list, tuple, set, frozenset)):
        return {str(value) for value in result if value in file_ids}
    if isinstance(result, int):
        return set(file_ids[: max(0, min(result, len(file_ids)))])
    return set()


def _record_trash(session: Session, file_ids: list[str], user_hash: str | None = None) -> None:
    if not file_ids:
        return
    existing = set(
        session.scalars(
            select(AdvisorFileTrash.file_id).where(AdvisorFileTrash.file_id.in_(file_ids))
        )
    )
    for file_id in dict.fromkeys(file_ids):
        if file_id not in existing:
            session.add(AdvisorFileTrash(file_id=file_id, user_hash=user_hash))


def _delete_and_record(
    session: Session,
    file_ids: list[str],
    *,
    settings: Settings,
    client: httpx.Client | None,
    user_hash: str | None = None,
) -> set[str]:
    """删除对象存储文件；失败的 ID 留在待删表，调用方只清成功引用。"""
    if not file_ids:
        return set()
    cloud_env = getattr(settings, "cloud_env", None)
    successful: set[str] = set()
    if cloud_env:
        try:
            result = delete_result_files(client, cloud_env, file_ids)
            successful = _successful_ids(result, file_ids)
        except Exception:  # noqa: BLE001 - 存储删除失败必须留下待删记录
            successful = set()
    failed = [file_id for file_id in file_ids if file_id not in successful]
    _record_trash(session, failed, user_hash)
    return successful


def recover_stale(session: Session, stamp: datetime | None = None, condition=None) -> int:
    """实例重启后把运行超过十分钟的任务放回队列，不增加 attempts。"""
    stamp = _utc(stamp or datetime.now(UTC))
    stmt = update(AdvisorTask)
    if condition is not None:
        stmt = stmt.where(condition)
    count = session.execute(
        stmt.where(
            AdvisorTask.status == "running",
            AdvisorTask.started_at.is_not(None),
            AdvisorTask.started_at < stamp - STALE_RUNNING,
        ).values(status="queued", started_at=None, finished_at=None)
    ).rowcount
    session.commit()
    return count or 0


def _fail_revoked_interpret(session: Session, settings: Settings, stamp: datetime) -> int:
    """退回失去授权的任务额度；未超时的运行任务仍由领取它的进程负责。返回退回了几个。

    admin 这边处理不了 AI 任务（开关关着、白名单是空的，或对外 AI 的密钥没配；和心跳用同一个
    判断 heartbeat.worker_capable）时，排队的 AI 任务全部退回，用户看到"AI 暂时不可用"，个数和
    退回在同一个事务里记进心跳，给运维提醒（D-075）；只是这个人不在白名单里了（被移出），
    写"功能未开放"，不提醒。"""
    paused = not heartbeat.worker_capable(settings)
    condition = AdvisorTask.kind.in_(AI_KINDS) & (
        (AdvisorTask.status == "queued")
        | (
            (AdvisorTask.status == "running")
            & AdvisorTask.started_at.is_not(None)
            & (AdvisorTask.started_at < stamp - STALE_RUNNING)
        )
    )
    if not paused:
        condition = condition & AdvisorTask.user_id.not_in(interpret.worker_user_ids(settings))
    tasks = session.scalars(
        select(AdvisorTask)
        .where(condition)
        .order_by(AdvisorTask.id)
        .with_for_update(skip_locked=True)
    ).all()
    for task in tasks:
        task.status, task.finished_at = "failed", stamp
        task.error = heartbeat.UNAVAILABLE if paused else AI_OFF_MESSAGE
        release_usage(session, task)
    if paused:
        heartbeat.add_refunds(session, len(tasks), stamp)
    session.commit()
    return len(tasks)


def _run_engine(session: Session, task: AdvisorTask, output: Path) -> int:
    from youju.native_results import build

    if not isinstance(task.input_rows, list):
        raise CompareInputError("任务输入已过期，请重新提交")
    if task.kind == "compare":
        report = compare_lists(session, task.input_rows)
        task.result_document_ids = sorted(
            {m.document_id for row in report.rows for m in row.matches}
        )
        task.result_json = build(task.kind, report)
        return export_compare_excel(report, output)
    if task.kind == "track":
        report = track_clients(session, task.input_rows)
        task.result_document_ids = sorted({item.document_id for item in report.items})
        task.result_json = build(task.kind, report)
        return export_track_excel(report, output)
    raise CompareInputError("不支持的工具任务类型")


def legacy_result_document_ids(session: Session, client: httpx.Client, url: str) -> list[int]:
    """旧结果没有依赖清单：只在存在下架文件时从既有 Excel 的原文链接补齐。

    不用当前输入重算，避免漏掉历史结果中的文件；只取文件引用，不记录客户内容。
    """
    data = bytearray()
    with client.stream("GET", url, timeout=30) as response:
        response.raise_for_status()
        for chunk in response.iter_bytes():
            data.extend(chunk)
            if len(data) > 4 * 1024 * 1024:
                raise ValueError("结果文件过大")
    with ZipFile(BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
            raise ValueError("结果文件过大")
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        rows = sheet.iter_rows(min_row=2, values_only=True)
        headers = next(rows)
        column = headers.index("原文链接")
        urls = {str(row[column]) for row in rows if len(row) > column and row[column]}
        return list(session.scalars(select(Document.id).where(Document.url.in_(urls))))
    finally:
        book.close()


# 同一份内容没通过内容安全以后，多久内再提交直接按没通过处理
BLOCK_CACHE = timedelta(days=1)


def _recently_blocked(
    session: Session, task: AdvisorTask, skill, model: str, key: str, stamp: datetime
) -> bool:
    return (
        session.scalar(
            select(AdvisorTask.id)
            .where(
                AdvisorTask.kind == skill.kind,
                AdvisorTask.id != task.id,
                AdvisorTask.prompt_version == skill.prompt_version,
                AdvisorTask.model == model,
                AdvisorTask.cache_key == key,
                AdvisorTask.status == "failed",
                AdvisorTask.error == result_safety.UNSAFE_MESSAGE,
                AdvisorTask.finished_at > stamp - BLOCK_CACHE,
            )
            .limit(1)
        )
        is not None
    )


def sweep_results(session: Session, stamp: datetime) -> int:
    """词库更新以后，把还在保存期内、显示着的 AI 结果用现在的本地词库再查一遍（在本机比，
    不调微信）。命中的撤下、写处置记录；没命中的把标记改成现在的版本。没有标记的（检查上线以前
    生成的）读的时候本来就不显示，到期删除。每天的清理里调，不提交。返回撤下了几份。"""
    version = safety_words.LIBRARY.version
    removed = 0
    rows = session.scalars(
        select(AdvisorTask)
        .where(
            AdvisorTask.kind.in_(AI_KINDS),
            AdvisorTask.status == "done",
            AdvisorTask.result_json.is_not(None),
            # 只扫还在保存期内的（过期的在前面已经整行删了）
            or_(AdvisorTask.expires_at.is_(None), AdvisorTask.expires_at > stamp),
        )
        .with_for_update(skip_locked=True)
    ).all()
    for task in rows:
        result = task.result_json or {}
        mark = result.get("safety")
        if not isinstance(mark, dict) or mark.get("words") == version:
            continue
        field = get_skill(task.kind).result_field
        found = result_safety.local_hit(result_safety.quotes(result, field))
        if found is None:
            task.result_json = {**result, "safety": {**mark, "words": version}}
            continue
        result_safety.record_local(session, task, found, stamp=stamp, extra={"sweep": True})
        task.status, task.error = "failed", result_safety.REMOVED_MESSAGE
        task.result_json, task.result_rows, task.cache_key = None, 0, None
        removed += 1
    return removed


def _remove_copies(session: Session, task: AdvisorTask, source_id: int) -> None:
    """撤下没通过内容安全的那份缓存结果和它的复制件：复制结果不刷新生成时间，同类型、同缓存键、
    生成时间相同的就是同一份。"""
    source = session.get(AdvisorTask, source_id)
    generated = ((source.result_json or {}).get("generated_at")) if source else None
    if source is None or not source.cache_key or not generated:
        return
    for row in session.scalars(
        select(AdvisorTask)
        .where(
            AdvisorTask.kind == source.kind,
            AdvisorTask.cache_key == source.cache_key,
            AdvisorTask.status == "done",
            AdvisorTask.id != task.id,
        )
        .with_for_update()
    ):
        if (row.result_json or {}).get("generated_at") != generated:
            continue
        row.status, row.error = "failed", result_safety.REMOVED_MESSAGE
        row.result_json, row.result_rows, row.cache_key = None, 0, None


def process_ai(
    session: Session,
    task: AdvisorTask,
    *,
    settings: Settings,
    stamp: datetime,
    timeout: timedelta = TASK_TIMEOUT,
    chat_client: ChatClient | None = None,
    live_time: bool = False,
) -> bool:
    """拒答是有原因的完成结果；暂时性调用失败最多重排两次。"""
    if task.status not in ("running", "queued") or not interpret.enabled(
        settings, session.get(User, task.user_id)
    ):
        return False
    skill = get_skill(task.kind)
    started = monotonic_clock.monotonic()
    own_client = chat_client is None
    # 对外 AI 的模型（YOUJU_AI_ADVISOR_MODEL，没单独配就是 YOUJU_AI_MODEL）；也是结果缓存键的一部分
    advisor_url, advisor_key, model, advisor_extra = advisor_endpoint(settings)
    call_reserved = False
    call_stamp = stamp
    cached_from = None  # 结果是从哪个任务的缓存复制来的（内容安全没通过时要撤下那份）
    task.model, task.prompt_version = model, skill.prompt_version
    try:
        if not task.input_json:
            raise interpret.InterpretRefusal("任务输入已过期，请重新提交")
        snapshot, verified, key = interpret.prepare(
            session, task.input_json["doc_id"], model, skill=skill
        )
        task.snapshot, task.cache_key = snapshot, key
        task.result_document_ids = snapshot["document_ids"]
        if _recently_blocked(session, task, skill, model, key, stamp):
            # 同一份内容 1 天内没通过内容安全：直接按没通过处理，不再调模型和微信（不另写记录）
            raise result_safety.ResultBlocked
        cached = session.scalar(
            select(AdvisorTask)
            .where(
                AdvisorTask.kind == skill.kind,
                AdvisorTask.id != task.id,
                AdvisorTask.prompt_version == skill.prompt_version,
                AdvisorTask.model == model,
                AdvisorTask.cache_key == key,
                AdvisorTask.status == "done",
                AdvisorTask.result_json.is_not(None),
                AdvisorTask.finished_at > stamp - FILE_RETENTION,
                AdvisorTask.expires_at > stamp,
            )
            .order_by(AdvisorTask.finished_at.desc())
            .limit(1)
        )
        reusable = False
        if cached is not None and cached.result_json:
            generated = _utc(datetime.fromisoformat(cached.result_json["generated_at"]))
            # 复制结果不会刷新模型生成时间，避免连续命中把 30 天期限无限延长。
            # 模型拒答（逐字核对后一条不剩）只复用 1 天：挡住反复提交同一份触发模型调用，
            # 又不让一次偶然的拒答把这份通知锁 30 天（模型输出会波动，见 AI 评测第八节）。
            window = REFUSAL_CACHE if cached.result_json.get("refusal") else FILE_RETENTION
            cached_doc = (cached.snapshot or {}).get("doc_id")
            if generated > stamp - window and cached_doc:
                try:
                    _, _, cached_current_key = interpret.prepare(
                        session, cached_doc, model, skill=skill
                    )
                    # 缓存键相同只说明原文和核对一样；发给模型的脱敏后段落也要一样才复用，
                    # 免得脱敏规则改了以后旧摘录被新任务复制出来；旧缓存缺正文指纹或段落的不复用，
                    # 重新调用模型（D-073，Codex 第二、三轮）
                    reusable = cached_current_key == key and not interpret.body_changed(
                        cached.snapshot, snapshot
                    )
                except interpret.InterpretRefusal:
                    pass
        if reusable:
            cached_from = cached.id
            result = (
                clause.notification_result(cached.result_json)
                if skill.kind == "clause"
                else copy.deepcopy(cached.result_json)
            )
        else:
            if chat_client is None and not advisor_key:
                raise AIError("AI 服务尚未配置")
            if chat_client is None:
                chat_client = ChatClient(
                    advisor_url,
                    advisor_key,
                    model,
                    timeout=min(
                        getattr(settings, "ai_advisor_timeout_seconds", 180),
                        timeout.total_seconds(),
                    ),
                    max_retries=0,
                    extra_body=advisor_extra,
                    stream=bool(getattr(settings, "ai_advisor_stream", False)),
                )
            call_stamp = datetime.now(UTC) if live_time else stamp
            reserve_ai_call(session, skill.kind, model, call_stamp, settings.ai_advisor_daily_cap)
            call_reserved = True
            payload = chat_client.complete_json(
                skill.system_prompt,
                json.dumps({"paragraphs": snapshot["paragraphs"]}, ensure_ascii=False),
                max_tokens=12000,
            )
            if monotonic_clock.monotonic() - started > timeout.total_seconds():
                raise TimeoutError
            result = interpret.result_metadata(
                interpret.validate(payload, snapshot["paragraphs"], skill=skill),
                stamp=stamp,
                model=model,
                verified=verified,
                attachments=snapshot["attachments"],
                model_disclosure=interpret.disclosure(settings),
                skill=skill,
            )
            # 调用期间原文或核对状态可能变更；刷新后重新校验。D-073：正文变了不发布；
            # 只是核对或附件变了照常发布，"已核对"和附件提醒换成现在的
            session.flush()
            session.expire_all()
            current, current_verified, current_key = interpret.prepare(
                session, task.input_json["doc_id"], model, skill=skill
            )
            if key != current_key:
                if interpret.body_changed(snapshot, current):
                    raise interpret.InterpretRefusal(interpret.BODY_CHANGED)
                result = interpret.refresh_verification(
                    result, current_verified, current["attachments"]
                )
                task.snapshot, task.cache_key = current, current_key
        # 结果存进去之前过两层内容安全（D-093 门槛第 4 条）：复用缓存的也查；通过的记上标记
        safety = result_safety.check(
            session, task, result, field=skill.result_field, stamp=stamp, cached_from=cached_from
        )
        if safety is not None:
            result = {**result, "safety": safety}
        task.result_json = result
        task.result_rows = len(result[skill.result_field])
        task.status, task.finished_at, task.error = "done", stamp, None
        if result["refusal"]:
            release_usage(session, task)
        session.commit()
        return True
    except result_safety.ResultBlocked:
        # 没通过内容安全：不显示、不存、退回次数；处置记录已经加进会话，和任务一起提交
        if cached_from:
            # 复用的缓存没通过：这份结果以前给别人显示过，撤下它和它的复制件（生成时间相同的），
            # 也不再拿来复用（暂行办法第十四条：发现违法内容要消除）
            _remove_copies(session, task, cached_from)
        task.status, task.finished_at, task.error = "failed", stamp, result_safety.UNSAFE_MESSAGE
        task.result_json = None
        task.result_rows = 0
        # 留着缓存键：1 天内同一份内容再提交，直接按没通过处理（_recently_blocked）
        release_usage(session, task)
        session.commit()
        return False
    except AICapReached:
        task.status, task.finished_at, task.error = "failed", stamp, AI_CAP_MESSAGE
        task.result_json = None
        task.cache_key = None
        release_usage(session, task)
        session.commit()
        return False
    except interpret.InterpretRefusal as exc:
        task.result_json = interpret.result_metadata(
            {
                skill.result_field: [],
                **({"missing": list(skill.types)} if skill.kind == "interpret" else {}),
                "dropped": 0,
                "refusal": str(exc),
            },
            stamp=stamp,
            model=model,
            verified=[],
            model_disclosure=interpret.disclosure(settings),
            skill=skill,
            attachments=(task.snapshot or {}).get("attachments", 0),
        )
        task.cache_key = None
        task.status, task.finished_at, task.error = "done", stamp, None
        task.result_rows = 0
        release_usage(session, task)
        session.commit()
        return True
    except Exception as exc:  # noqa: BLE001 - 不向用户暴露模型原始响应或配置
        session.rollback()
        task = session.get(AdvisorTask, task.id)
        if task is None:
            return False
        task.model, task.prompt_version = model, skill.prompt_version
        task.result_json = None
        task.cache_key = None
        # 内容安全接口出错不放行，也不重排：暂时性的错误已经原地重试过（result_safety），
        # 重排要重新调模型、占全站当天的次数
        retryable = (isinstance(exc, AIError) and exc.retryable) or isinstance(
            exc, (httpx.TransportError, TimeoutError)
        )
        task.error = (
            "AI 返回内容格式无效"
            if isinstance(exc, ValueError)
            else (
                "AI 条款对照失败，请稍后重试"
                if skill.kind == "clause"
                else "AI 解读失败，请稍后重试"
            )
        )
        if isinstance(exc, AIError) and not exc.retryable:
            task.error = "AI 返回内容无效或请求被拒绝，请稍后重试"
        if isinstance(exc, wx_security.CheckDown):
            task.error = result_safety.CHECK_DOWN_MESSAGE
            result_safety.alert_check_down(exc)
        if retryable and task.attempts <= 2:
            task.status, task.started_at, task.finished_at = "queued", None, None
        else:
            task.status, task.started_at, task.finished_at = "failed", None, stamp
            release_usage(session, task)
        session.commit()
        return False
    finally:
        if chat_client is not None:
            try:
                record_usage(
                    session,
                    chat_client,
                    skill.kind,
                    now=call_stamp,
                    precounted_calls=int(call_reserved),
                )
            finally:
                if own_client:
                    chat_client.close()


def process_interpret(session: Session, task: AdvisorTask, **kwargs) -> bool:
    """兼容原有政策解读调用入口。"""
    return process_ai(session, task, **kwargs)


def process_clause(session: Session, task: AdvisorTask, **kwargs) -> bool:
    """条款对照共用同一处理流程。"""
    return process_ai(session, task, **kwargs)


def record_pass_completion(session: Session, task: AdvisorTask) -> None:
    if task.pass_id is None:
        return
    # 只更新最近提交任务的凭据，不让较早任务的晚完成覆盖最新任务。
    session.execute(
        update(WorkbenchPass)
        .where(
            WorkbenchPass.id == task.pass_id,
            WorkbenchPass.last_task_id == task.id,
        )
        .values(last_task_completed_at=task.finished_at)
    )


def process_task(
    session: Session,
    task: AdvisorTask,
    *,
    settings: Settings | None = None,
    client: httpx.Client | None = None,
    stamp: datetime | None = None,
    timeout: timedelta = TASK_TIMEOUT,
    live_time: bool = False,
) -> bool:
    """处理一条已锁定为 running 的任务；失败按最多两次重试回队列。"""
    settings = settings or get_settings()
    stamp = _utc(stamp or datetime.now(UTC))
    if task.kind in AI_KINDS:
        return process_ai(
            session, task, settings=settings, stamp=stamp, timeout=timeout, live_time=live_time
        )
    own_client = client is None
    client = client or httpx.Client(trust_env=False)
    started = monotonic_clock.monotonic()
    uploaded_file_id: str | None = None
    # 处理途中任务行可能被删（用户注销、过期清理）：rollback 之后再读 task.id
    # 会报 ObjectDeletedError，先记下编号，失败分支用它重新取
    task_id = task.id
    try:
        cloud_env = getattr(settings, "cloud_env", None)
        if not cloud_env:
            raise BackupError("未配置对象存储环境")
        user = session.get(User, task.user_id)
        if user is None:
            raise CompareInputError("用户不存在，任务无法处理")
        with tempfile.TemporaryDirectory() as tmp:
            local = Path(tmp) / f"{task.id}.xlsx"
            rows = _run_engine(session, task, local)
            if monotonic_clock.monotonic() - started > timeout.total_seconds():
                raise TimeoutError
            file_id = upload(
                client,
                cloud_env,
                result_path(user, task.id, engine=session.get_bind()),
                local,
                what="工具结果",
            )
            uploaded_file_id = file_id
        if monotonic_clock.monotonic() - started > timeout.total_seconds():
            raise TimeoutError
        # 和提交任务、注销一样先锁用户，再写任务完成（首次批量奖励已删，D-061、D-070）。
        # 引擎已修改任务的结果引用；自动刷新会先写任务再等待用户，和注销互锁。
        with session.no_autoflush:
            user = session.scalar(
                select(User)
                .where(User.id == task.user_id)
                .with_for_update()
                .execution_options(populate_existing=True)
            )
        if user is None:
            raise CompareInputError("用户不存在，任务无法处理")
        task.status = "done"
        task.result_file_id = file_id
        task.result_rows = rows
        task.finished_at = stamp
        task.error = None
        record_pass_completion(session, task)
        session.commit()
        return True
    except Exception as exc:  # noqa: BLE001 - 用户只需看到稳定的中文原因
        session.rollback()
        if uploaded_file_id:
            _delete_and_record(
                session,
                [uploaded_file_id],
                settings=settings,
                client=client,
            )
        task = session.get(AdvisorTask, task_id)
        if task is None:
            session.commit()
            return False
        task.error = _friendly_error(exc)
        if task.attempts <= 2:
            task.status = "queued"
            task.started_at = None
            task.finished_at = None
        else:
            task.status = "failed"
            task.started_at = None
            task.finished_at = stamp
            release_usage(session, task)
            record_pass_completion(session, task)
        session.commit()
        log.warning("工具任务 #%s 失败（第 %s 次）：%s", task.id, task.attempts, task.error)
        return False
    finally:
        if own_client:
            client.close()


def run_advisor_tasks(
    session: Session,
    *,
    settings: Settings | None = None,
    stamp: datetime | None = None,
    timeout: timedelta = TASK_TIMEOUT,
    max_tasks: int | None = None,
    poll_window: float = 0.0,
    poll_interval: float = POLL_INTERVAL,
    claim_limit: float = CLAIM_LIMIT,
    clock=monotonic_clock.monotonic,
    sleep=monotonic_clock.sleep,
) -> dict[str, int]:
    """定时任务入口；先终结失去授权的解读任务，再按开关领取任务。

    默认（poll_window 0）把现在排着的处理完就返回（测试、手动 --once）。poll_window 大于 0 时
    （每分钟的 advisor-run 用 POLL_WINDOW），AI 任务排空了不马上结束：开始后 poll_window 秒以内、
    或别的线程还在做任务时，每隔 poll_interval 秒再领一次；开始满 claim_limit 秒不再领新的，
    手上的做完就返回。数据工具不轮询，每轮只领一次。"""
    start = clock()
    deadline = start + poll_window if poll_window > 0 else None
    claim_deadline = start + claim_limit if poll_window > 0 else None
    settings = settings or get_settings()
    live_timing = stamp is None
    stamp = _utc(stamp or datetime.now(UTC))
    revoked = _fail_revoked_interpret(session, settings, stamp)
    data_enabled = getattr(settings, "advisor_enabled", False)
    # 和心跳、上面的退回同一个判断：缺对外 AI 密钥时也不领 AI 任务（排队的已经退回）
    ai_enabled = heartbeat.worker_capable(settings)
    data_allowed = AdvisorTask.kind.in_(("compare", "track"))
    ai_allowed = AdvisorTask.kind.in_(AI_KINDS) & AdvisorTask.user_id.in_(
        interpret.worker_user_ids(settings)
    )
    if not data_enabled and not ai_enabled:
        return {"processed": 0, "done": 0, "failed": 0, "requeued": 0, "revoked": revoked}
    allowed = data_allowed if data_enabled else ai_allowed
    if data_enabled and ai_enabled:
        allowed = data_allowed | ai_allowed
    requeued = recover_stale(session, stamp, allowed)
    result = {"processed": 0, "done": 0, "failed": 0, "requeued": requeued, "revoked": revoked}
    claimed: set[int] = set()
    claim_lock = Lock()
    # 正在处理的任务数（领到加 1，处理完减 1；数据工具也算）：空闲的 AI 线程据此判断这一轮还结束不了
    busy = {"in_flight": 0}

    def claim(work_session: Session, condition) -> AdvisorTask | None:
        # 临界区只负责领取；耗时的模型请求在锁外并行运行。
        with claim_lock:
            if max_tasks is not None and result["processed"] >= max_tasks:
                return None
            stmt = (
                select(AdvisorTask)
                .where(AdvisorTask.status == "queued", condition)
                .order_by(AdvisorTask.id)
                .limit(1)
                .with_for_update(skip_locked=True)
            )
            if claimed:
                stmt = stmt.where(AdvisorTask.id.not_in(claimed))
            task = work_session.scalar(stmt)
            if task is None:
                work_session.rollback()
                return None
            task.status = "running"
            task.started_at = datetime.now(UTC) if live_timing else stamp
            task.attempts += 1
            task.error = None
            work_session.commit()
            claimed.add(task.id)
            result["processed"] += 1
            busy["in_flight"] += 1
            return task

    def run_one(work_session: Session, task: AdvisorTask) -> None:
        # 任务行可能在处理途中被删（注销、过期清理）：之后只按编号重新取，不读 task 的属性
        task_id = task.id
        try:
            try:
                done = process_task(
                    work_session,
                    task,
                    settings=settings,
                    stamp=datetime.now(UTC) if live_timing else stamp,
                    timeout=timeout,
                    live_time=live_timing,
                )
            except Exception:
                work_session.rollback()
                log.exception("顾问任务 #%s 处理异常", task_id)
                failed_task = work_session.get(AdvisorTask, task_id)
                if failed_task is not None:
                    failed_task.status, failed_task.finished_at = "failed", stamp
                    failed_task.error = "系统繁忙，请稍后重试"
                    release_usage(work_session, failed_task)
                    work_session.commit()
                done = False
            current = work_session.get(AdvisorTask, task_id) if live_timing else None
            if current is not None and current.status in ("done", "failed"):
                current.finished_at = datetime.now(UTC)
                work_session.commit()
            with claim_lock:
                result["done" if done else "failed"] += 1
        finally:
            with claim_lock:
                busy["in_flight"] -= 1

    def run_data() -> None:
        # 数据工具始终由入口会话逐个处理，同样开始满 claim_limit 秒不再领新的。处理数据工具的
        # 整段时间都算"有人在忙"：两条数据任务之间计数不掉到 0，空闲的 AI 线程不会这时一起退出
        if not data_enabled:
            return
        with claim_lock:
            busy["in_flight"] += 1
        try:
            while claim_deadline is None or clock() < claim_deadline:
                task = claim(session, data_allowed)
                if task is None:
                    break
                run_one(session, task)
        finally:
            with claim_lock:
                busy["in_flight"] -= 1

    if ai_enabled:
        factory = sessionmaker(bind=session.get_bind(), expire_on_commit=False)

        def ai_worker() -> None:
            with factory() as work_session:
                # 开始满 claim_limit 秒不再领新的（手上的已经做完）：剩下的留给下一轮
                while claim_deadline is None or clock() < claim_deadline:
                    task = claim(work_session, ai_allowed)
                    if task is not None:
                        run_one(work_session, task)
                        continue
                    # 队里空了：不轮询（测试、--once）或领满 max_tasks 就结束
                    if deadline is None or (
                        max_tasks is not None and result["processed"] >= max_tasks
                    ):
                        break
                    now = clock()
                    with claim_lock:
                        others_busy = busy["in_flight"] > 0
                    # 过了 poll_window、别的线程也都空着：这一轮结束
                    if now >= deadline and not others_busy:
                        break
                    # poll_window 以内睡到下一次看（最后一次正好在到点时看）；过了 poll_window
                    # 但别的线程还在做任务（这一轮反正结束不了），照常隔几秒看
                    wait = poll_interval if now >= deadline else min(poll_interval, deadline - now)
                    if now + wait >= claim_deadline:
                        break
                    sleep(wait)

        concurrency = getattr(settings, "ai_advisor_concurrency", 3)
        with ThreadPoolExecutor(max_workers=concurrency) as pool:
            futures = [pool.submit(ai_worker) for _ in range(concurrency)]
            run_data()
            for future in futures:
                future.result()
    else:
        run_data()
    session.expire_all()
    return result


def download_url(client: httpx.Client, env: str, file_id: str) -> str | None:
    """结果文件的临时下载链接（只在服务端用）。

    请求和返回的格式同附件副本（attachments.download_link，线上一直在用）：
    file_list 里每项是 {fileid, max_age}，返回的 file_list 每项带 status 和 download_url。
    原来这里把文件编号直接当字符串传，10/10 第一次连真实环境时取不到链接。
    """
    data = _open_api(
        client,
        "/tcb/batchdownloadfile",
        {"env": env, "file_list": [{"fileid": file_id, "max_age": 600}]},
    )
    items = data.get("file_list") or data.get("download_list") or []
    item = next((i for i in items if i.get("fileid") in (None, file_id)), None)
    if not item or item.get("status") not in (None, 0):
        return None
    url = item.get("download_url") or item.get("url") or ""
    return url if url.startswith("https://") else None


# 结果文件由 api 自己从对象存储分段读出来再发给小程序（和附件副本同一种做法）：
# 不把存储的下载地址交给小程序——手机直接下载要在小程序后台配 downloadFile 合法域名，
# 10/10 真机试用前发现结果因此打不开。每段大小留在 callContainer 返回包上限（1000 KiB）以内。
RESULT_PART_SIZE = 900 * 1024
RESULT_MAX_BYTES = 20 * 1024 * 1024
_FILE_ENV = re.compile(r"^cloud://([^./]+)\.")
_RANGE_TOTAL = re.compile(r"^bytes \d+-\d+/(\d+)$")


def result_env(settings, file_id: str) -> str | None:
    """对象存储环境：优先用配置；api 服务没配时从文件编号里取（cloud://环境ID.存储桶/路径）。"""
    configured = getattr(settings, "cloud_env", None)
    if configured:
        return configured
    match = _FILE_ENV.match(file_id or "")
    return match.group(1) if match else None


def result_size(client: httpx.Client, url: str) -> int:
    """结果文件的字节数：只请求第一个字节，从 Content-Range 里读总长度。"""
    with client.stream("GET", url, headers={"Range": "bytes=0-0"}, timeout=10) as res:
        if res.status_code == 206:
            match = _RANGE_TOTAL.match(res.headers.get("content-range", ""))
            if not match:
                raise BackupError("读不出结果文件的大小")
            total = int(match.group(1))
        elif res.status_code == 200:
            total = int(res.headers.get("content-length") or 0)
        else:
            raise BackupError(f"读取结果文件失败（HTTP {res.status_code}）")
    if not 0 < total <= RESULT_MAX_BYTES:
        raise BackupError("结果文件大小不对")
    return total


def read_result_part(client: httpx.Client, url: str, index: int, total: int) -> bytes:
    """读结果文件的第 index 段（从 0 起）。"""
    from youju.attachments import read_range

    start = index * RESULT_PART_SIZE
    if index < 0 or start >= total:
        raise ValueError("没有这一段")
    return read_range(client, url, start, min(total, start + RESULT_PART_SIZE), total)


def delete_result_files(client: httpx.Client, env: str, file_ids: list[str]) -> list[str]:
    """返回对象存储确认删除成功的文件 ID。"""
    if not file_ids:
        return []
    data = _open_api(client, "/tcb/batchdeletefile", {"env": env, "fileid_list": file_ids})
    return [
        item["fileid"]
        for item in data.get("delete_list", [])
        if item.get("status") == 0 and item.get("fileid") in file_ids
    ]


def cleanup_advisor(
    session: Session,
    *,
    settings: Settings | None = None,
    stamp: datetime | None = None,
    client: httpx.Client | None = None,
) -> dict[str, int]:
    """清理七天输入和三十天导出；对象存储删除成功后才清空 file id，
    过了三十天的数据工具任务行随后整行删除。"""
    settings = settings or get_settings()
    stamp = _utc(stamp or datetime.now(UTC))
    input_count = (
        session.execute(
            update(AdvisorTask)
            .where(
                AdvisorTask.input_rows.is_not(None),
                AdvisorTask.created_at <= stamp - INPUT_RETENTION,
            )
            .values(input_rows=None)
        ).rowcount
        or 0
    )
    session.execute(
        update(AdvisorTask)
        .where(AdvisorTask.kind.in_(("compare", "track")), AdvisorTask.expires_at <= stamp)
        .values(result_json=None)
    )
    session.execute(
        update(AdvisorTask)
        .where(AdvisorTask.kind.in_(AI_KINDS), AdvisorTask.created_at <= stamp - INPUT_RETENTION)
        .values(input_json=None)
    )
    session.execute(
        update(AdvisorTask)
        .where(AdvisorTask.kind.in_(AI_KINDS), AdvisorTask.expires_at <= stamp)
        .values(
            result_json=None,
            snapshot=None,
            cache_key=None,
            model=None,
            prompt_version=None,
            result_document_ids=None,
        )
    )
    old = list(
        session.scalars(
            select(AdvisorTask).where(
                AdvisorTask.result_file_id.is_not(None),
                AdvisorTask.expires_at.is_not(None),
                AdvisorTask.expires_at <= stamp,
            )
        )
    )
    deleted_count = 0
    own_client = client is None
    trash = list(session.scalars(select(AdvisorFileTrash)))
    file_ids = list(
        dict.fromkeys(
            [task.result_file_id for task in old if task.result_file_id]
            + [row.file_id for row in trash]
        )
    )
    if file_ids:
        client = client or httpx.Client(trust_env=False)
        try:
            successful = _delete_and_record(
                session,
                file_ids,
                settings=settings,
                client=client,
            )
        finally:
            if own_client:
                client.close()
        for task in old:
            if task.result_file_id in successful:
                task.result_file_id = None
                deleted_count += 1
        for row in trash:
            if row.file_id in successful:
                session.delete(row)
    # AI 任务过了结果保存期（30 天）整行删除，任务和用户的关联不再保留（隐私保护指引一·14）。
    ai_tasks = (
        session.execute(
            delete(AdvisorTask).where(
                AdvisorTask.kind.in_(AI_KINDS), AdvisorTask.expires_at <= stamp
            )
        ).rowcount
        or 0
    )
    # 数据工具的任务记录和结果文件一起保存 30 天（隐私保护指引一·13，个保法第十九条）：
    # 文件删成以后整行删除；这次没删成的（result_file_id 还在）留到下次删成再删，免得丢了文件编号。
    # 付费核对看使用期上的 last_used_at 等最少记录，不靠任务行（用户本来也能随时删任务）
    session.flush()
    expired_data = (
        AdvisorTask.kind.not_in(AI_KINDS),
        AdvisorTask.expires_at <= stamp,
        AdvisorTask.result_file_id.is_(None),
    )
    # 批量工具关着时排队的不会再处理、处理中的也不会被放回队列：过了保存期先退回占用的次数
    # （按原提交日退，账本已清掉的不退），标成已取消，和已结束的一起删。worker 领取时锁住排队的行、
    # 立刻改成处理中，所以这里跳过被锁的就不会抢；处理中的行在处理期间并不一直锁着，只动开始超过
    # ABANDONED 的（换版本时旧实例还会跑十几分钟，一天以上不可能还在跑），免得删掉正在上传结果的
    for stale in session.scalars(
        select(AdvisorTask)
        .where(
            *expired_data,
            or_(
                AdvisorTask.status == "queued",
                and_(
                    AdvisorTask.status == "running",
                    AdvisorTask.started_at <= stamp - ABANDONED,
                ),
            ),
        )
        .with_for_update(skip_locked=True)
    ).all():
        release_usage(session, stale)
        stale.status = "cancelled"
    session.flush()
    data_tasks = (
        session.execute(
            delete(AdvisorTask).where(
                AdvisorTask.kind.not_in(AI_KINDS),
                AdvisorTask.expires_at <= stamp,
                AdvisorTask.result_file_id.is_(None),
                AdvisorTask.status.in_(("done", "failed", "cancelled")),
            )
        ).rowcount
        or 0
    )
    # 每人每天的使用次数只用于每日、每月上限和自助退款条件，保留本月和上个月，更早的删除。
    keep_from = (usage_day(stamp).replace(day=1) - timedelta(days=1)).replace(day=1)
    usage_rows = (
        session.execute(delete(AdvisorUsage).where(AdvisorUsage.day < keep_from)).rowcount or 0
    )
    # 词库更新以后再查一遍还显示着的 AI 结果（门槛第 4 道）；处置记录 190 天到期删除，
    # 删任务、注销都不提前删
    swept = sweep_results(session, stamp)
    safety_hits = (
        session.execute(delete(AISafetyHit).where(AISafetyHit.expires_at <= stamp)).rowcount or 0
    )
    session.commit()
    return {
        "advisor_input_rows": input_count,
        "advisor_files": deleted_count,
        "advisor_ai_tasks": ai_tasks,
        "advisor_data_tasks": data_tasks,
        "advisor_usage_rows": usage_rows,
        "ai_safety_hits": safety_hits,
        "ai_results_removed": swept,
    }


def delete_user_tasks(
    session: Session,
    user_id: int,
    *,
    settings: Settings | None = None,
    client: httpx.Client | None = None,
) -> int:
    """注销或用户删除任务时删除结果文件，再删除数据库任务记录。"""
    settings = settings or get_settings()
    tasks = list(session.scalars(select(AdvisorTask).where(AdvisorTask.user_id == user_id)))
    if tasks:
        # 锁定读取最新文件引用，避免注销的旧快照漏清理刚完成的结果。
        tasks = list(
            session.scalars(
                select(AdvisorTask)
                .where(
                    AdvisorTask.user_id == user_id,
                )
                .with_for_update()
                .execution_options(populate_existing=True)
            )
        )
    file_ids = [task.result_file_id for task in tasks if task.result_file_id]
    own_client = client is None
    if file_ids:
        client = client or httpx.Client(trust_env=False)
        try:
            _delete_and_record(session, file_ids, settings=settings, client=client)
        finally:
            if own_client:
                client.close()
    # 没有任务就不发 DELETE：正式库给 youju_api 授权顾问表之前（功能关闭期间），
    # 注销不能因为缺这张表的 DELETE 权限而失败（MySQL 按语句检查权限，删 0 行也会报错）。
    if tasks:
        session.execute(delete(AdvisorTask).where(AdvisorTask.user_id == user_id))
    return len(tasks)
