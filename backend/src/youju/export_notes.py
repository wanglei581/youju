"""各类导出共用的来源说明、可选小程序码和公众号名称。"""

import re
from pathlib import Path

from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet

from youju.config import get_settings

DISCLAIMER = "第三方按公开发布的信息整理，仅供参考，以发文机关原文为准。"
DEADLINE_HEADER = "截止日期（以原文为准）"
# 小程序全称"有据申报与行业助手"（10/18 验收：只搜"有据"不一定找得到）
SEARCH_NOTICE = "在微信里搜索小程序「有据申报与行业助手」可以查看原文和更新。"
SCAN_NOTICE = "微信扫码查看原文和更新"
MINI_CODE_PATH = Path(__file__).parent / "assets" / "mini-code.png"
# 名称里只留中文、字母、数字和常见标点，别的（如换行、引号、网址符号）不进导出文件。
_ACCOUNT_NAME = re.compile(r"[\u4e00-\u9fffA-Za-z0-9·（）()\- ]{1,30}")


def official_account_notice() -> str | None:
    """页脚的公众号一行（PR #287）：配置了公众号名称才有，没填返回 None。"""
    name = (get_settings().official_account_name or "").strip()
    if not name or not _ACCOUNT_NAME.fullmatch(name):
        return None
    # 陈述句、不用"关注"祈使：导出文件常转给领导和客户（agy 合规意见）
    return f"微信公众号「{name}」：每周一发申报周报。"


def add_excel_notes(sheet: Worksheet) -> None:
    """追加说明，不改变现有表头、数据行和核对日期的位置。"""
    sheet.append(["使用说明", DISCLAIMER])
    sheet.append(["查看原文和更新", SEARCH_NOTICE])
    rows = 2
    notice = official_account_notice()
    if notice:
        sheet.append(["公众号", notice])
        rows += 1
    for row in sheet.iter_rows(min_row=sheet.max_row - rows + 1, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 42
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 65
    if not MINI_CODE_PATH.is_file():
        return
    # 可选资源晚于代码部署，缺失或损坏时仍能导出文字说明。
    try:
        image = Image(str(MINI_CODE_PATH))
    except OSError:
        return
    scale = 160 / max(image.width, image.height)
    image.width, image.height = image.width * scale, image.height * scale
    row_number = sheet.max_row + 2
    sheet.add_image(image, f"B{row_number}")
    sheet.cell(row_number, 3, SCAN_NOTICE)
    sheet.column_dimensions["C"].width = max(sheet.column_dimensions["C"].width, 30)
    sheet.row_dimensions[row_number].height = 130
