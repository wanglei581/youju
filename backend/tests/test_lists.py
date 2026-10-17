from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from youju.crawler.fetcher import FetchResult
from youju.lists.parser import (
    MAX_PDF_PAGES,
    parse_attachment,
    parse_attachment_result,
    parse_html_tables,
    parse_table_rows,
)
from youju.lists.rules import (
    candidate_key,
    classify_title,
    is_person_query,
    normalize_credit_code,
    normalize_name,
    valid_org_name,
)
from youju.lists.service import _should_attempt, parse_lists
from youju.models import Document, DocumentAttachment, DocumentDetail, ListEntry, ListParse, Source

FIXTURES = Path(__file__).parent / "fixtures" / "lists"


@pytest.mark.parametrize(
    ("filename", "count", "first"),
    [
        ("panjin_table.html", 30, "智源（辽宁）碳科技有限公司"),
        ("heilongjiang_30.xlsx", 30, "哈尔滨力强药业有限责任公司"),
        ("heilongjiang_30.xls", 30, "哈尔滨力强药业有限责任公司"),
        ("shanghai_page1.pdf", 27, "中卉泽科技（上海）有限公司"),
    ],
)
def test_real_government_list_fixtures(filename, count, first) -> None:
    path = FIXTURES / filename
    assert path.stat().st_size <= 300 * 1024
    entries = (
        parse_html_tables(path.read_text())
        if path.suffix == ".html"
        else parse_attachment(path.read_bytes(), path.suffix)
    )
    assert len(entries) == count
    assert entries[0].name == first
    assert entries[0].sequence == "1"
    if filename.endswith((".xlsx", ".xls")):
        assert entries[0].region == "黑龙江省/哈尔滨市/依兰县"
    assert all(valid_org_name(entry.name, entry.credit_code) for entry in entries)


def test_docx_and_personal_columns_are_not_saved() -> None:
    doc = DocxDocument()
    table = doc.add_table(rows=3, cols=5)
    values = [
        ["序号", "企业名称", "联系人", "手机号", "所在地区"],
        ["1", "北京星河科技有限公司", "张三", "13800138000", "北京市"],
        ["2", "李四", "王五", "13900139000", "北京市"],
    ]
    for row, values_row in zip(table.rows, values, strict=True):
        for cell, value in zip(row.cells, values_row, strict=True):
            cell.text = value
    stream = BytesIO()
    doc.save(stream)
    entries = parse_attachment(stream.getvalue(), "docx")
    assert [(entry.name, entry.region) for entry in entries] == [("北京星河科技有限公司", "北京市")]
    assert "张三" not in str(entries)
    assert "13800138000" not in str(entries)
    assert parse_table_rows([["姓名", "电话"], ["张三", "13800138000"]]) == []


def test_names_codes_and_title_rules() -> None:
    assert normalize_name("ＡＢＣ（上海） 科技有限 公司") == "ABC(上海)科技有限公司"
    assert candidate_key("北京星河（上海）有限公司") == candidate_key("北京星河有限公司")
    assert is_person_query("张 三")
    assert not is_person_query("北京研究院")
    assert not valid_org_name("李四")
    assert not valid_org_name("老茶厂、铸铁锅等设备和档案资料。")
    assert valid_org_name("洞口县华宇机械制造厂（普通合伙）")
    assert normalize_credit_code("９１３１０１１７５８８６５７２５６ｃ") == "91310117588657256C"
    assert normalize_credit_code("123456789012345679") is None
    assert classify_title("关于2026年第4批拟入库科技型中小企业名单的公示").is_list
    info = classify_title("关于公布2026年第4批专精特新中小企业名单的公告")
    assert (info.category, info.batch) == ("专精特新中小企业", "第4批")
    assert classify_title("第六批专精特新“小巨人”企业名单公示").category == '专精特新"小巨人"'
    assert classify_title("高新技术企业名单").is_list
    # 与全站分类一致：招聘、职称等无关公示不解析
    assert classify_title("关于2025年度工程系列科技管理专业高级职称评审结果的公示").is_list is False
    assert classify_title("国家统计局所属单位2026年度公开招聘第三批拟聘用人员公示").is_list is False
    assert classify_title("关于2026年高新技术企业认定结果的通知").is_list
    assert classify_title("关于开展高新技术企业认定申报工作的通知").is_list is False
    assert classify_title("关于高企认定评审专家培训的预通知").is_list is False
    assert classify_title("采购结果公示").is_list is False
    assert classify_title("取消有关企业专精特新中小企业认定的公告").is_list is False
    assert classify_title("关于公布实验动物行政许可注销名单的通知").is_list is False
    assert classify_title("关于拟对高新技术企业给予奖励的公示").is_list is False


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("关于2026年工业和信息化部科技型企业孵化器拟推荐申报名单的公示", True),
        ("广东省科学技术厅关于2026年工业和信息化部科技型企业孵化器推荐申报名单的公示", True),
        ("四川省科学技术厅关于拟推荐申报2026年度工业和信息化部科技型企业孵化器名单的公示", True),
        ("重庆市科学技术局关于公示2026年工业和信息化部科技型企业孵化器拟推荐申报名单的公告", True),
        ("江西省2026年第二批拟申请备案高新技术企业名单公示", True),
        ("吉林省工业和信息化厅关于组织开展国家级零碳工厂建设名单申报工作的通知", False),
        ("关于受理山东省自然科学基金依托单位注册申请的通知", False),
        ("关于组织征集2026年度四川省消费名品培育名单的通知", False),
        ("关于征集推荐第五批陕西省“工信贷”白名单企业的通知", False),
        ("关于2026年黑龙江省科技成果转移转化服务体系奖励拟支持单位名单的公示", False),
        ("省企业技术中心拟认定、撤销名单公示", False),
    ],
)
def test_real_titles_affected_by_negative_words(title: str, expected: bool) -> None:
    assert classify_title(title).is_list is expected


def test_unsupported_and_scanned_file() -> None:
    with pytest.raises(ValueError, match="暂不支持"):
        parse_attachment(b"old doc", "doc")
    with pytest.raises(ValueError, match="不是 PDF"):
        parse_attachment(b"not a pdf", "pdf")


def _fake_pdf(page_count: int, visited: list[int]):
    class Page:
        def __init__(self, number: int):
            self.number = number

        def close(self):
            pass

        def extract_text(self):
            visited.append(self.number)
            return "企业名称" if self.number == MAX_PDF_PAGES else ""

        def extract_tables(self, table_settings=None):
            if self.number == MAX_PDF_PAGES:
                return [[["企业名称"], ["北京星河科技有限公司"]]]
            return []

    class Pdf:
        def __init__(self):
            self.pages = [Page(i) for i in range(1, page_count + 1)]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

    return Pdf()


@pytest.mark.parametrize("page_count", [MAX_PDF_PAGES, MAX_PDF_PAGES + 1])
def test_pdf_page_limit_keeps_first_pages(monkeypatch, page_count: int) -> None:
    visited: list[int] = []
    monkeypatch.setattr(
        "youju.lists.parser.pdfplumber.open", lambda _: _fake_pdf(page_count, visited)
    )
    result = parse_attachment_result(b"%PDF-fake", "pdf")
    assert [entry.name for entry in result.entries] == ["北京星河科技有限公司"]
    assert visited == list(range(1, MAX_PDF_PAGES + 1))
    assert result.failure_reason == (
        f"PDF 共{page_count}页，仅解析前{MAX_PDF_PAGES}页" if page_count > MAX_PDF_PAGES else None
    )


def test_pdf_page_limit_is_recorded_for_scanned_file(monkeypatch) -> None:
    pdf = _fake_pdf(MAX_PDF_PAGES + 1, [])
    for page in pdf.pages:
        page.extract_text = lambda: ""
        page.extract_tables = lambda table_settings=None: []
    monkeypatch.setattr("youju.lists.parser.pdfplumber.open", lambda _: pdf)
    over = MAX_PDF_PAGES + 1
    with pytest.raises(
        ValueError, match=f"暂不支持扫描件；PDF 共{over}页，仅解析前{MAX_PDF_PAGES}页"
    ):
        parse_attachment_result(b"%PDF-fake", "pdf")


class FakeFetcher:
    def __init__(self, responses):
        self.responses = responses
        self.calls = []

    def fetch(self, url):
        self.calls.append(url)
        return self.responses[url]


def test_pdf_page_limit_saves_partial_parse_reason(session, monkeypatch) -> None:
    visited: list[int] = []
    over = MAX_PDF_PAGES + 1
    monkeypatch.setattr("youju.lists.parser.pdfplumber.open", lambda _: _fake_pdf(over, visited))
    source = Source(
        level="省级",
        province="北京",
        dept="科技",
        dept_full_name="北京市科学技术委员会",
        column_name="通知公告",
        url="https://kw.beijing.gov.cn/list/",
        fetch_mode="static",
        enabled=True,
        parser_config={},
        rights_level=1,
    )
    session.add(source)
    session.flush()
    now = datetime.now(UTC)
    url = "https://kw.beijing.gov.cn/partial-list.pdf"
    document = Document(
        source_id=source.id,
        url=url,
        url_key="kw.beijing.gov.cn/partial-list.pdf",
        title="2026年高新技术企业名单公告",
        publish_date=date(2026, 8, 1),
        first_seen_at=now,
        last_seen_at=now,
        is_baseline=False,
    )
    session.add(document)
    session.commit()
    fake = FakeFetcher({url: FetchResult(url, "ok", final_url=url, content=b"%PDF-fake")})

    report = parse_lists(session, fetcher=fake)
    record = session.query(ListParse).filter_by(document_id=document.id, source_url=url).one()
    assert record.status == "ok"
    assert record.entry_count == 1
    assert record.failure_reason == f"PDF 共{over}页，仅解析前{MAX_PDF_PAGES}页"
    assert report.names == 1
    assert report.failures == {record.failure_reason: 1}
    assert session.query(ListEntry).filter_by(
        document_id=document.id, source_url=url
    ).one().name == ("北京星河科技有限公司")
    assert visited == list(range(1, MAX_PDF_PAGES + 1))
    assert parse_lists(session, fetcher=fake).names == 1
    assert fake.calls == [url]

    entry = session.query(ListEntry).filter_by(document_id=document.id, source_url=url).one()
    entry.name = "李四"
    session.commit()
    assert parse_lists(session, fetcher=fake).names == 0
    assert record.status == "empty"
    assert record.failure_reason == f"PDF 共{over}页，仅解析前{MAX_PDF_PAGES}页"


def _parse_record(status: str, reason: str | None, hours_ago: float = 24) -> ListParse:
    return ListParse(
        document_id=1,
        source_url="https://a.gov.cn/f/名单.pdf",
        source_kind="attachment",
        status=status,
        failure_reason=reason,
        attempted_at=datetime.now(UTC) - timedelta(hours=hours_ago),
    )


def test_pdfs_cut_at_an_older_page_limit_are_reread_once() -> None:
    """10/17 PDF 上限由 300 调到 1000：原因恰好是以前上限下的截断原话（ok、empty，或扫描件）的重读一次；
    重读成功后原因变了，重读失败会加上"；重读失败"，都不再是原话，所以只重读一次（Codex 审阅 #274）。"""
    assert MAX_PDF_PAGES == 1000
    assert _should_attempt(_parse_record("ok", "PDF 共567页，仅解析前300页"))
    assert _should_attempt(_parse_record("empty", "PDF 共680页，仅解析前300页"))
    assert not _should_attempt(_parse_record("ok", f"PDF 共1200页，仅解析前{MAX_PDF_PAGES}页"))
    # 扫描件：前 300 页没有文字，后面的页可能有名单
    assert _should_attempt(
        _parse_record("unsupported", "暂不支持扫描件；PDF 共400页，仅解析前300页")
    )
    assert not _should_attempt(
        _parse_record("unsupported", f"暂不支持扫描件；PDF 共1200页，仅解析前{MAX_PDF_PAGES}页")
    )
    assert not _should_attempt(_parse_record("unsupported", "暂不支持扫描件"))
    assert not _should_attempt(_parse_record("ok", None))
    assert not _should_attempt(_parse_record("empty", "未识别到企业或机构名称表格"))
    # 只认 parser 写的整句原话：重读失败记过的、别的原因里带着相近文字的都不重读
    assert not _should_attempt(
        _parse_record("ok", "PDF 共567页，仅解析前300页；重读失败：抓取失败：error")
    )
    assert not _should_attempt(_parse_record("ok", "附件 PDF 共567页，仅解析前300页"))
    # failed 本来就按 6 小时规则重试，这里只看截断规则：1 小时内不因截断而重读
    assert not _should_attempt(_parse_record("failed", "PDF 共567页，仅解析前300页", hours_ago=1))
    # 和 failed 一样 6 小时内不重复试（临时抓取失败后留着原话）
    assert not _should_attempt(_parse_record("ok", "PDF 共567页，仅解析前300页", hours_ago=1))


def test_reread_replaces_entries_from_the_old_page_limit(session, monkeypatch) -> None:
    """以前只读了前 300 页的名单：重读后条目整份换成新结果（不重复叠加），原因清掉；再跑不再重读。"""
    visited: list[int] = []

    class Page:
        def __init__(self, number: int):
            self.number = number

        def close(self):
            pass

        def extract_text(self):
            return "名单"

        def extract_tables(self, table_settings=None):
            if table_settings is None:
                visited.append(self.number)
            # 企业名称表格在第 400 页：旧的 300 页上限读不到
            return [[["企业名称"], ["北京星河科技有限公司"]]] if self.number == 400 else []

    class Pdf:
        pages = [Page(i) for i in range(1, 568)]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

    monkeypatch.setattr("youju.lists.parser.pdfplumber.open", lambda _: Pdf())
    source = Source(
        level="省级",
        province="河南",
        dept="科技",
        dept_full_name="河南省科学技术厅",
        column_name="通知公告",
        url="https://kjt.henan.gov.cn/list/",
        fetch_mode="static",
        enabled=True,
        parser_config={},
        rights_level=1,
    )
    session.add(source)
    session.flush()
    now = datetime.now(UTC)
    url = "https://kjt.henan.gov.cn/2026-list.pdf"
    document = Document(
        source_id=source.id,
        url=url,
        url_key="kjt.henan.gov.cn/2026-list.pdf",
        title="2026年第一批入库科技型中小企业名单公告",
        publish_date=date(2026, 9, 2),
        first_seen_at=now,
        last_seen_at=now,
        is_baseline=False,
    )
    session.add(document)
    session.flush()
    session.add(
        ListParse(
            document_id=document.id,
            source_url=url,
            source_kind="document",
            status="ok",
            failure_reason="PDF 共567页，仅解析前300页",
            entry_count=1,
            attempted_at=now - timedelta(days=1),
        )
    )
    session.add(
        ListEntry(
            document_id=document.id,
            source_url=url,
            name="旧条目有限公司",
            normalized_name="旧条目有限公司",
            category="科技型中小企业",
        )
    )
    session.commit()
    fake = FakeFetcher({url: FetchResult(url, "ok", final_url=url, content=b"%PDF-fake")})

    parse_lists(session, fetcher=fake)
    record = session.query(ListParse).filter_by(document_id=document.id, source_url=url).one()
    names = [e.name for e in session.query(ListEntry).filter_by(document_id=document.id)]
    assert names == ["北京星河科技有限公司"]
    assert record.status == "ok" and record.failure_reason is None and record.entry_count == 1
    assert visited == list(range(1, 568))
    parse_lists(session, fetcher=fake)
    assert fake.calls == [url]


def _henan_list(session, url: str):
    source = Source(
        level="省级",
        province="河南",
        dept="科技",
        dept_full_name="河南省科学技术厅",
        column_name="通知公告",
        url="https://kjt.henan.gov.cn/list/",
        fetch_mode="static",
        enabled=True,
        parser_config={},
        rights_level=1,
    )
    session.add(source)
    session.flush()
    now = datetime.now(UTC)
    document = Document(
        source_id=source.id,
        url=url,
        url_key=url.split("://", 1)[1],
        title="2026年第一批入库科技型中小企业名单公告",
        publish_date=date(2026, 9, 2),
        first_seen_at=now,
        last_seen_at=now,
        is_baseline=False,
    )
    session.add(document)
    session.flush()
    return document


def _cut_record(session, url: str) -> Document:
    """以前只读了前 300 页、已有 1 条条目的名单（一天前解析的）。"""
    document = _henan_list(session, url)
    session.add(
        ListParse(
            document_id=document.id,
            source_url=url,
            source_kind="document",
            status="ok",
            failure_reason="PDF 共567页，仅解析前300页",
            entry_count=1,
            attempted_at=datetime.now(UTC) - timedelta(days=1),
        )
    )
    session.add(
        ListEntry(
            document_id=document.id,
            source_url=url,
            name="已读到有限公司",
            normalized_name="已读到有限公司",
            category="科技型中小企业",
        )
    )
    session.commit()
    return document


def test_reread_fetch_failure_keeps_entries_and_retries_later(session) -> None:
    """重读时网站临时出错：已读到的条目、状态、原因都照旧，只记下这次尝试时间；
    6 小时内不再试，之后再试（Codex 审阅 #274：一次临时失败不能清掉条目，也不能从此不再重读）。"""
    url = "https://kjt.henan.gov.cn/2026-list-b.pdf"
    document = _cut_record(session, url)
    fake = FakeFetcher({url: FetchResult(url, "network_error", final_url=url, error="503")})

    parse_lists(session, fetcher=fake)
    record = session.query(ListParse).filter_by(document_id=document.id, source_url=url).one()
    names = [e.name for e in session.query(ListEntry).filter_by(document_id=document.id)]
    assert names == ["已读到有限公司"]
    assert record.status == "ok" and record.entry_count == 1
    assert record.failure_reason == "PDF 共567页，仅解析前300页"
    assert record.attempted_at > datetime.now(UTC) - timedelta(minutes=5)
    parse_lists(session, fetcher=fake)
    assert fake.calls == [url]
    record.attempted_at = datetime.now(UTC) - timedelta(hours=7)
    session.commit()
    parse_lists(session, fetcher=fake)
    assert fake.calls == [url, url]


def test_reread_parse_failure_keeps_entries_and_stops(session) -> None:
    """重读时解析出错（确定的结果）：条目和状态照旧，原因后面记上这次失败，之后不再重读。"""
    url = "https://kjt.henan.gov.cn/2026-list-d.pdf"
    document = _cut_record(session, url)
    fake = FakeFetcher({url: FetchResult(url, "ok", final_url=url, content=b"not a pdf")})

    parse_lists(session, fetcher=fake)
    record = session.query(ListParse).filter_by(document_id=document.id, source_url=url).one()
    names = [e.name for e in session.query(ListEntry).filter_by(document_id=document.id)]
    assert names == ["已读到有限公司"]
    assert record.status == "ok" and record.entry_count == 1
    assert record.failure_reason.startswith("PDF 共567页，仅解析前300页；重读失败：解析失败")
    record.attempted_at = datetime.now(UTC) - timedelta(days=1)
    session.commit()
    parse_lists(session, fetcher=fake)
    assert fake.calls == [url]


def test_scanned_pdf_cut_at_old_limit_is_reread(session, monkeypatch) -> None:
    """以前记成扫描件、只看了前 300 页的：后面的页有名单，重读后能读出来（Codex 审阅 #274）。"""

    class Page:
        def __init__(self, number: int):
            self.number = number

        def close(self):
            pass

        def extract_text(self):
            return "名单" if self.number >= 350 else ""

        def extract_tables(self, table_settings=None):
            return [[["企业名称"], ["郑州星河科技有限公司"]]] if self.number == 400 else []

    class Pdf:
        pages = [Page(i) for i in range(1, 451)]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

    monkeypatch.setattr("youju.lists.parser.pdfplumber.open", lambda _: Pdf())
    url = "https://kjt.henan.gov.cn/2026-list-c.pdf"
    document = _henan_list(session, url)
    session.add(
        ListParse(
            document_id=document.id,
            source_url=url,
            source_kind="document",
            status="unsupported",
            failure_reason="暂不支持扫描件；PDF 共450页，仅解析前300页",
            entry_count=0,
            attempted_at=datetime.now(UTC) - timedelta(days=1),
        )
    )
    session.commit()
    fake = FakeFetcher({url: FetchResult(url, "ok", final_url=url, content=b"%PDF-fake")})

    parse_lists(session, fetcher=fake)
    record = session.query(ListParse).filter_by(document_id=document.id, source_url=url).one()
    assert record.status == "ok" and record.failure_reason is None
    names = [e.name for e in session.query(ListEntry).filter_by(document_id=document.id)]
    assert names == ["郑州星河科技有限公司"]


def test_parse_lists_backfills_detail_and_is_idempotent(session) -> None:
    source = Source(
        level="省级",
        province="黑龙江",
        dept="科技",
        dept_full_name="黑龙江省科学技术厅",
        column_name="公示公告",
        url="https://kjt.hlj.gov.cn/notices/",
        fetch_mode="static",
        enabled=True,
        parser_config={},
        rights_level=1,
    )
    session.add(source)
    session.flush()
    now = datetime.now(UTC)
    url = "https://kjt.hlj.gov.cn/notice.html"
    attachment = "https://kjt.hlj.gov.cn/list.xlsx"
    document = Document(
        source_id=source.id,
        url=url,
        url_key="kjt.hlj.gov.cn/notice.html",
        title="关于黑龙江省2026年第2批拟入库科技型中小企业名单的公示",
        publish_date=date(2026, 8, 6),
        first_seen_at=now,
        last_seen_at=now,
        is_baseline=False,
    )
    session.add(document)
    session.commit()
    page = (
        "<html><body><div id='zoom'><p>现将拟入库科技型中小企业名单予以公示。"
        "有关单位可查看附件名单。</p><a href='/list.xlsx'>附件：名单.xlsx</a>"
        "</div></body></html>"
    )
    fake = FakeFetcher(
        {
            url: FetchResult(url, "ok", final_url=url, http_status=200, text=page),
            attachment: FetchResult(
                attachment,
                "ok",
                final_url=attachment,
                http_status=200,
                content=(FIXTURES / "heilongjiang_30.xlsx").read_bytes(),
            ),
        }
    )
    result = parse_lists(session, fetcher=fake)
    assert (
        result.documents,
        result.parsed_documents,
        result.names,
        result.attachments_fetched,
    ) == (1, 1, 30, 1)
    assert result.categories == {"科技型中小企业": 30}
    assert result.samples[0].status == "公示"
    assert len(fake.calls) == 3  # 补详情、读 HTML 表格、读附件
    session.add(
        ListEntry(
            document_id=document.id,
            source_url=attachment,
            name="老茶厂、铸铁锅等设备和档案资料。",
            normalized_name="老茶厂、铸铁锅等设备和档案资料。",
            category="其他",
        )
    )
    session.commit()
    assert parse_lists(session, fetcher=fake).names == 30
    assert len(fake.calls) == 3
    assert session.query(ListEntry).count() == 30
    assert session.get(DocumentDetail, document.id).attachments[0].url == attachment
    document.title = "关于开展2026年科技型中小企业申报的通知"
    session.commit()
    assert parse_lists(session, fetcher=fake).documents == 0
    assert session.query(ListEntry).count() == 0


def test_attachment_limit(session) -> None:
    source = Source(
        level="省级",
        province="北京",
        dept="科技",
        dept_full_name="北京市科学技术委员会",
        column_name="通知公告",
        url="https://kw.beijing.gov.cn/list/",
        fetch_mode="static",
        enabled=True,
        parser_config={},
        rights_level=1,
    )
    session.add(source)
    session.flush()
    now = datetime.now(UTC)
    doc = Document(
        source_id=source.id,
        url="https://kw.beijing.gov.cn/notice.html",
        url_key="kw.beijing.gov.cn/notice.html",
        title="2026年高新技术企业名单公告",
        publish_date=date(2026, 8, 1),
        first_seen_at=now,
        last_seen_at=now,
        is_baseline=False,
    )
    session.add(doc)
    session.flush()
    links = [f"https://kw.beijing.gov.cn/{i}.xlsx" for i in range(3)]
    session.add(
        DocumentDetail(
            document_id=doc.id,
            body_text="名单附件",
            fetched_at=now,
            status="ok",
            final_url=doc.url,
            attachments=[
                DocumentAttachment(position=i, name=f"名单{i}", url=link, extension="xlsx")
                for i, link in enumerate(links)
            ],
        )
    )
    session.commit()
    fake = FakeFetcher(
        {
            doc.url: FetchResult(
                doc.url, "ok", final_url=doc.url, text="<html><body></body></html>"
            ),
            **{
                link: FetchResult(
                    link,
                    "ok",
                    final_url=link,
                    content=(FIXTURES / "heilongjiang_30.xlsx").read_bytes(),
                )
                for link in links
            },
        }
    )
    report = parse_lists(session, attachment_limit=2, fetcher=fake)
    assert report.attachments_fetched == 2
    assert report.attachments_deferred == 1
    assert report.failures == {}
    assert links[2] not in fake.calls
    fake.responses[links[2]].content = b"damaged spreadsheet"
    retry = parse_lists(session, attachment_limit=1, fetcher=fake)
    assert retry.attachments_fetched == 1
    assert retry.failures["解析失败：BadZipFile"] == 1
    assert retry.parsed_documents == 1


def _direct_list_document(session, url: str) -> None:
    source = Source(
        level="省级",
        province="北京",
        dept="科技",
        dept_full_name="北京市科学技术委员会",
        column_name="通知公告",
        url="https://kw.beijing.gov.cn/list/",
        fetch_mode="static",
        enabled=True,
        parser_config={},
        rights_level=1,
    )
    session.add(source)
    session.flush()
    now = datetime.now(UTC)
    session.add(
        Document(
            source_id=source.id,
            url=url,
            url_key=url.split("://", 1)[1],
            title="2026年高新技术企业名单公告",
            publish_date=date(2026, 8, 1),
            first_seen_at=now,
            last_seen_at=now,
            is_baseline=False,
        )
    )
    session.commit()


def test_direct_unsupported_format_is_not_fetched(session) -> None:
    _direct_list_document(session, "https://kw.beijing.gov.cn/old-list.et")
    fake = FakeFetcher({})
    report = parse_lists(session, fetcher=fake)
    assert report.failures == {"暂不支持旧版或专有格式": 1}
    assert fake.calls == []


def test_direct_old_doc_is_parsed(session) -> None:
    """10/16 起 Word 97 二进制（.doc/.wps）也能读：文件本身就是 .doc 的通知照常下载解析。"""
    url = "https://kw.beijing.gov.cn/old-list.doc"
    _direct_list_document(session, url)
    content = (FIXTURES / "shanxi_green_factories.doc").read_bytes()
    fake = FakeFetcher({url: FetchResult(url, "ok", final_url=url, content=content)})
    report = parse_lists(session, fetcher=fake)
    assert fake.calls == [url]
    assert report.failures == {}
    assert session.query(ListEntry).count() == 79


def test_cells_drop_line_breaks_inside_chinese_names() -> None:
    """表格格子里名称折行（PDF 表格常见），读出来不要留空格；英文名称里的空格保留。"""
    from youju.lists.parser import _cell

    assert _cell("广睿创业服务有限\n公司") == "广睿创业服务有限公司"
    assert _cell("襄阳捌捌零网络科技有限 公司") == "襄阳捌捌零网络科技有限公司"
    assert _cell("ABC Technology Co., Ltd.") == "ABC Technology Co., Ltd."
    assert _cell(" 北京\xa0某某 科技 （中国） 有限公司 ") == "北京某某科技 （中国） 有限公司"
    assert _cell(12.0) == "12" and _cell(None) == ""
