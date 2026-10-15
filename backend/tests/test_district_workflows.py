"""区县条件跨输入、查询、任务、导出的回归；只用合成来源。"""

import base64
from datetime import UTC, date, datetime

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select

from tests.test_advisor import advisor_settings as advisor_settings
from tests.test_advisor import agree_advisor
from tests.test_advisor import client as client
from tests.test_city_workflows import city_docs as city_docs
from youju.ai import clause
from youju.api import advisor, me, v1
from youju.api.app import SourceCreateIn, SourceOut
from youju.classify import applicable_region
from youju.compare import CompareInputError, compare_lists, export_compare_excel
from youju.models import Document, DocumentDeadline, ListEntry, Source
from youju.regions import covered_district_scopes, resolve_district
from youju.sources.importer import import_sources
from youju.track import export_track_excel, parse_customer_profile, track_clients


@pytest.fixture
def district_docs(session, city_docs):
    stamp = datetime(2026, 10, 15, 10, tzinfo=UTC)
    docs = {}
    for index, (province, city, district) in enumerate(
        [
            ("山东", "青岛", "市南区"),
            ("山东", "青岛", "市北区"),
            ("山东", "济南", "市南区"),
            ("北京", "北京", "朝阳区"),
            ("吉林", "长春", "朝阳区"),
        ]
    ):
        source = Source(
            level="区县级",
            province=province,
            city=city,
            district=district,
            dept="科技",
            dept_full_name=f"合成{city}{district}科技部门",
            column_name="通知",
            url=f"https://example.gov.cn/district-source/{index}",
            rights_level=1,
            last_success_at=stamp,
        )
        session.add(source)
        session.flush()
        doc = Document(
            source_id=source.id,
            title=f"关于组织申报合成区县科技项目的通知{index}",
            url=f"https://example.gov.cn/district-doc/{index}",
            url_key=f"district/{index}",
            publish_date=date(2026, 10, 14),
            first_seen_at=stamp,
            last_seen_at=stamp,
            kind_override="申报通知",
        )
        session.add(doc)
        session.flush()
        session.add(
            DocumentDeadline(
                document_id=doc.id,
                fingerprint=f"district-deadline-{index}",
                kind="企业申报截止",
                candidate_date=date(2026, 10, 20),
                approved_date=date(2026, 10, 20),
                evidence="合成截止",
                confidence="高",
                status="approved",
            )
        )
        docs[f"{province}/{city}/{district}"] = doc
    session.commit()
    return docs


def test_search_follow_and_coverage_keep_district_without_parent_city_source(
    client, session, district_docs
):
    agree_advisor(client)
    for path in ["山东/青岛/市南区", "北京/北京/朝阳区", "吉林/长春/朝阳区"]:
        res = client.get("/api/v1/search", params={"districts": path, "time": "不限"})
        assert res.status_code == 200, res.text
        assert [row["id"] for row in res.json()["items"]] == [str(district_docs[path].id)]
        assert res.json()["items"][0]["district"] == path.split("/")[2]
        follow = client.post(
            "/api/v1/me/follows",
            json={"name": "区县关注", "conditions": {"districts": path, "time": "不限"}},
        )
        assert follow.status_code == 201, follow.text
        assert follow.json()["conditions"]["districts"] == path
        ids = {
            row[0].id
            for row in session.execute(me._conditions_stmt(session, follow.json()["conditions"]))
        }
        assert ids == {district_docs[path].id}
    coverage = client.get("/api/v1/coverage").json()
    assert any(
        row["province"] == "吉林" and row["city"] == "长春" and row["district"] == "朝阳区"
        for row in coverage["districts"]
    )
    # Coverage never fabricates a successful city-level source for a district-only parent.
    assert not any(row["city"] == "长春" for row in coverage["cities"])
    missing = client.get("/api/v1/search", params={"districts": "山东/青岛/不存在区"})
    assert missing.status_code == 422 and "不会改为全市" in missing.text
    for invalid in ["山东/青岛", "山东//市南区", "山东/青岛/市南", "北京/长春/朝阳区"]:
        assert client.get("/api/v1/search", params={"districts": invalid}).status_code == 422


@pytest.mark.parametrize("excerpts", [None, 1])
def test_deadline_calendar_exact_district(session, district_docs, excerpts):
    for province, city, district in [("山东", "青岛市", "市南区"), ("北京", "北京市", "朝阳区")]:
        report = v1.deadline_report(
            session, excerpts=excerpts, province=province, city=city, district=district
        )
        assert report["total"] == 1
        assert report["items"][0]["district"] == district
        calendar = v1.calendar(
            session, excerpts=excerpts, province=province, city=city, district=district
        )
        assert len(calendar["items"]) == 1
        assert calendar["items"][0]["district"] == district
        assert district in calendar["items"][0]["region"]
    with pytest.raises(HTTPException, match="不会改为全市"):
        v1.calendar(session, province="山东", city="青岛", district="未知区")


def test_track_six_columns_json_export_and_uncovered_scope(
    session, district_docs, city_docs, tmp_path
):
    data = "客户名称,省份,关键词,支持方式,城市,区县\n合成科技有限公司,山东,,,青岛市,市南区"
    body = advisor.TaskIn(
        kind="track",
        file={"name": "synthetic.csv", "content_base64": base64.b64encode(data.encode()).decode()},
    )
    rows = advisor._file_rows(body)
    assert rows[0]["district"] == "市南区"
    report = track_clients(session, rows)
    row = report.clients[0]
    assert row.profile.district == "市南区"
    assert {i.document_id for i in row.items} == {
        district_docs["山东/青岛/市南区"].id,
        *(city_docs[k].id for k in ["国家", "山东", "青岛"]),
    }
    output = tmp_path / "track.xlsx"
    export_track_excel(report, output)
    book = load_workbook(output)
    sheet = book["客户相关通知"]
    columns = {cell.value: i for i, cell in enumerate(sheet[2])}
    assert all(r[columns["客户区县"]].value == "市南区" for r in sheet.iter_rows(min_row=3))
    assert any(r[columns["区县"]].value == "市南区" for r in sheet.iter_rows(min_row=3))
    missing = track_clients(
        session, [{"name": "合成客户", "province": "山东", "city": "青岛", "district": "未知区"}]
    ).clients[0]
    assert missing.profile.invalid_district and not missing.items
    assert (
        parse_customer_profile(
            {"name": "合成客户", "province": "北京", "city": "北京市", "district": "朝阳区"}
        ).city
        == "北京"
    )
    assert (
        parse_customer_profile({"name": "合成客户", "province": "北京", "city": "北京市"}).city
        == ""
    )


@pytest.mark.parametrize("value", [42, ["市南区"], "市南区/市北区", "市南区;市北区", "区" * 33])
def test_invalid_district_rejected_before_task_quota(value):
    with pytest.raises(CompareInputError):
        parse_customer_profile(
            {"name": "合成客户", "province": "山东", "city": "青岛", "district": value}
        )
    with pytest.raises(ValidationError):
        clause.Company(region="山东", city="青岛", district=value)


def test_district_needs_parents_and_company_display():
    with pytest.raises(CompareInputError):
        parse_customer_profile({"name": "合成客户", "province": "山东", "district": "市南区"})
    with pytest.raises(ValidationError):
        clause.Company(region="山东", district="市南区")
    company = clause.Company(region="北京", city="市辖区", district="朝阳区")
    assert company.city == "北京"
    shown = clause.with_company(
        {"rows": [{"type": "条件", "text": "注册地", "keys": ["region"], "refs": []}]},
        {"company": company.model_dump(mode="json")},
    )
    assert shown["rows"][0]["company"][0]["value"] == "北京 · 朝阳区"
    assert applicable_region("北京", "区县级", "北京", "朝阳区") == "北京·朝阳区"


def test_compare_export_preserves_source_district_not_entry_region(
    session, district_docs, tmp_path
):
    doc = district_docs["山东/青岛/市南区"]
    doc.title = "关于公布合成科技企业认定名单的通知"
    doc.kind_override = "名单"
    session.add(
        ListEntry(
            document_id=doc.id,
            source_url=doc.url,
            name="合成区县科技有限公司",
            normalized_name="合成区县科技有限公司",
            region="其他地区",
            category="高新技术企业",
        )
    )
    session.commit()
    report = compare_lists(session, ["合成区县科技有限公司"])
    match = report.rows[0].matches[0]
    assert match.district == "市南区" and match.region == "其他地区"
    assert match.as_dict()["district"] == "市南区"
    output = tmp_path / "compare.xlsx"
    export_compare_excel(report, output)
    sheet = load_workbook(output)["明细"]
    row = dict(zip([c.value for c in sheet[2]], [c.value for c in sheet[3]], strict=True))
    assert row["区县"] == "市南区" and row["名单中企业地区"] == "其他地区"


def test_source_district_create_csv_and_exact_resolution(session, tmp_path):
    body = SourceCreateIn(
        url="https://example.gov.cn/area",
        level="区县级",
        province="北京",
        city="北京市",
        district="朝阳区",
        dept="科技",
        dept_full_name="合成科技局",
        column_name="通知",
        trial_token="synthetic-token",
        reviewed_by="合成测试",
    )
    assert body.city == "北京"
    with pytest.raises(ValidationError):
        SourceCreateIn(**{**body.model_dump(), "level": "市级"})
    path = tmp_path / "sources.csv"
    path.write_text(
        "层级,省份,城市,区县,部门,部门全称,栏目名称,栏目URL,验证状态\n区县级,北京,北京市,朝阳区,科技,合成科技局,通知,https://example.gov.cn/area,已验证",
        encoding="utf-8",
    )
    import_sources(session, path)
    source = session.scalar(select(Source))
    assert source.city == "北京" and source.district == "朝阳区"
    assert SourceOut.model_validate(source).district == "朝阳区"
    # A prepared but never successfully crawled source is not coverage.
    assert not covered_district_scopes(session)
    assert resolve_district("山东", "青岛市", "市南区", {("山东", "青岛", "市南区")}) == (
        "青岛",
        "市南区",
    )
    assert resolve_district("山东", "青岛", "市南", {("山东", "青岛", "市南区")}) is None


def test_district_migration_preserves_old_source_values(session, db_engine):
    from alembic import command
    from sqlalchemy import text

    from youju.migrations import alembic_config, upgrade

    source = Source(
        level="市级",
        province="山东",
        city="青岛",
        dept="科技",
        dept_full_name="合成科技局",
        column_name="通知",
        url="https://example.gov.cn/old-source",
    )
    session.add(source)
    session.commit()
    source_id = source.id
    session.close()
    url = db_engine.url.render_as_string(hide_password=False)
    try:
        command.downgrade(alembic_config(url), "0028")
        with db_engine.connect() as conn:
            assert (
                conn.execute(
                    text("SELECT city FROM sources WHERE id=:id"), {"id": source_id}
                ).scalar()
                == "青岛"
            )
        upgrade(url)
        with db_engine.connect() as conn:
            row = conn.execute(
                text("SELECT province, city, district FROM sources WHERE id=:id"), {"id": source_id}
            ).one()
            assert tuple(row) == ("山东", "青岛", None)
    finally:
        upgrade(url)


@pytest.mark.parametrize("excerpts", [False, True])
@pytest.mark.parametrize(
    "status,evidence",
    [
        ("approved", "请于2026年10月20日前提交。"),
        ("changed", "原定2026年10月20日前提交。"),
        ("approved", ""),
    ],
)
def test_actual_excel_word_keep_verified_original_quote(
    session, district_docs, excerpts, status, evidence
):
    from io import BytesIO

    from docx import Document as WordDocument

    from youju.api import exports

    doc = district_docs["山东/青岛/市南区"]
    deadline = session.scalar(
        select(DocumentDeadline).where(DocumentDeadline.document_id == doc.id)
    )
    deadline.evidence = evidence
    if status == "changed":
        deadline.change_reason = "延期"
        changed = Document(
            source_id=doc.source_id,
            title="关于合成区县项目延期的通知",
            url="https://example.gov.cn/district-extension",
            url_key="district-extension",
            publish_date=date(2026, 10, 15),
            first_seen_at=doc.first_seen_at,
            last_seen_at=doc.last_seen_at,
            kind_override="申报通知",
        )
        session.add(changed)
        session.flush()
        deadline.changed_to_document_id = changed.id
    deadline.status = status
    session.commit()
    items = exports._items(session, [str(doc.id)], excerpts=excerpts)
    assert items[0].deadline_quote == evidence
    assert items[0].region == "山东·青岛·市南区"
    sheet = load_workbook(BytesIO(exports._xlsx(items, "2026-10-15", excerpts)))["文件清单"]
    row = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]], strict=True))
    assert (row["截止原文句子"] or "") == evidence
    assert row["地区"] == "山东·青岛·市南区"
    word = WordDocument(BytesIO(exports._docx(items, "2026-10-15", excerpts)))
    fields = {r.cells[0].text: r.cells[1].text for r in word.tables[1].rows}
    assert fields["截止原文句子"] == evidence
    assert fields["地区"] == "山东·青岛·市南区"


def test_district_changes_do_not_share_repost_identity_or_deadline_change_target(
    session, district_docs
):
    from youju.api.app import _check_change_target
    from youju.repost import Candidate, member_facts

    first, second = [district_docs[f"山东/青岛/{d}"] for d in ["市南区", "市北区"]]
    args = dict(
        title="关于合成项目申报的通知",
        kind="申报通知",
        province="山东",
        city="青岛",
        dept="科技",
        publish_date=date(2026, 10, 14),
    )
    assert (
        Candidate(id=1, district="市南区", **args).agency
        != Candidate(id=2, district="市北区", **args).agency
    )
    facts, fingerprints = member_facts(session, [first.id, second.id])
    assert facts[first.id][0] != facts[second.id][0]
    deadline = session.scalar(
        select(DocumentDeadline).where(DocumentDeadline.document_id == first.id)
    )
    with pytest.raises(HTTPException, match="不是同一地区"):
        _check_change_target(session, deadline, second.id)
    old = fingerprints[first.id]
    first.source.district = "新区"
    session.commit()
    assert member_facts(session, [first.id])[1][first.id] != old


@pytest.mark.parametrize(
    "city", ["青岛/济南", "青岛\\济南", "青岛、济南", "青岛|济南", "青岛\t济南"]
)
def test_company_single_city_validation_is_not_conditional_on_district(city):
    with pytest.raises(ValidationError):
        clause.Company(region="山东", city=city)


def test_whitespace_parents_cannot_expand_district_scope():
    with pytest.raises(CompareInputError):
        parse_customer_profile(
            {"name": "合成客户", "province": "山东", "city": " ", "district": "市南区"}
        )
    with pytest.raises(ValidationError):
        clause.Company(region="山东", city=" ", district="市南区")
    assert clause.Company(region="山东", district=" ").district == ""
    assert clause.Company(region="山东").district is None
