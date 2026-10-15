"""公开导出只包含归并后的题录、原文网址和人工核对字段。"""

import hashlib
import json
from collections import deque
from datetime import UTC, date, datetime, time
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from youju.api import exports
from youju.api.app import app
from youju.api.exports import MAX_BYTES
from youju.db import get_session
from youju.models import (
    Document,
    DocumentDeadline,
    DocumentDetail,
    DocumentField,
    DocumentMerge,
    Source,
)


@pytest.fixture
def client(session_factory):
    def override():
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_session] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture
def docs(session):
    source = Source(
        level="省级",
        province="辽宁",
        dept="工信",
        dept_full_name="辽宁省工业和信息化厅",
        column_name="通知",
        url="https://gxt.ln.gov.cn/tz/",
        enabled=True,
        parser_config={},
        verified_on=date(2026, 9, 25),
    )
    session.add(source)
    session.flush()

    def add_doc(number, title, day=date(2026, 9, 24)):
        now = datetime.now(UTC)
        doc = Document(
            source_id=source.id,
            url=f"https://gxt.ln.gov.cn/{number}.html",
            url_key=f"gxt.ln.gov.cn/{number}.html",
            title=title,
            publish_date=day,
            first_seen_at=now,
            last_seen_at=now,
        )
        session.add(doc)
        session.flush()
        return doc

    title = "辽宁省关于组织申报科技项目的通知"
    original = add_doc(1, title)
    duplicate = add_doc(2, title)
    listed = add_doc(3, "辽宁省科技型中小企业拟入库名单公示")
    changed_to = add_doc(4, "关于延长科技项目申报时间的通知", date(2026, 9, 26))
    session.add(
        DocumentMerge(
            duplicate_id=duplicate.id, canonical_id=original.id, active=True, match_key="test"
        )
    )
    session.add(
        DocumentDetail(
            document_id=duplicate.id,
            body_text="",
            issuer="辽宁省工业和信息化厅",
            doc_number="辽工信〔2026〕1号",
            fetched_at=datetime.now(UTC),
            status="ok",
        )
    )
    reviewed = datetime(2026, 9, 26, 3, tzinfo=UTC)

    def deadline(doc, fingerprint, day, status, kind="企业申报截止", **extra):
        return DocumentDeadline(
            document_id=doc.id,
            fingerprint=fingerprint,
            kind=kind,
            candidate_date=day,
            evidence="核对记录",
            confidence="高",
            status=status,
            approved_date=day if status != "pending" else None,
            reviewed_at=reviewed if status != "pending" else None,
            **extra,
        )

    session.add_all(
        [
            deadline(
                original,
                "a",
                date(2026, 10, 18),
                "changed",
                change_reason="延期",
                changed_to_document_id=changed_to.id,
                changed_at=reviewed,
            ),
            deadline(duplicate, "b", date(2026, 10, 18), "approved"),
            deadline(duplicate, "c", date(2026, 10, 20), "approved", kind="推荐或审核截止"),
            deadline(listed, "d", date(2026, 10, 5), "pending"),
        ]
    )
    session.commit()
    return original.id, duplicate.id, listed.id


def _export(client, docs, format):
    original, duplicate, listed = docs
    response = client.post(
        "/api/v1/exports",
        json={"ids": [str(duplicate), str(original), str(listed)], "format": format},
    )
    assert response.status_code == 200
    assert len(response.content) < MAX_BYTES
    assert (
        "filename*=UTF-8''%E6%9C%89%E6%8D%AE%E5%AF%BC%E5%87%BA_"
        in response.headers["content-disposition"]
    )
    return response


def test_xlsx_export(client, docs):
    response = _export(client, docs, "xlsx")
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["文件清单", "数据说明"]
    sheet = book["文件清单"]
    assert sheet.max_row == 3  # 两个原文网址归并为一行，请求顺序仍保留
    assert [sheet.cell(1, i).value for i in range(1, 13)] == [
        "序号",
        "标题",
        "类型",
        "名单状态",
        "发文机关",
        "文号",
        "发布日期",
        "地区",
        "截止日期（以原文为准）",
        "日期角色",
        "支持方式",
        "未核项",
    ]
    assert sheet["M1"].value == "截止原文句子"
    assert sheet["N1"].value == "原文链接" and sheet["O1"].value == "原文链接 2"
    assert sheet["B2"].value == "辽宁省关于组织申报科技项目的通知"
    assert sheet["D2"].value is None  # 申报通知没有名单状态，status 契约保持不变
    assert sheet["D3"].value == "公示"
    assert "2026-10-20" in sheet["I2"].value
    assert "原定 企业申报截止 2026-10-18 已延期" in sheet["I2"].value
    assert "2026-10-05" not in sheet["I3"].value and sheet["I3"].value == "见原文"
    # 日期角色按截止日期那一格的顺序写类型；已变更的注明；没有核对过的截止不写日期
    assert sheet["J2"].value == "推荐或审核截止；企业申报截止（已延期）"
    assert sheet["J3"].value == "截止时间见原文"
    assert sheet["K2"].value == "见原文"  # 申报通知还没有核对通过的支持方式
    assert sheet["K3"].value == "—"  # 名单没有支持方式
    assert sheet["L2"].value == "支持方式"
    assert sheet["L3"].value == "截止日期（以原文为准）、名单状态"
    assert sheet["N2"].hyperlink.target.endswith("/1.html")
    assert sheet["O2"].hyperlink.target.endswith("/2.html")
    note = book["数据说明"]
    assert "第三方整理，非政府平台" in note["B1"].value
    assert "“截止日期（以原文为准）”列" in note["B1"].value
    assert note["D4"].value == "2026-09-25"
    assert note["E4"].value == "2026-09-26"


def test_docx_export(client, docs):
    response = _export(client, docs, "docx")
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    book = WordDocument(BytesIO(response.content))
    headings = [p.text for p in book.paragraphs if p.style.name == "Heading 1"]
    assert len(headings) == 2
    assert headings[0].startswith("1. 辽宁省关于组织申报")
    assert "第三方整理，非政府平台" in book.paragraphs[1].text
    text = "\n".join(p.text for p in book.paragraphs)
    assert "https://gxt.ln.gov.cn/1.html" in text
    assert "https://gxt.ln.gov.cn/2.html" in text
    assert len(book.tables) == 3
    assert len(book.tables[0].rows) == 3
    assert book.tables[0].rows[1].cells[4].text == "2026-09-26"
    fields = {row.cells[0].text: row.cells[1].text for row in book.tables[1].rows}
    assert "2026-10-20" in fields["截止日期（以原文为准）"]
    assert "原定 企业申报截止 2026-10-18 已延期" in fields["截止日期（以原文为准）"]
    assert fields["截止日期核对日期"] == "2026-09-26"
    assert fields["日期角色"] == "推荐或审核截止；企业申报截止（已延期）"
    assert fields["支持方式"] == "见原文"
    assert fields["未核项"] == "支持方式"
    listed_fields = {row.cells[0].text: row.cells[1].text for row in book.tables[2].rows}
    assert listed_fields["截止日期（以原文为准）"] == "见原文"
    assert listed_fields["日期角色"] == "截止时间见原文"
    assert listed_fields["支持方式"] == "—"
    assert listed_fields["未核项"] == "截止日期（以原文为准）、名单状态"


@pytest.fixture
def role_docs(session):
    """日期角色、未核项用：全部核对过 / 截止只有待核对的候选 / 部分核对过。"""
    source = Source(
        level="省级",
        province="辽宁",
        dept="科技",
        dept_full_name="辽宁省科学技术厅",
        column_name="通知公告",
        url="https://kjt.ln.gov.cn/tz/",
        enabled=True,
        parser_config={},
        verified_on=date(2026, 9, 25),
    )
    session.add(source)
    session.flush()
    reviewed = datetime(2026, 9, 27, 2, tzinfo=UTC)
    body = "本项目给予资金支持。对符合条件的贷款给予贴息。申报时间以通知为准。"
    body_hash = hashlib.sha256(body.encode()).hexdigest()

    def add_doc(slug, title):
        now = datetime.now(UTC)
        doc = Document(
            source_id=source.id,
            url=f"https://kjt.ln.gov.cn/{slug}.html",
            url_key=f"kjt.ln.gov.cn/{slug}.html",
            title=title,
            publish_date=date(2026, 9, 24),
            first_seen_at=now,
            last_seen_at=now,
        )
        session.add(doc)
        session.flush()
        session.add(
            DocumentDetail(
                document_id=doc.id,
                body_text=body,
                issuer="辽宁省科学技术厅",
                fetched_at=now,
                status="ok",
            )
        )
        return doc

    def deadline(doc, fingerprint, kind, day, status, at=None, origin="rule"):
        return DocumentDeadline(
            document_id=doc.id,
            fingerprint=fingerprint,
            kind=kind,
            candidate_date=day,
            candidate_time=at,
            evidence="核对记录",
            confidence="高",
            status=status,
            origin=origin,
            approved_date=day if status == "approved" else None,
            approved_time=at if status == "approved" else None,
            reviewed_by="审阅人" if status != "pending" else None,
            reviewed_at=reviewed if status != "pending" else None,
        )

    def support(doc, value, status):
        return DocumentField(
            document_id=doc.id,
            field="support",
            value=value,
            evidence=body[:10],
            evidence_offset=0,
            body_hash=body_hash,
            status=status,
            approved_value=value if status == "approved" else None,
            reviewed_by="审阅人" if status == "approved" else None,
            reviewed_at=reviewed if status == "approved" else None,
        )

    # 标题会被按标题归类打上筛选标签（贷款贴息、资金支持），导出的支持方式不能用这些标签
    checked = add_doc("a", "辽宁省关于组织申报智能制造贷款贴息项目的通知")
    unchecked = add_doc("b", "辽宁省关于组织申报数字经济项目资金补助的通知")
    partial = add_doc("c", "辽宁省关于组织申报新材料项目的通知")
    session.add_all(
        [
            deadline(checked, "a1", "企业申报截止", date(2026, 11, 10), "approved"),
            support(checked, "资格认定", "approved"),
            support(checked, "资金支持", "approved"),
            # 只有待核对、已驳回的候选：任何格子里都不能出现这些日期
            deadline(unchecked, "b1", "企业申报截止", date(2026, 11, 23), "pending", time(17, 0)),
            deadline(unchecked, "b2", "材料报送截止", date(2026, 11, 27), "pending", origin="ai"),
            deadline(unchecked, "b3", "推荐或审核截止", date(2026, 11, 25), "rejected"),
            support(unchecked, "资金支持", "pending"),
            # 已核对一部分：列出的已核对，另有待核对的截止和支持方式
            deadline(partial, "c1", "材料报送截止", date(2026, 12, 1), "approved"),
            deadline(partial, "c2", "推荐或审核截止", date(2026, 12, 8), "pending"),
            support(partial, "资金支持", "approved"),
            support(partial, "贷款贴息", "pending"),
        ]
    )
    session.commit()
    return checked, unchecked, partial


ROW_COLUMNS = ("截止日期（以原文为准）", "日期角色", "支持方式", "未核项")


def _export_rows(client, docs, format_):
    """导出后按标题取出（截止日期, 日期角色, 支持方式, 未核项），以及整个文件的全部文字。"""
    response = client.post(
        "/api/v1/exports", json={"ids": [str(d.id) for d in docs], "format": format_}
    )
    assert response.status_code == 200, response.text
    if format_ == "xlsx":
        book = load_workbook(BytesIO(response.content))
        sheet = book["文件清单"]
        header = [cell.value for cell in sheet[1]]
        col = {name: header.index(name) for name in ("标题", *ROW_COLUMNS)}
        rows = {
            row[col["标题"]].value: tuple(row[col[name]].value for name in ROW_COLUMNS)
            for row in sheet.iter_rows(min_row=2)
        }
        text = "\n".join(str(cell.value) for ws in book for row in ws for cell in row)
        book.close()
        return rows, text
    book = WordDocument(BytesIO(response.content))
    rows = {}
    for table in book.tables[1:]:
        fields = {row.cells[0].text: row.cells[1].text for row in table.rows}
        rows[fields["标题"]] = tuple(fields[name] for name in ROW_COLUMNS)
    text = "\n".join(
        [p.text for p in book.paragraphs]
        + [cell.text for table in book.tables for row in table.rows for cell in row.cells]
    )
    return rows, text


@pytest.mark.parametrize("format_", ["xlsx", "docx"])
def test_date_role_and_unchecked_columns(client, role_docs, format_):
    checked, unchecked, partial = role_docs
    rows, text = _export_rows(client, role_docs, format_)
    assert rows[checked.title] == (
        "企业申报截止 2026-11-10",
        "企业申报截止",
        "资金支持、资格认定",  # 只写人工核对通过的，顺序和详情页一致
        "无",
    )
    assert rows[unchecked.title] == (
        "见原文",
        "截止时间见原文",
        "见原文",
        "截止日期（以原文为准）、支持方式",
    )
    assert rows[partial.title] == (
        "材料报送截止 2026-12-01",
        "材料报送截止",
        "资金支持",  # 待核对的贷款贴息不写
        "其他截止日期（以原文为准）、支持方式",
    )
    # 按标题归类的筛选标签不进支持方式这一列
    assert checked.rule_supports == "贷款贴息" and unchecked.rule_supports == "资金支持"
    assert "贷款贴息" not in rows[checked.title][2]
    # 没有人工核对通过的截止日期，不管哪种写法都不能出现在文件里（含日期角色、数据说明）
    for day in ("2026-11-23", "2026-11-25", "2026-11-27", "2026-12-08"):
        y, m, d = day.split("-")
        for shown in (day, f"{y}/{m}/{d}", f"{int(m)}月{int(d)}日", f"{m}-{d}", f"{m}月{d}日"):
            assert shown not in text, shown
    assert "以发文机关原文为准" in text and "未核项" in text


def test_word_and_excel_columns_match(client, role_docs, docs, session):
    ids = [*role_docs, *(session.get(Document, i) for i in docs)]
    xlsx, _ = _export_rows(client, ids, "xlsx")
    docx, _ = _export_rows(client, ids, "docx")
    assert len(xlsx) == 5  # 归并的两个网址算一行
    assert xlsx == docx


def test_export_rejects_missing_ids_and_over_200(client, docs):
    original, _, _ = docs
    missing = client.post(
        "/api/v1/exports", json={"ids": [str(original), "999999", "bad"], "format": "xlsx"}
    )
    assert missing.status_code == 422
    assert missing.json()["detail"]["不存在的文件编号"] == ["999999", "bad"]
    too_many = client.post("/api/v1/exports", json={"ids": [str(original)] * 201, "format": "docx"})
    assert too_many.status_code == 422


def test_export_rejects_file_at_size_limit(client, docs, monkeypatch):
    monkeypatch.setattr(exports, "_xlsx", lambda items, stamp: b"x" * MAX_BYTES)
    response = client.post("/api/v1/exports", json={"ids": [str(docs[0])], "format": "xlsx"})
    assert response.status_code == 413


def test_restricted_sources_are_not_exported_or_listed(client, docs, session):
    """只有 2 级来源（只存题录）、没有 1 级副本的文件：搜索、详情、导出都当作没有。"""
    restricted = Source(
        level="省级",
        province="辽宁",
        dept="工信",
        dept_full_name="某授权平台",
        column_name="资讯",
        url="https://example.gov.cn/zx/",
        enabled=True,
        parser_config={},
        rights_level=2,
    )
    session.add(restricted)
    session.flush()
    now = datetime.now(UTC)
    doc = Document(
        source_id=restricted.id,
        url="https://example.gov.cn/zx/1.html",
        url_key="example.gov.cn/zx/1.html",
        title="辽宁省关于组织申报智能制造项目的通知",
        publish_date=date(2026, 9, 24),
        first_seen_at=now,
        last_seen_at=now,
    )
    session.add(doc)
    session.commit()
    exported = client.post("/api/v1/exports", json={"ids": [str(doc.id)], "format": "xlsx"})
    assert exported.status_code == 422
    assert exported.json()["detail"]["不存在的文件编号"] == [str(doc.id)]
    assert client.get(f"/api/v1/documents/{doc.id}").status_code == 404
    found = client.get("/api/v1/search", params={"q": "智能制造", "time": "不限"}).json()
    assert found["total"] == 0


def test_exports_are_rate_limited_per_user(client, docs, monkeypatch):
    from youju.api import exports

    original, _, _ = docs
    monkeypatch.setattr(exports, "get_settings", lambda: SimpleNamespace(exports_per_10min=2))
    monkeypatch.setattr(
        "youju.api.me.get_settings",
        lambda: SimpleNamespace(trust_wx_headers=True, dev_openid=None),
    )
    body = {"ids": [str(original)], "format": "xlsx"}

    def post(openid):
        headers = {"X-WX-OPENID": openid, "X-WX-SOURCE": "wx"}
        return client.post("/api/v1/exports", json=body, headers=headers).status_code

    assert [post("o-a"), post("o-a"), post("o-a")] == [200, 200, 429]
    assert post("o-b") == 200  # 别的用户不受影响


def test_export_limits_evict_idle_callers_and_cap_each_instance(monkeypatch):
    from youju.api import exports

    monkeypatch.setattr(exports, "get_settings", lambda: SimpleNamespace(exports_per_10min=2))
    monkeypatch.setattr(exports, "_EXPORT_KEYS_MAX", 3)
    monkeypatch.setattr(exports, "_EXPORTS_PER_INSTANCE", 5)
    monkeypatch.setattr(exports, "_instance_times", deque())
    now = 1000.0
    assert [exports._allow_export("u:a", now), exports._allow_export("u:a", now)] == [None, None]
    assert exports._allow_export("u:a", now) == "user"
    for caller in ("u:b", "u:c", "u:d"):  # 记录超过上限：淘汰最久没导出的，不整个清空
        assert exports._allow_export(caller, now + 1) is None
    assert list(exports._export_times) == ["u:b", "u:c", "u:d"]
    assert exports._allow_export("u:e", now + 2) == "busy"  # 这个实例 10 分钟已导出 5 次
    assert exports._allow_export("u:e", now + 602) is None  # 过了 10 分钟又能导出


def test_home_and_calendar_skip_restricted_sources(client, session, monkeypatch):
    """首页的本周计数、最新文件和日历都不算只有 2 级来源的文件。"""
    from youju.api import v1

    monkeypatch.setattr(v1, "now", lambda: datetime(2026, 9, 29, 10, 0, tzinfo=UTC))
    restricted = Source(
        level="省级",
        province="辽宁",
        dept="工信",
        dept_full_name="某授权平台",
        column_name="资讯",
        url="https://example.gov.cn/zx/",
        enabled=True,
        parser_config={},
        rights_level=2,
    )
    session.add(restricted)
    session.flush()
    now = datetime.now(UTC)
    doc = Document(
        source_id=restricted.id,
        url="https://example.gov.cn/zx/9.html",
        url_key="example.gov.cn/zx/9.html",
        title="辽宁省关于组织申报智能制造项目的通知",
        publish_date=date(2026, 9, 29),
        first_seen_at=now,
        last_seen_at=now,
    )
    session.add(doc)
    session.flush()
    session.add(
        DocumentDeadline(
            document_id=doc.id,
            fingerprint="r" * 64,
            kind="企业申报截止",
            candidate_date=date(2026, 10, 10),
            approved_date=date(2026, 10, 10),
            evidence="请于10月10日前申报",
            confidence="高",
            status="approved",
            reviewed_by="审阅人",
            reviewed_at=now,
        )
    )
    session.commit()
    home = client.get("/api/v1/home").json()
    assert home["stats"]["week_total"] == 0
    assert all(item["id"] != str(doc.id) for item in home["items"])
    calendar = json.dumps(client.get("/api/v1/calendar").json(), ensure_ascii=False)
    assert "智能制造" not in calendar


def test_excerpt_export_marks_machine_excerpts(client, session, docs, monkeypatch):
    """D-061：小程序在请求体里带 excerpts=1（网址参数也认）。8–30 天的机器摘录在截止日期后写
    "（机器摘录）"，旁边一列放原句；说明里写明未经人工核对。不带时和以前一样，不出现候选日期。"""
    from zoneinfo import ZoneInfo

    from sqlalchemy import select

    from youju.api import v1

    monkeypatch.setattr(
        v1, "now", lambda: datetime(2026, 10, 12, 9, tzinfo=ZoneInfo("Asia/Shanghai"))
    )
    source = session.scalar(select(Source))
    stamp = datetime.now(UTC)
    doc = Document(
        source_id=source.id,
        url="https://gxt.ln.gov.cn/9.html",
        url_key="gxt.ln.gov.cn/9.html",
        title="辽宁省关于组织申报2026年度工业项目的通知",
        publish_date=date(2026, 10, 9),
        first_seen_at=stamp,
        last_seen_at=stamp,
    )
    session.add(doc)
    session.flush()
    quote = "请于2026年10月25日前将申报材料报送至省工业和信息化厅。"
    session.add(
        DocumentDetail(
            document_id=doc.id, body_text=f"一、申报要求\n{quote}", status="ok", fetched_at=stamp
        )
    )
    session.add(
        DocumentDeadline(
            document_id=doc.id,
            fingerprint="excerpt",
            kind="企业申报截止",
            candidate_date=date(2026, 10, 25),
            evidence=quote,
            confidence="高",
            status="pending",
        )
    )
    session.commit()

    def sheet(**request):
        response = client.post("/api/v1/exports", **request)
        assert response.status_code == 200, response.text
        book = load_workbook(BytesIO(response.content))
        rows = [[cell.value for cell in row] for row in book["文件清单"].iter_rows()]
        return dict(zip(rows[0], rows[1], strict=False)), book["数据说明"]["B1"].value

    body = {"ids": [str(doc.id)], "format": "xlsx"}
    for request in ({"json": {**body, "excerpts": 1}}, {"json": body, "params": {"excerpts": 1}}):
        row, notice = sheet(**request)
        assert row["截止日期（以原文为准）"] == "2026-10-25（机器摘录）"
        assert row["日期角色"] == "企业申报截止（机器摘录）"
        assert row["未核项"] == "截止日期（以原文为准）、支持方式"
        assert row["截止原文句子"] == quote
        assert "机器摘录" in notice and "未经人工核对" in notice
        assert "机器摘录的“日期角色”也标有“（机器摘录）”" in notice
        assert "“日期角色”是人工核对时确认的截止类型" not in notice
        assert "“截止日期（以原文为准）”列" in notice
    row, notice = sheet(json=body)
    assert row["截止原文句子"] is None and "2026-10-25" not in str(row["截止日期（以原文为准）"])
    assert "机器摘录" not in notice


def test_excerpt_export_keeps_all_verified_deadlines_and_cleans_legacy_fields(
    client, session, docs, monkeypatch
):
    from zoneinfo import ZoneInfo

    from sqlalchemy import select

    from youju.api import v1

    monkeypatch.setattr(
        v1, "now", lambda: datetime(2026, 10, 12, 9, tzinfo=ZoneInfo("Asia/Shanghai"))
    )
    original, duplicate, _ = docs
    session.get(Document, original).url = "http://gxt.ln.gov.cn:443/1.html"
    detail = session.scalar(select(DocumentDetail).where(DocumentDetail.document_id == duplicate))
    detail.issuer = "安 徽 省 财 政 厅 福建省职称改革工作办公室"
    session.commit()
    response = client.post(
        "/api/v1/exports",
        json={"ids": [str(original)], "format": "xlsx", "excerpts": 1},
    )
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    rows = [[cell.value for cell in row] for row in book["文件清单"].iter_rows()]
    row = dict(zip(rows[0], rows[1], strict=False))
    assert "推荐或审核截止 2026-10-20" in row["截止日期（以原文为准）"]
    assert "原定 企业申报截止 2026-10-18 已延期" in row["截止日期（以原文为准）"]
    assert row["日期角色"] == "推荐或审核截止；企业申报截止（已延期）"
    assert row["发文机关"] == "安徽省财政厅、福建省职称改革工作办公室"
    assert row["原文链接"] == "https://gxt.ln.gov.cn/1.html"


def test_excerpt_export_only_falls_back_when_no_verified_deadline(
    client, session, role_docs, monkeypatch
):
    from zoneinfo import ZoneInfo

    from sqlalchemy import select

    from youju.api import v1

    monkeypatch.setattr(
        v1, "now", lambda: datetime(2026, 10, 12, 9, tzinfo=ZoneInfo("Asia/Shanghai"))
    )
    _, unchecked, partial = role_docs
    quote = "推荐或审核截止：2026年10月25日。"
    pending = session.scalar(
        select(DocumentDeadline).where(
            DocumentDeadline.document_id == partial.id,
            DocumentDeadline.status == "pending",
        )
    )
    pending.candidate_date = date(2026, 10, 25)
    pending.evidence = quote
    detail = session.scalar(select(DocumentDetail).where(DocumentDetail.document_id == partial.id))
    detail.body_text += quote
    session.commit()
    response = client.post(
        "/api/v1/exports",
        json={"ids": [str(partial.id), str(unchecked.id)], "format": "xlsx", "excerpts": 1},
    )
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    rows = [[cell.value for cell in row] for row in book["文件清单"].iter_rows()]
    records = {row[1]: dict(zip(rows[0], row, strict=False)) for row in rows[1:]}
    assert records[partial.title]["截止日期（以原文为准）"] == "材料报送截止 2026-12-01"
    assert records[partial.title]["日期角色"] == "材料报送截止"
    assert "其他截止日期（以原文为准）" in records[partial.title]["未核项"]
    assert records[unchecked.title]["截止日期（以原文为准）"] == "截止未核对，以原文为准"
    assert "截止日期（以原文为准）" in records[unchecked.title]["未核项"]
    assert "其他截止日期" not in records[unchecked.title]["未核项"]
