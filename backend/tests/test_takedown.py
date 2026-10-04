"""真实 MySQL 上逐接口验证下架、归并组、历史、缓存和恢复。"""

from datetime import datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import httpx
import pytest
from alembic import command
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import inspect, select

from youju import advisor_tasks, track
from youju.api import advisor, exports, me, v1
from youju.api import app as admin
from youju.config import Settings
from youju.db import get_session
from youju.migrations import alembic_config
from youju.models import (
    AccessLog,
    AdvisorTask,
    Document,
    DocumentDeadline,
    DocumentMerge,
    DocumentTakedownEvent,
    Feedback,
    ListEntry,
    Source,
    Subscription,
    User,
)
from youju.notify import run_notify
from youju.takedown import is_taken_down

NOW = datetime(2026, 9, 30, 10, tzinfo=ZoneInfo("Asia/Shanghai"))
HEADERS = {"X-Admin-Token": "takedown-test-only"}
NOTE = {"note": "已核实，按要求处理", "operated_by": "测试管理员"}
REMOVE = {**NOTE, "reason": "侵权投诉"}


@pytest.fixture
def settings(monkeypatch):
    value = Settings(
        _env_file=None,
        admin_token=HEADERS["X-Admin-Token"],
        dev_openid="takedown-user",
        advisor_enabled=True,
        cloud_env="test-env",
        agreement_version="2026-09",
        send_wx_messages=True,
        template_deadline={"id": "due", "fields": {"thing1": "title"}},
        template_newdocs={"id": "new", "fields": {"thing1": "summary"}},
        exports_per_10min=100,
    )
    for module in (admin, advisor, me, exports):
        monkeypatch.setattr(module, "get_settings", lambda: value)
    monkeypatch.setattr(v1, "now", lambda: NOW)
    monkeypatch.setattr(track, "now", lambda: NOW)
    return value


@pytest.fixture
def client(session_factory, settings):
    def override():
        with session_factory() as s:
            yield s

    admin.app.dependency_overrides[get_session] = override
    with TestClient(admin.app) as c:
        yield c
    admin.app.dependency_overrides.clear()


@pytest.fixture
def docs(session, client):
    src = Source(
        level="省级",
        province="辽宁",
        dept="科技",
        dept_full_name="辽宁省科技厅",
        column_name="通知",
        url="https://example.gov.cn/tz",
    )
    session.add(src)
    session.flush()
    result = []
    for n, title in enumerate(
        [
            "辽宁省科技项目申报通知",
            "辽宁省科技项目申报通知",
            "辽宁省2026年科技型中小企业名单公示",
            "辽宁省2026年科技型中小企业名单公示",
            "辽宁省科技政策管理办法",
        ],
        1,
    ):
        doc = Document(
            source_id=src.id,
            url=f"https://example.gov.cn/{n}.html",
            url_key=f"example.gov.cn/{n}.html",
            title=title,
            publish_date=NOW.date(),
            first_seen_at=NOW + timedelta(minutes=1),
            last_seen_at=NOW + timedelta(minutes=1),
        )
        session.add(doc)
        session.flush()
        result.append(doc)
    for master, twin in ((result[0], result[1]), (result[2], result[3])):
        session.add(
            DocumentMerge(
                duplicate_id=twin.id, canonical_id=master.id, active=True, match_key=master.title
            )
        )
    session.add(
        DocumentDeadline(
            document_id=result[1].id,
            fingerprint="test-due",
            kind="企业申报截止",
            candidate_date=NOW.date() + timedelta(days=2),
            approved_date=NOW.date() + timedelta(days=2),
            evidence="截止原句",
            status="approved",
            confidence="高",
            reviewed_by="核对人",
            reviewed_at=NOW,
        )
    )
    session.add(
        ListEntry(
            document_id=result[3].id,
            source_url=result[3].url,
            name="辽宁创新科技有限公司",
            normalized_name="辽宁创新科技有限公司",
            category="企业",
            sequence="1",
        )
    )
    session.commit()
    assert client.post("/api/v1/me/agreement", json={"version": "2026-09"}).status_code == 200
    return result


def remove(client, doc):
    response = client.post(f"/api/admin/takedown/{doc.id}/remove", headers=HEADERS, json=REMOVE)
    assert response.status_code == 200, response.text
    return response.json()


def restore(client, doc):
    response = client.post(f"/api/admin/takedown/{doc.id}/restore", headers=HEADERS, json=NOTE)
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.parametrize(
    "path,index",
    [
        ("/api/v1/home", 0),
        ("/api/v1/search", 0),
        ("/api/v1/search?kind=名单", 2),
        ("/api/documents", 0),
        ("/api/documents?q=科技", 2),
        ("/api/v1/calendar", 0),
        ("/api/v1/reports/deadlines", 0),
        ("/api/v1/reports/lists", 2),
        ("/api/v1/lists/recent", 2),
        ("/api/v1/lists/lookup?q=辽宁创新科技有限公司", 2),
        ("/api/lists/search?q=辽宁创新科技有限公司", 2),
    ],
)
def test_each_public_listing_and_warmed_cache(client, docs, path, index):
    target = docs[index]
    before = client.get(path)
    assert before.status_code == 200, before.text
    marker = target.url if path.startswith("/api/lists/search") else target.title
    assert marker in before.text
    # 不清缓存：公开服务要从数据库历史版本识别另一进程的下架。
    remove(client, docs[index + 1])
    hidden = client.get(path)
    assert hidden.status_code == 200, hidden.text
    assert target.title not in hidden.text
    assert target.url not in hidden.text
    assert client.get(f"/api/v1/documents/{docs[4].id}").status_code == 200
    restore(client, target)
    assert marker in client.get(path).text


@pytest.mark.parametrize("prefix", ["/api/documents", "/api/v1/documents"])
@pytest.mark.parametrize("index", [0, 1, 2, 3])
def test_detail_returns_specific_404_and_restores(client, docs, prefix, index):
    path = f"{prefix}/{docs[index].id}"
    assert client.get(path).status_code == 200
    remove(client, docs[index])
    response = client.get(path)
    assert response.status_code == 404
    assert response.json()["detail"] == "该文件已下架"
    restore(client, docs[index])
    assert client.get(path).status_code == 200


@pytest.mark.parametrize("format_", ["xlsx", "docx"])
@pytest.mark.parametrize("index", [0, 1, 2, 3])
def test_exports_filter_each_member(client, docs, format_, index):
    payload = {"ids": [str(docs[index].id)], "format": format_}
    assert client.post("/api/v1/exports", json=payload).status_code == 200
    remove(client, docs[index])
    assert client.post("/api/v1/exports", json=payload).status_code == 422
    assert (
        client.post("/api/v1/exports", json={**payload, "ids": [str(docs[4].id)]}).status_code
        == 200
    )
    restore(client, docs[index])
    assert client.post("/api/v1/exports", json=payload).status_code == 200


def test_legacy_stats_and_lookup_metadata(client, docs):
    assert client.get("/api/stats").json()["documents_total"] == 3
    path = "/api/v1/lists/lookup?q=辽宁创新科技股份有限公司"
    # 相近名称、覆盖统计都不能泄露下架名单。
    assert client.get(path).json()["similar"]
    remove(client, docs[3])
    data = client.get(path).json()
    assert data["similar"] == [] and data["coverage"]["lists"] == 0
    assert client.get("/api/stats").json()["documents_total"] == 2
    restore(client, docs[2])
    assert client.get("/api/stats").json()["documents_total"] == 3


@pytest.mark.parametrize(
    "path,method",
    [
        ("/api/admin/takedown", "get"),
        ("/api/admin/takedown?removed=true", "get"),
        ("/api/admin/takedown/1/history", "get"),
        ("/api/admin/takedown/1/remove", "post"),
        ("/api/admin/takedown/1/restore", "post"),
    ],
)
def test_admin_requires_password(client, path, method):
    kwargs = {"json": REMOVE if path.endswith("remove") else NOTE} if method == "post" else {}
    assert getattr(client, method)(path, **kwargs).status_code == 401


@pytest.mark.parametrize(
    "payload",
    [
        {**REMOVE, "note": ""},
        {**REMOVE, "note": " \n "},
        {**REMOVE, "note": "字" * 301},
        {**REMOVE, "reason": "随意"},
        {**REMOVE, "operated_by": " "},
        {"reason": "侵权投诉", "operated_by": "测试"},
    ],
)
def test_required_reason_note_operator(client, docs, payload):
    assert (
        client.post(
            f"/api/admin/takedown/{docs[0].id}/remove", headers=HEADERS, json=payload
        ).status_code
        == 422
    )
    assert client.get(f"/api/v1/documents/{docs[0].id}").status_code == 200


def test_group_history_search_overview_and_audit(client, docs, session):
    for _ in range(2):
        assert remove(client, docs[1])["document_ids"] == [docs[0].id, docs[1].id]
    for member in docs[:2]:
        session.refresh(member)
        assert member.taken_down_by == NOTE["operated_by"] and member.takedown_note == NOTE["note"]
        assert member.taken_down_at is not None
    for value in (str(docs[1].id), docs[0].title, docs[1].url):
        data = client.get("/api/admin/takedown", params={"q": value}, headers=HEADERS).json()
        assert data["items"] and all(row["taken_down"] for row in data["items"])
        assert data["items"][0]["reason"] == "侵权投诉"
    removed = client.get("/api/admin/takedown?removed=true", headers=HEADERS).json()
    assert removed["total"] == 1 and removed["items"][0]["id"] == docs[0].id
    assert client.get("/api/admin/overview", headers=HEADERS).json()["todos"]["takedown"] == 1
    for _ in range(2):
        restore(client, docs[0])
    remove(client, docs[0])
    restore(client, docs[1])
    for member in docs[:2]:
        history = client.get(f"/api/admin/takedown/{member.id}/history", headers=HEADERS).json()
        assert [r["action"] for r in history] == [
            "restore",
            "takedown",
            "restore",
            "restore",
            "takedown",
            "takedown",
        ]
        session.refresh(member)
        assert member.restored_at and member.restored_by == NOTE["operated_by"]
    events = session.scalars(select(DocumentTakedownEvent)).all()
    assert len(events) == 12
    logs = [row for row in session.scalars(select(AccessLog)).all() if row.detail]
    actions = [row.detail["action"] for row in logs]
    assert actions.count("document_takedown") == 3 and actions.count("document_restore") == 3
    record = next(row.detail for row in logs if row.detail["action"] == "document_takedown")
    assert record["by"] == NOTE["operated_by"] and record["document_ids"] == [1, 2]
    assert record["reason"] == "侵权投诉" and record["note"] == NOTE["note"]
    assert client.get("/api/admin/takedown?removed=true", headers=HEADERS).json()["total"] == 0


def test_later_group_member_cannot_reappear(client, docs, session):
    remove(client, docs[0])
    # 把原来独立的文件归进已下架通知，模拟采集后的自动归并。
    session.add(
        DocumentMerge(
            duplicate_id=docs[4].id, canonical_id=docs[0].id, active=True, match_key="later"
        )
    )
    session.commit()
    assert is_taken_down(session, docs[4].id)
    assert client.get(f"/api/v1/documents/{docs[4].id}").json()["detail"] == "该文件已下架"
    restore(client, docs[4])
    assert client.get(f"/api/v1/documents/{docs[4].id}").status_code == 200


@pytest.mark.parametrize("path,index", [("deadlines", 0), ("lists", 2)])
def test_warmed_report_hides_later_group_member(client, docs, session, path, index):
    # 先下架独立的政策文件；缓存仍含公开的申报/名单组。
    remove(client, docs[4])
    endpoint = f"/api/v1/reports/{path}"
    assert docs[index].title in client.get(endpoint).text
    session.add(
        DocumentMerge(
            duplicate_id=docs[4].id, canonical_id=docs[index].id, active=True, match_key="later"
        )
    )
    session.commit()
    assert docs[index].title not in client.get(endpoint).text
    restore(client, docs[index])
    assert docs[index].title in client.get(endpoint).text


def test_personal_data_and_subscription_references(client, docs, session):
    sub = client.post(
        "/api/v1/me/subscriptions",
        json={"kind": "deadline", "templateId": "due", "docId": docs[0].id},
    )
    assert sub.status_code == 201
    user = session.scalar(select(User))
    session.add(
        Feedback(
            user_id=user.id,
            document_id=docs[1].id,
            content="截止需要核对",
            check_result="unchecked",
        )
    )
    session.commit()
    assert client.get("/api/v1/me/data").json()["feedback"][0]["docTitle"] == docs[0].title
    remove(client, docs[1])
    assert client.get("/api/v1/me/subscriptions").json() == {"items": []}
    assert client.get("/api/v1/me").json()["subscriptions"]["deadline"] == 0
    data = client.get("/api/v1/me/data").json()
    for key in ("subscriptions", "feedback"):
        assert data[key][0]["docId"] is None and data[key][0]["docTitle"] is None
    restore(client, docs[0])
    assert client.get("/api/v1/me/subscriptions").json()["items"] == [sub.json()]
    assert client.get("/api/v1/me").json()["subscriptions"]["deadline"] == 1
    assert docs[0].title in client.get("/api/v1/me/data").text


@pytest.mark.parametrize("format_", ["xlsx", "docx"])
def test_hidden_change_notice_reference(client, docs, session, format_):
    due = session.scalar(select(DocumentDeadline))
    due.status, due.change_reason = "changed", "延期"
    due.changed_to_document_id = docs[4].id
    due.changed_by, due.changed_at = "测试管理员", NOW
    session.commit()
    payload = {"ids": [str(docs[0].id)], "format": format_}

    def export_text():
        response = client.post("/api/v1/exports", json=payload)
        assert response.status_code == 200
        if format_ == "xlsx":
            book = load_workbook(BytesIO(response.content))
            try:
                return " ".join(str(cell.value) for sheet in book for row in sheet for cell in row)
            finally:
                book.close()
        book = WordDocument(BytesIO(response.content))
        return " ".join(
            cell.text for table in book.tables for row in table.rows for cell in row.cells
        )

    assert docs[4].title in export_text()
    remove(client, docs[4])
    detail = client.get(f"/api/v1/documents/{docs[0].id}").json()["deadline_changes"][0]
    assert detail["newDocId"] is None and detail["newTitle"] is None
    assert docs[4].title not in export_text()
    restore(client, docs[4])
    assert docs[4].title in export_text()


def test_follow_markers_and_both_notifications(client, docs, session, settings):
    follow = client.post(
        "/api/v1/me/follows", json={"name": "科技", "conditions": {"time": "不限"}}
    ).json()
    path = f"/api/v1/me/follows/{follow['id']}/new"
    initial = client.get(path).json()["total"]
    user = session.scalar(select(User))
    subscriptions = [
        Subscription(user_id=user.id, kind="deadline", template_id="due", document_id=docs[1].id),
        Subscription(user_id=user.id, kind="newDocs", template_id="new", follow_id=follow["id"]),
    ]
    session.add_all(subscriptions)
    session.commit()
    for doc in (docs[0], docs[2], docs[4]):
        remove(client, doc)
    assert client.get(path).json() == {"total": 0, "items": []}
    assert client.get("/api/v1/me/follows").json()["items"][0]["newCount"] == 0
    sent = []
    run_notify(
        session, settings=settings, sender=lambda p: sent.append(p) or {"errcode": 0}, stamp=NOW
    )
    assert sent == [] and all(s.status == "unused" for s in subscriptions)
    assert (
        client.post(
            "/api/v1/me/subscriptions",
            json={"kind": "deadline", "templateId": "due", "docId": docs[0].id},
        ).status_code
        == 422
    )
    for doc in (docs[0], docs[2], docs[4]):
        restore(client, doc)
    assert client.get(path).json()["total"] == initial
    result = run_notify(
        session, settings=settings, sender=lambda p: sent.append(p) or {"errcode": 0}, stamp=NOW
    )
    assert result["sent"] == 2 and len(sent) == 2


@pytest.mark.parametrize(
    "kind,index,rows",
    [
        ("compare", 2, ["辽宁创新科技有限公司"]),
        ("track", 0, [{"name": "辽宁创新科技有限公司", "province": "辽宁"}]),
    ],
)
def test_tool_generation_and_existing_download(
    client, docs, session, settings, monkeypatch, tmp_path, kind, index, rows
):
    monkeypatch.setattr(advisor_tasks, "upload", lambda *a, **kw: "cloud://test-file")
    monkeypatch.setattr(advisor_tasks, "download_url", lambda *a: "https://example.test/file")
    task_id = client.post(
        "/api/v1/advisor/tasks",
        headers={"Idempotency-Key": "before"},
        json={"kind": kind, "rows": rows},
    ).json()["id"]
    task = session.get(AdvisorTask, task_id)
    task.status = "running"
    assert advisor_tasks.process_task(session, task, settings=settings)
    assert docs[index].id in task.result_document_ids
    path = f"/api/v1/advisor/tasks/{task_id}/download"
    assert client.get(path).status_code == 200
    remove(client, docs[index + 1])
    assert client.get(path).status_code == 410
    user = session.scalar(select(User))
    hidden = AdvisorTask(
        user_id=user.id,
        kind=kind,
        status="running",
        input_rows=rows,
        row_count=1,
        idempotency_key="after",
        created_at=NOW,
        expires_at=NOW + timedelta(days=30),
    )
    session.add(hidden)
    session.commit()
    output = tmp_path / "filtered.xlsx"
    advisor_tasks._run_engine(session, hidden, output)
    assert hidden.result_document_ids == []
    cells = [str(cell.value) for sheet in load_workbook(output) for row in sheet for cell in row]
    assert docs[index].title not in cells and docs[index].url not in cells
    hidden.result_file_id, hidden.status = "cloud://clean-file", "done"
    session.commit()
    assert client.get(f"/api/v1/advisor/tasks/{hidden.id}/download").status_code == 200
    restore(client, docs[index])
    assert client.get(path).status_code == 200


@pytest.mark.parametrize("kind,index", [("compare", 2), ("track", 0)])
@pytest.mark.parametrize("state", ["hidden", "clean", "corrupt"])
def test_legacy_tool_download_checks_actual_excel(
    client, docs, session, monkeypatch, tmp_path, kind, index, state
):
    user = session.scalar(select(User))
    rows = ["辽宁创新科技有限公司"] if kind == "compare" else [{"name": "客户", "province": "辽宁"}]
    task = AdvisorTask(
        user_id=user.id,
        kind=kind,
        status="done",
        input_rows=rows,
        row_count=1,
        result_file_id="cloud://legacy",
        idempotency_key="legacy",
        created_at=NOW,
        expires_at=NOW + timedelta(days=30),
    )
    session.add(task)
    session.commit()
    output = tmp_path / "legacy.xlsx"
    if state == "clean":
        remove(client, docs[index])
        advisor_tasks._run_engine(session, task, output)
    else:
        advisor_tasks._run_engine(session, task, output)
        remove(client, docs[index])
    # 模拟迁移前生成的结果：没有依赖清单，甚至输入已经清理。
    task.result_document_ids = None
    task.input_rows = None
    session.commit()
    content = b"broken excel" if state == "corrupt" else output.read_bytes()
    real_client = httpx.Client
    transport = httpx.MockTransport(lambda request: httpx.Response(200, content=content))
    monkeypatch.setattr(advisor.httpx, "Client", lambda **kw: real_client(transport=transport))
    monkeypatch.setattr(advisor_tasks, "download_url", lambda *a: "https://example.test/file")
    endpoint = f"/api/v1/advisor/tasks/{task.id}/download"
    response = client.get(endpoint)
    assert response.status_code == {"hidden": 410, "clean": 200, "corrupt": 503}[state]
    session.refresh(task)
    if state == "corrupt":
        assert task.result_document_ids is None
    else:
        assert task.result_document_ids == ([docs[index].id] if state == "hidden" else [])
        restore(client, docs[index])
        assert client.get(endpoint).status_code == 200


def test_admin_missing_document_and_numeric_title_search(client, docs):
    for action, payload in (("remove", REMOVE), ("restore", NOTE)):
        assert (
            client.post(
                f"/api/admin/takedown/99999/{action}", headers=HEADERS, json=payload
            ).status_code
            == 404
        )
    assert client.get("/api/admin/takedown/99999/history", headers=HEADERS).status_code == 404
    results = client.get("/api/admin/takedown?q=2026", headers=HEADERS).json()["items"]
    assert {row["id"] for row in results} == {docs[2].id, docs[3].id}


def test_admin_page_embeds_and_portal_count(client):
    for suffix in ("", "?embed=1"):
        response = client.get(f"/admin/takedown{suffix}")
        assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
        assert "/admin/embed.js" in response.text and "youjuAdminToken()" in response.text
        assert "confirm(" in response.text and 'maxlength="300"' in response.text
    assert 'data-menu="takedown"' in client.get("/admin").text


def test_migration_0012_downgrade_upgrade_preserves_documents(db_engine, session):
    # 仅使用当前 db_engine（YOUJU_TEST_DATABASE_URL），不触及开发库。
    url = db_engine.url.render_as_string(hide_password=False)
    assert db_engine.url.database == "youju_test_codex" or db_engine.url.database.startswith(
        "youju_test"
    )
    src = Source(
        level="省级",
        province="辽宁",
        dept="科技",
        dept_full_name="辽宁省科技厅",
        column_name="通知",
        url="https://example.gov.cn/migration",
    )
    session.add(src)
    session.flush()
    doc = Document(
        source_id=src.id,
        title="迁移前题录",
        url=src.url + "/doc",
        url_key="migration",
        publish_date=NOW.date(),
        first_seen_at=NOW,
        last_seen_at=NOW,
    )
    session.add(doc)
    session.commit()
    cfg = alembic_config(url)
    try:
        command.downgrade(cfg, "0011")
        assert "document_takedown_events" not in inspect(db_engine).get_table_names()
        assert "taken_down_at" not in {
            c["name"] for c in inspect(db_engine).get_columns("documents")
        }
    finally:
        command.upgrade(cfg, "head")
    session.expire_all()
    assert session.get(Document, doc.id).title == "迁移前题录"
    assert session.get(Document, doc.id).taken_down_at is None
