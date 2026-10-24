"""名单比对后台引擎的文件、匹配和导出契约。"""

from datetime import UTC, date, datetime
from time import perf_counter

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import insert

from youju.compare import (
    CompareInputError,
    CompareRow,
    compare_lists,
    export_compare_excel,
    read_customer_file,
)
from youju.lists.rules import normalize_name
from youju.models import Document, DocumentDetail, DocumentMerge, ListEntry, Source

CODE_A = "91310117588657256C"


def _add_list(session, source, number, title, name, code, stamp, url_suffix=""):
    now = datetime.now(UTC)
    doc = Document(
        source_id=source.id,
        url=f"https://example.gov.cn/{number}{url_suffix}",
        url_key=f"example.gov.cn/{number}{url_suffix}",
        title=title,
        publish_date=date(2026, 9, number),
        first_seen_at=now,
        last_seen_at=stamp,
        rule_kind="名单",
    )
    session.add(doc)
    session.flush()
    session.add(
        ListEntry(
            document_id=doc.id,
            source_url=doc.url,
            name=name,
            normalized_name=normalize_name(name),
            credit_code=code,
            sequence="1",
            category="高新技术企业",
            region=source.province,
        )
    )
    session.add(
        DocumentDetail(
            document_id=doc.id,
            body_text="名单正文",
            issuer="示例科技厅",
            doc_number="示科函〔2026〕1号",
            fetched_at=now,
            status="ok",
        )
    )
    return doc


def test_file_reading_headers_and_encodings(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.append(["序号", "企业名称"])
    sheet.append([1, "北京星河科技有限公司"])
    xlsx = tmp_path / "customers.xlsx"
    book.save(xlsx)
    assert read_customer_file(xlsx) == ["北京星河科技有限公司"]

    utf8 = tmp_path / "customers.csv"
    utf8.write_text("信用代码\n" + CODE_A + "\n", encoding="utf-8")
    assert read_customer_file(utf8) == [CODE_A]
    gbk = tmp_path / "customers-gbk.csv"
    gbk.write_bytes("企业名称\n北京星河科技有限公司\n".encode("gbk"))
    assert read_customer_file(gbk) == ["北京星河科技有限公司"]

    no_header = tmp_path / "no-header.csv"
    no_header.write_text("北京星河科技有限公司\n上海星河科技有限公司\n", encoding="utf-8")
    assert read_customer_file(no_header) == ["北京星河科技有限公司", "上海星河科技有限公司"]


def test_file_limits(tmp_path):
    path = tmp_path / "too-many.csv"
    path.write_text("企业名称\n" + "\n".join(f"公司{i}" for i in range(501)), encoding="utf-8")
    try:
        read_customer_file(path)
    except CompareInputError as exc:
        assert "500" in str(exc)
    else:
        raise AssertionError("应拒绝超过 500 行")


def test_compare_deduplicates_groups_and_excludes_private(session):
    public = Source(
        level="省级",
        province="北京",
        dept="科技",
        dept_full_name="示例科技厅",
        column_name="通知公告",
        url="https://example.gov.cn/source",
        rights_level=1,
    )
    private = Source(
        level="省级",
        province="上海",
        dept="科技",
        dept_full_name="不公开来源",
        column_name="通知公告",
        url="https://private.example.gov.cn/source",
        rights_level=2,
    )
    session.add_all([public, private])
    session.flush()
    stamp = datetime(2026, 9, 30, tzinfo=UTC)
    canonical = _add_list(
        session,
        public,
        1,
        "关于公布2026年高新技术企业名单的通知",
        "北京星河（上海）科技有限公司",
        CODE_A,
        stamp,
    )
    duplicate = _add_list(
        session,
        public,
        2,
        "关于公布2026年高新技术企业名单的通知",
        "北京星河（上海）科技有限公司",
        CODE_A,
        stamp,
        "-copy",
    )
    session.add(
        DocumentMerge(
            duplicate_id=duplicate.id, canonical_id=canonical.id, active=True, match_key="test"
        )
    )
    _add_list(
        session,
        private,
        3,
        "关于公布2026年高新技术企业名单的通知",
        "北京星河（上海）科技有限公司",
        CODE_A,
        stamp,
        "-private",
    )
    session.commit()

    report = compare_lists(
        session,
        ["ＡＢＣ", "北京星河（上海）科技有限公司", CODE_A, "", "北京星河（上海）科技有限公司"],
    )
    assert [row.status for row in report.rows] == ["未命中", "命中", "命中"]
    assert len(report.rows[1].matches) == 1
    assert report.rows[1].matches[0].status == "公布"
    assert report.rows[1].matches[0].doc_number == "示科函〔2026〕1号"
    assert {item.reason for item in report.removed} == {"空行", "重复行"}
    assert report.data_version == stamp


def test_multiple_same_name_and_export_formula_escape(session, tmp_path):
    source = Source(
        level="省级",
        province="北京",
        dept="科技",
        dept_full_name="示例科技厅",
        column_name="通知公告",
        url="https://example.gov.cn/source2",
        rights_level=1,
    )
    session.add(source)
    session.flush()
    # 第二个代码按规则算出，确保两个同名主体都能入库。
    from youju.lists.rules import _CREDIT_ALPHABET, _CREDIT_WEIGHTS

    body = "91310117588657257"
    checksum = (
        31
        - sum(_CREDIT_ALPHABET.index(c) * w for c, w in zip(body, _CREDIT_WEIGHTS, strict=True))
        % 31
    ) % 31
    code_b = body + _CREDIT_ALPHABET[checksum]
    _add_list(
        session,
        source,
        4,
        "2026年高新技术企业名单公示",
        "同名科技有限公司",
        CODE_A,
        datetime.now(UTC),
        "-a",
    )
    _add_list(
        session,
        source,
        5,
        "2025年高新技术企业名单公示",
        "同名科技有限公司",
        code_b,
        datetime.now(UTC),
        "-b",
    )
    session.commit()
    report = compare_lists(session, ["同名科技有限公司"])
    assert report.rows[0].status == "多个同名"
    assert len(report.rows[0].matches) == 2

    output = tmp_path / "out.xlsx"
    report.rows.append(CompareRow("=客户", "=客户", "企业名称", "未命中", matches=[]))
    assert export_compare_excel(report, output) == report.detail_row_count
    exported = load_workbook(output, data_only=False)
    assert exported["明细"]["A5"].value == "'=客户"
    assert exported["明细"]["A1"].value.startswith("第三方整理")
    # 用词 D-090：表头和说明写"发布阶段"，不再写"名单状态"
    rows = [[cell.value for cell in row] for row in exported["明细"].iter_rows(max_row=4)]
    header = next(row for row in rows if "客户输入" in row)
    assert "发布阶段" in header and "名单状态" not in header
    assert "发布阶段照录标题用词" in exported["明细"]["A1"].value


def test_compare_and_export_allow_missing_year_and_publish_date(session, tmp_path):
    source = Source(
        level="省级",
        province="北京",
        dept="科技",
        dept_full_name="示例科技厅",
        column_name="通知公告",
        url="https://example.gov.cn/source-none-date",
        rights_level=1,
    )
    session.add(source)
    session.flush()
    stamp = datetime.now(UTC)
    without_year = _add_list(
        session,
        source,
        6,
        "关于公布高新技术企业名单的通知",
        "无年份科技有限公司",
        CODE_A,
        stamp,
        "-without-year",
    )
    without_year.publish_date = None
    _add_list(
        session,
        source,
        7,
        "关于公布2026年高新技术企业名单的通知",
        "无年份科技有限公司",
        CODE_A,
        stamp,
        "-with-year",
    )
    session.commit()

    report = compare_lists(session, ["无年份科技有限公司"])
    assert report.rows[0].status == "命中"
    assert [match.year for match in report.rows[0].matches] == [2026, None]
    output = tmp_path / "missing-date.xlsx"
    assert export_compare_excel(report, output) == 2
    book = load_workbook(output, data_only=False)
    assert book["明细"]["E4"].value in (None, "")
    assert book["明细"]["L4"].value in (None, "")


def test_compare_skips_person_names(session):
    report = compare_lists(session, ["张三"])
    assert report.rows[0].status == "像个人姓名，未查询"
    assert report.rows[0].matches == []


def test_compare_50_customers_with_20000_entries_under_five_seconds(session):
    source = Source(
        level="省级",
        province="北京",
        dept="科技",
        dept_full_name="示例科技厅",
        column_name="通知公告",
        url="https://example.gov.cn/source-performance",
        rights_level=1,
    )
    session.add(source)
    session.flush()
    now = datetime.now(UTC)
    doc = Document(
        source_id=source.id,
        url="https://example.gov.cn/performance",
        url_key="example.gov.cn/performance",
        title="关于公布2026年高新技术企业名单的通知",
        publish_date=date(2026, 9, 30),
        first_seen_at=now,
        last_seen_at=now,
        rule_kind="名单",
    )
    session.add(doc)
    session.flush()
    entries = [
        {
            "document_id": doc.id,
            "source_url": doc.url,
            "name": f"性能测试公司{i}有限公司",
            "normalized_name": f"性能测试公司{i}有限公司",
            "credit_code": None,
            "sequence": str(i),
            "category": "高新技术企业",
            "batch": None,
            "region": "北京",
        }
        for i in range(20_000)
    ]
    session.execute(insert(ListEntry), entries)
    session.commit()

    values = [f"性能测试公司{i}有限公司" for i in range(50)]
    started = perf_counter()
    report = compare_lists(session, values)
    elapsed = perf_counter() - started
    assert elapsed < 5, f"名单比对耗时 {elapsed:.2f}s"
    assert len(report.rows) == 50
    assert all(row.status == "命中" and len(row.matches) == 1 for row in report.rows)


def test_csv_preserves_quoted_newlines():
    content = '企业名称\r\n"第一行\r\n第二行公司"\r\n'.encode()
    assert read_customer_file("clients.csv", content=content) == ["第一行\r\n第二行公司"]


@pytest.mark.parametrize("text", ['企业名称\n"未闭合公司', '企业名称\n"甲公司"多余内容'])
def test_csv_malformed_quotes_return_chinese_reason(text):
    with pytest.raises(CompareInputError, match="CSV 格式有误"):
        read_customer_file("clients.csv", content=text.encode())


def test_xlsx_rejects_overlong_cell_reference(tmp_path):
    """超长列坐标直接报错，不做大整数计算。"""
    import zipfile

    import pytest
    from openpyxl import Workbook

    from youju.compare import CompareInputError, read_customer_file

    good = tmp_path / "good.xlsx"
    book = Workbook()
    book.active["A1"] = "企业名称"
    book.active["A2"] = "演示科技有限公司"
    book.save(good)
    bad = tmp_path / "bad.xlsx"
    with zipfile.ZipFile(good) as src, zipfile.ZipFile(bad, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                data = data.replace(b'r="A2"', b'r="' + b"A" * 5000 + b'2"')
            dst.writestr(item, data)
    with pytest.raises(CompareInputError, match="坐标无效"):
        read_customer_file(bad)
