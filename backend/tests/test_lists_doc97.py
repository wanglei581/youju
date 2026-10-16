"""Word 97 二进制附件（.doc / .wps）里的名单表格（10/16）。"""

import struct
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from youju.lists import doc97
from youju.lists.parser import parse_attachment, parse_table_rows
from youju.lists.service import WORD97_SINCE, _should_attempt
from youju.models import ListParse

FIXTURES = Path(__file__).parent / "fixtures" / "lists"


def test_real_word97_tables() -> None:
    """两份真实政府附件：山西绿色工厂（.doc，表头"工厂名称"）、福建省级绿色制造（.wps，带地区）。"""
    shanxi = parse_attachment((FIXTURES / "shanxi_green_factories.doc").read_bytes(), "doc")
    # 附件 1 的 77 家和附件 3 的 2 家；附件 2 是园区（表头"园区名称"），不收
    assert len(shanxi) == 79
    assert [shanxi[0].name, shanxi[0].sequence] == ["山西华豹新材料有限公司", "1"]
    assert "沁水经济技术开发区" not in {entry.name for entry in shanxi}
    fujian = parse_attachment((FIXTURES / "fujian_green_manufacturing.wps").read_bytes(), "wps")
    assert len(fujian) == 179
    assert [fujian[0].name, fujian[0].region] == ["东南（福建）汽车工业股份有限公司", "福州"]
    assert [fujian[-1].name, fujian[-1].sequence] == ["福建品品香茶业有限公司", "179"]


R = doc97.ROW_END


def test_rows_from_cell_marks() -> None:
    """格子以 \\x07 结束，行尾段落（读出时换成 ROW_END）断行；空格子、表格前后的段落、
    表格前只有标记的空行都要处理对。"""
    text = (
        "附件1\r2026年度名单\r"
        f"\x07{R}"  # 表格前一个只有标记的空行
        f"序号\x07企业名称\x07地区\x07{R}"
        f"1\x07甲科技有限公司\x07\x07{R}"  # 地区空着
        f"2\x07乙有限公司\x07福州\x07{R}"
        "附件2\r移出名单\r"
        f"序号\x07单位名称\x07{R}"
        f"1\x07丙研究院\x07{R}"
        "后记\r"
    )
    assert doc97.doc_tables(text) == [
        [["序号", "企业名称", "地区"], ["1", "甲科技有限公司", ""], ["2", "乙有限公司", "福州"]],
        [["序号", "单位名称"], ["1", "丙研究院"]],
    ]


def test_extra_empty_cell_does_not_shift_columns() -> None:
    """一行多出一个空格子（Word 允许各行格数不同）：按行尾段落断行，下一行的推荐单位
    不会落进企业名称列（Codex 审查：按表头格数猜行尾会错位）。"""
    text = (
        f"推荐单位\x07企业名称\x07{R}"
        f"甲市工信局\x07乙科技有限公司\x07\x07{R}"
        f"丙市工信局\x07丁科技有限公司\x07{R}"
    )
    rows = doc97.doc_tables(text)[0]
    assert rows[2] == ["丙市工信局", "丁科技有限公司"]
    assert [e.name for e in parse_table_rows(rows)] == ["乙科技有限公司", "丁科技有限公司"]


def test_field_codes_keep_only_results() -> None:
    text = '甲\x13 HYPERLINK "https://example.gov.cn" \x14乙有限公司\x15丙'
    assert doc97._drop_field_codes(text) == "甲乙有限公司丙"
    nested = "\x13 IF \x13 PAGE \x141\x15 \x14结果\x15"
    assert doc97._drop_field_codes(nested) == "结果"
    with pytest.raises(doc97.Doc97Error, match="域嵌套过深"):
        doc97._drop_field_codes("\x13" * (doc97.MAX_FIELD_DEPTH + 1) + "甲")


def _fib(pieces_fc: int, piece_chars: int, ccp_text: int) -> tuple[bytes, bytes]:
    """最小的 WordDocument 流和表格流：FIB 只填用到的字段，片段表一段，段落属性表一页（空）。"""
    wd = bytearray(2048)
    struct.pack_into("<HH", wd, 0, 0xA5EC, 0x00C1)
    struct.pack_into("<H", wd, 32, 0)  # csw
    struct.pack_into("<H", wd, 34, 4)  # cslw：ccpText 是第 3 个
    struct.pack_into("<i", wd, 36 + 3 * 4, ccp_text)
    rgfclcb = 36 + 4 * 4 + 2
    table = bytearray()
    # 片段表：Pcdt(0x02) + lcb + 两个字符位置 + 一个片段描述
    plc = struct.pack("<II", 0, piece_chars) + struct.pack("<HIH", 0, pieces_fc, 0)
    clx = b"\x02" + struct.pack("<I", len(plc)) + plc
    struct.pack_into("<II", wd, rgfclcb + 33 * 8, len(table), len(clx))
    table += clx
    # 段落属性表：一页，页号 3（偏移 1536），页里一段、没有属性
    struct.pack_into("<II", wd, rgfclcb + 13 * 8, len(table), 12)
    table += struct.pack("<III", 0, 2048, 3)
    page = bytearray(512)
    struct.pack_into("<II", page, 0, 0, 2048)
    page[511] = 1
    wd[1536:2048] = page
    return bytes(wd), bytes(table)


def test_pieces_out_of_range_are_unsupported() -> None:
    """片段指到流外、正文长度超预算：报"暂不支持"，不当成空名单（Codex 审查）。"""
    wd, table = _fib(pieces_fc=600, piece_chars=10, ccp_text=10)
    assert doc97._main_text(wd, table) == bytes(wd[600:620]).decode("utf-16-le")
    wd, table = _fib(pieces_fc=2040, piece_chars=10, ccp_text=10)
    with pytest.raises(doc97.Doc97Error, match="文字片段越界"):
        doc97._main_text(wd, table)
    wd, table = _fib(pieces_fc=600, piece_chars=10, ccp_text=doc97.MAX_CHARS + 1)
    with pytest.raises(doc97.Doc97Error, match="正文过长"):
        doc97._main_text(wd, table)


def test_broken_streams_are_unsupported(monkeypatch) -> None:
    """复合文件扇区链损坏时 olefile 读流报 OSError：记"暂不支持"，不记"解析失败"（Codex 审查）。"""

    class Broken:
        def exists(self, name):
            return True

        def openstream(self, name):
            raise OSError("incorrect OLE sector index")

        def close(self):
            pass

    monkeypatch.setattr(doc97.olefile, "OleFileIO", lambda _data: Broken())
    with pytest.raises(doc97.Doc97Error, match="暂不支持旧版或专有格式"):
        doc97.doc_text(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")


def test_misnamed_and_unreadable_files() -> None:
    """扩展名是 .doc/.wps，内容按实际格式处理：DOCX 照读，网页读表格，RTF 和其他不支持。"""
    docx = DocxDocument()
    table = docx.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "序号", "企业名称"
    table.cell(1, 0).text, table.cell(1, 1).text = "1", "北京星河科技有限公司"
    buffer = BytesIO()
    docx.save(buffer)
    assert [e.name for e in parse_attachment(buffer.getvalue(), "doc")] == ["北京星河科技有限公司"]
    page = "<html><body><table><tr><td>企业名称</td></tr><tr><td>上海云帆科技有限公司</td></tr>"
    assert [e.name for e in parse_attachment(page.encode(), "wps")] == ["上海云帆科技有限公司"]
    with pytest.raises(ValueError, match="暂不支持 RTF"):
        parse_attachment(b"{\\rtf1\\ansi table}", "doc")
    with pytest.raises(ValueError, match="暂不支持旧版或专有格式"):
        parse_attachment(b"old doc", "doc")
    # OLE 文件但不是 Word 文档（旧版 Excel）
    with pytest.raises(ValueError, match="不是 Word 97 文档"):
        parse_attachment((FIXTURES / "heilongjiang_30.xls").read_bytes(), "doc")


def test_header_written_as_single_word() -> None:
    """整格写作"企业"的表头也认（福建水效领跑者名单）；"企业类型"不算；单写"单位"不收。"""
    rows = [["序号", "企业", "地区"], ["1", "福建联合石油化工有限公司", "泉州"]]
    assert [e.name for e in parse_table_rows(rows)] == ["福建联合石油化工有限公司"]
    assert parse_table_rows([["序号", "企业类型"], ["1", "福建联合石油化工有限公司"]]) == []
    awards = [["排名", "姓名", "单位"], ["1", "张三", "福建省科学技术馆"]]
    assert parse_table_rows(awards) == []


def _record(url: str, status: str, reason: str | None, attempted: datetime) -> ListParse:
    return ListParse(
        document_id=1,
        source_url=url,
        source_kind="attachment",
        status=status,
        failure_reason=reason,
        attempted_at=attempted,
    )


def test_old_word_records_are_retried_once() -> None:
    """以前记作"暂不支持旧版或专有格式"的 .doc/.wps 补解析一次；补过的、别的格式照旧不重试。"""
    before = WORD97_SINCE - timedelta(days=3)
    after = WORD97_SINCE + timedelta(minutes=5)
    reason = "暂不支持旧版或专有格式"
    assert _should_attempt(_record("https://a.gov.cn/f/名单.doc", "unsupported", reason, before))
    assert _should_attempt(_record("https://a.gov.cn/f/名单.WPS", "unsupported", reason, before))
    assert not _should_attempt(_record("https://a.gov.cn/f/名单.doc", "unsupported", reason, after))
    assert not _should_attempt(_record("https://a.gov.cn/f/名单.et", "unsupported", reason, before))
    assert not _should_attempt(
        _record("https://a.gov.cn/f/名单.doc", "unsupported", "文件超过20MB", before)
    )
    assert not _should_attempt(_record("https://a.gov.cn/f/名单.doc", "empty", None, before))
    # 网址不带扩展名、附件题录里认出是 .doc（链接文字"名单.doc"）的也补
    no_ext = _record("https://a.gov.cn/download?id=7", "unsupported", reason, before)
    assert _should_attempt(no_ext, "doc")
    assert not _should_attempt(no_ext)
    # 临时失败照旧 6 小时后再试
    assert _should_attempt(
        _record(
            "https://a.gov.cn/f/名单.doc",
            "failed",
            "抓取失败",
            datetime.now(UTC) - timedelta(hours=7),
        )
    )


def test_pieces_must_cover_main_text() -> None:
    """片段要从正文开头连续盖满正文；少一段要报"暂不支持"，不能悄悄少字（Codex 复核）。"""
    wd, table = _fib(pieces_fc=600, piece_chars=10, ccp_text=20)
    with pytest.raises(doc97.Doc97Error, match="文字片段不完整"):
        doc97._main_text(wd, table)


def _ttp_document(text: str, ttp_at: int) -> tuple[bytes, bytes]:
    """正文 text（UTF-16）放在 600 字节处；字节位置 ttp_at 起的 2 字节是一个行尾段落。"""
    raw = text.encode("utf-16-le")
    wd, table = _fib(pieces_fc=600, piece_chars=len(raw) // 2, ccp_text=len(raw) // 2)
    wd = bytearray(wd)
    wd[600 : 600 + len(raw)] = raw
    page = bytearray(512)
    struct.pack_into("<III", page, 0, 0, ttp_at, ttp_at + 2)
    struct.pack_into("<I", page, 12, 2048)
    page[511] = 3
    rgbx = 4 * 4
    page[rgbx + 13] = 400 // 2  # 第 2 段的段落属性在页内 400 字节处
    page[400:406] = bytes([3, 0, 0]) + struct.pack("<HB", 0x2417, 1)  # cb=3：样式 + sprmPFTtp=1
    wd[1536:2048] = page
    return bytes(wd), table


def test_row_end_after_supplementary_character() -> None:
    """生僻字（如𠀀）在 UTF-16 里占 4 字节：按原始字节找标记，行尾不会偏（Codex 复核）。"""
    wd, table = _ttp_document("𠀀\x07甲\x07", ttp_at=604)
    assert doc97._main_text(wd, table) == f"𠀀{doc97.ROW_END}甲\x07"


def test_row_end_flag_in_data_stream_and_after_special_tabs() -> None:
    """行尾标记可能放在 Data 流里（sprmPHugePapx），或排在特殊写法的制表位后面（Codex 复核）。"""
    data = b"\x00" * 10 + struct.pack("<h", 3) + struct.pack("<HB", 0x2417, 1)
    assert doc97._grpprl_has_ttp(struct.pack("<HI", 0x6646, 10), data)
    with pytest.raises(doc97.Doc97Error, match="段落属性指针越界"):
        doc97._grpprl_has_ttp(struct.pack("<HI", 0x6646, 999), data)
    tabs = struct.pack("<H", 0xC615) + bytes([255, 1]) + b"\x01\x00\x02\x00" + bytes([1]) + b"abc"
    assert doc97._grpprl_has_ttp(tabs + struct.pack("<HB", 0x2417, 1), b"")
    assert not doc97._grpprl_has_ttp(tabs + struct.pack("<HB", 0x2417, 0), b"")
