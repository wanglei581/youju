import json
import re
from datetime import UTC, date, datetime, time, timedelta
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from youju.api import v1
from youju.api.app import app
from youju.db import get_session
from youju.dedup import reconcile_documents
from youju.models import (
    Document,
    DocumentAttachment,
    DocumentDeadline,
    DocumentDetail,
    DocumentMerge,
    JobRun,
    ListEntry,
    Source,
)

# 北京时间 2026-09-27（周日）上午 10 点；本周是 9/21（周一）到 9/28（下周一，不含）
NOW = datetime(2026, 9, 27, 10, 0, tzinfo=ZoneInfo("Asia/Shanghai"))


@pytest.fixture
def client(session_factory, monkeypatch):
    def override():
        with session_factory() as s:
            yield s

    monkeypatch.setattr(v1, "now", lambda: NOW)
    app.dependency_overrides[get_session] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


def _doc(src, n, title, day):
    now = datetime.now(UTC)
    return Document(
        source_id=src.id,
        url=f"https://{src.url.split('/')[2]}/{n}.html",
        url_key=f"{src.url.split('/')[2]}/{n}.html",
        title=title,
        publish_date=day,
        first_seen_at=now,
        last_seen_at=now,
    )


@pytest.fixture
def seeded(session):
    ln = Source(
        level="省级",
        province="辽宁",
        dept="工信",
        dept_full_name="辽宁省工业和信息化厅",
        column_name="通知",
        url="https://gxt.ln.gov.cn/tz/",
        enabled=True,
        parser_config={},
    )
    sd = Source(
        level="省级",
        province="山东",
        dept="科技",
        dept_full_name="山东省科学技术厅",
        column_name="通知",
        url="https://kjt.shandong.gov.cn/tz/",
        enabled=True,
        parser_config={},
    )
    gj = Source(
        level="国家级",
        province="国家",
        dept="统计",
        dept_full_name="国家统计局",
        column_name="发布",
        url="https://www.stats.gov.cn/fb/",
        enabled=True,
        parser_config={},
    )
    session.add_all([ln, sd, gj])
    session.flush()
    apply_doc = _doc(
        ln,
        1,
        "辽宁省工业和信息化厅关于组织开展2026年三季度贷款贴息项目申报工作的通知",
        date(2026, 9, 24),
    )
    list_doc = _doc(sd, 2, "山东省2026年第4批拟入库科技型中小企业名单公示", date(2026, 9, 24))
    policy_doc = _doc(
        sd, 3, "山东省科学技术厅关于印发《科技计划项目管理办法》的通知", date(2026, 8, 1)
    )
    stat_doc = _doc(gj, 4, "2026年9月中旬流通领域重要生产资料市场价格变动情况", date(2026, 9, 26))
    future_doc = _doc(gj, 5, "2026年9月下旬流通领域重要生产资料市场价格变动情况", date(2026, 9, 28))
    session.add_all([apply_doc, list_doc, policy_doc, stat_doc, future_doc])
    session.flush()
    session.add(
        DocumentDetail(
            document_id=apply_doc.id,
            body_text="……",
            issuer="辽宁省工业和信息化厅",
            doc_number="辽工信投资〔2026〕130号",
            fetched_at=datetime.now(UTC),
            status="ok",
            attachments=[
                DocumentAttachment(
                    position=0,
                    name="附件1 实施细则",
                    url="https://gxt.ln.gov.cn/a1.pdf",
                    extension="pdf",
                )
            ],
        )
    )
    reviewed = datetime(2026, 9, 26, 3, 0, tzinfo=UTC)
    session.add_all(
        [
            DocumentDeadline(
                document_id=apply_doc.id,
                fingerprint="a",
                kind="企业申报截止",
                candidate_date=date(2026, 10, 18),
                evidence="企业申报截止时间为10月18日",
                confidence="高",
                status="approved",
                reviewed_by="创始人",
                reviewed_at=reviewed,
                approved_date=date(2026, 10, 18),
            ),
            DocumentDeadline(
                document_id=apply_doc.id,
                fingerprint="b",
                kind="推荐或审核截止",
                candidate_date=date(2026, 10, 20),
                evidence="市级推荐截止10月20日",
                confidence="高",
                status="approved",
                reviewed_by="创始人",
                reviewed_at=reviewed,
                approved_date=date(2026, 10, 20),
                approved_time=time(17, 0),
            ),
            # 当天 9 点已截止（现在 10 点）：日历里不能出现
            DocumentDeadline(
                document_id=policy_doc.id,
                fingerprint="d",
                kind="材料报送截止",
                candidate_date=date(2026, 9, 27),
                evidence="9月27日9:00前报送",
                confidence="高",
                status="approved",
                reviewed_by="创始人",
                reviewed_at=reviewed,
                approved_date=date(2026, 9, 27),
                approved_time=time(9, 0),
            ),
            # 当天 17 点截止：还没过，日历里要有
            DocumentDeadline(
                document_id=policy_doc.id,
                fingerprint="e",
                kind="企业申报截止",
                candidate_date=date(2026, 9, 27),
                evidence="9月27日17:00前提交",
                confidence="高",
                status="approved",
                reviewed_by="创始人",
                reviewed_at=reviewed,
                approved_date=date(2026, 9, 27),
                approved_time=time(17, 0),
            ),
            # 没核对的候选：任何公开接口都不能出现
            DocumentDeadline(
                document_id=list_doc.id,
                fingerprint="c",
                kind="企业申报截止",
                candidate_date=date(2026, 10, 5),
                evidence="10月5日前",
                confidence="中",
                status="pending",
            ),
        ]
    )
    session.commit()
    return {
        "apply": apply_doc.id,
        "list": list_doc.id,
        "policy": policy_doc.id,
        "stat": stat_doc.id,
    }


def test_home_counts_week_and_skips_other_kinds_in_latest(client, session, seeded):
    # 栏目数、省份数与收录范围同一口径：只算成功检查过的栏目
    for src in session.scalars(select(Source)).all():
        src.last_success_at = datetime(2026, 9, 28, 1, 30, tzinfo=UTC)
    session.commit()
    body = client.get("/api/v1/home").json()
    assert body["stats"]["week_total"] == 3  # 9/24、9/24、9/26；8/1 的政策不算本周
    assert (
        body["stats"]["apply"] == 1 and body["stats"]["lists"] == 1 and body["stats"]["other"] == 1
    )
    assert body["stats"]["provinces"] == 2 and body["stats"]["sources"] == 3
    kinds = [i["kind"] for i in body["items"]]
    assert "其他" not in kinds
    apply_item = next(i for i in body["items"] if i["id"] == str(seeded["apply"]))
    assert apply_item["deadline"] == {"date": "2026-10-18", "time": "", "verified": True}


def test_home_unverified_count_excludes_verified_due_soon(client, session, seeded):
    approved = session.scalar(
        select(DocumentDeadline).where(
            DocumentDeadline.document_id == seeded["apply"],
            DocumentDeadline.kind == "企业申报截止",
        )
    )
    approved.approved_date = date(2026, 9, 30)
    session.add(
        DocumentDeadline(
            document_id=seeded["apply"],
            fingerprint="other-pending",
            kind="材料报送截止",
            candidate_date=date(2026, 10, 1),
            evidence="材料报送截止10月1日",
            confidence="高",
            status="pending",
        )
    )
    session.commit()
    v1._EXCERPT_CACHE.clear()
    v1._REPORT_CACHE.clear()
    body = client.get("/api/v1/home", params={"excerpts": 1}).json()
    assert [item["id"] for item in body["dueSoon"]] == [str(seeded["apply"])]
    assert body["dueSoonUnverifiedTotal"] == 0


def test_excerpt_report_uses_future_machine_date_after_verified_dates_expire(
    client, session, seeded
):
    for row in session.scalars(
        select(DocumentDeadline).where(DocumentDeadline.document_id == seeded["apply"])
    ):
        row.approved_date = date(2026, 9, 25)
    detail = session.scalar(
        select(DocumentDetail).where(DocumentDetail.document_id == seeded["apply"])
    )
    detail.body_text = "材料报送截止：2026年10月10日。"
    session.add(
        DocumentDeadline(
            document_id=seeded["apply"],
            fingerprint="future-excerpt",
            kind="材料报送截止",
            candidate_date=date(2026, 10, 10),
            evidence="材料报送截止：2026年10月10日。",
            confidence="高",
            status="pending",
        )
    )
    session.commit()
    v1._EXCERPT_CACHE.clear()
    v1._REPORT_CACHE.clear()
    body = client.get("/api/v1/reports/deadlines", params={"days": 30, "excerpts": 1}).json()
    item = next(item for item in body["items"] if item["id"] == str(seeded["apply"]))
    assert item["source"] == "excerpt"
    assert item["date"] == "2026-10-10"


def test_snapshot_latest_apply_has_no_time_limit(client, session, seeded):
    source = session.get(Document, seeded["apply"]).source_id
    old = Document(
        source_id=source,
        url="https://gxt.ln.gov.cn/older-apply.html",
        url_key="gxt.ln.gov.cn/older-apply.html",
        title="关于组织申报上一批科技项目的通知",
        publish_date=date(2026, 8, 1),
        first_seen_at=datetime.now(UTC),
        last_seen_at=datetime.now(UTC),
    )
    session.add(old)
    session.commit()
    v1._SNAPSHOT_CACHE.clear()
    body = client.get("/api/v1/snapshot").json()
    assert str(old.id) in {item["id"] for item in body["home"]["items"]}


def _new_apply(session, seeded, slug: str) -> Document:
    """在申报通知同一来源下加一份新的申报通知（快照首页"最新 8 份"里会有它）。"""
    source = session.get(Document, seeded["apply"]).source_id
    doc = Document(
        source_id=source,
        url=f"https://gxt.ln.gov.cn/{slug}.html",
        url_key=f"gxt.ln.gov.cn/{slug}.html",
        title=f"关于组织申报{slug}科技项目的通知",
        publish_date=date(2026, 9, 26),
        first_seen_at=datetime.now(UTC),
        last_seen_at=datetime.now(UTC),
    )
    session.add(doc)
    session.commit()
    return doc


def _snapshot_ids(body) -> set[str]:
    return {item["id"] for item in body["home"]["items"]}


def _age_snapshot(seconds: float) -> None:
    key, stamp, payload = v1._SNAPSHOT_CACHE["first"]
    v1._SNAPSHOT_CACHE["first"] = (key, stamp - seconds, payload)


def test_snapshot_stale_copy_first_then_refreshed_after_response(client, session, seeded):
    """S2：60 秒到期后、5 分钟内先给旧的，同时在后台刷新一份；刷新好以后给新的。"""
    v1._SNAPSHOT_CACHE.clear()
    first = client.get("/api/v1/snapshot").json()
    _age_snapshot(120)
    doc = _new_apply(session, seeded, "stale")
    stale = client.get("/api/v1/snapshot").json()
    assert stale == first and str(doc.id) not in _snapshot_ids(stale)
    # TestClient 等后台任务做完才返回：这时缓存已经换成新算的
    assert str(doc.id) in _snapshot_ids(v1._SNAPSHOT_CACHE["first"][2])
    assert v1._snapshot_refresh_at is None
    assert str(doc.id) in _snapshot_ids(client.get("/api/v1/snapshot").json())


def test_snapshot_rebuilt_in_request_when_no_usable_copy(client, session, seeded, monkeypatch):
    """旧的超过 5 分钟、跨天、算好以后有文件下架时当场重算，不先给旧的。"""
    from youju.takedown import change_state

    v1._SNAPSHOT_CACHE.clear()
    client.get("/api/v1/snapshot")
    _age_snapshot(v1.SNAPSHOT_STALE + 1)
    doc = _new_apply(session, seeded, "tooold")
    assert str(doc.id) in _snapshot_ids(client.get("/api/v1/snapshot").json())
    # 下架：过期后不再先给旧的（旧的里还有这份），当场重算
    change_state(session, doc.id, action="takedown", note="测试", operator="测试", reason="其他")
    _age_snapshot(120)
    assert str(doc.id) not in _snapshot_ids(client.get("/api/v1/snapshot").json())
    # 跨天：日期变了，当场按新的一天算
    monkeypatch.setattr(v1, "now", lambda: NOW + timedelta(days=1))
    assert client.get("/api/v1/snapshot").json()["generatedAt"][:10] == "2026-09-28"


def test_snapshot_not_served_stale_after_direct_takedown(client, session, seeded):
    """不经管理台、直接改库下架（没有下架记录）：过期后也不先给旧的，当场重算。"""
    v1._SNAPSHOT_CACHE.clear()
    doc = _new_apply(session, seeded, "direct")
    assert str(doc.id) in _snapshot_ids(client.get("/api/v1/snapshot").json())
    session.get(Document, doc.id).taken_down_at = datetime.now(UTC)
    session.commit()
    _age_snapshot(120)
    assert str(doc.id) not in _snapshot_ids(client.get("/api/v1/snapshot").json())


def test_snapshot_refresh_failure_keeps_old_copy(client, session, seeded, monkeypatch, caplog):
    """后台刷新出错：这次照样给旧的，记下错误，下一个请求可以再安排刷新。"""
    v1._SNAPSHOT_CACHE.clear()
    first = client.get("/api/v1/snapshot").json()
    _age_snapshot(120)

    def broken(*args, **kwargs):
        raise RuntimeError("数据库暂时连不上")

    monkeypatch.setattr(v1, "home", broken)
    response = client.get("/api/v1/snapshot")
    assert response.status_code == 200 and response.json() == first
    assert "首屏快照后台刷新失败" in caplog.text
    assert v1._snapshot_refresh_at is None
    assert v1._claim_snapshot_refresh()
    v1._snapshot_refresh_at = None


def _snapshot_race(monkeypatch):
    """当场算快照时第一个开始算的停住，等调用方放行；记下算了几份、查了几次下架版本。"""
    import threading

    release = threading.Event()
    state = {"builds": 0, "revisions": 0}
    real_home, real_revision = v1.home, v1.cache_revision

    def slow_home(*args, **kwargs):
        state["builds"] += 1
        assert release.wait(30)
        return real_home(*args, **kwargs)

    def revision(session):
        state["revisions"] += 1
        return real_revision(session)

    monkeypatch.setattr(v1, "home", slow_home)
    monkeypatch.setattr(v1, "cache_revision", revision)
    return state, release


def _wait_until(check, seconds: float = 30) -> bool:
    import threading

    for _ in range(int(seconds * 100)):
        if check():
            return True
        threading.Event().wait(0.01)
    return check()


def test_snapshot_concurrent_requests_build_once(
    client, session_factory, session, seeded, monkeypatch
):
    """缓存过期时同时来的请求只算一份（线上 10/17 18:03 同一秒 3 个请求各算了一遍）；
    等着的请求不占数据库连接。"""
    from concurrent.futures import ThreadPoolExecutor

    v1._SNAPSHOT_CACHE.clear()
    monkeypatch.setattr(v1, "SNAPSHOT_WAIT", 60.0)  # 慢机器上也不会等超时
    state, release = _snapshot_race(monkeypatch)
    pool = session_factory.kw["bind"].pool
    with ThreadPoolExecutor(3) as executor:
        futures = [executor.submit(client.get, "/api/v1/snapshot") for _ in range(3)]
        # 第一个在算（拿着连接）；第一份算完之前另外两个都看不到可用的，在等锁、已经把连接还回去
        assert _wait_until(lambda: state["builds"] == 1 and v1._snapshot_waiters == 3)
        assert _wait_until(lambda: pool.checkedout() == 1)
        release.set()
        bodies = [future.result().json() for future in futures]
    assert state["builds"] == 1
    assert state["revisions"] == 1  # 只有算的那个查了下架版本；等着的拿到锁后直接用
    assert bodies[0] == bodies[1] == bodies[2]
    assert v1._snapshot_waiters == 0


@pytest.mark.parametrize(
    ("waiters", "wait"),
    [
        (1, 60.0),  # 已经有一个在算、满了：马上拒，不等 60 秒
        (8, 0.2),  # 没满：等 0.2 秒还没算完就放弃
    ],
)
def test_snapshot_waiters_capped_and_wait_bounded(
    client, session_factory, session, seeded, monkeypatch, waiters, wait
):
    """当场算的时候，等着的请求有个数上限、等的时间有上限，超了回 503
    （小程序取快照只试一次，取不到就等实时数据）；被拒的请求一次数据库都不查。"""
    from concurrent.futures import ThreadPoolExecutor

    from sqlalchemy import event

    v1._SNAPSHOT_CACHE.clear()
    monkeypatch.setattr(v1, "SNAPSHOT_WAITERS", waiters)
    monkeypatch.setattr(v1, "SNAPSHOT_WAIT", wait)
    state, release = _snapshot_race(monkeypatch)
    engine = session_factory.kw["bind"]
    statements: list[str] = []

    def capture(connection, cursor, statement, parameters, context, executemany):
        statements.append(statement)

    event.listen(engine, "before_cursor_execute", capture)
    try:
        with ThreadPoolExecutor(2) as executor:
            building = executor.submit(client.get, "/api/v1/snapshot")
            assert _wait_until(lambda: state["builds"] == 1)
            before = len(statements)  # 算的那个停在这里，不再查数据库
            refused = executor.submit(client.get, "/api/v1/snapshot")
            assert _wait_until(refused.done, 10)
            assert refused.result().status_code == 503
            assert refused.result().json() == {"detail": "首屏数据正在生成，请稍后再试"}
            assert statements[before:] == []
            release.set()
            assert building.result().status_code == 200
    finally:
        event.remove(engine, "before_cursor_execute", capture)
    assert v1._snapshot_waiters == 0


def test_snapshot_fresh_copy_needs_no_database(client, session, seeded, monkeypatch):
    """60 秒内的直接给、不查数据库（数据库一时连不上也照样给）；过期后先核对下架版本。"""
    v1._SNAPSHOT_CACHE.clear()
    first = client.get("/api/v1/snapshot").json()

    def unreachable(session):
        raise RuntimeError("数据库暂时连不上")

    monkeypatch.setattr(v1, "cache_revision", unreachable)
    assert client.get("/api/v1/snapshot").json() == first


def test_api_cleans_legacy_issuer_and_source_url(client, session, seeded):
    doc = session.get(Document, seeded["apply"])
    doc.url = "http://gxt.ln.gov.cn:443/1.html"
    detail = session.scalar(select(DocumentDetail).where(DocumentDetail.document_id == doc.id))
    detail.issuer = "安 徽 省 财 政 厅 福建省职称改革工作办公室"
    session.commit()
    expected = "安徽省财政厅、福建省职称改革工作办公室"
    for item in (
        next(i for i in client.get("/api/v1/home").json()["items"] if i["id"] == str(doc.id)),
        client.get(f"/api/v1/documents/{doc.id}").json(),
    ):
        assert item["issuer"] == expected
        assert item["url"] == "https://gxt.ln.gov.cn/1.html"


def test_search_filters_and_totals(client, seeded):
    r = client.get("/api/v1/search", params={"q": "科技型中小企业"}).json()
    assert [i["id"] for i in r["items"]] == [str(seeded["list"])]
    assert r["items"][0]["status"] == "公示"
    assert r["items"][0]["deadline"] is None  # 只有 pending 候选，不能显示
    r = client.get("/api/v1/search", params={"types": "政策文件", "time": "近 30 天"}).json()
    assert r["total"] == 0 and r["total_all"] == 1  # 8/1 的政策超出近 30 天
    r = client.get("/api/v1/search", params={"regions": "国家部委", "time": "不限"}).json()
    assert r["total"] == 2 and r["items"][0]["region"] == "国家"
    r = client.get("/api/v1/search", params={"depts": "工信部门", "time": "不限"}).json()
    assert [i["id"] for i in r["items"]] == [str(seeded["apply"])]
    r = client.get("/api/v1/search", params={"q": "辽工信投资", "time": "不限"}).json()
    assert r["total"] == 1  # 文号也能搜到
    r = client.get("/api/v1/search", params={"status": "公布", "time": "不限"}).json()
    assert r["total"] == 0


@pytest.fixture
def city_search_docs(session):
    specs = [
        ("国家级", "国家", None),
        ("省级", "山东", None),
        ("市级", "山东", "青岛"),
        ("市级", "山东", "济南"),
        ("市级", "江苏", "青岛"),
        ("省级", "江苏", None),
    ]
    ids = []
    for n, (level, province, city) in enumerate(specs):
        source = Source(
            level=level,
            province=province,
            city=city,
            dept="科技",
            dept_full_name=f"{province}{city or ''}科技部门",
            column_name="通知",
            url=f"https://city{n}.gov.cn/tz/",
            enabled=True,
            parser_config={},
            last_success_at=NOW,
        )
        session.add(source)
        session.flush()
        doc = _doc(source, n, "关于组织数字经济项目申报的通知", date(2026, 1, 1))
        session.add(doc)
        session.flush()
        ids.append(str(doc.id))
        session.add(
            DocumentDeadline(
                document_id=doc.id,
                fingerprint=f"city-{n}",
                kind="企业申报截止",
                candidate_date=date(2026, 10, 18),
                evidence="10月18日截止",
                confidence="高",
                status="approved",
                approved_date=date(2026, 10, 18),
                reviewed_at=NOW,
                reviewed_by="创始人",
            )
        )
    session.commit()
    return ids


def test_search_city_union_national_and_provinces(client, city_search_docs):
    ids = city_search_docs
    params = {"time": "不限", "regions": "国家部委", "cities": "山东/青岛"}
    res = client.get("/api/v1/search", params=params)
    assert res.status_code == 200
    assert {v["id"] for v in res.json()["items"]} == {ids[0], ids[2]}
    # 市名和省份作为一对比较，不误列另一省同名城市。
    params["regions"] = "国家部委,江苏"
    res = client.get("/api/v1/search", params=params).json()
    assert {v["id"] for v in res["items"]} == {ids[0], ids[2], ids[4], ids[5]}
    assert res["total"] == res["total_all"] == 4
    params |= {"regions": "国家部委", "cities": "山东/青岛,山东/济南"}
    assert {v["id"] for v in client.get("/api/v1/search", params=params).json()["items"]} == {
        ids[0],
        ids[2],
        ids[3],
    }


def test_client_search_city_open_industry_unlimited_time(client, session, city_search_docs):
    ids = city_search_docs
    # 未截止条件仍只认人工核对的截止，不能由标题或正文推定。
    session.scalar(
        select(DocumentDeadline).where(DocumentDeadline.document_id == int(ids[2]))
    ).status = "pending"
    session.commit()
    params = {
        "kind": "申报通知",
        "time": "不限",
        "open": 1,
        "regions": "国家部委",
        "cities": "山东/青岛",
        "industries": "数字经济",
        "group": 1,
    }
    res = client.get("/api/v1/search", params=params)
    assert res.status_code == 200
    assert [v["id"] for v in res.json()["items"]] == [ids[0]]
    # 不限产业只移除该筛选，其他条件继续生效。
    params["industries"] = "现代农业"
    assert client.get("/api/v1/search", params=params).json()["total"] == 0
    del params["industries"]
    assert [v["id"] for v in client.get("/api/v1/search", params=params).json()["items"]] == [
        ids[0]
    ]


@pytest.mark.parametrize("cities", ["山东/未收录", "青岛", "山东/青岛/其他", "江苏/济南"])
def test_search_rejects_unknown_city(client, city_search_docs, cities):
    res = client.get("/api/v1/search", params={"cities": cities})
    assert res.status_code == 422
    assert res.json()["detail"] == "城市未收录，请重新选择地区"


def test_link_gone_date_in_detail_and_cards(client, session, seeded):
    """原网页确认失效（youju/link_check.py）：详情和卡片都带上失效日期（北京时间），能打开的为 None。"""
    doc = session.get(Document, seeded["apply"])
    assert client.get(f"/api/v1/documents/{doc.id}").json()["link_gone"] is None
    doc.link_gone_at = datetime(2026, 10, 15, 20, 0, tzinfo=UTC)  # 北京时间 10 月 16 日 4:00
    session.commit()
    assert client.get(f"/api/v1/documents/{doc.id}").json()["link_gone"] == "2026-10-16"
    items = client.get("/api/v1/search", params={"time": "不限"}).json()["items"]
    card = next(item for item in items if item["id"] == str(doc.id))
    assert card["link_gone"] == "2026-10-16"


def test_document_detail_only_approved_deadlines(client, seeded):
    d = client.get(f"/api/v1/documents/{seeded['apply']}").json()
    assert d["status"] is None and d["app_status"]["code"] == "open"
    assert d["deadline"]["label"] == "企业申报截止" and d["deadline"]["date"] == "2026-10-18"
    assert d["deadline"]["notes"] == ["推荐或审核截止 10月20日 周二 17:00"]
    assert ["文号", "辽工信投资〔2026〕130号"] in d["points"]
    # 还没有副本、还没取文字：不能在小程序里打开、读，给原因（副本见 test_attachments.py；
    # 能不能读和阅读页同一个判断，见 test_readers.py）
    [attachment] = d["attachments"]
    assert attachment["id"].isdigit()
    assert {k: v for k, v in attachment.items() if k != "id"} == {
        "ext": "PDF",
        "name": "附件1 实施细则",
        "url": "https://gxt.ln.gov.cn/a1.pdf",
        "preview": False,
        "reason": "pending",
        "readable": False,
        "state": "pending",
        "stateText": "新附件，一般半天内可读",
        "scannedPages": 0,
    }
    assert d["checked"] == "2026-09-26" and d["host"] == "gxt.ln.gov.cn"
    listed = client.get(f"/api/v1/documents/{seeded['list']}").json()
    assert listed["deadline"] is None and listed["excerpt"] is None
    assert client.get("/api/v1/documents/999999").status_code == 404
    # 正文只有"……"（像网页上的字，没有一句话）不给，列着附件的说内容在附件里；名单类没有正文
    # 也没有附件的说正文还取不到（有正文的挡掉名单内容后给，D-092，见 test_list_body.py）
    assert d["body"] == {"status": "none", "reason": "attachments"}
    assert listed["body"] == {"status": "none", "reason": "site"}


def test_body_for_reading_cleans_only_known_strings():
    from youju.api.v1 import BODY_MAX_CHARS, body_for_reading

    text = (
        "各有关单位：\n\n\n\n请于10月18日前报送。联系人：[姓名已脱敏]，电话：[电话已脱敏]，"
        "[手机号已脱敏]，邮箱 [邮箱已脱敏]。【关闭窗口】【字体：大 中 小】扫一扫在手机打开当前页"
        + "正文"
        * 100
    )
    body = body_for_reading(text, "申报通知")
    assert body["status"] == "ok" and body["truncated"] is False
    assert "各有关单位：\n\n请于10月18日前报送" in body["text"]
    assert (
        "联系人：（这里隐去了联系人），电话：（这里隐去了电话），（这里隐去了电话），"
        "邮箱 （这里隐去了邮箱）" in body["text"]
    )
    for gone in ("【关闭窗口】", "【字体", "扫一扫", "已脱敏"):
        assert gone not in body["text"]
    assert body["total"] == len(body["text"])
    long = body_for_reading("字" * (BODY_MAX_CHARS + 5), "政策")
    assert long["truncated"] is True and long["total"] == BODY_MAX_CHARS + 5
    assert len(long["text"]) == BODY_MAX_CHARS
    # 像网页导航的（没有一句话、不长）不给；真短的通知照常给（D-097 去掉"不到 200 字不给"）
    assert body_for_reading("首页\n政务公开\n科技资讯\n搜 索", "申报通知") == {"status": "none"}
    assert body_for_reading("短" * 199, "申报通知") == {"status": "none"}
    assert body_for_reading(None, "申报通知") == {"status": "none"}
    short = body_for_reading("现将《管理办法》印发给你们，请遵照执行。", "政策")
    assert short == {
        "status": "ok",
        "text": "现将《管理办法》印发给你们，请遵照执行。",
        "total": 20,
        "truncated": False,
    }
    # 没有标点的长文（表格）照常给
    assert body_for_reading("序号\n企业名称\n" * 100, "政策")["status"] == "ok"
    assert body_for_reading("字" * 500, "名单") == {"status": "list"}
    # 资讯与指南、其他：标题写着个人事项（招聘、职称评审等）的不给，其余给全文（D-097 ②）；
    # 不给标题按个人事项算
    personal = {"status": "original", "reason": "personal"}
    assert body_for_reading("字" * 500, "其他") == personal
    assert body_for_reading("字" * 500, "其他", title="2026年职称评审结果公示") == personal
    assert (
        body_for_reading("字。" * 250, "资讯与指南", title="关于召开宣讲会的通知")["status"] == "ok"
    )
    # 给出前按现行规则再隐去一遍；Windows 换行也按换行处理
    again = body_for_reading("咨询电话：13812345678\r\n\r\n\r\n" + "内容。" * 80, "政策")
    assert "13812345678" not in again["text"] and "（这里隐去了电话）" in again["text"]
    assert "\r" not in again["text"] and "\n\n\n" not in again["text"]


def test_detail_body_prefers_own_level1_text(client, session, seeded):
    """同一通知几个网址：自己有正文就用自己的；没有才用别的网址里最长的；只取一级来源的正文。"""
    own = session.get(Document, seeded["policy"])
    src = session.get(Source, own.source_id)
    level2 = Source(
        level="省级",
        province="山东",
        dept="科技",
        dept_full_name="山东省科学技术厅",
        column_name="转载",
        url="https://example.gov.cn/zz/",
        enabled=True,
        parser_config={},
        rights_level=2,
    )
    session.add(level2)
    session.flush()
    twin1 = _doc(src, 31, own.title, own.publish_date)
    twin2 = _doc(level2, 32, own.title, own.publish_date)
    session.add_all([twin1, twin2])
    session.flush()
    now = datetime.now(UTC)
    for doc, text in ((twin1, "一级副本" * 100), (twin2, "二级来源" * 300)):
        session.add(DocumentDetail(document_id=doc.id, body_text=text, fetched_at=now, status="ok"))
        session.add(
            DocumentMerge(duplicate_id=doc.id, canonical_id=own.id, active=True, match_key="t")
        )
    session.commit()
    body = client.get(f"/api/v1/documents/{own.id}").json()["body"]
    assert body["status"] == "ok" and body["text"].startswith("一级副本")  # 二级的更长也不用
    session.add(
        DocumentDetail(document_id=own.id, body_text="本文。" * 100, fetched_at=now, status="ok")
    )
    session.commit()
    body = client.get(f"/api/v1/documents/{own.id}").json()["body"]
    assert body["text"].startswith("本文")  # 自己有就用自己的


def test_calendar_lists_only_approved_upcoming(client, seeded):
    items = client.get("/api/v1/calendar").json()["items"]
    assert [i["due"] for i in items] == ["2026-09-27", "2026-10-18", "2026-10-20"]
    assert all(i["verified"] for i in items)
    assert items[2]["time"] == "17:00" and items[1]["region"] == "辽宁 · 工信"


def test_calendar_and_home_merge_duplicate_notice(client, seeded, session):
    original = session.get(Document, seeded["apply"])
    duplicate = _doc(original.source, 88, original.title, original.publish_date)
    session.add(duplicate)
    session.flush()
    session.add(
        DocumentDeadline(
            document_id=duplicate.id,
            fingerprint="duplicate",
            kind="企业申报截止",
            candidate_date=date(2026, 10, 18),
            approved_date=date(2026, 10, 18),
            evidence="申报截止10月18日",
            confidence="高",
            status="approved",
        )
    )
    reconcile_documents(session)
    session.commit()
    body = client.get("/api/v1/calendar").json()
    assert body["total"] == 3
    assert len([i for i in body["items"] if i["docId"] == str(original.id)]) == 2
    assert client.get("/api/v1/home").json()["stats"]["week_total"] == 3


def test_recent_lists(client, seeded):
    body = client.get("/api/v1/lists/recent").json()
    assert body["total"] == 1 and body["week_total"] == 1
    assert body["items"][0]["kind"] == "名单"


def test_home_week_has_upper_bound_and_followed_new_is_null(client, seeded):
    stats = client.get("/api/v1/home").json()["stats"]
    assert stats["week_total"] == 3  # 9/28（下周一）的文件不算本周
    assert stats["followed_new"] is None


def test_home_last_checked_at_is_latest_successful_crawl(client, session, seeded):
    # 没有运行记录时为空
    assert client.get("/api/v1/home").json()["stats"]["last_checked_at"] is None
    session.add_all(
        [
            # 最近一次成功结束的抓取：9/26 22:30 UTC = 9/27 06:30 北京时间
            JobRun(
                job="crawl",
                started_at=datetime(2026, 9, 26, 22, 0, tzinfo=UTC),
                finished_at=datetime(2026, 9, 26, 22, 30, tzinfo=UTC),
                status="ok",
            ),
            JobRun(
                job="crawl",
                started_at=datetime(2026, 9, 25, 22, 0, tzinfo=UTC),
                finished_at=datetime(2026, 9, 25, 22, 40, tzinfo=UTC),
                status="ok",
            ),
            # 更晚的失败、还在跑的抓取和其他任务都不算
            JobRun(
                job="crawl",
                started_at=datetime(2026, 9, 27, 0, 0, tzinfo=UTC),
                finished_at=datetime(2026, 9, 27, 0, 30, tzinfo=UTC),
                status="failed",
            ),
            JobRun(job="crawl", started_at=datetime(2026, 9, 27, 1, 0, tzinfo=UTC)),
            JobRun(
                job="round",
                started_at=datetime(2026, 9, 27, 1, 0, tzinfo=UTC),
                finished_at=datetime(2026, 9, 27, 1, 50, tzinfo=UTC),
                status="ok",
            ),
        ]
    )
    session.commit()
    stats = client.get("/api/v1/home").json()["stats"]
    assert stats["last_checked_at"] == "2026-09-27T06:30:00+08:00"


def test_home_deadline_reviewed_at_is_latest_human_review(client, session, seeded):
    """截止核对断档（定稿方案第八节）：首页统计给最近一次人工核对截止的时间，待核对的不算。"""
    from sqlalchemy import func, update

    stats = client.get("/api/v1/home").json()["stats"]
    latest = session.scalar(
        select(func.max(DocumentDeadline.reviewed_at)).where(DocumentDeadline.status != "pending")
    )
    assert latest is not None
    assert stats["deadline_reviewed_at"] == latest.astimezone(ZoneInfo("Asia/Shanghai")).isoformat(
        timespec="seconds"
    )
    # 一条都没人工核对过：为空
    session.execute(
        update(DocumentDeadline).values(
            status="pending",
            reviewed_at=None,
            reviewed_by=None,
            approved_date=None,
            approved_time=None,
            changed_to_document_id=None,
        )
    )
    session.commit()
    assert client.get("/api/v1/home").json()["stats"]["deadline_reviewed_at"] is None


def test_calendar_skips_deadlines_already_passed_today(client, seeded):
    body = client.get("/api/v1/calendar").json()
    dues = [(i["due"], i["time"]) for i in body["items"]]
    assert ("2026-09-27", "09:00") not in dues
    assert dues[0] == ("2026-09-27", "17:00")
    assert body["total"] == len(body["items"]) == 3
    assert all(i["docId"] and i["id"].startswith("dl") for i in body["items"])


@pytest.mark.parametrize(
    "params",
    [
        {"types": "公告"},
        {"regions": "火星"},
        {"status": "已撤销"},
        {"depts": "财政部门"},
        {"kind": "新闻"},
        {"time": "近 180 天"},
        {"offset": 5001},
        {"regions": ",".join(["北京"] * 41)},
        {"levels": "县级"},
    ],
)
def test_search_rejects_invalid_filters(client, seeded, params):
    assert client.get("/api/v1/search", params=params).status_code == 422


def test_lists_lookup_exact_similar_and_coverage(client, seeded, session):
    list_doc = session.query(Document).filter(Document.title.contains("拟入库科技型")).one()
    common = {"category": "科技型中小企业", "batch": "第4批"}
    session.add_all(
        [
            ListEntry(
                document_id=list_doc.id,
                source_url=list_doc.url,
                name="山东星河科技有限公司",
                normalized_name="山东星河科技有限公司",
                credit_code="91310117588657256C",
                sequence="12",
                **common,
            ),
            # 同一份名单的附件里又出现一次：只列一次
            ListEntry(
                document_id=list_doc.id,
                source_url="https://kjt.shandong.gov.cn/a.xlsx",
                name="山东星河科技有限公司",
                normalized_name="山东星河科技有限公司",
                sequence="12",
                **common,
            ),
            ListEntry(
                document_id=list_doc.id,
                source_url=list_doc.url,
                name="山东星河科技（青岛）有限公司",
                normalized_name="山东星河科技(青岛)有限公司",
                sequence="13",
                **common,
            ),
        ]
    )
    session.commit()
    url = "/api/v1/lists/lookup"
    body = client.get(url, params={"q": "山东星河科技有限公司"}).json()
    assert [i["id"] for i in body["items"]] == [str(list_doc.id)]
    item = body["items"][0]
    assert item["kind"] == "名单" and item["status"] == "公示"
    # 查"原名单里的那一行"用的编号带签名（D-088）
    assert re.fullmatch(r"[A-Za-z0-9_-]{32}", item["matched"].pop("entryToken"))
    assert item["matched"] == {
        "name": "山东星河科技有限公司",
        "category": "科技型中小企业",
        "batch": "第4批",
        "sequence": "12",
        "creditCode": "91310117588657256C",
        "region": None,
        "via": "name",
    }
    # 名称相近的另列，不当作同一家
    assert body["similar"] == [{"name": "山东星河科技（青岛）有限公司", "lists": 1}]
    assert body["coverage"] == {
        "lists": 1,
        "since": "2026-09-24",
        "unreadLists": 0,
        "partialLists": 0,
    }
    # 全角括号、空格不同也能查到；信用代码不分大小写
    assert len(client.get(url, params={"q": "山东星河科技（青岛） 有限公司"}).json()["items"]) == 1
    assert len(client.get(url, params={"q": "91310117588657256c"}).json()["items"]) == 1
    # 查不到：结果为空，收录范围照给，由小程序说明"已收录的名单中未找到"
    missing = client.get(url, params={"q": "不存在的科技有限公司"}).json()
    assert missing["items"] == [] and missing["coverage"]["lists"] == 1
    # 只查企业和机构
    refused = client.get(url, params={"q": "张三"})
    assert refused.status_code == 422 and "只查企业和机构" in refused.json()["detail"]
    assert body["ambiguous"] is False and body["truncated"] is False
    assert (
        item["matched"]["via"] == "name" and item["matched"]["creditCode"] == "91310117588657256C"
    )


def test_lists_lookup_rename_same_name_and_duplicate_notice(client, seeded, session):
    """更名：信用代码相同、名称不同的也列出；同名不同代码：提示可能不是同一家；重复通知只列一次。"""
    list_doc = session.get(Document, seeded["list"])
    src = list_doc.source
    old_doc = _doc(src, 11, "山东省2025年第9批入库科技型中小企业名单公告", date(2025, 11, 3))
    dup_doc = _doc(src, 12, list_doc.title, list_doc.publish_date)  # 同一通知，网址不同
    other = _doc(src, 13, "2026年度高新技术企业认定名单公告", date(2026, 9, 1))
    # 别的机关同一天发了同名通知：不是同一份，不能合并
    ln_src = session.query(Source).filter_by(province="辽宁").one()
    same_title_other_agency = _doc(ln_src, 14, list_doc.title, list_doc.publish_date)
    session.add_all([old_doc, dup_doc, other, same_title_other_agency])
    session.flush()
    code = "91310117588657256C"
    rows = [
        (list_doc, "山东星河科技有限公司", code),
        (dup_doc, "山东星河科技有限公司", code),
        (old_doc, "济南星河科技有限公司", code),  # 更名前
        (other, "山东星河科技有限公司", "91370100MA3C5LTN6Q"),  # 同名、代码不同
        (same_title_other_agency, "山东星河科技有限公司", code),
    ]
    for i, (doc, name, credit) in enumerate(rows):
        session.add(
            ListEntry(
                document_id=doc.id,
                source_url=doc.url,
                name=name,
                normalized_name=name,
                credit_code=credit,
                sequence=str(i + 1),
                category="科技型中小企业",
            )
        )
    session.commit()
    reconcile_documents(session)  # 按内容归并：同一机关同名同日的两个网址归为一份
    session.commit()
    body = client.get("/api/v1/lists/lookup", params={"q": "山东星河科技有限公司"}).json()
    titles = [i["title"] for i in body["items"]]
    # 同一机关的重复通知只列一次；别的机关的同名通知单独列出
    assert titles.count(list_doc.title) == 2
    regions = sorted(i["region"] for i in body["items"] if i["title"] == list_doc.title)
    assert regions == ["山东 · 省级", "辽宁 · 省级"]
    assert old_doc.title in titles and other.title in titles
    renamed = next(i for i in body["items"] if i["title"] == old_doc.title)
    assert renamed["matched"]["name"] == "济南星河科技有限公司"
    assert renamed["matched"]["via"] == "creditCode"
    assert body["ambiguous"] is True
    # 更名前的名称已在结果里，不再算"名称相近"
    assert body["similar"] == []


def test_lists_lookup_shows_master_even_if_its_source_is_not_parsed(client, seeded, session):
    """名单只从可解析的来源（rights_level=1）抽取；条目所在副本归并到别的来源的主文件时，仍列出主文件。"""
    from youju.models import DocumentMerge

    list_doc = session.get(Document, seeded["list"])
    other_src = Source(
        level="省级",
        province="山东",
        dept="科技",
        dept_full_name="山东省科学技术厅（转载）",
        column_name="转载",
        url="https://www.shandong.gov.cn/kj/",
        enabled=True,
        parser_config={},
        rights_level=0,
    )
    session.add(other_src)
    session.flush()
    master = _doc(other_src, 51, list_doc.title, list_doc.publish_date)
    session.add(master)
    session.flush()
    # 主文件编号更小时才会成为主文件；这里直接写归并关系：list_doc 是副本，master 是主文件
    session.add(
        DocumentMerge(duplicate_id=list_doc.id, canonical_id=master.id, active=True, match_key="k")
    )
    session.add(
        ListEntry(
            document_id=list_doc.id,
            source_url=list_doc.url,
            name="山东星河科技有限公司",
            normalized_name="山东星河科技有限公司",
            sequence="1",
            category="科技型中小企业",
        )
    )
    session.commit()
    body = client.get("/api/v1/lists/lookup", params={"q": "山东星河科技有限公司"}).json()
    assert [i["id"] for i in body["items"]] == [str(master.id)]
    assert body["items"][0]["matched"]["name"] == "山东星河科技有限公司"


def test_deadline_change_hides_old_deadline_and_links_new_notice(
    client, seeded, session, monkeypatch
):
    """验收清单 A7：发了延期通知后，原通知的截止标为已变更，公开端不再当作有效截止。"""
    from types import SimpleNamespace

    monkeypatch.setattr("youju.api.app.get_settings", lambda: SimpleNamespace(admin_token="secret"))
    headers = {"X-Admin-Token": "secret"}
    old = session.get(Document, seeded["apply"])
    main = (
        session.query(DocumentDeadline)
        .filter_by(document_id=old.id, kind="企业申报截止", status="approved")
        .one()
    )
    new = _doc(old.source, 9, "关于延长2026年三季度贷款贴息项目申报时间的通知", date(2026, 9, 26))
    session.add(new)
    session.commit()

    found = client.get(
        "/api/admin/deadlines/approved", params={"q": "贷款贴息 三季度"}, headers=headers
    ).json()
    assert main.id in [f["id"] for f in found]
    url = f"/api/admin/deadlines/{main.id}/change"
    body = {"reviewed_by": "审阅人", "reason": "延期", "new_document_id": new.id}
    assert client.post(url, json=body).status_code == 401
    # 必须关联发布变更的新通知；原因只能是三种之一；新通知不能是自己
    no_link = {"reviewed_by": "审阅人", "reason": "撤销"}
    assert client.post(url, json=no_link, headers=headers).status_code == 422
    bad_reason = {**body, "reason": "推迟"}
    assert client.post(url, json=bad_reason, headers=headers).status_code == 422
    self_link = {**body, "new_document_id": old.id}
    assert client.post(url, json=self_link, headers=headers).status_code == 422
    # 防止关联到无关文件：别的地区、或发布早于原通知的，不行
    other_region = session.get(Document, seeded["list"])  # 山东的文件，原通知在辽宁
    wrong = {**body, "new_document_id": other_region.id}
    assert client.post(url, json=wrong, headers=headers).status_code == 422
    earlier = _doc(old.source, 10, "关于调整贷款贴息项目有关事项的通知", date(2026, 9, 1))
    session.add(earlier)
    session.commit()
    too_early = {**body, "new_document_id": earlier.id}
    assert client.post(url, json=too_early, headers=headers).status_code == 422
    undated = _doc(old.source, 11, "关于延长贷款贴息项目申报时间的补充通知", None)
    session.add(undated)
    session.commit()
    no_date = {**body, "new_document_id": undated.id}
    assert client.post(url, json=no_date, headers=headers).status_code == 422
    assert client.post(url, json=body, headers=headers).status_code == 200
    assert client.post(url, json=body, headers=headers).status_code == 409

    detail = client.get(f"/api/v1/documents/{old.id}").json()
    # 原来的企业申报截止不再作为主截止，剩下的推荐截止照常显示
    assert detail["deadline"]["label"] == "推荐或审核截止"
    assert detail["deadline_changes"] == [
        {
            "label": "企业申报截止",
            "date": "2026-10-18",
            "time": "",
            "reason": "延期",
            "newDocId": str(new.id),
            "newTitle": new.title,
        }
    ]
    calendar = client.get("/api/v1/calendar").json()["items"]
    assert not [i for i in calendar if i["docId"] == str(old.id) and i["due"] == "2026-10-18"]

    # 标错了可以撤回：恢复为已通过，变更记录清空
    back = f"/api/admin/deadlines/{main.id}/unchange"
    assert client.post(back, json={"reviewed_by": "审阅人"}, headers=headers).status_code == 200
    assert client.get(f"/api/v1/documents/{old.id}").json()["deadline_changes"] == []
    session.refresh(main)
    assert main.status == "approved" and main.changed_by is None and main.changed_at is None
    # 撤销也要关联撤销通知
    revoke = {"reviewed_by": "审阅人", "reason": "撤销", "new_document_id": new.id}
    assert client.post(url, json=revoke, headers=headers).status_code == 200
    changes = client.get(f"/api/v1/documents/{old.id}").json()["deadline_changes"]
    assert changes[0]["reason"] == "撤销" and changes[0]["newDocId"] == str(new.id)
    docs = client.get("/api/admin/documents", params={"q": "延长 贷款贴息"}, headers=headers)
    assert [d["id"] for d in docs.json()][0] == new.id  # 按发布日期从新到旧，日期不详的排后


def test_deadline_change_covers_merged_twins(client, seeded, session, monkeypatch):
    """同一通知的两个网址归并后，各有一条相同的已通过截止：标记变更时两条一并处理，公开端不再显示。"""
    from types import SimpleNamespace

    from youju.models import DocumentMerge

    monkeypatch.setattr("youju.api.app.get_settings", lambda: SimpleNamespace(admin_token="secret"))
    headers = {"X-Admin-Token": "secret"}
    old = session.get(Document, seeded["apply"])
    twin = _doc(old.source, 21, old.title, old.publish_date)
    new = _doc(old.source, 22, "关于延长2026年三季度贷款贴息项目申报时间的通知", date(2026, 9, 26))
    session.add_all([twin, new])
    session.flush()
    session.add(
        DocumentMerge(duplicate_id=twin.id, canonical_id=old.id, active=True, match_key="k")
    )
    main = (
        session.query(DocumentDeadline)
        .filter_by(document_id=old.id, kind="企业申报截止", status="approved")
        .one()
    )
    copy = DocumentDeadline(
        document_id=twin.id,
        fingerprint="twin",
        kind=main.kind,
        candidate_date=main.candidate_date,
        evidence=main.evidence,
        confidence=main.confidence,
        status="approved",
        reviewed_by="创始人",
        reviewed_at=main.reviewed_at,
        approved_date=main.approved_date,
        approved_time=main.approved_time,
    )
    session.add(copy)
    session.commit()
    # 核对台搜索只列一条
    found = client.get("/api/admin/deadlines/approved", params={"q": "贷款贴息"}, headers=headers)
    assert len([f for f in found.json() if f["kind"] == "企业申报截止"]) == 1
    # 新通知不能是原通知的另一个网址
    body = {"reviewed_by": "审阅人", "reason": "延期", "new_document_id": twin.id}
    url = f"/api/admin/deadlines/{copy.id}/change"
    assert client.post(url, json=body, headers=headers).status_code == 422
    body["new_document_id"] = new.id
    assert client.post(url, json=body, headers=headers).status_code == 200
    session.refresh(main)
    session.refresh(copy)
    assert main.status == copy.status == "changed"
    detail = client.get(f"/api/v1/documents/{old.id}").json()
    assert detail["deadline"]["label"] == "推荐或审核截止"
    assert [c["reason"] for c in detail["deadline_changes"]] == ["延期"]
    # 撤回也一并恢复
    back = f"/api/admin/deadlines/{main.id}/unchange"
    assert client.post(back, json={"reviewed_by": "审阅人"}, headers=headers).status_code == 200
    session.refresh(copy)
    assert copy.status == "approved"


def test_changed_deadline_does_not_come_back_via_new_duplicate_url(
    client, seeded, session, monkeypatch
):
    """截止已标\"已延期\"后，又抓到同一通知的另一个网址：它的候选不能再作为有效截止出现。"""
    from types import SimpleNamespace

    from youju.dedup import reconcile_documents
    from youju.extract.service import extract_existing

    monkeypatch.setattr("youju.api.app.get_settings", lambda: SimpleNamespace(admin_token="secret"))
    headers = {"X-Admin-Token": "secret"}
    old = session.get(Document, seeded["apply"])
    main = (
        session.query(DocumentDeadline)
        .filter_by(document_id=old.id, kind="企业申报截止", status="approved")
        .one()
    )
    new = _doc(old.source, 31, "关于延长2026年三季度贷款贴息项目申报时间的通知", date(2026, 9, 26))
    session.add(new)
    session.commit()
    body = {"reviewed_by": "审阅人", "reason": "延期", "new_document_id": new.id}
    assert (
        client.post(
            f"/api/admin/deadlines/{main.id}/change", json=body, headers=headers
        ).status_code
        == 200
    )
    # 情形一：新网址先抽出候选（还没归并），之后归并——待核对副本沿用\"已变更\"
    twin = _doc(old.source, 32, old.title, old.publish_date)
    session.add(twin)
    session.flush()
    pending = DocumentDeadline(
        document_id=twin.id,
        fingerprint="dup-url",
        kind=main.kind,
        candidate_date=main.candidate_date,
        candidate_time=main.candidate_time,
        evidence=main.evidence,
        confidence="高",
        status="pending",
    )
    session.add(pending)
    session.commit()
    reconcile_documents(session)
    session.commit()
    session.refresh(pending)
    assert pending.status == "changed" and pending.changed_to_document_id == new.id
    detail = client.get(f"/api/v1/documents/{old.id}").json()
    assert detail["deadline"]["label"] == "推荐或审核截止"  # 企业申报截止不再作为有效截止
    # 情形二：已归并后再抽取这个网址的正文：同组已核对过的截止不再送审
    session.add(
        DocumentDetail(
            document_id=twin.id,
            body_text="企业申报截止时间为10月18日。",
            fetched_at=datetime.now(UTC),
            status="ok",
        )
    )
    session.commit()
    before = session.query(DocumentDeadline).filter_by(status="pending").count()
    extract_existing(session, document_ids=[twin.id])
    assert session.query(DocumentDeadline).filter_by(status="pending").count() == before


def test_reviewed_search_prefers_still_valid_deadline(client, seeded, session, monkeypatch):
    """核对台搜索：同组同时有已通过和已变更的，列出已通过那条（仍有效，需要处理）。"""
    from types import SimpleNamespace

    from youju.models import DocumentMerge

    monkeypatch.setattr("youju.api.app.get_settings", lambda: SimpleNamespace(admin_token="secret"))
    headers = {"X-Admin-Token": "secret"}
    old = session.get(Document, seeded["apply"])
    main = (
        session.query(DocumentDeadline)
        .filter_by(document_id=old.id, kind="企业申报截止", status="approved")
        .one()
    )
    new = _doc(old.source, 41, "关于延长2026年三季度贷款贴息项目申报时间的通知", date(2026, 9, 26))
    twin = _doc(old.source, 42, old.title, old.publish_date)
    session.add_all([new, twin])
    session.flush()
    session.add(
        DocumentMerge(duplicate_id=twin.id, canonical_id=old.id, active=True, match_key="k")
    )
    main.status, main.change_reason, main.changed_to_document_id = "changed", "延期", new.id
    session.add(
        DocumentDeadline(
            document_id=twin.id,
            fingerprint="still-valid",
            kind=main.kind,
            candidate_date=main.candidate_date,
            evidence=main.evidence,
            confidence="高",
            status="approved",
            reviewed_by="创始人",
            reviewed_at=main.reviewed_at,
            approved_date=main.approved_date,
        )
    )
    session.commit()
    found = client.get("/api/admin/deadlines/approved", params={"q": "贷款贴息"}, headers=headers)
    rows = [f for f in found.json() if f["kind"] == "企业申报截止"]
    assert len(rows) == 1 and rows[0]["status"] == "approved"
    # 公开端：同组同一截止已标"已变更"，仍为"已通过"的副本也不当作有效截止
    detail = client.get(f"/api/v1/documents/{old.id}").json()
    assert detail["deadline"]["label"] == "推荐或审核截止"
    assert [c["reason"] for c in detail["deadline_changes"]] == ["延期"]


def test_twin_sync_rules_for_confirmed_and_rejected(client, seeded, session):
    """已通过、已变更的：同组同类型同日期时间的待核对副本沿用结果（证据句写法不同也算）；
    已驳回的：证据句也相同才沿用，免得把真正的截止当成已驳回的那句。"""
    from youju.dedup import sync_reviewed_twins
    from youju.models import DocumentMerge

    old = session.get(Document, seeded["apply"])
    twin = _doc(old.source, 61, old.title, old.publish_date)
    session.add(twin)
    session.flush()
    session.add(
        DocumentMerge(duplicate_id=twin.id, canonical_id=old.id, active=True, match_key="k")
    )
    rejected = DocumentDeadline(
        document_id=old.id,
        fingerprint="rej",
        kind="材料报送截止",
        candidate_date=date(2026, 11, 30),
        evidence="证书有效期至11月30日",
        confidence="低",
        status="rejected",
        reviewed_by="创始人",
        reviewed_at=datetime.now(UTC),
    )
    session.add(rejected)
    twins = {
        "confirmed": DocumentDeadline(
            document_id=twin.id,
            fingerprint="t1",
            kind="企业申报截止",
            candidate_date=date(2026, 10, 18),
            evidence="申报截止时间：10月18日",  # 与已通过那条写法不同
            confidence="高",
            status="pending",
        ),
        "real": DocumentDeadline(
            document_id=twin.id,
            fingerprint="t2",
            kind="材料报送截止",
            candidate_date=date(2026, 11, 30),
            evidence="纸质材料11月30日前报送",  # 同日期同类型，但不是被驳回的那句
            confidence="高",
            status="pending",
        ),
    }
    session.add_all(twins.values())
    session.commit()
    sync_reviewed_twins(session, {old.id})
    session.commit()
    session.refresh(twins["confirmed"])
    session.refresh(twins["real"])
    assert twins["confirmed"].status == "approved"
    assert twins["confirmed"].approved_date == date(2026, 10, 18)
    assert twins["real"].status == "pending"
    # 同一日期既有通过又有驳回时，证据句与驳回那条完全相同的副本沿用驳回
    approved_same_day = DocumentDeadline(
        document_id=old.id,
        fingerprint="ok-1130",
        kind="材料报送截止",
        candidate_date=date(2026, 11, 30),
        evidence="纸质材料于11月30日前寄送",
        confidence="高",
        status="approved",
        reviewed_by="创始人",
        reviewed_at=datetime.now(UTC),
        approved_date=date(2026, 11, 30),
    )
    copy_of_rejected = DocumentDeadline(
        document_id=twin.id,
        fingerprint="t3",
        kind="材料报送截止",
        candidate_date=date(2026, 11, 30),
        evidence="证书有效期至11月30日",
        confidence="低",
        status="pending",
    )
    session.add_all([approved_same_day, copy_of_rejected])
    session.commit()
    sync_reviewed_twins(session, {old.id})
    session.commit()
    session.refresh(copy_of_rejected)
    assert copy_of_rejected.status == "rejected"


def test_coverage_lists_covered_regions_and_uncovered_provinces(client, session, seeded):
    # seeded：国家、辽宁、山东各一个启用的栏目。已收录要求至少成功检查过一次
    ok = datetime(2026, 9, 28, 1, 30, tzinfo=UTC)
    for province in ("国家", "山东"):
        session.scalar(select(Source).where(Source.province == province)).last_success_at = ok
    session.add_all(
        [
            Source(
                level="省级",
                province="浙江",
                dept="科技",
                dept_full_name="浙江省科学技术厅",
                column_name="通知",
                url="https://kjt.zj.gov.cn/tz/",
                enabled=False,
                fetch_mode="browser",
                parser_config={},
            ),
            Source(
                level="省级",
                province="广西",
                dept="科技",
                dept_full_name="广西壮族自治区科学技术厅",
                column_name="通知",
                url="https://kjt.gxzf.gov.cn/tz/",
                enabled=False,
                fetch_mode="blocked",
                parser_config={},
            ),
        ]
    )
    session.commit()

    body = client.get("/api/v1/coverage").json()
    assert [r["name"] for r in body["regions"]] == ["国家部委", "山东"]
    assert body["sources"] == 2
    assert body["provinces"] == 1
    assert body["national"] is True
    assert body["regions"][1] == {
        "name": "山东",
        "sources": 1,
        "lastUpdatedAt": "2026-09-28T09:30:00+08:00",
        "columns": [
            {"name": "山东省科学技术厅 · 通知", "lastUpdatedAt": "2026-09-28T09:30:00+08:00"}
        ],
    }
    assert body["updatedAt"] == "2026-09-28T09:30:00+08:00"
    reasons = {item["name"]: item["reason"] for item in body["notCovered"]}
    assert reasons["辽宁"] == "正在接入"  # 已启用但还没成功检查过
    # 需要浏览器渲染的没有接入计划（浙江 9/29 决定不绕过网站的访问限制，#16），不写"正在接入"（10/13）
    assert reasons["浙江"] == "暂未接入"
    assert reasons["广西"] == "网站不允许程序访问"
    assert reasons["北京"] == "暂未接入"
    assert "山东" not in reasons and "国家部委" not in reasons
    assert len(body["notCovered"]) == 30


def test_coverage_reports_national_when_no_national_source_succeeded(client, session, seeded):
    body = client.get("/api/v1/coverage").json()
    assert body["regions"] == []
    assert body["national"] is False
    assert {"name": "国家部委", "reason": "正在接入"} in body["notCovered"]


def test_coverage_counts_only_successfully_checked_columns(client, session, seeded):
    # 山东两个启用的栏目，只有一个成功检查过：栏目数算 1，另一个列出但时间为空（正在接入）
    sd = session.scalar(select(Source).where(Source.province == "山东"))
    sd.last_success_at = datetime(2026, 9, 28, 1, 30, tzinfo=UTC)
    session.add(
        Source(
            level="省级",
            province="山东",
            dept="工信",
            dept_full_name="山东省工业和信息化厅",
            column_name="公示",
            url="https://gxt.shandong.gov.cn/gs/",
            enabled=True,
            parser_config={},
        )
    )
    session.commit()
    shandong = client.get("/api/v1/coverage").json()["regions"][0]
    assert shandong["sources"] == 1
    assert [c["lastUpdatedAt"] for c in shandong["columns"]] == [
        None,
        "2026-09-28T09:30:00+08:00",
    ]


def _candidate(doc_id, fingerprint, due, evidence, status="pending", clock=None):
    """截止候选；status 为 approved 时同时填核定日期。"""
    return DocumentDeadline(
        document_id=doc_id,
        fingerprint=fingerprint,
        kind="企业申报截止",
        candidate_date=due,
        candidate_time=clock,
        evidence=evidence,
        confidence="中",
        status=status,
        approved_date=due if status == "approved" else None,
        approved_time=clock if status == "approved" else None,
    )


def test_deadline_state_verified_unverified_none(client, session, seeded):
    """三种截止状态：有核对通过的 verified；只有待核对候选的 unverified；没有候选的 none。"""
    sd = session.get(Document, seeded["list"]).source
    day = date(2026, 9, 25)
    waiting = _doc(sd, 41, "山东省科学技术厅关于组织申报2027年度重点研发计划项目的通知", day)
    plain = _doc(sd, 42, "山东省科学技术厅关于组织申报2027年度软科学研究计划项目的通知", day)
    rejected = _doc(sd, 43, "山东省科学技术厅关于组织申报2027年度科技示范工程项目的通知", day)
    master = _doc(sd, 44, "山东省科学技术厅关于组织申报2027年度技术创新引导项目的通知", day)
    twin = _doc(sd, 45, master.title, day)
    session.add_all([waiting, plain, rejected, master, twin])
    session.flush()
    session.add_all(
        [
            _candidate(waiting.id, "w", date(2026, 11, 23), "11月23日前提交"),
            _candidate(rejected.id, "r", date(2026, 11, 23), "11月23日前提交", status="rejected"),
            # 待核对候选在同一通知的另一个网址上：主文件也算"截止未核对"
            _candidate(twin.id, "t", date(2026, 11, 23), "11月23日前提交"),
            DocumentMerge(duplicate_id=twin.id, canonical_id=master.id, active=True, match_key="t"),
        ]
    )
    session.commit()

    body = client.get("/api/v1/search", params={"time": "不限", "limit": 50}).json()
    items = {i["id"]: i for i in body["items"]}
    assert all(i["deadline_state"] in v1.DEADLINE_STATES for i in body["items"])
    assert all(
        (i["deadline"] is not None) == (i["deadline_state"] == "verified") for i in items.values()
    )
    expected = {
        seeded["apply"]: "verified",
        waiting.id: "unverified",
        plain.id: "none",
        rejected.id: "none",
        master.id: "unverified",
    }
    for doc_id, state in expected.items():
        assert items[str(doc_id)]["kind"] == "申报通知"
        assert items[str(doc_id)]["deadline_state"] == state
        detail = client.get(f"/api/v1/documents/{doc_id}").json()
        assert detail["deadline_state"] == state
        assert (detail["deadline"] is not None) == (state == "verified")
    assert str(twin.id) not in items
    # 首页、最近名单用同一个列表格式
    home = {i["id"]: i for i in client.get("/api/v1/home").json()["items"]}
    assert home[str(waiting.id)]["deadline_state"] == "unverified"
    assert client.get("/api/v1/lists/recent").json()["items"][0]["deadline_state"] == "unverified"

    # 已有核对通过的截止，再来新的待核对候选：仍按已核对算，日期只给核对过的
    session.add(_candidate(seeded["apply"], "p", date(2026, 10, 30), "10月30日前报送"))
    session.commit()
    detail = client.get(f"/api/v1/documents/{seeded['apply']}").json()
    assert detail["deadline_state"] == "verified" and detail["deadline"]["date"] == "2026-10-18"


def test_empty_states_get_verified_total(client, session, seeded):
    """截止雷达和申报日历返回已核对条数：有核对通过截止的公开文件份数，同一通知算一份。"""
    v1._REPORT_CACHE.clear()
    apply_doc = session.get(Document, seeded["apply"])
    twin = _doc(apply_doc.source, 71, apply_doc.title, apply_doc.publish_date)
    hidden_source = Source(
        level="省级",
        province="辽宁",
        dept="工信",
        dept_full_name="辽宁省工业和信息化厅",
        column_name="转载",
        url="https://example.gov.cn/zz/",
        enabled=True,
        parser_config={},
        rights_level=2,
    )
    session.add_all([twin, hidden_source])
    session.flush()
    hidden = _doc(
        hidden_source, 72, "辽宁省关于组织申报2026年度技术改造项目的通知", date(2026, 9, 20)
    )
    session.add(hidden)
    session.flush()
    session.add_all(
        [
            _candidate(twin.id, "twin", date(2026, 10, 18), "10月18日", status="approved"),
            DocumentMerge(
                duplicate_id=twin.id, canonical_id=apply_doc.id, active=True, match_key="t"
            ),
            # 只存题录的二级来源：不公开，不算
            _candidate(hidden.id, "hidden", date(2026, 10, 9), "10月9日", status="approved"),
        ]
    )
    session.commit()

    # 申报通知（两个截止、两个网址）和政策各一份；名单只有待核对候选，不算
    report = client.get("/api/v1/reports/deadlines", params={"days": 30}).json()
    assert report["verified_total"] == 2
    filtered = client.get("/api/v1/reports/deadlines", params={"province": "北京"}).json()
    assert filtered["items"] == [] and filtered["verified_total"] == 2  # 不随筛选变化
    assert client.get("/api/v1/calendar").json()["verifiedTotal"] == 2

    # 线上现状：全部还在待核对。两处都没有内容，已核对条数为 0
    for row in session.scalars(select(DocumentDeadline)).all():
        row.status, row.approved_date, row.approved_time = "pending", None, None
    session.commit()
    report = client.get("/api/v1/reports/deadlines", params={"days": 30}).json()
    assert report["items"] == [] and report["total"] == 0 and report["verified_total"] == 0
    calendar = client.get("/api/v1/calendar").json()
    assert calendar["items"] == [] and calendar["past"] == [] and calendar["verifiedTotal"] == 0
    assert isinstance(calendar["verifiedTotal"], int)


def test_search_level_filter_and_region_merge_rule(client, session, seeded):
    """层级和地区同时用：组内取并集、组间取交集；两组里的"国家部委"是同一个条件。"""
    city = Source(
        level="市级",
        province="山东",
        dept="科技",
        dept_full_name="济南市科学技术局",
        column_name="通知",
        url="https://jnsti.jinan.gov.cn/tz/",
        enabled=True,
        parser_config={},
    )
    session.add(city)
    session.flush()
    city_doc = _doc(
        city, 81, "济南市科学技术局关于组织申报2026年度科技计划项目的通知", date(2026, 9, 20)
    )
    session.add(city_doc)
    session.commit()

    def ids(**params):
        response = client.get("/api/v1/search", params={"time": "不限", "limit": 50, **params})
        assert response.status_code == 200, response.text
        return {int(i["id"]) for i in response.json()["items"]}

    everything = ids()
    national = ids(regions="国家部委")
    provincial = {seeded["apply"], seeded["list"], seeded["policy"]}
    shandong_provincial = {seeded["list"], seeded["policy"]}
    assert len(national) == 2 and city_doc.id in everything
    # 现有行为不变：不传层级或传空串，和以前一样
    assert ids(levels="") == everything
    assert ids(regions="山东") == shandong_provincial | {city_doc.id}
    # 国家部委：只在地区选、只在层级选、两边都选，结果相同
    assert ids(levels="国家部委") == national
    assert ids(regions="国家部委", levels="国家部委") == national
    assert ids(regions="不限地区", levels="国家部委") == national
    assert ids(regions="国家部委,山东", levels="国家部委") == national
    # 省级不含国家级，也不含市级
    assert ids(levels="市级") == {city_doc.id}
    assert ids(levels="国家部委,省级,市级") == everything
    assert ids(regions="山东", levels="市级") == {city_doc.id}
    assert ids(regions="国家部委,山东", levels="国家部委,市级") == national | {city_doc.id}
    assert ids(levels="省级") == provincial
    assert ids(levels="国家部委,省级") == everything - {city_doc.id}
    assert ids(regions="山东", levels="省级") == shandong_provincial
    assert ids(regions="国家部委,山东", levels="省级") == shandong_provincial
    assert ids(regions="国家部委,山东", levels="国家部委,省级") == national | shandong_provincial
    # 没有交集的选法：结果为空，不报错
    assert ids(regions="国家部委", levels="省级") == set()
    assert ids(regions="山东", levels="国家部委") == set()
    # 和其他筛选、页签、不限时间总数一起用
    assert ids(levels="省级", kind="申报通知") == {seeded["apply"]}
    body = client.get("/api/v1/search", params={"levels": "省级"}).json()
    assert body["total"] == 2 and body["total_all"] == 3  # 8/1 的政策超出近 30 天
    for bad in ("县级", "国家级", "省级,火星"):
        assert client.get("/api/v1/search", params={"levels": bad}).status_code == 422


def test_search_national_level_with_regions_as_miniapp_sends_it(client, session, seeded):
    """UI 第二轮：小程序的地区里不再有"国家部委"，层级里的国家部委不受地区限制、和所选地区的文件
    一起列出。服务端不改：小程序把请求拼成 regions=国家部委,省份（层级同时含省级时再带
    levels=国家部委,省级），见 apps/miniapp/utils/conditions.js 的 requestScope。这里核对这样拼出来的
    结果，以及旧条件打开后换成的新请求（migrateNational）结果不变。"""
    city = Source(
        level="市级",
        province="山东",
        dept="科技",
        dept_full_name="济南市科学技术局",
        column_name="通知",
        url="https://jnsti.jinan.gov.cn/tz/",
        enabled=True,
        parser_config={},
    )
    session.add(city)
    session.flush()
    city_doc = _doc(
        city, 82, "济南市科学技术局关于组织申报2026年度科技计划项目的通知", date(2026, 9, 20)
    )
    session.add(city_doc)
    session.commit()

    def ids(**params):
        response = client.get("/api/v1/search", params={"time": "不限", "limit": 50, **params})
        assert response.status_code == 200, response.text
        return {int(i["id"]) for i in response.json()["items"]}

    national = ids(levels="国家部委")
    provincial = ids(levels="省级")
    shandong = {seeded["list"], seeded["policy"]}
    assert len(national) == 2 and shandong <= provincial
    # 界面：地区山东 + 层级国家部委 → 国家部委和山东的文件（包括已收录的市级文件）
    assert ids(regions="国家部委,山东") == national | shandong | {city_doc.id}
    # 地区山东 + 层级国家部委、省级 → 国家部委和山东的省级文件
    assert ids(regions="国家部委,山东", levels="国家部委,省级") == national | shandong
    # 地区山东 + 层级省级、不限地区 + 两级：和以前一样
    assert ids(regions="山东", levels="省级") == shandong
    assert ids(regions="不限地区", levels="国家部委,省级") == national | provincial

    # 旧条件 → 打开后小程序发的请求（与 apps/miniapp/tests/search-r2.test.js 的转换表一致），结果相同
    same = [
        ({"regions": "国家部委"}, {"regions": "不限地区", "levels": "国家部委"}),
        ({"regions": "不限地区,国家部委"}, {"regions": "不限地区", "levels": "国家部委"}),
        (
            {"regions": "国家部委", "levels": "国家部委,省级"},
            {"regions": "不限地区", "levels": "国家部委"},
        ),
        (
            {"regions": "国家部委,山东", "levels": "国家部委"},
            {"regions": "不限地区", "levels": "国家部委"},
        ),
        (
            {"regions": "国家部委,山东,北京", "levels": "省级"},
            {"regions": "山东,北京", "levels": "省级"},
        ),
        ({"regions": "山东", "levels": "国家部委,省级"}, {"regions": "山东", "levels": "省级"}),
    ]
    for old, new in same:
        assert ids(**old), old
        assert ids(**old) == ids(**new), (old, new)
    # 原来必然查不到的两种，打开后按字面意思列出
    assert ids(regions="国家部委", levels="省级") == set()
    assert ids(regions="不限地区", levels="国家部委,省级") == national | provincial
    assert ids(regions="山东", levels="国家部委") == set()
    assert ids(regions="国家部委,山东") == national | shandong | {city_doc.id}


def test_no_public_output_contains_unverified_dates(client, session, seeded):
    """红线：没核对的截止日期不能从任何公开接口出来（含时刻和原文证据）。

    正文（D-030）是原文全文、页面写明未经核对，这里的正文不含候选日期，只检查结构化字段。
    """
    ln = session.get(Document, seeded["apply"]).source
    notice = _doc(
        ln, 91, "辽宁省工业和信息化厅关于组织申报2026年度数字化转型项目的通知", date(2026, 9, 25)
    )
    session.add(notice)
    session.flush()
    session.add_all(
        [
            DocumentDetail(
                document_id=notice.id,
                body_text="各有关单位：" + "请按要求组织申报。" * 40,
                issuer="辽宁省工业和信息化厅",
                fetched_at=datetime.now(UTC),
                status="ok",
            ),
            _candidate(
                notice.id, "x", date(2026, 11, 23), "请于11月23日16:30前提交", clock=time(16, 30)
            ),
            # 已核对的文件上又有一条待核对候选
            _candidate(seeded["apply"], "y", date(2026, 11, 24), "材料11月24日前报送"),
        ]
    )
    session.commit()
    v1._REPORT_CACHE.clear()
    found = client.get("/api/v1/search", params={"time": "不限", "kind": "申报通知"}).json()
    assert {i["id"]: i["deadline_state"] for i in found["items"]}[str(notice.id)] == "unverified"

    forbidden = ("2026-11-23", "11月23日", "11-23", "16:30", "2026-11-24", "11月24日", "11-24")
    forbidden += ("2026-10-05", "10月5日", "10-05")  # seeded 名单上的待核对候选
    paths = [
        "/api/v1/home",
        "/api/v1/search?time=不限&limit=50",
        "/api/v1/search?time=不限&kind=申报通知&levels=省级",
        "/api/v1/search?time=不限&open=1",
        "/api/v1/calendar",
        "/api/v1/calendar?days=366&limit=200",
        "/api/v1/reports/deadlines?days=30",
        "/api/v1/lists/recent",
        f"/api/v1/documents/{notice.id}",
        f"/api/v1/documents/{seeded['apply']}",
        f"/api/v1/documents/{seeded['list']}",
    ]
    for path in paths:
        response = client.get(path)
        assert response.status_code == 200, path
        text = json.dumps(response.json(), ensure_ascii=False)
        leaked = [value for value in forbidden if value in text]
        assert not leaked, f"{path} 出现了未核对的日期：{leaked}"


@pytest.fixture
def city_document(session, seeded):
    source = Source(
        level="市级",
        province="山东",
        city="青岛",
        dept="科技",
        dept_full_name="测试市科学技术局",
        column_name="通知公告",
        url="https://city.example.gov.cn/list/",
        enabled=True,
        parser_config={},
        last_success_at=NOW,
    )
    session.add(source)
    session.flush()
    doc = _doc(source, 92, "关于组织申报科技计划项目的通知", NOW.date())
    session.add(doc)
    session.commit()
    return doc


def test_city_display_coverage_and_home(client, session, city_document):
    doc = city_document
    payload = client.get("/api/v1/search", params={"regions": "山东", "levels": "市级"}).json()
    assert payload["total"] == 1
    item = payload["items"][0]
    assert (item["region"], item["city"], item["level"], item["province"]) == (
        "青岛 · 市级",
        "青岛",
        "市级",
        "山东",
    )
    detail = client.get(f"/api/v1/documents/{doc.id}").json()
    assert (detail["region"], detail["city"], detail["level"]) == ("青岛 · 市级", "青岛", "市级")
    assert ["适用地区", "山东·青岛"] in detail["points"]
    assert detail["deadline"] is None
    body = client.get("/api/v1/coverage").json()
    assert body["cities"] == [{"province": "山东", "city": "青岛", "sources": 1}]
    assert body["provinces"] == 1
    assert client.get("/api/v1/home").json()["stats"]["cities"] == 1
    # 同城多栏目只算一个城市，停用和未成功检查的栏目不计数。
    for index, enabled, success in [(1, True, NOW), (2, False, NOW), (3, True, None)]:
        session.add(
            Source(
                level="市级",
                province="山东",
                city="青岛" if index == 1 else "济南",
                dept="科技",
                dept_full_name="测试机关",
                column_name="通知",
                url=f"https://city.example.gov.cn/list{index}/",
                enabled=enabled,
                last_success_at=success,
            )
        )
    session.commit()
    assert client.get("/api/v1/coverage").json()["cities"] == [
        {"province": "山东", "city": "青岛", "sources": 2}
    ]
    assert client.get("/api/v1/home").json()["stats"]["cities"] == 1


def test_no_city_counts_and_labels_unchanged(client, seeded):
    assert client.get("/api/v1/coverage").json()["cities"] == []
    assert client.get("/api/v1/home").json()["stats"]["cities"] == 0
    doc = client.get(f"/api/v1/documents/{seeded['apply']}").json()
    assert doc["region"] == "辽宁 · 省级"
    assert doc["city"] is None and doc["level"] == "省级"
    assert ["适用地区", "辽宁"] in doc["points"]


def test_city_calendar_radar_and_exports_require_approved_deadlines(client, session, city_document):
    from io import BytesIO

    from docx import Document as WordDocument
    from openpyxl import load_workbook

    doc = city_document
    dl = DocumentDeadline(
        document_id=doc.id,
        fingerprint="city-deadline",
        kind="企业申报截止",
        candidate_date=date(2026, 9, 30),
        evidence="测试截止证据",
        confidence="高",
        status="pending",
    )
    session.add(dl)
    session.commit()

    def calendar_ids():
        return {item["docId"] for item in client.get("/api/v1/calendar").json()["items"]}

    def radar():
        return client.get("/api/v1/reports/deadlines", params={"province": "山东"}).json()

    assert str(doc.id) not in calendar_ids()
    assert radar()["total"] == 0
    assert all(item["id"] != str(doc.id) for item in radar()["items"])
    dl.status, dl.approved_date, dl.reviewed_at = "approved", dl.candidate_date, NOW
    session.commit()
    v1._REPORT_CACHE.clear()
    assert str(doc.id) in calendar_ids()
    assert any(item["id"] == str(doc.id) for item in radar()["items"])
    for fmt in ("xlsx", "docx"):
        response = client.post("/api/v1/exports", json={"ids": [str(doc.id)], "format": fmt})
        assert response.status_code == 200
        if fmt == "xlsx":
            sheet = load_workbook(BytesIO(response.content))["文件清单"]
            assert sheet.cell(2, 8).value == "山东·青岛"
        else:
            book = WordDocument(BytesIO(response.content))
            text = " ".join(
                cell.text for table in book.tables for row in table.rows for cell in row.cells
            )
            assert "山东·青岛" in text


def test_city_deadline_change_requires_same_city(session, city_document):
    from fastapi import HTTPException

    from youju.api.app import _check_change_target

    old = city_document
    source = session.get(Source, old.source_id)
    new = _doc(source, 94, "关于延长科技计划项目申报时间的通知", NOW.date())
    other_source = Source(
        level="市级",
        province="山东",
        city="济南",
        dept="科技",
        dept_full_name="测试另一市科技局",
        column_name="通知",
        url="https://another.example.gov.cn/list/",
    )
    session.add(other_source)
    session.flush()
    other = _doc(other_source, 95, new.title, NOW.date())
    session.add_all([new, other])
    session.flush()
    deadline = DocumentDeadline(document_id=old.id)
    _check_change_target(session, deadline, new.id)
    with pytest.raises(HTTPException, match="不是同一地区") as error:
        _check_change_target(session, deadline, other.id)
    assert error.value.status_code == 422


def test_client_search_90_days_excludes_only_verified_closed(client, seeded, session):
    source = session.query(Source).filter_by(province="山东").one()
    docs = [
        _doc(source, 100 + i, f"第{i}项先进制造项目申报通知", day)
        for i, day in enumerate(
            [NOW.date(), date(2026, 6, 29), date(2026, 6, 28), NOW.date(), NOW.date(), NOW.date()]
        )
    ]
    session.add_all(docs)
    session.flush()
    for i, state in [(0, "pending"), (3, "approved"), (4, "approved"), (5, "changed")]:
        session.add(
            DocumentDeadline(
                document_id=docs[i].id,
                fingerprint=str(i),
                kind="企业申报截止",
                status=state,
                candidate_date=date(2026, 9, 26),
                evidence="原文",
                confidence="高",
                approved_date=date(2026, 9, 26) if state != "pending" else None,
                # 已变更的要写原因并关联发布变更的通知（数据库约束）
                change_reason="延期" if state == "changed" else None,
                changed_to_document_id=docs[0].id if state == "changed" else None,
            )
        )
    # 第4项今天的截止已过，同样排除；无日期、候选过期和变更仍列出。
    session.flush()
    dl = session.query(DocumentDeadline).filter_by(document_id=docs[4].id).one()
    dl.approved_date = NOW.date()
    dl.approved_time = time(9, 0)
    session.commit()
    params = {"q": "先进制造", "kind": "申报通知", "time": "近 90 天"}
    before = client.get("/api/v1/search", params=params).json()
    assert before["total"] == 5
    body = client.get("/api/v1/search", params={**params, "exclude_closed": 1}).json()
    assert {i["id"] for i in body["items"]} == {str(docs[i].id) for i in [0, 1, 5]}
    assert body["total_all"] == 4
    assert all(i["deadline"] is None for i in body["items"])
    assert body["verified_open_total"] == 1  # 全站已核对未截止，不受关键词影响
    strict = client.get("/api/v1/search", params={**params, "open": 1}).json()
    assert strict["total"] == 0
    assert client.get("/api/v1/search", params={**params, "exclude_closed": 0}).json()["total"] == 5


def test_client_search_zero_verified_still_lists_unverified(client, seeded, session):
    for dl in session.query(DocumentDeadline):
        dl.status = "pending"
    session.commit()
    body = client.get(
        "/api/v1/search", params={"kind": "申报通知", "exclude_closed": 1, "time": "近 90 天"}
    ).json()
    assert body["total"] == 1 and body["verified_open_total"] == 0
    assert body["items"][0]["deadline"] is None
    assert client.get("/api/v1/search", params={"open": 1}).json()["total"] == 0


def test_lists_shorthand_candidates_are_orgs_sorted_unique_and_limited(client, seeded, session):
    from youju.lists.lookup import lookup_name_candidates
    from youju.lists.rules import normalize_name

    doc = session.query(Document).filter(Document.title.contains("拟入库科技型")).one()
    names = ["海尔集团", "海尔研究院", "青岛海尔有限公司", "海尔集团", "张三", "海尔小王"]
    names += [f"海尔{i:02d}技术有限公司" for i in range(25)]
    for i, name in enumerate(names):
        session.add(
            ListEntry(
                document_id=doc.id,
                source_url=doc.url,
                name=name,
                normalized_name=normalize_name(name),
                sequence=str(900 + i),
                category="科技型中小企业",
            )
        )
    session.commit()
    body = client.get("/api/v1/lists/lookup", params={"q": "海尔"}).json()
    # 以"海尔"开头的排前面（"青岛海尔有限公司"排在后面），再按名称长度、字典序
    expected = sorted(
        set(n for n in names if n not in {"张三", "海尔小王"}),
        key=lambda n: (not n.startswith("海尔"), len(n), n),
    )[:20]
    assert "青岛海尔有限公司" not in expected
    assert body["items"] == [] and body["candidates"] == expected
    exact = client.get("/api/v1/lists/lookup", params={"q": expected[0]}).json()
    assert len(exact["items"]) == 1 and exact["candidates"] == []
    assert client.get("/api/v1/lists/lookup", params={"q": "张三"}).status_code == 422
    assert client.get("/api/v1/lists/lookup", params={"q": "海"}).status_code == 422
    for query in ["海尔%", "海尔_", "' OR 1=1 --", "海尔!"]:
        assert client.get("/api/v1/lists/lookup", params={"q": query}).json()["candidates"] == []
    # 查全部条目，数量在数据库里限住
    assert lookup_name_candidates(session, "海尔", limit=3) == expected[:3]
    # 像个人姓名的输入不做精确比对：即使库里有同名的遗留条目也不返回
    assert client.get("/api/v1/lists/lookup", params={"q": "海尔小王"}).status_code == 422
    doc.kind_override = "其他"
    session.commit()
    assert client.get("/api/v1/lists/lookup", params={"q": "海尔集团"}).json()["candidates"] == []


def test_lists_lookup_rate_limited_and_recovers(client, monkeypatch):
    monkeypatch.setattr(v1._time, "monotonic", lambda: 1000)
    for _ in range(30):
        assert client.get("/api/v1/lists/lookup", params={"q": "测试公司"}).status_code == 200
    assert client.get("/api/v1/lists/lookup", params={"q": "测试公司"}).status_code == 429
    monkeypatch.setattr(v1._time, "monotonic", lambda: 1061)
    assert client.get("/api/v1/lists/lookup", params={"q": "测试公司"}).status_code == 200


def test_search_window_per_address_and_recovers():
    """S3 计数：每个来源地址一分钟滑动窗口；超了拒绝，被拒的不记次；不同地址互不影响。"""
    for _ in range(v1.SEARCH_PER_MINUTE):
        assert v1._allow_search("203.0.113.1", 1000.0)
    assert not v1._allow_search("203.0.113.1", 1059.9)
    assert len(v1._SEARCH_TIMES["ip:203.0.113.1"]) == v1.SEARCH_PER_MINUTE  # 被拒的不记
    assert v1._allow_search("203.0.113.2", 1059.9)
    assert v1._allow_search("203.0.113.1", 1060.0)  # 最早那批满 60 秒，恢复


def test_search_rate_limited_per_address(client, monkeypatch):
    """S3：同一来源地址每分钟搜索有上限，超了返回 429"搜索太频繁，请稍后再试"；别的地址照常；
    快照里取最新 8 份不计次；计数只按来源地址，不保存 openid（隐私保护指引一·2 没写搜索计次）。"""
    from types import SimpleNamespace

    from youju.api import me

    monkeypatch.setattr(
        me, "get_settings", lambda: SimpleNamespace(trust_wx_headers=True, dev_openid=None)
    )
    monkeypatch.setattr(v1, "SEARCH_PER_MINUTE", 5)
    user = {
        "X-WX-SOURCE": "wx_client",
        "X-WX-OPENID": "o-search-a",
        "X-Original-Forwarded-For": "203.0.113.7",
    }
    for _ in range(5):
        assert client.get("/api/v1/search", headers=user).status_code == 200
    refused = client.get("/api/v1/search", headers=user, params={"q": "申报"})
    assert refused.status_code == 429
    assert refused.json() == {"detail": "搜索太频繁，请稍后再试"}
    # 换个 openid、同一出口：照样被拦（不按 openid 算）
    other_user = {**user, "X-WX-OPENID": "o-search-b"}
    assert client.get("/api/v1/search", headers=other_user).status_code == 429
    elsewhere = {**user, "X-Original-Forwarded-For": "203.0.113.8"}
    assert client.get("/api/v1/search", headers=elsewhere).status_code == 200
    v1._SNAPSHOT_CACHE.clear()
    assert client.get("/api/v1/snapshot", headers=elsewhere).status_code == 200
    assert len(v1._SEARCH_TIMES["ip:203.0.113.8"]) == 1
    assert set(v1._SEARCH_TIMES) == {"ip:203.0.113.7", "ip:203.0.113.8"}
    assert not any("o-search" in key for key in v1._SEARCH_TIMES)


def test_search_limit_on_admin_by_address_before_db(session_factory, monkeypatch):
    """admin 上同样按来源地址计次（换着 openid 刷绕不过去）；超限的请求在取数据库会话之前就拒绝。"""
    from types import SimpleNamespace

    from youju.api import app as app_module

    opened = 0

    def override():
        nonlocal opened
        opened += 1
        with session_factory() as s:
            yield s

    monkeypatch.setattr(
        app_module, "get_settings", lambda: SimpleNamespace(role="admin", admin_token="t" * 32)
    )
    monkeypatch.setattr(v1, "SEARCH_PER_MINUTE", 3)
    app.dependency_overrides[get_session] = override
    try:
        client = TestClient(app)
        office = {"X-WX-SOURCE": "wx_client", "X-Original-Forwarded-For": "203.0.113.7"}
        for n in range(3):
            headers = {**office, "X-WX-OPENID": f"o-{n}"}
            assert client.get("/api/v1/search", headers=headers).status_code == 200
        assert opened == 3
        refused = client.get("/api/v1/search", headers={**office, "X-WX-OPENID": "o-new"})
        assert refused.status_code == 429 and opened == 3
        elsewhere = {**office, "X-Original-Forwarded-For": "203.0.113.8"}
        assert client.get("/api/v1/search", headers=elsewhere).status_code == 200
        assert set(v1._SEARCH_TIMES) == {"ip:203.0.113.7", "ip:203.0.113.8"}
    finally:
        app.dependency_overrides.clear()
