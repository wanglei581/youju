from datetime import date, datetime, timedelta
from hashlib import sha256
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook

from youju.api import v1
from youju.models import Document, DocumentDeadline, DocumentDetail, DocumentField, Source
from youju.track import (
    CustomerProfile,
    CustomerTrack,
    TrackReport,
    export_track_excel,
    read_customer_profiles,
    track_clients,
)

BJ = ZoneInfo("Asia/Shanghai")


def test_read_profiles_and_mark_unknown_province(tmp_path: Path):
    path = tmp_path / "clients.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["客户名称", "省份", "关键词", "支持方式"])
    sheet.append(["甲公司", "北京", "专精特新; 数字化", "资金支持;未知标签"])
    sheet.append(["乙公司", "火星", "", ""])
    book.save(path)

    profiles = read_customer_profiles(path)
    assert profiles[0].keywords == ("专精特新", "数字化")
    assert profiles[0].supports == ("资金支持",)
    assert profiles[1].invalid_province


def test_export_keeps_empty_clients_and_escapes_formula(tmp_path: Path):
    output = tmp_path / "track.xlsx"
    report = TrackReport(
        [CustomerTrack(CustomerProfile("=客户", "北京"))],
        datetime(2026, 9, 30, 1, 2, tzinfo=ZoneInfo("UTC")),
    )
    assert export_track_excel(report, output) == 1
    book = load_workbook(output, data_only=False)
    assert book["客户相关通知"]["A3"].value == "'=客户"
    assert book["客户相关通知"]["B3"].value == "暂无"
    assert book["汇总"]["B3"].value == 0
    # 没有通知的客户：不写“未核对”，也不写 0 以外的数
    assert [cell.value for cell in book["汇总"][3]][2:] == [0, None, None, None, 0]
    assert "2026-09-30 09:02" in book["客户相关通知"]["A1"].value


def test_track_uses_radar_scope_and_customer_filters(session, monkeypatch, tmp_path):
    stamp = datetime(2026, 9, 30, 10, 0, tzinfo=BJ)
    monkeypatch.setattr(v1, "now", lambda: stamp)
    national = Source(
        level="国家级",
        province="国家",
        dept="科技",
        dept_full_name="国家科技部",
        column_name="通知",
        url="https://example.gov.cn/national",
        rights_level=1,
    )
    beijing = Source(
        level="省级",
        province="北京",
        dept="科技",
        dept_full_name="北京市科学技术委员会",
        column_name="通知",
        url="https://example.gov.cn/beijing",
        rights_level=1,
    )
    session.add_all([national, beijing])
    session.flush()

    def add(source, number, title, days, support="资金支持", status="approved"):
        document = Document(
            source_id=source.id,
            url=f"https://example.gov.cn/{number}",
            url_key=f"example/{number}",
            title=title,
            publish_date=date(2026, 9, 1),
            first_seen_at=stamp,
            last_seen_at=stamp,
        )
        session.add(document)
        session.flush()
        body = f"正文 {number}"
        session.add(
            DocumentDetail(
                document_id=document.id,
                body_text=body,
                issuer=source.dept_full_name,
                fetched_at=stamp,
                status="ok",
            )
        )
        session.add(
            DocumentDeadline(
                document_id=document.id,
                fingerprint=number,
                kind="企业申报截止",
                candidate_date=stamp.date() + timedelta(days=days),
                approved_date=stamp.date() + timedelta(days=days) if status == "approved" else None,
                evidence="企业申报截止",
                confidence="高",
                status=status,
            )
        )
        if support:
            session.add(
                DocumentField(
                    document_id=document.id,
                    field="support",
                    value=support,
                    approved_value=support,
                    evidence=support,
                    evidence_offset=0,
                    body_hash=sha256(body.encode()).hexdigest(),
                    status="approved",
                )
            )

    add(national, "n", "国家数字化申报通知", 5)
    add(beijing, "b", "北京专精特新申报通知", 10, "贷款贴息")
    add(beijing, "late", "北京远期申报通知", 31)
    add(beijing, "pending", "北京未核对申报通知", 5, status="pending")
    session.commit()

    radar = v1._report_payload(session, 30, "", "", "", 1, 100)
    report = track_clients(
        session,
        [
            CustomerProfile("北京客户", "北京"),
            CustomerProfile("关键词客户", "北京", ("专精特新",), ("贷款贴息",)),
            CustomerProfile("上海客户", "上海"),
        ],
    )
    by_name = {client.customer_name: client for client in report.clients}
    beijing_items = by_name["北京客户"].items
    # 跟进表列全部未截止项目；30 天内的部分与截止雷达同口径（验收 A2）
    assert (
        len(
            [
                item
                for item in beijing_items
                if item.remaining_days is not None and item.remaining_days <= 30
            ]
        )
        == radar["total"]
    )
    assert [item.title for item in beijing_items][-1] == "北京未核对申报通知"
    # 每一行的剩余天数是它自己的，不是最后一份通知的
    assert [item.remaining_days for item in beijing_items] == [5, 10, 31, None]
    assert [item.title for item in by_name["关键词客户"].items] == ["北京专精特新申报通知"]
    assert [item.title for item in by_name["上海客户"].items] == ["国家数字化申报通知"]
    # 导出：核对过的写日期和“已核对”；汇总只在已核对的里面数 7/14/30 天，另列没核对的条数
    output = tmp_path / "mixed.xlsx"
    export_track_excel(report, output)
    book = load_workbook(output)
    rows = [
        [cell.value for cell in row]
        for row in book["客户相关通知"].iter_rows(min_row=3)
        if row[0].value == "北京客户"
    ]
    assert [row[5] for row in rows] == ["已核对", "已核对", "已核对", "未核对，请以原文为准"]
    assert [row[7] is not None for row in rows] == [True, True, True, False]
    assert [row[8] for row in rows] == [5, 10, 31, None]
    summary = {
        row[0].value: [cell.value for cell in row] for row in book["汇总"].iter_rows(min_row=3)
    }
    assert summary["北京客户"] == ["北京客户", 4, 3, 1, 2, 2, 1]
    assert summary["上海客户"] == ["上海客户", 1, 1, 1, 1, 1, 0]


def test_track_without_verified_deadlines_and_excludes_only_closed(session, monkeypatch, tmp_path):
    stamp = datetime(2026, 10, 9, 12, tzinfo=BJ)
    monkeypatch.setattr(v1, "now", lambda: stamp)
    source = Source(
        level="省级",
        province="北京",
        dept="科技",
        dept_full_name="北京科技部门",
        column_name="通知",
        url="https://example.gov.cn/notices",
        rights_level=1,
    )
    session.add(source)
    session.flush()

    def add(key, age):
        document = Document(
            source_id=source.id,
            title=f"{key}项目申报通知",
            url=f"https://example.gov.cn/{key}",
            url_key=f"example/{key}",
            publish_date=stamp.date() - timedelta(days=age),
            first_seen_at=stamp,
            last_seen_at=stamp,
        )
        session.add(document)
        session.flush()
        return document

    recent = add("近期待核对", 1)
    boundary = add("90天", 90)
    add("91天", 91)
    closed = add("已截止", 1)
    session.commit()
    report = track_clients(session, [CustomerProfile("客户", "北京")])
    assert {item.document_id for item in report.items} == {recent.id, boundary.id, closed.id}
    assert all(
        item.as_dict()["deadline_status"] == "截止未核对，请以原文为准" for item in report.items
    )
    # 没核对的排在后面，新发布的先（发布 1 天的两份在前，90 天的在后）
    assert {item.document_id for item in report.items[:2]} == {recent.id, closed.id}
    assert report.items[-1].document_id == boundary.id
    output = tmp_path / "unknown.xlsx"
    export_track_excel(report, output)
    book = load_workbook(output)
    sheet = book["客户相关通知"]
    assert [cell.value for cell in sheet[2]] == [
        "客户名称",
        "通知标题",
        "发文机关",
        "省份",
        "发布日期",
        "截止核对情况",
        "截止类型",
        "截止日期",
        "剩余天数",
        "支持方式",
        "原文链接",
    ]
    # 截止日期一列只写日期：没核对的留空，核对情况另写一列
    assert sheet["E3"].value == (stamp.date() - timedelta(days=1)).isoformat()
    assert sheet["F3"].value == "未核对，请以原文为准"
    assert sheet["G3"].value is None and sheet["H3"].value is None and sheet["I3"].value is None
    # 汇总：一条都没核对时，“N 天内截止”写“未核对”，不写 0（0 会被看成没有快截止的）
    summary = book["汇总"]
    assert [cell.value for cell in summary[2]] == [
        "客户名称",
        "通知数",
        "已核对截止日期的",
        "7 天内截止",
        "14 天内截止",
        "30 天内截止",
        "截止未核对的",
    ]
    assert [cell.value for cell in summary[3]] == ["客户", 3, 0, "未核对", "未核对", "未核对", 3]
    assert "只数已核对截止日期的通知" in summary["A1"].value
    session.add(
        DocumentDeadline(
            document_id=closed.id,
            fingerprint="closed",
            kind="企业申报截止",
            candidate_date=stamp.date() - timedelta(days=1),
            approved_date=stamp.date() - timedelta(days=1),
            evidence="截止",
            confidence="高",
            status="approved",
        )
    )
    session.commit()
    report = track_clients(session, [CustomerProfile("客户", "北京")])
    assert {item.document_id for item in report.items} == {recent.id, boundary.id}
