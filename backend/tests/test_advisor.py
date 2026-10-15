"""工具任务接口、队列状态和名单保存期限。"""

import base64
from datetime import UTC, datetime, timedelta
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile
from zoneinfo import ZoneInfo

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select

from youju import advisor_tasks as service
from youju.api import advisor as advisor_api
from youju.api import me, v1
from youju.api.app import app
from youju.db import get_session
from youju.models import AdvisorFileTrash, AdvisorTask, AdvisorUsage, User

BJ = ZoneInfo("Asia/Shanghai")
NOW = datetime(2026, 9, 30, 23, 59, tzinfo=BJ)


@pytest.fixture
def advisor_settings(monkeypatch):
    settings = SimpleNamespace(
        advisor_enabled=True,
        cloud_env="test-env",
        dev_openid="advisor-user",
        trust_wx_headers=False,
        agreement_version="2026-09",
        retain_events_days=90,
        retain_feedback_days=180,
        retain_used_subscriptions_days=30,
        retain_unused_subscriptions_days=365,
    )
    monkeypatch.setattr(advisor_api, "get_settings", lambda: settings)
    monkeypatch.setattr(me, "get_settings", lambda: settings)
    monkeypatch.setattr(v1, "now", lambda: NOW)
    return settings


@pytest.fixture
def client(session_factory, advisor_settings):
    def override():
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_session] = override
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def agree(client):
    response = client.post("/api/v1/me/agreement", json={"version": "2026-09"})
    assert response.status_code == 200


def agree_advisor(client):
    agree(client)
    response = client.post(
        "/api/v1/me/agreement",
        json={"scope": "advisor", "version": service.ADVISOR_AGREEMENT_VERSION},
    )
    assert response.status_code == 200


def upload_body(content, *, name="clients.csv", kind="compare"):
    return {
        "kind": kind,
        "file": {"name": name, "content_base64": base64.b64encode(content).decode()},
    }


def xlsx_content(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_advisor_status_public_quota_and_switch(client, advisor_settings, monkeypatch):
    advisor_settings.advisor_enabled = False
    assert client.get("/api/v1/advisor/status").status_code == 404
    advisor_settings.advisor_enabled = True
    expected = {"used": 0, "limit": 3, "maxRows": 50}
    assert client.get("/api/v1/advisor/status").json() == {
        "enabled": True,
        "pass": {"active": False, "expiresAt": None},
        "quota": {"compare": {**expected, "limit": 3}, "track": expected},
    }
    agree_advisor(client)
    created = create(client).json()
    assert client.get("/api/v1/advisor/status").json()["quota"]["compare"]["used"] == 1
    client.post(f"/api/v1/advisor/tasks/{created['id']}/cancel")
    assert client.get("/api/v1/advisor/status").json()["quota"]["compare"]["used"] == 0
    create(client, key="another")
    monkeypatch.setattr(v1, "now", lambda: datetime(2026, 10, 1, 0, 1, tzinfo=BJ))
    assert client.get("/api/v1/advisor/status").json()["quota"]["compare"]["used"] == 0


def test_advisor_agreement_is_separate_and_versioned(client, advisor_settings):
    agree(client)
    assert (
        client.post("/api/v1/me/agreement", json={"scope": "advisor", "version": "old"}).status_code
        == 422
    )
    agree_advisor(client)
    profile = client.get("/api/v1/me").json()
    assert profile["agreed"]["version"] == "2026-09"
    assert profile["advisorAgreed"]["version"] == service.ADVISOR_AGREEMENT_VERSION
    assert client.get("/api/v1/me/data").json()["user"]["advisorAgreed"] == profile["advisorAgreed"]
    advisor_settings.advisor_enabled = False
    assert (
        client.post(
            "/api/v1/me/agreement",
            json={"scope": "advisor", "version": service.ADVISOR_AGREEMENT_VERSION},
        ).status_code
        == 404
    )


def test_compare_monthly_quota_keeps_daily_ledger_and_refunds_original_month(
    client, session, monkeypatch
):
    agree_advisor(client)
    stamps = [
        datetime(2026, 9, 1, 0, 0, tzinfo=BJ),
        datetime(2026, 9, 2, 0, 0, tzinfo=BJ),
        datetime(2026, 9, 30, 23, 59, 59, tzinfo=BJ),
    ]
    created = []
    for n, stamp in enumerate(stamps):
        monkeypatch.setattr(v1, "now", lambda stamp=stamp: stamp.astimezone(UTC))
        response = create(client, key=f"monthly-{n}", rows=["甲公司"] * 50)
        assert response.status_code == 202
        created.append(response.json()["id"])
        assert client.get("/api/v1/advisor/status").json()["quota"]["compare"]["used"] == n + 1
        assert (
            client.get("/api/v1/advisor/capabilities").json()["data_tools"]["remaining"]["compare"]
            == 2 - n
        )
    assert create(client, key="monthly-fourth").status_code == 429
    assert len(list(session.scalars(select(AdvisorUsage)))) == 3

    # UTC 尚在九月，北京时间已到十月；取消旧月任务不增加新月额度。
    stamp = datetime(2026, 9, 30, 16, 0, tzinfo=UTC)
    monkeypatch.setattr(v1, "now", lambda: stamp)
    assert client.get("/api/v1/advisor/status").json()["quota"]["compare"] == {
        "used": 0,
        "limit": 3,
        "maxRows": 50,
    }
    assert create(client, key="new-month").status_code == 202
    assert client.post(f"/api/v1/advisor/tasks/{created[0]}/cancel").status_code == 200
    assert client.get("/api/v1/advisor/status").json()["quota"]["compare"]["used"] == 1
    assert client.post(f"/api/v1/advisor/tasks/{created[0]}/cancel").status_code == 200
    session.expire_all()
    assert (
        service.usage_counts(session, session.scalar(select(User.id)), stamps[0].date())["compare"]
        == 2
    )


def test_compare_reservation_reads_latest_usage_after_user_lock(client, session_factory):
    agree_advisor(client)
    assert create(client, key="first").status_code == 202
    assert create(client, key="second").status_code == 202
    with session_factory() as stale:
        stale.connection(execution_options={"isolation_level": "REPEATABLE READ"})
        user_id = stale.scalar(select(User.id))
        day = service.usage_day(NOW)
        assert service.usage_counts(stale, user_id, day)["compare"] == 2
        # 另一个请求先占了最后一次；旧快照不能在拿到用户锁后继续放行。
        assert create(client, key="third").status_code == 202
        stale.execute(select(User.id).where(User.id == user_id).with_for_update()).one()
        assert not service.reserve_usage(stale, user_id, "compare", day)


@pytest.mark.parametrize(
    "content,name",
    [
        ("企业名称\n甲公司\n乙公司\n".encode(), "clients.csv"),
        ("企业名称\n甲公司\n乙公司\n".encode("gbk"), "clients.xlsx"),
        (xlsx_content([["企业名称"], ["甲公司"], ["乙公司"]]), "clients.csv"),
    ],
)
def test_upload_content_detection_preview_and_idempotency(client, session, content, name):
    agree(client)
    body = upload_body(content, name=name)
    assert client.post("/api/v1/advisor/preview", json=body).status_code == 403
    assert (
        client.post(
            "/api/v1/advisor/tasks", headers={"Idempotency-Key": "file"}, json=body
        ).status_code
        == 403
    )
    agree_advisor(client)
    assert client.post("/api/v1/advisor/preview", json=body).json() == {"rowCount": 2}
    assert client.get("/api/v1/advisor/status").json()["quota"]["compare"]["used"] == 0
    first = client.post("/api/v1/advisor/tasks", headers={"Idempotency-Key": "file"}, json=body)
    assert first.status_code == 202 and first.json()["rowCount"] == 2
    row = session.get(AdvisorTask, first.json()["id"])
    assert row.input_rows == ["甲公司", "乙公司"]
    assert name not in str(row.input_rows) and "content_base64" not in first.text
    duplicate = client.post(
        "/api/v1/advisor/tasks", headers={"Idempotency-Key": "file"}, json=upload_body(b"PKinvalid")
    )
    assert duplicate.status_code == 202 and duplicate.json()["id"] == row.id
    assert client.get("/api/v1/advisor/status").json()["quota"]["compare"]["used"] == 1


@pytest.mark.parametrize("xlsx", [False, True])
def test_upload_track_profiles(client, session, xlsx):
    agree_advisor(client)
    rows = [["客户名称", "省份", "关键词", "支持方式"], ["甲公司", "辽宁", "高新;科技", "资金支持"]]
    content = xlsx_content(rows) if xlsx else "\n".join(",".join(row) for row in rows).encode()
    body = upload_body(content, kind="track")
    assert client.post("/api/v1/advisor/preview", json=body).json()["rowCount"] == 1
    response = client.post("/api/v1/advisor/tasks", headers={"Idempotency-Key": "track"}, json=body)
    assert response.status_code == 202
    assert session.get(AdvisorTask, response.json()["id"]).input_rows == [
        {
            "name": "甲公司",
            "province": "辽宁",
            "city": "",
            "keywords": ["高新", "科技"],
            "supports": ["资金支持"],
        }
    ]


@pytest.mark.parametrize(
    "body,reason",
    [
        (upload_body(b"test", name="clients.txt"), "只支持"),
        ({"kind": "compare", "file": {"name": "a.csv", "content_base64": "!"}}, "编码无效"),
        (upload_body(b"x" * (service.MAX_UPLOAD_BYTES + 1)), "60 KB"),
        (upload_body(b"PK\x03\x04broken"), "不是 CSV"),
        (upload_body(b"\x00\x01binary"), "不是 CSV"),
        (upload_body(b""), "没有可提交"),
        (upload_body(b"no header", kind="track"), "表头"),
    ],
)
def test_bad_upload_does_not_use_quota(client, body, reason):
    agree_advisor(client)
    response = client.post("/api/v1/advisor/tasks", headers={"Idempotency-Key": "bad"}, json=body)
    assert response.status_code == 422
    assert reason in response.json()["detail"]
    assert client.get("/api/v1/advisor/status").json()["quota"]["compare"]["used"] == 0
    assert client.get("/api/v1/advisor/tasks").json()["items"] == []


def test_upload_limits_exclusive_inputs_and_chunked_body(client):
    agree_advisor(client)
    body = upload_body(b"\n".join([b"company"] * 51))
    assert (
        client.post(
            "/api/v1/advisor/tasks", headers={"Idempotency-Key": "51"}, json=body
        ).status_code
        == 429
    )
    assert (
        client.post(
            "/api/v1/advisor/tasks",
            headers={"Idempotency-Key": "both"},
            json=body | {"rows": ["x"]},
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/v1/advisor/tasks", headers={"Idempotency-Key": "none"}, json={"kind": "compare"}
        ).status_code
        == 422
    )
    content = b'{"kind":"compare","rows":["' + b"x" * service.MAX_BODY_BYTES + b'"]}'
    response = client.post(
        "/api/v1/advisor/tasks",
        headers={"Idempotency-Key": "chunked", "content-type": "application/json"},
        content=iter([content]),
    )
    assert response.status_code == 413


def test_upload_exact_byte_boundary_and_xlsx_expansion(client):
    agree_advisor(client)
    boundary = upload_body(b"a" * service.MAX_UPLOAD_BYTES)
    assert client.post("/api/v1/advisor/preview", json=boundary).json() == {"rowCount": 1}
    assert (
        client.post(
            "/api/v1/advisor/tasks", headers={"Idempotency-Key": "boundary"}, json=boundary
        ).status_code
        == 202
    )
    archive = BytesIO()
    with ZipFile(archive, "w", compression=8) as book:
        book.writestr("xl/worksheets/sheet1.xml", b" " * (2 * 1024 * 1024 + 1))
    response = client.post("/api/v1/advisor/preview", json=upload_body(archive.getvalue()))
    assert response.status_code == 422 and "解压后过大" in response.json()["detail"]


@pytest.mark.parametrize("sparse", [False, True])
def test_xlsx_untrusted_dimensions_and_sparse_rows(client, sparse):
    agree_advisor(client)
    source = xlsx_content([["企业名称"], ["甲公司"]])
    archive = BytesIO()
    with ZipFile(BytesIO(source)) as old, ZipFile(archive, "w", compression=8) as new:
        for name in old.namelist():
            content = old.read(name)
            if name == "xl/worksheets/sheet1.xml":
                content = content.replace(b'ref="A1:A2"', b'ref="A1:XFD1048576"')
                if sparse:
                    content = content.replace(b'<row r="2">', b'<row r="1000000">').replace(
                        b'r="A2"', b'r="A1000000"'
                    )
            new.writestr(name, content)
    response = client.post("/api/v1/advisor/preview", json=upload_body(archive.getvalue()))
    if sparse:
        assert response.status_code == 422 and "行号过大" in response.json()["detail"]
    else:
        assert response.json() == {"rowCount": 1}


def create(client, key="k1", rows=None):
    return client.post(
        "/api/v1/advisor/tasks",
        headers={"Idempotency-Key": key},
        json={"kind": "compare", "rows": rows or ["甲公司"]},
    )


def task(session, user_id, **values):
    status = values.pop("status", "queued")
    row = AdvisorTask(
        user_id=user_id,
        kind=values.pop("kind", "compare"),
        status=status,
        input_rows=["客户秘密"],
        row_count=1,
        idempotency_key=values.pop("idempotency_key", f"direct-{user_id}"),
        created_at=values.pop("created_at", NOW.astimezone(UTC)),
        expires_at=values.pop("expires_at", NOW.astimezone(UTC) + timedelta(days=30)),
        **values,
    )
    session.add(row)
    session.commit()
    return row


class FakeResponse:
    status_code = 200

    def __init__(self, body):
        self._body = body

    def json(self):
        return self._body

    def raise_for_status(self):
        return None


class FakeObjectStore:
    def __init__(self):
        self.calls = []

    def post(self, url, **kwargs):
        self.calls.append((url, kwargs))
        if url.endswith("/tcb/uploadfile"):
            return FakeResponse(
                {
                    "errcode": 0,
                    "url": "https://cos.test/upload",
                    "authorization": "sig",
                    "token": "token",
                    "cos_file_id": "cos-file",
                    "file_id": "cloud://advisor-result",
                }
            )
        if url.endswith("/tcb/batchdownloadfile"):
            # 按微信接口的真实格式：请求里每项是 {fileid, max_age}，返回的 file_list 每项带 fileid
            asked = kwargs["json"]["file_list"]
            assert all(set(item) == {"fileid", "max_age"} for item in asked)
            return FakeResponse(
                {
                    "errcode": 0,
                    "file_list": [
                        {
                            "fileid": asked[0]["fileid"],
                            "status": 0,
                            "download_url": "https://download",
                        }
                    ],
                }
            )
        return FakeResponse({"errcode": 0})


def test_switch_idempotency_quota_rollover_and_cancel(
    client, session, advisor_settings, monkeypatch
):
    advisor_settings.advisor_enabled = False
    assert client.get("/api/v1/advisor/tasks").status_code == 404
    advisor_settings.advisor_enabled = True
    agree_advisor(client)
    assert create(client, key="missing").status_code == 202
    duplicate = create(client, key="missing", rows=["不同内容"])
    assert duplicate.status_code == 202
    assert duplicate.json()["id"] == 1 and duplicate.json()["rowCount"] == 1
    assert (
        client.post("/api/v1/advisor/tasks", json={"kind": "compare", "rows": ["x"]}).status_code
        == 422
    )
    assert create(client, key="second").status_code == 202
    assert create(client, key="third").status_code == 202
    blocked = create(client, key="fourth")
    assert blocked.status_code == 429
    assert "本月" in blocked.json()["detail"]["message"]

    cancelled = client.post("/api/v1/advisor/tasks/1/cancel")
    assert cancelled.status_code == 200 and cancelled.json()["status"] == "cancelled"
    assert create(client, key="after-cancel").status_code == 202

    # 处理失败的任务不算用掉额度。
    failed = session.get(AdvisorTask, 4)
    failed.status = "failed"
    service.release_usage(session, failed)
    session.commit()
    assert create(client, key="after-failed").status_code == 202

    # 北京时间自然月切换后恢复三次免费额度。
    monkeypatch.setattr(v1, "now", lambda: datetime(2026, 10, 1, 0, 1, tzinfo=BJ))
    assert create(client, key="next-month").status_code == 202


def test_row_and_body_limits_and_personal_data(client):
    agree_advisor(client)
    assert create(client, key="too-many", rows=["x"] * 51).status_code == 429
    huge = "x" * (service.MAX_BODY_BYTES + 1)
    response = create(client, key="huge", rows=[huge])
    assert response.status_code == 413
    data = client.get("/api/v1/me/data").json()
    assert data["advisorTasksTotal"] == 0
    assert data["advisorTasks"] == []
    created = create(client, key="private").json()
    data = client.get("/api/v1/me/data").json()
    assert data["advisorTasksTotal"] == 1
    assert "inputRows" not in str(data) and "客户秘密" not in str(data)
    assert data["advisorTasks"][0]["id"] == created["id"]


def test_task_ownership_for_read_cancel_delete_download(
    client, advisor_settings, session, monkeypatch
):
    agree_advisor(client)
    created = create(client, key="owned").json()
    task_id = created["id"]
    user = session.scalar(select(User).where(User.openid == "advisor-user"))
    row = session.get(AdvisorTask, task_id)
    row.status, row.result_file_id = "done", "cloud://owned"
    session.commit()
    deleted = []
    monkeypatch.setattr(
        service, "delete_result_files", lambda client, env, ids: deleted.extend(ids) or 1
    )

    advisor_settings.dev_openid = "other-user"
    agree_advisor(client)
    assert client.get(f"/api/v1/advisor/tasks/{task_id}").status_code == 404
    assert client.post(f"/api/v1/advisor/tasks/{task_id}/cancel").status_code == 404
    assert client.get(f"/api/v1/advisor/tasks/{task_id}/download").status_code == 404
    assert client.delete(f"/api/v1/advisor/tasks/{task_id}").status_code == 404

    advisor_settings.dev_openid = "advisor-user"
    monkeypatch.setattr(service, "download_url", lambda client, env, file_id: "https://download")
    row.expires_at = NOW.astimezone(UTC) + timedelta(minutes=1)
    session.commit()
    # 结果文件由接口自己分段读出来发给小程序，不把对象存储的下载地址交出去
    content = bytes(range(256)) * 4000  # 1 024 000 字节：分两段
    monkeypatch.setattr(service, "result_size", lambda client, url: len(content))
    monkeypatch.setattr(
        service,
        "read_result_part",
        lambda client, url, index, total: (
            content[index * service.RESULT_PART_SIZE : (index + 1) * service.RESULT_PART_SIZE]
            or (_ for _ in ()).throw(ValueError("没有这一段"))
        ),
    )
    downloaded = client.get(f"/api/v1/advisor/tasks/{task_id}/download")
    assert downloaded.status_code == 200
    assert downloaded.json() == {
        "size": len(content),
        "parts": 2,
        "partSize": service.RESULT_PART_SIZE,
    }
    assert "url" not in downloaded.text and "https://" not in downloaded.text
    parts = [client.get(f"/api/v1/advisor/tasks/{task_id}/file?part={n}") for n in (0, 1)]
    assert [p.status_code for p in parts] == [200, 200]
    assert b"".join(p.content for p in parts) == content
    assert all(len(p.content) <= service.RESULT_PART_SIZE for p in parts)
    assert client.get(f"/api/v1/advisor/tasks/{task_id}/file?part=2").status_code == 404
    # 别人取不到；api 服务没配存储环境时，从文件编号里取环境（线上 10/10 的情况）
    advisor_settings.dev_openid = "other-user"
    assert client.get(f"/api/v1/advisor/tasks/{task_id}/file?part=0").status_code == 404
    advisor_settings.dev_openid = "advisor-user"
    assert service.result_env(SimpleNamespace(cloud_env=None), "cloud://prod-x.bucket/a.xlsx") == (
        "prod-x"
    )
    assert service.result_env(SimpleNamespace(cloud_env="configured"), "cloud://prod-x.b/a") == (
        "configured"
    )
    assert service.result_env(SimpleNamespace(cloud_env=None), "not-a-file-id") is None
    row.expires_at = NOW.astimezone(UTC)
    session.commit()
    assert client.get(f"/api/v1/advisor/tasks/{task_id}/download").status_code == 410
    assert client.get(f"/api/v1/advisor/tasks/{task_id}/file?part=0").status_code == 410
    assert client.delete(f"/api/v1/advisor/tasks/{task_id}").status_code == 204
    assert deleted == ["cloud://owned"]
    session.expire_all()
    assert session.get(AdvisorTask, task_id) is None
    assert user is not None


def test_queue_state_stale_recovery_and_retry(session, advisor_settings, monkeypatch):
    user = User(openid="queue-user", agreement_version="2026-09", agreement_at=NOW)
    session.add(user)
    session.commit()
    queued = task(session, user.id)
    monkeypatch.setattr(
        service, "_run_engine", lambda session, task, output: output.write_bytes(b"x") or 1
    )
    monkeypatch.setattr(service, "upload", lambda *args, **kwargs: "cloud://done")
    report = service.run_advisor_tasks(
        session, settings=advisor_settings, stamp=NOW.astimezone(UTC), max_tasks=1
    )
    session.refresh(queued)
    assert report["done"] == 1 and queued.status == "done" and queued.attempts == 1

    stale = task(
        session,
        user.id,
        idempotency_key="stale",
        status="running",
        started_at=(NOW - timedelta(minutes=11)).astimezone(UTC),
    )
    assert service.recover_stale(session, NOW.astimezone(UTC)) == 1
    session.refresh(stale)
    assert stale.status == "queued" and stale.started_at is None
    stale.status = "cancelled"
    session.commit()

    failing = task(session, user.id, idempotency_key="failing")
    monkeypatch.setattr(
        service, "_run_engine", lambda *args: (_ for _ in ()).throw(RuntimeError("boom"))
    )
    for attempts in (1, 2):
        service.run_advisor_tasks(
            session, settings=advisor_settings, stamp=NOW.astimezone(UTC), max_tasks=1
        )
        session.refresh(failing)
        assert failing.attempts == attempts and failing.status == "queued"
    service.run_advisor_tasks(
        session, settings=advisor_settings, stamp=NOW.astimezone(UTC), max_tasks=1
    )
    session.refresh(failing)
    assert failing.attempts == 3 and failing.status == "failed"


def test_data_tool_worker_stays_serial(session, advisor_settings, monkeypatch):
    user = User(openid="serial-data", agreement_version="2026-09", agreement_at=NOW)
    session.add(user)
    session.commit()
    tasks = [task(session, user.id, idempotency_key=kind) for kind in ("compare", "track")]
    tasks[1].kind = "track"
    session.commit()
    seen = []

    def process(work, row, **kwargs):
        assert work is session
        seen.append(row.id)
        row.status, row.finished_at = "done", NOW
        work.commit()
        return True

    monkeypatch.setattr(service, "process_task", process)
    report = service.run_advisor_tasks(session, settings=advisor_settings, stamp=NOW)
    assert seen == [row.id for row in tasks]
    assert report == {"processed": 2, "done": 2, "failed": 0, "requeued": 0}


def test_fake_object_store_upload_download_and_expiry(session, advisor_settings, monkeypatch):
    user = User(openid="file-user", agreement_version="2026-09", agreement_at=NOW)
    session.add(user)
    session.commit()
    row = task(session, user.id, idempotency_key="file")
    monkeypatch.setattr(
        service, "_run_engine", lambda session, task, output: output.write_bytes(b"x") or 2
    )
    store = FakeObjectStore()
    assert (
        service.process_task(session, row, settings=advisor_settings, client=store, stamp=NOW)
        is True
    )
    session.refresh(row)
    assert row.status == "done" and row.result_file_id == "cloud://advisor-result"
    assert service.download_url(store, "test-env", row.result_file_id) == "https://download"

    row.expires_at = NOW.astimezone(UTC)
    session.commit()
    assert row.expires_at <= NOW.astimezone(UTC)


def test_cleanup_ignores_switch_and_deletes_expired_files(session, advisor_settings):
    user = User(openid="cleanup-user")
    session.add(user)
    session.commit()
    old = task(
        session,
        user.id,
        idempotency_key="old",
        created_at=(NOW - timedelta(days=30)).astimezone(UTC),
        expires_at=NOW.astimezone(UTC),
        result_file_id="cloud://old",
    )
    advisor_settings.advisor_enabled = False
    deleted = []
    client = SimpleNamespace()
    monkeypatch = pytest.MonkeyPatch()
    monkeypatch.setattr(
        service, "delete_result_files", lambda client, env, ids: deleted.extend(ids) or len(ids)
    )
    result = service.cleanup_advisor(
        session,
        settings=advisor_settings,
        stamp=NOW.astimezone(UTC),
        client=client,
    )
    monkeypatch.undo()
    assert result["advisor_input_rows"] == 1 and result["advisor_files"] == 1
    session.refresh(old)
    assert old.input_rows is None and old.result_file_id is None and deleted == ["cloud://old"]


def test_cleanup_deletes_expired_ai_tasks_and_old_usage(session, advisor_settings):
    """AI 任务过了 30 天结果保存期整行删除；使用次数只留本月和上个月。"""
    user = User(openid="retention-user")
    session.add(user)
    session.commit()
    stamp = NOW.astimezone(UTC)
    expired = task(
        session,
        user.id,
        kind="interpret",
        idempotency_key="ai-expired",
        created_at=stamp - timedelta(days=31),
        expires_at=stamp - timedelta(days=1),
    )
    fresh = task(
        session,
        user.id,
        kind="clause",
        idempotency_key="ai-fresh",
        created_at=stamp - timedelta(days=1),
        expires_at=stamp + timedelta(days=29),
    )
    data_tool = task(
        session,
        user.id,
        idempotency_key="data-expired",
        created_at=stamp - timedelta(days=31),
        expires_at=stamp - timedelta(days=1),
    )
    today = service.usage_day(stamp)
    this_month = today.replace(day=1)
    last_month = (this_month - timedelta(days=1)).replace(day=1)
    two_months_ago = (last_month - timedelta(days=1)).replace(day=1)
    for day in (today, last_month, two_months_ago):
        session.add(AdvisorUsage(user_id=user.id, kind="interpret", day=day, used_count=1))
    session.commit()
    expired_id, fresh_id, data_id = expired.id, fresh.id, data_tool.id
    result = service.cleanup_advisor(session, settings=advisor_settings, stamp=stamp)
    assert result["advisor_ai_tasks"] == 1 and result["advisor_usage_rows"] == 1
    session.expire_all()
    assert session.get(AdvisorTask, expired_id) is None
    assert session.get(AdvisorTask, fresh_id) is not None
    assert session.get(AdvisorTask, data_id) is not None
    days = sorted(session.scalars(select(AdvisorUsage.day)))
    assert days == [last_month, today]


def test_delete_account_removes_advisor_tasks_and_files(
    client, advisor_settings, session, monkeypatch
):
    agree_advisor(client)
    created = create(client, key="logout").json()
    row = session.get(AdvisorTask, created["id"])
    row.status, row.result_file_id = "done", "cloud://logout"
    session.commit()
    deleted = []
    monkeypatch.setattr(
        service, "delete_result_files", lambda client, env, ids: deleted.extend(ids) or 1
    )
    assert client.delete("/api/v1/me").status_code == 204
    assert deleted == ["cloud://logout"]
    session.expire_all()
    assert session.get(AdvisorTask, created["id"]) is None


def test_delete_terminal_task_keeps_usage_and_queued_delete_refunds(
    client, advisor_settings, session
):
    agree_advisor(client)
    created = create(client, key="ledger").json()
    usage = session.scalar(select(AdvisorUsage))
    assert usage is not None and usage.used_count == 1
    queued_delete = client.delete(f"/api/v1/advisor/tasks/{created['id']}")
    # 排队中的任务一步取消并删除，额度退回
    assert queued_delete.status_code == 204
    session.expire_all()
    assert session.get(AdvisorTask, created["id"]) is None
    assert session.scalar(select(AdvisorUsage.used_count)) == 0

    created = create(client, key="ledger-done").json()
    row = session.get(AdvisorTask, created["id"])
    row.status = "done"
    session.commit()
    assert client.delete(f"/api/v1/advisor/tasks/{created['id']}").status_code == 204
    assert session.scalar(select(AdvisorUsage.used_count)) == 1
    assert create(client, key="ledger-second").status_code == 202
    assert create(client, key="ledger-third").status_code == 202
    assert create(client, key="ledger-blocked").status_code == 429


def test_storage_failure_is_queued_and_retried(session, advisor_settings, monkeypatch):
    user = User(openid="trash-user", agreement_version="2026-09", agreement_at=NOW)
    session.add(user)
    session.commit()
    row = task(
        session,
        user.id,
        idempotency_key="trash",
        expires_at=NOW.astimezone(UTC),
        result_file_id="cloud://trash",
    )
    attempts = iter([RuntimeError("暂时失败"), ["cloud://trash"]])
    monkeypatch.setattr(service, "delete_result_files", lambda *args: next(attempts))
    first = service.cleanup_advisor(
        session, settings=advisor_settings, stamp=NOW.astimezone(UTC), client=SimpleNamespace()
    )
    assert first["advisor_files"] == 0
    session.refresh(row)
    assert row.result_file_id == "cloud://trash"
    assert session.scalar(select(AdvisorFileTrash.file_id)) == "cloud://trash"

    second = service.cleanup_advisor(
        session,
        settings=advisor_settings,
        stamp=(NOW + timedelta(days=1)).astimezone(UTC),
        client=SimpleNamespace(),
    )
    assert second["advisor_files"] == 1
    session.refresh(row)
    assert row.result_file_id is None
    assert session.scalar(select(AdvisorFileTrash.file_id)) is None


def test_delete_and_logout_succeed_when_storage_delete_fails(
    client, advisor_settings, session, monkeypatch
):
    agree_advisor(client)
    created = create(client, key="api-trash").json()
    row = session.get(AdvisorTask, created["id"])
    row.status, row.result_file_id = "done", "cloud://api-trash"
    session.commit()
    monkeypatch.setattr(
        service, "delete_result_files", lambda *args: (_ for _ in ()).throw(RuntimeError("down"))
    )
    assert client.delete(f"/api/v1/advisor/tasks/{created['id']}").status_code == 204
    assert session.scalar(select(AdvisorFileTrash.file_id)) == "cloud://api-trash"

    created = client.post(
        "/api/v1/advisor/tasks",
        headers={"Idempotency-Key": "logout-trash"},
        json={"kind": "track", "rows": [{"name": "客户"}]},
    ).json()
    row = session.get(AdvisorTask, created["id"])
    row.status, row.result_file_id = "done", "cloud://logout-trash"
    session.commit()
    assert client.delete("/api/v1/me").status_code == 204
    assert (
        session.scalar(
            select(AdvisorFileTrash.file_id).where(
                AdvisorFileTrash.file_id == "cloud://logout-trash"
            )
        )
        == "cloud://logout-trash"
    )


def test_uploaded_file_is_deleted_or_queued_when_task_times_out(
    session, advisor_settings, monkeypatch
):
    user = User(openid="timeout-user", agreement_version="2026-09", agreement_at=NOW)
    session.add(user)
    session.commit()
    row = task(session, user.id, idempotency_key="timeout", status="running")
    monkeypatch.setattr(
        service, "_run_engine", lambda session, task, output: output.write_bytes(b"x") or 1
    )
    monkeypatch.setattr(service, "upload", lambda *args, **kwargs: "cloud://timeout")
    monkeypatch.setattr(
        service, "delete_result_files", lambda *args: (_ for _ in ()).throw(RuntimeError("down"))
    )
    clock = iter((0.0, 0.0, 2.0))
    monkeypatch.setattr(service.monotonic_clock, "monotonic", lambda: next(clock))
    assert (
        service.process_task(
            session,
            row,
            settings=advisor_settings,
            client=SimpleNamespace(),
            stamp=NOW.astimezone(UTC),
            timeout=timedelta(seconds=1),
        )
        is False
    )
    assert session.scalar(select(AdvisorFileTrash.file_id)) == "cloud://timeout"


def relocate_sheet(source, transform=lambda data: data):
    """保留有效工作簿关系，将工作表移到无 .xml 后缀的非标准路径。"""
    output = BytesIO()
    old_path = "xl/worksheets/sheet1.xml"
    new_path = "custom/customers.data"
    with ZipFile(BytesIO(source)) as old, ZipFile(output, "w", compression=8) as new:
        for name in old.namelist():
            data = old.read(name)
            if name == old_path:
                name, data = new_path, transform(data)
            elif name in ("xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
                data = data.replace(("/" + old_path).encode(), ("/" + new_path).encode())
            new.writestr(name, data)
    return output.getvalue()


def test_upload_nonstandard_worksheet_path_and_first_sheet(client):
    agree_advisor(client)
    book = Workbook()
    book.active.append(["企业名称"])
    book.active.append(["首表公司"])
    book.create_sheet("当前表").append(["其他公司"])
    book.active = 1
    output = BytesIO()
    book.save(output)
    body = upload_body(relocate_sheet(output.getvalue()))
    assert client.post("/api/v1/advisor/preview", json=body).json() == {"rowCount": 1}
    response = client.post("/api/v1/advisor/tasks", headers={"Idempotency-Key": "first"}, json=body)
    assert response.status_code == 202
    assert client.get(f"/api/v1/advisor/tasks/{response.json()['id']}").json()["rowCount"] == 1


@pytest.mark.parametrize(
    "payload,reason",
    [
        (b'<row r="1000000"><c r="A1000000"/></row>', "行号过大"),
        (b'<row r="1"><c r="A1000000"/></row>', "行号过大"),
        (b'<row r="1"><c r="BM1"/></row>', "列数过多"),
        (b'<row r="1"><c r="A1"/></row>' * 503, "行号过大"),
        (b'<row r="1">' + b'<c r="A1"/>' * 32129 + b"</row>", "单元格过多"),
        (b"<x/>" * 100001, "内容过多"),
    ],
    ids=["row-index", "cell-index", "columns", "rows", "cells", "elements"],
)
@pytest.mark.parametrize("kind", ["compare", "track"])
def test_nonstandard_worksheet_limits(client, payload, reason, kind):
    agree_advisor(client)
    sheet = (
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        + payload
        + b"</sheetData></worksheet>"
    )
    content = relocate_sheet(xlsx_content([["企业名称"], ["甲公司"]]), lambda _: sheet)
    response = client.post("/api/v1/advisor/preview", json=upload_body(content, kind=kind))
    assert response.status_code == 422 and reason in response.json()["detail"]
    assert client.get("/api/v1/advisor/status").json()["quota"][kind]["used"] == 0


def test_xlsx_cumulative_expansion_limit(client):
    agree_advisor(client)
    output = BytesIO()
    with ZipFile(output, "w", compression=8) as archive:
        archive.writestr("a", b"x" * (1024 * 1024 + 1))
        archive.writestr("b", b"x" * (1024 * 1024 + 1))
    response = client.post("/api/v1/advisor/preview", json=upload_body(output.getvalue()))
    assert response.status_code == 422 and "解压后过大" in response.json()["detail"]


@pytest.mark.parametrize("content", [b"PK company\nOther company", b"PKbroken"])
def test_pk_prefix_csv_is_not_zip(client, content):
    agree_advisor(client)
    response = client.post("/api/v1/advisor/preview", json=upload_body(content))
    assert response.status_code == 200
    assert response.json()["rowCount"] == len(content.splitlines())


@pytest.mark.parametrize("disabled,outdated", [(True, False), (False, True), (True, True)])
def test_delete_requires_authorization_but_logout_ignores_switch_and_agreement(
    client, advisor_settings, session, disabled, outdated
):
    agree_advisor(client)
    created = create(client).json()
    advisor_settings.advisor_enabled = not disabled
    if outdated:
        advisor_settings.agreement_version = "new-general-version"
        user = session.scalar(select(User))
        user.advisor_agreement_version = "old-advisor-version"
        session.commit()
    # 我的数据中的删除与我的任务使用同一个删除接口。
    data = client.get("/api/v1/me/data")
    assert data.status_code == 200
    assert data.json()["advisorTasks"][0]["id"] == created["id"]
    advisor_settings.dev_openid = "somebody-else"
    expected = 404 if disabled else 403
    assert client.delete(f"/api/v1/advisor/tasks/{created['id']}").status_code == expected
    advisor_settings.dev_openid = "advisor-user"
    assert client.delete(f"/api/v1/advisor/tasks/{created['id']}").status_code == expected
    assert client.get("/api/v1/me/data").json()["advisorTasks"][0]["id"] == created["id"]
    assert client.delete("/api/v1/me").status_code == 204
    session.expire_all()
    assert session.get(AdvisorTask, created["id"]) is None
    assert client.get("/api/v1/me").json()["agreed"] is None


def test_xlsx_cell_limit_across_worksheets(client):
    agree_advisor(client)
    book = Workbook()
    book.create_sheet("第二表")
    output = BytesIO()
    book.save(output)
    sheet = (
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1">'
        + b'<c r="A1"/>' * 16100
        + b"</row></sheetData></worksheet>"
    )
    result = BytesIO()
    with ZipFile(BytesIO(output.getvalue())) as old, ZipFile(result, "w", compression=8) as new:
        for name in old.namelist():
            new.writestr(name, sheet if name.startswith("xl/worksheets/") else old.read(name))
    response = client.post("/api/v1/advisor/preview", json=upload_body(result.getvalue()))
    assert response.status_code == 422 and "单元格过多" in response.json()["detail"]


def test_result_file_is_read_from_storage_in_ranges():
    """结果文件按段从对象存储读：先用 Range: bytes=0-0 拿总长度，再按 900 KiB 一段读。"""
    content = bytes(range(256)) * 4000
    seen = []

    def handler(request: httpx.Request) -> httpx.Response:
        start, end = (int(v) for v in request.headers["range"].removeprefix("bytes=").split("-"))
        seen.append((start, end))
        return httpx.Response(
            206,
            content=content[start : end + 1],
            headers={"content-range": f"bytes {start}-{end}/{len(content)}"},
        )

    with httpx.Client(transport=httpx.MockTransport(handler)) as http:
        assert service.result_size(http, "https://cos.test/a.xlsx") == len(content)
        first = service.read_result_part(http, "https://cos.test/a.xlsx", 0, len(content))
        second = service.read_result_part(http, "https://cos.test/a.xlsx", 1, len(content))
        with pytest.raises(ValueError):
            service.read_result_part(http, "https://cos.test/a.xlsx", 2, len(content))
    assert first + second == content and len(first) == service.RESULT_PART_SIZE
    assert seen[0] == (0, 0) and seen[1] == (0, service.RESULT_PART_SIZE - 1)

    too_big = httpx.MockTransport(
        lambda request: httpx.Response(
            206, content=b"x", headers={"content-range": "bytes 0-0/99999999"}
        )
    )
    with httpx.Client(transport=too_big) as http, pytest.raises(service.BackupError):
        service.result_size(http, "https://cos.test/big.xlsx")


def test_download_url_handles_missing_file_and_bad_links():
    """文件不存在（status 非 0）、返回的不是 https 链接、返回别的文件：都当作取不到。"""

    class Store:
        def __init__(self, items):
            self.items = items

        def post(self, url, **kwargs):
            return FakeResponse({"errcode": 0, "file_list": self.items})

    file_id = "cloud://env.bucket/advisor/1.xlsx"
    ok = [{"fileid": file_id, "status": 0, "download_url": "https://cos.test/1.xlsx"}]
    assert service.download_url(Store(ok), "env", file_id) == "https://cos.test/1.xlsx"
    gone = [{"fileid": file_id, "status": -503002, "download_url": ""}]
    assert service.download_url(Store(gone), "env", file_id) is None
    plain = [{"fileid": file_id, "status": 0, "download_url": "http://cos.test/1.xlsx"}]
    assert service.download_url(Store(plain), "env", file_id) is None
    other = [
        {"fileid": "cloud://env.bucket/advisor/2.xlsx", "status": 0, "download_url": "https://x"}
    ]
    assert service.download_url(Store(other), "env", file_id) is None
    assert service.download_url(Store([]), "env", file_id) is None


@pytest.mark.parametrize(
    "kind,rows",
    [
        ("compare", ["合成城市科技有限公司"]),
        ("track", [{"name": "合成城市科技有限公司", "province": "山东", "city": "青岛"}]),
    ],
)
def test_pasted_rows_require_independent_tool_agreement(client, kind, rows):
    agree(client)
    body = {"kind": kind, "rows": rows}
    response = client.post(
        "/api/v1/advisor/tasks", json=body, headers={"Idempotency-Key": "paste-consent"}
    )
    assert response.status_code == 403
    assert "工具使用说明" in response.json()["detail"]
    assert client.get("/api/v1/advisor/status").json()["quota"][kind]["used"] == 0
    agree_advisor(client)
    accepted = client.post(
        "/api/v1/advisor/tasks", json=body, headers={"Idempotency-Key": "paste-consent"}
    )
    assert accepted.status_code == 202
    assert client.get("/api/v1/advisor/status").json()["quota"][kind]["used"] == 1
