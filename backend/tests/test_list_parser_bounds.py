"""只使用内存合成文件；资源预算不能靠真的耗尽机器内存来验证。"""

from io import BytesIO

import pytest
from openpyxl import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from youju.lists import parser


def _xlsx(workbook):
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_sparse_xlsx_rejected_before_expanding_rows(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["企业名称"])
    workbook.active.append(["示例科技有限公司"])
    workbook.active.cell(row=100_000, column=16_384, value="末尾")
    payload = _xlsx(workbook)
    assert len(payload) < 6_000

    def unexpected_rows(*args, **kwargs):
        pytest.fail("稀疏声明范围必须在展开空行之前拒绝")

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", unexpected_rows)
    with pytest.raises(ValueError, match="暂不支持.*安全解析预算"):
        parser.parse_attachment_result(payload, "xlsx")


@pytest.mark.parametrize(
    ("setting", "limit", "rows"),
    [
        ("_MAX_ROWS", 2, [["企业名称"], ["示例科技有限公司"], ["另一科技有限公司"]]),
        ("_MAX_COLUMNS", 1, [["企业名称", "序号"]]),
        ("_MAX_CELLS", 3, [["企业名称", "序号"], ["示例科技有限公司", "1"]]),
        ("_MAX_TEXT_CHARS", 8, [["企业名称"], ["示例科技有限公司"]]),
    ],
)
def test_budget_rejects_whole_input_instead_of_returning_partial(monkeypatch, setting, limit, rows):
    monkeypatch.setattr(parser, setting, limit)
    with pytest.raises(ValueError, match="暂不支持.*安全解析预算"):
        parser.parse_table_rows(iter(rows))


def test_budget_accepts_exact_boundary(monkeypatch):
    rows = [["企业名称"], ["示例科技有限公司"]]
    monkeypatch.setattr(parser, "_MAX_ROWS", 2)
    monkeypatch.setattr(parser, "_MAX_CELLS", 2)
    monkeypatch.setattr(parser, "_MAX_COLUMNS", 1)
    monkeypatch.setattr(parser, "_MAX_TEXT_CHARS", sum(len(row[0]) for row in rows))
    assert [entry.name for entry in parser.parse_table_rows(rows)] == ["示例科技有限公司"]


def test_xlsx_budget_shared_between_sheets(monkeypatch):
    workbook = Workbook()
    for sheet in (workbook.active, workbook.create_sheet()):
        sheet.append(["企业名称"])
        sheet.append(["示例科技有限公司"])
    monkeypatch.setattr(parser, "_MAX_CELLS", 3)
    with pytest.raises(ValueError, match="安全解析预算"):
        parser.parse_attachment_result(_xlsx(workbook), "xlsx")


def test_html_budget_shared_between_tables(monkeypatch):
    table = "<table><tr><td>企业名称</td></tr><tr><td>示例科技有限公司</td></tr></table>"
    monkeypatch.setattr(parser, "_MAX_ROWS", 3)
    with pytest.raises(ValueError, match="安全解析预算"):
        parser.parse_html_tables("<html>" + table * 2 + "</html>")


@pytest.mark.parametrize("failure", [None, "budget", "parser"])
def test_pdf_releases_each_page_before_next_and_on_failure(monkeypatch, failure):
    closed = []
    visited = []

    class Page:
        def __init__(self, number):
            self.number = number

        def extract_text(self):
            return "企业名称"

        def extract_tables(self, table_settings=None):
            assert closed == list(range(self.number))
            visited.append(self.number)
            if failure == "parser" and self.number == 1:
                raise RuntimeError("synthetic parser failure")
            return [[["企业名称"], ["示例科技有限公司"]]]

        def close(self):
            closed.append(self.number)

    class Pdf:
        pages = [Page(0), Page(1), Page(2)]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

    monkeypatch.setattr(parser.pdfplumber, "open", lambda _: Pdf())
    if failure == "budget":
        monkeypatch.setattr(parser, "_MAX_CELLS", 3)
        with pytest.raises(ValueError, match="安全解析预算"):
            parser.parse_attachment_result(b"%PDF-fake", "pdf")
    elif failure == "parser":
        with pytest.raises(RuntimeError, match="synthetic"):
            parser.parse_attachment_result(b"%PDF-fake", "pdf")
    else:
        assert len(parser.parse_attachment_result(b"%PDF-fake", "pdf").entries) == 1
    assert closed == visited == ([0, 1] if failure else [0, 1, 2])


def test_oversize_source_is_recorded_without_aborting_batch(monkeypatch):
    from types import SimpleNamespace

    from youju.lists import service

    workbook = Workbook()
    workbook.active.cell(row=100_000, column=16_384, value="末尾")
    response = SimpleNamespace(status="ok", final_url=None, content=_xlsx(workbook), text="")
    fetcher = SimpleNamespace(fetch=lambda _: response)
    recorded = []
    monkeypatch.setattr(service, "_government_url", lambda _: True)
    monkeypatch.setattr(service, "_save", lambda *args, **kwargs: recorded.append(args))
    service._parse_source(
        None, None, "https://test.gov.cn/list.xlsx", "xlsx", "attachment", None, fetcher
    )
    assert len(recorded) == 1
    assert recorded[0][-2] is None
    assert recorded[0][-1].startswith("暂不支持超过安全解析预算")


def test_dimensionless_xlsx_is_rejected_before_unbounded_row_allocation(monkeypatch):
    import re
    import zipfile

    workbook = Workbook()
    workbook.active.append(["企业名称"])
    output = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(_xlsx(workbook))) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for name in source.namelist():
            data = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"<dimension[^>]*/>", b"", data)
            target.writestr(name, data)
    monkeypatch.setattr(
        ReadOnlyWorksheet, "iter_rows", lambda *args, **kwargs: pytest.fail("必须先拒绝")
    )
    with pytest.raises(ValueError, match="暂不支持未声明行列范围"):
        parser.parse_attachment_result(output.getvalue(), "xlsx")
