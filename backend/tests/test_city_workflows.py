"""省市从客户输入/来源到真实查询及 Excel 的合成数据回归，不读取生产数据。"""

import base64
from datetime import UTC, date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from youju.api import advisor, me, v1
from youju.compare import compare_lists, export_compare_excel
from youju.models import Document, DocumentDeadline, ListEntry, Source
from youju.track import export_track_excel, read_customer_profiles, track_clients


@pytest.fixture
def city_docs(session, monkeypatch):
    stamp = datetime(2026, 10, 15, 10, tzinfo=UTC)
    monkeypatch.setattr(v1, "now", lambda: stamp)
    docs = {}
    for index, (province, city, level) in enumerate(
        [
            ("国家", None, "国家级"),
            ("山东", None, "省级"),
            ("山东", "青岛", "市级"),
            ("山东", "济南", "市级"),
            ("江苏", "苏州", "市级"),
        ]
    ):
        source = Source(
            level=level,
            province=province,
            city=city,
            dept="科技",
            dept_full_name=f"合成{city or province}科技部门",
            column_name="通知",
            url=f"https://example.gov.cn/city-source/{index}",
            rights_level=1,
            last_success_at=stamp,
        )
        session.add(source)
        session.flush()
        doc = Document(
            source_id=source.id,
            title=f"关于组织申报合成科技项目的通知{index}",
            url=f"https://example.gov.cn/city-doc/{index}",
            url_key=f"city/{index}",
            publish_date=date(2026, 10, 14),
            first_seen_at=stamp,
            last_seen_at=stamp,
            kind_override="申报通知",
        )
        session.add(doc)
        session.flush()
        session.add(
            DocumentDeadline(
                document_id=doc.id,
                fingerprint=f"city-deadline-{index}",
                kind="企业申报截止",
                candidate_date=date(2026, 10, 20),
                approved_date=date(2026, 10, 20),
                evidence="合成截止",
                confidence="高",
                status="approved",
            )
        )
        docs[city or province] = doc
    session.commit()
    return docs


@pytest.mark.parametrize("suffix", ["csv", "xlsx"])
def test_track_city_file_api_query_and_export(session, city_docs, tmp_path, suffix):
    path = tmp_path / f"synthetic.{suffix}"
    headers = ["客户名称", "省份", "关键词", "支持方式", "城市"]
    rows = [
        ["合成青岛科技有限公司", "山东省", "", "", "青岛市"],
        ["合成省级科技有限公司", "山东", "", "", ""],
        ["合成未收录科技有限公司", "山东", "", "", "苏州"],
    ]
    if suffix == "xlsx":
        book = Workbook()
        book.active.append(headers)
        for row in rows:
            book.active.append(row)
        book.save(path)
    else:
        path.write_text("\n".join(",".join(row) for row in [headers, *rows]), encoding="utf-8")
    profiles = read_customer_profiles(path)
    assert profiles[0].province == "山东"
    assert profiles[0].city == "青岛市"
    body = advisor.TaskIn(
        kind="track",
        file={"name": path.name, "content_base64": base64.b64encode(path.read_bytes()).decode()},
    )
    payload = advisor._file_rows(body)
    assert payload[0]["city"] == "青岛市"
    assert payload[0]["province"] == "山东"
    report = track_clients(session, payload)
    by_name = {row.customer_name: row for row in report.clients}
    qd = by_name[rows[0][0]]
    assert {item.document_id for item in qd.items} == {
        city_docs[key].id for key in ["国家", "山东", "青岛"]
    }
    assert {item.city for item in qd.items} == {"", "青岛"}
    assert len(by_name[rows[1][0]].items) == 4
    missing = by_name[rows[2][0]]
    assert missing.profile.invalid_city and not missing.items
    assert (
        next(
            row for row in report.as_dict()["clients"] if row["customer_name"] == qd.customer_name
        )["city"]
        == "青岛"
    )
    output = tmp_path / "result.xlsx"
    export_track_excel(report, output)
    book = load_workbook(output)
    sheet = book["客户相关通知"]
    columns = {cell.value: index for index, cell in enumerate(sheet[2])}
    results = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=3)]
    qd_row = next(
        row for row in results if row[columns["城市"]] == "青岛" and row[0] == qd.customer_name
    )
    assert [qd_row[columns[key]] for key in ["省份", "城市", "层级", "客户省份", "客户城市"]] == [
        "山东",
        "青岛",
        "市级",
        "山东",
        "青岛",
    ]
    assert (
        "未查询"
        in next(row for row in results if row[0] == missing.customer_name)[columns["查询说明"]]
    )


def test_compare_keeps_source_city_and_level_in_json_and_excel(session, city_docs, tmp_path):
    doc = city_docs["青岛"]
    doc.title = "关于公布合成科技企业认定名单的通知"
    doc.kind_override = "名单"
    session.add(
        ListEntry(
            document_id=doc.id,
            source_url=doc.url,
            name="合成城市科技有限公司",
            normalized_name="合成城市科技有限公司",
            sequence="1",
            category="高新技术企业",
            region="青岛",
            credit_code="91310117588657256C",
            batch="第三批",
        )
    )
    session.commit()
    report = compare_lists(session, ["合成城市科技有限公司"])
    match = report.rows[0].matches[0]
    assert (match.province, match.city, match.level) == ("山东", "青岛", "市级")
    assert match.as_dict()["city"] == "青岛"
    assert (match.name, match.credit_code, match.batch, match.sequence, match.region) == (
        "合成城市科技有限公司",
        "91310117588657256C",
        "第三批",
        "1",
        "青岛",
    )
    output = tmp_path / "compare.xlsx"
    export_compare_excel(report, output)
    book = load_workbook(output)
    fields = dict(
        zip(
            [cell.value for cell in book["明细"][2]],
            [cell.value for cell in book["明细"][3]],
            strict=True,
        )
    )
    assert (fields["省份"], fields["城市"], fields["层级"]) == ("山东", "青岛", "市级")
    assert [
        fields[key]
        for key in ["名单中名称", "统一社会信用代码", "批次", "名单序号", "名单中企业地区"]
    ] == ["合成城市科技有限公司", "91310117588657256C", "第三批", "1", "青岛"]


def test_follow_city_contract_preserves_exact_area(session, city_docs):
    conditions = me.Conditions(regions="国家部委", cities="山东/青岛", time="不限")
    statement = me._conditions_stmt(session, conditions.model_dump())
    ids = {row[0].id for row in session.execute(statement)}
    assert ids == {city_docs["国家"].id, city_docs["青岛"].id}


@pytest.mark.parametrize("excerpts", [None, 1])
def test_deadline_and_calendar_city_filter_and_metadata(session, city_docs, excerpts):
    report = v1.deadline_report(session, excerpts=excerpts, province="山东", city="青岛")
    assert report["total"] == 1
    assert (
        report["items"][0]["province"],
        report["items"][0]["city"],
        report["items"][0]["level"],
    ) == ("山东", "青岛", "市级")
    calendar = v1.calendar(session, excerpts=excerpts, province="山东", city="青岛")
    assert len(calendar["items"]) == 1
    assert calendar["items"][0]["province"] == "山东"
    assert calendar["items"][0]["city"] == "青岛"
    assert "山东" in calendar["items"][0]["region"] and "青岛" in calendar["items"][0]["region"]
    with pytest.raises(Exception, match="城市未收录"):
        v1.deadline_report(session, province="山东", city="苏州")


def test_company_city_is_optional_local_display_and_does_not_change_model_keys():
    from pydantic import ValidationError

    from youju.ai import clause

    company = clause.Company(region="山东", city="青岛")
    result = clause.with_company(
        {"rows": [{"type": "条件", "text": "在青岛注册", "keys": ["region"], "refs": []}]},
        {"company": company.model_dump(mode="json")},
    )
    assert result["rows"][0]["company"] == [
        {"key": "region", "label": "注册地", "value": "山东 · 青岛"}
    ]
    assert "city" not in clause.KEYS
    assert clause.Company(region="山东").city is None
    with pytest.raises(ValidationError, match="提供省份"):
        clause.Company(city="青岛")


def test_track_strings_and_lists_already_share_multiple_condition_contract():
    from youju.track import _profile

    for keywords, supports in [
        ("科技;创新", "资金支持;资格认定"),
        (["科技", "创新"], ["资金支持", "资格认定"]),
    ]:
        profile = _profile(
            {
                "name": "合成科技有限公司",
                "province": "山东",
                "keywords": keywords,
                "supports": supports,
            }
        )
        assert profile.keywords == ("科技", "创新")
        assert profile.supports == ("资金支持", "资格认定")


def test_track_json_full_names_match_only_same_province_city(session, city_docs):
    report = track_clients(session, [{"name": "合成企业", "province": "山东省", "city": "青岛市"}])
    row = report.clients[0]
    assert row.profile.province == "山东" and row.profile.city == "青岛"
    assert {item.document_id for item in row.items} == {
        city_docs[k].id for k in ["国家", "山东", "青岛"]
    }
    # 真实来源可能保留“市”后缀，短名输入应仍回到该来源的明确城市标识。
    source = session.get(Source, city_docs["青岛"].source_id)
    source.city = "青岛市"
    session.commit()
    report = track_clients(session, [{"name": "合成企业", "province": "山东省", "city": "青岛"}])
    assert report.clients[0].profile.city == "青岛市"
    assert len(report.clients[0].items) == 3


@pytest.mark.parametrize(
    "province, city, expected",
    [
        ("吉林省", "延边朝鲜族自治州", "吉林"),
        ("新疆维吾尔自治区", "阿克苏地区", "新疆"),
        ("广西壮族自治区", "南宁市", "广西"),
    ],
)
def test_track_profile_preserves_autonomous_city_and_region_names(province, city, expected):
    from youju.track import _profile

    profile = _profile({"name": "合成企业", "province": province, "city": city})
    assert profile.province == expected and profile.city == city and not profile.invalid_province


@pytest.mark.parametrize("excerpts", [None, 1])
def test_calendar_and_report_resolve_full_short_city_to_real_source(session, city_docs, excerpts):
    assert (
        v1.deadline_report(session, excerpts=excerpts, province="山东", city="青岛市")["total"] == 1
    )
    source = session.get(Source, city_docs["青岛"].source_id)
    source.city = "青岛市"
    session.commit()
    assert (
        v1.calendar(session, excerpts=excerpts, province="山东", city="青岛")["items"][0]["city"]
        == "青岛市"
    )
