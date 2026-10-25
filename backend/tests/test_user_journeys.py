"""按用户操作顺序验证跨接口闭环，仅使用 conftest 提供的测试库。

微信消息、内容安全、对象存储及网页抓取使用离线替身，业务查询和文件生成用真实实现。
原 9 条问题取证用例保留为回归断言，验证对应问题已修复。
"""

import re
from datetime import UTC, datetime, timedelta
from hashlib import sha256
from io import BytesIO
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import httpx
import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event, func, select, update
from sqlalchemy.sql import Select, visitors

from youju import advisor_tasks, track
from youju.agreements import TOOLS_AGREEMENT_VERSION
from youju.api import advisor, exports, me, v1
from youju.api import app as admin
from youju.config import Settings
from youju.crawler import service as crawler
from youju.crawler.fetcher import FetchResult
from youju.crawler.urls import url_key
from youju.db import get_session
from youju.lists.parser import ParsedEntry
from youju.lists.rules import classify_title, normalize_name
from youju.lists.service import _save as save_list_parse
from youju.models import (
    AdvisorTask,
    AdvisorUsage,
    Document,
    DocumentDeadline,
    DocumentDetail,
    DocumentField,
    DocumentMerge,
    Event,
    Feedback,
    Follow,
    ListEntry,
    ListParse,
    Source,
    Subscription,
    User,
)
from youju.notify import run_notify
from youju.report_revision import publication_revision
from youju.takedown import change_state

NOW = datetime(2026, 10, 3, 10, tzinfo=ZoneInfo("Asia/Shanghai"))
IDENTITY = {"X-WX-SOURCE": "miniapp", "X-WX-OPENID": "journey-test-user"}
ADMIN = {"X-Admin-Token": "journey-test-only"}
REVIEWER = "旅程测试管理员"
COMPANY = "北京旅程（科技）有限公司"


@pytest.fixture
def journey(session_factory, session, monkeypatch):
    # 身份通过可信微信头传入，禁止退回开发身份；配置显式禁用 .env。
    settings = Settings(
        _env_file=None,
        database_url=session_factory.kw["bind"].url.render_as_string(hide_password=False),
        trust_wx_headers=True,
        dev_openid=None,
        admin_token=ADMIN["X-Admin-Token"],
        agreement_version="2026-10",
        template_deadline={
            "id": "journey-deadline",
            "fields": {"thing1": "title", "date2": "date", "thing3": "kind"},
        },
        template_newdocs={
            "id": "journey-newdocs",
            "fields": {"thing1": "name", "number2": "count"},
        },
        send_wx_messages=True,
        advisor_enabled=False,
        cloud_env="journey-offline",
        exports_per_10min=100,
    )
    clock = SimpleNamespace(stamp=NOW)
    for module in (admin, advisor, exports, me):
        monkeypatch.setattr(module, "get_settings", lambda: settings)
    monkeypatch.setattr(v1, "now", lambda: clock.stamp)
    monkeypatch.setattr(track, "now", lambda: clock.stamp)
    monkeypatch.setattr(me, "check_feedback", lambda openid, content: True)
    v1._REPORT_CACHE.clear()

    def override():
        with session_factory() as request_session:
            yield request_session

    admin.app.dependency_overrides[get_session] = override
    try:
        with TestClient(admin.app, headers=IDENTITY) as client:
            yield SimpleNamespace(client=client, settings=settings, clock=clock, session=session)
    finally:
        admin.app.dependency_overrides.clear()
        v1._REPORT_CACHE.clear()


def _json(response, status=200):
    assert response.status_code == status, response.text
    return response.json()


def _agree(j):
    agreed = _json(
        j.client.post("/api/v1/me/agreement", json={"version": j.settings.agreement_version})
    )
    assert _json(j.client.get("/api/v1/me"))["agreed"] == agreed
    return agreed


def _tools_ready(j):
    """批量工具要求同意过写进任务记录的那一版指引（2026-11）；旅程按 2026-10 配置，
    用批量工具前先把本人的同意补到 2026-11，再同意工具使用说明。"""
    j.session.execute(
        update(User)
        .where(User.openid == IDENTITY["X-WX-OPENID"])
        .values(agreement_version=TOOLS_AGREEMENT_VERSION)
    )
    j.session.commit()


def _source(session, province="北京", dept="科技", suffix="default"):
    source = Source(
        level="省级",
        province=province,
        dept=dept,
        dept_full_name=f"{province}省{'科学技术厅' if dept == '科技' else '工业和信息化厅'}",
        column_name="通知公告",
        url=f"https://journey.example.gov.cn/{suffix}/notice/",
        rights_level=1,
        enabled=True,
        last_success_at=NOW - timedelta(days=1),
        verified_on=NOW.date(),
    )
    session.add(source)
    session.flush()
    return source


def _document(session, source, title, suffix, days=None, status="pending", seen=None):
    url = f"https://journey.example.gov.cn/{suffix}.html"
    doc = Document(
        source_id=source.id,
        title=title,
        url=url,
        url_key=url_key(url),
        publish_date=NOW.date(),
        first_seen_at=seen or NOW - timedelta(days=1),
        last_seen_at=seen or NOW - timedelta(days=1),
    )
    session.add(doc)
    session.flush()
    body = (
        f"{title}。本项目给予资金支持。"
        + "申报单位应按主管部门要求准备材料并对材料真实性负责。" * 12
    )
    if days is not None:
        due = NOW.date() + timedelta(days=days)
        body += f"企业申报截止时间为{due.isoformat()}。"
        deadline = DocumentDeadline(
            document_id=doc.id,
            fingerprint=sha256(suffix.encode()).hexdigest(),
            kind="企业申报截止",
            candidate_date=due,
            approved_date=due if status == "approved" else None,
            evidence=f"企业申报截止时间为{due.isoformat()}。",
            confidence="高",
            status=status,
            reviewed_by=REVIEWER if status == "approved" else None,
            reviewed_at=NOW if status == "approved" else None,
        )
        session.add(deadline)
    else:
        deadline = None
    session.add(
        DocumentDetail(
            document_id=doc.id,
            body_text=body,
            issuer=source.dept_full_name,
            doc_number=f"旅测〔2026〕{doc.id}号",
            fetched_at=NOW,
            status="ok",
        )
    )
    session.flush()
    return doc, deadline


def _support(j, doc, label="资金支持", approve=True):
    body = j.session.get(DocumentDetail, doc.id).body_text
    evidence = "本项目给予资金支持。"
    field = DocumentField(
        document_id=doc.id,
        field="support",
        value=label,
        evidence=evidence,
        evidence_offset=body.index(evidence),
        body_hash=sha256(body.encode()).hexdigest(),
        status="pending",
    )
    j.session.add(field)
    j.session.commit()
    if approve:
        assert _json(
            j.client.post(
                f"/api/admin/fields/{field.id}/approve",
                headers=ADMIN,
                json={"reviewed_by": REVIEWER},
            )
        ) == {"status": "approved"}
    return field


def _review(j, deadline, action="approve", **values):
    return _json(
        j.client.post(
            f"/api/admin/deadlines/{deadline.id}/{action}",
            headers=ADMIN,
            json={"reviewed_by": REVIEWER, **values},
        )
    )


def _follow(j, name, conditions):
    saved = _json(
        j.client.post("/api/v1/me/follows", json={"name": name, "conditions": conditions}), 201
    )
    assert saved in _json(j.client.get("/api/v1/me/follows"))["items"]
    return saved


def _subscribe(j, kind, **target):
    template = j.settings.template_deadline if kind == "deadline" else j.settings.template_newdocs
    saved = _json(
        j.client.post(
            "/api/v1/me/subscriptions",
            json={"kind": kind, "templateId": template["id"], **target},
        ),
        201,
    )
    assert saved in _json(j.client.get("/api/v1/me/subscriptions"))["items"]
    return saved


def _notify(j, sent, errcode=0):
    # 后台独立 session，验证不是靠请求 session 中未提交的对象完成闭环。
    with j.session.get_bind().connect() as connection:
        from sqlalchemy.orm import Session

        with Session(connection) as worker:
            return run_notify(
                worker,
                settings=j.settings,
                sender=lambda payload: sent.append(payload) or {"errcode": errcode},
                stamp=j.clock.stamp,
            )


def _entry(session, doc, name=COMPANY, code="91110108123456789A", sequence="1"):
    entry = ListEntry(
        document_id=doc.id,
        source_url=doc.url,
        name=name,
        normalized_name=normalize_name(name),
        credit_code=code,
        sequence=sequence,
        category="企业",
    )
    session.add(entry)
    return entry


def _export(client, docs, format_):
    response = client.post(
        "/api/v1/exports", json={"ids": [str(d.id) for d in docs], "format": format_}
    )
    assert response.status_code == 200, response.text
    assert format_ in response.headers["content-disposition"]
    if format_ == "xlsx":
        book = load_workbook(BytesIO(response.content))
        try:
            text = "\n".join(str(cell.value) for sheet in book for row in sheet for cell in row)
            links = [
                cell.hyperlink.target for row in book["文件清单"] for cell in row if cell.hyperlink
            ]
            assert links == [d.url for d in docs]
        finally:
            book.close()
    else:
        book = WordDocument(BytesIO(response.content))
        text = "\n".join(
            [p.text for p in book.paragraphs]
            + [cell.text for table in book.tables for row in table.rows for cell in row.cells]
        )
    for doc in docs:
        assert doc.title in text and doc.url in text
    assert "第三方整理" in text and "以发文机关原文为准" in text
    return text


def test_a_application_owner_search_detail_calendar_radar_reminder_and_delete(journey):
    j = journey
    source = _source(j.session)
    target, _ = _document(j.session, source, "高新技术企业项目申报通知", "target", 7, "approved")
    middle, _ = _document(j.session, source, "高新技术企业项目申报通知二", "middle", 14, "approved")
    far, _ = _document(j.session, source, "高新技术企业项目申报通知三", "far", 30, "approved")
    _document(j.session, source, "高新技术企业已截止申报通知", "past", -1, "approved")
    _document(j.session, source, "高新技术企业待核对申报通知", "pending", 2)
    _document(j.session, source, "高新技术企业远期申报通知", "future", 31, "approved")
    other = _source(j.session, "辽宁", "工信", "other")
    elsewhere, _ = _document(j.session, other, "高新技术企业申报通知", "elsewhere", 5, "approved")
    j.session.commit()
    _support(j, target)

    home = _json(j.client.get("/api/v1/home"))
    assert str(target.id) in {row["id"] for row in home["items"]}
    assert home["stats"]["apply"] == 7 and home["stats"]["sources"] == 2
    search = _json(
        j.client.get(
            "/api/v1/search",
            params={"q": "高新技术企业", "regions": "北京", "types": "申报通知", "open": 1},
        )
    )
    assert search["total"] == 4
    result = next(row for row in search["items"] if row["id"] == str(target.id))
    assert all(row["deadline"]["verified"] for row in search["items"])
    detail = _json(j.client.get(f"/api/v1/documents/{result['id']}"))
    assert detail["title"] == result["title"] and detail["url"] == target.url
    assert detail["deadline"]["date"] == result["deadline"]["date"] == "2026-10-10"
    assert detail["excerpt"]["text"] in detail["body"]["text"]
    assert detail["fields"]["support"] == ["资金支持"]
    assert detail["fields"]["support_evidence"]["资金支持"] in detail["body"]["text"]
    calendar = _json(j.client.get("/api/v1/calendar"))
    event = next(row for row in calendar["items"] if row["docId"] == result["id"])
    assert event["due"] == detail["deadline"]["date"] and event["verified"]
    for days, expected in (
        (7, {target.id}),
        (14, {target.id, middle.id}),
        (30, {target.id, middle.id, far.id}),
    ):
        radar = _json(
            j.client.get(
                "/api/v1/reports/deadlines",
                params={"days": days, "province": "北京", "dept": "科技部门"},
            )
        )
        assert {int(row["id"]) for row in radar["items"]} == expected
        assert radar["summary"] == {"7": 1, "14": 2, "30": 3}
    supported = _json(
        j.client.get("/api/v1/reports/deadlines", params={"days": 30, "support": "资金支持"})
    )
    assert [row["id"] for row in supported["items"]] == [str(target.id)]
    assert _json(j.client.get("/api/v1/reports/deadlines", params={"dept": "工信部门"}))["items"][
        0
    ]["id"] == str(elsewhere.id)
    assert (
        _json(j.client.get("/api/v1/reports/deadlines", params={"support": "贷款贴息"}))["items"]
        == []
    )

    _agree(j)
    subscription = _subscribe(j, "deadline", docId=target.id)
    assert (
        _json(
            j.client.post(
                "/api/v1/me/subscriptions",
                json={"kind": "deadline", "templateId": "journey-deadline", "docId": target.id},
            ),
            201,
        )
        == subscription
    )
    sent = []
    assert _notify(j, sent)["skipped"] == 1 and sent == []
    j.clock.stamp = NOW + timedelta(days=4)
    assert _notify(j, sent)["sent"] == 1
    assert _notify(j, sent)["sent"] == 0 and len(sent) == 1
    assert sent[0]["touser"] == IDENTITY["X-WX-OPENID"]
    assert sent[0]["page"] == f"pages/detail/index?id={target.id}"
    assert sent[0]["data"]["date2"]["value"] == detail["deadline"]["date"]
    data = _json(j.client.get("/api/v1/me/data"))
    assert data["subscriptions"][0]["id"] == subscription["id"]
    assert data["subscriptions"][0]["status"] == "sent" and data["subscriptions"][0]["attemptedAt"]
    assert _json(j.client.get("/api/v1/me"))["subscriptions"]["deadline"] == 0
    assert _json(j.client.get("/api/v1/me/subscriptions"))["items"] == []
    # 已发送授权是历史记录，产品只允许删除未用授权；再订阅一次后取消这次提醒。
    again = _subscribe(j, "deadline", docId=target.id)
    assert again["id"] != subscription["id"]
    assert j.client.delete(f"/api/v1/me/subscriptions/{again['id']}").status_code == 204
    assert _json(j.client.get("/api/v1/me/subscriptions"))["items"] == []
    assert _json(j.client.get("/api/v1/me/data"))["subscriptionsTotal"] == 1


def test_b_advisor_crawl_follow_unread_send_export_lookup_and_board(
    journey, session_factory, monkeypatch
):
    j = journey
    source = _source(j.session)
    old, _ = _document(j.session, source, "高新技术企业申报通知", "existing")
    listed, _ = _document(j.session, source, "2026年高新技术企业名单公示", "listed")
    same_name, _ = _document(j.session, source, "2026年专精特新企业名单公布", "same-name")
    _entry(j.session, listed)
    _entry(j.session, same_name, code="91110108987654321B")
    _entry(j.session, listed, name="北京旅程股份有限公司", code="91110108000000000C", sequence="2")
    j.session.commit()
    searched = _json(
        j.client.get(
            "/api/v1/search", params={"q": "高新技术企业", "regions": "北京", "kind": "申报通知"}
        )
    )
    assert [row["id"] for row in searched["items"]] == [str(old.id)]
    _agree(j)
    follow = _follow(
        j,
        "北京高企申报",
        {"q": "高新技术企业", "regions": "北京", "kind": "申报通知", "time": "近 30 天"},
    )
    other_follow = _follow(j, "专精特新申报", {"q": "专精特新", "kind": "申报通知"})
    assert follow["newCount"] == other_follow["newCount"] == 0
    j.clock.stamp = NOW + timedelta(minutes=1)
    title = "关于新增高新技术企业项目申报的通知"
    html = (
        f'<ul class="list"><li><a href="/fresh.html">{title}</a><span>2026-10-03</span></li></ul>'
    )
    calls = []

    class OfflineFetcher:
        def fetch(self, url):
            calls.append(url)
            assert url == source.url
            return FetchResult(
                url, "ok", final_url=url, http_status=200, content_type="text/html", text=html
            )

    monkeypatch.setattr(crawler, "_now", lambda: j.clock.stamp.astimezone(UTC))
    report = crawler.crawl_sources(session_factory, [source.id], OfflineFetcher(), max_workers=1)[0]
    assert (report.status, report.baseline, report.items_new) == ("ok", False, 1)
    assert calls == [source.url]
    j.session.expire_all()
    fresh = j.session.scalar(select(Document).where(Document.title == title))
    assert fresh is not None and not fresh.is_baseline
    markers = _json(j.client.get("/api/v1/me/follows"))["items"]
    assert {row["id"]: row["newCount"] for row in markers} == {
        follow["id"]: 1,
        other_follow["id"]: 0,
    }
    new = _json(j.client.get(f"/api/v1/me/follows/{follow['id']}/new"))
    assert new["total"] == 1 and new["items"][0]["id"] == str(fresh.id)
    first = _subscribe(j, "newDocs", followId=follow["id"])
    other_subscription = _subscribe(j, "newDocs", followId=other_follow["id"])
    j.clock.stamp += timedelta(minutes=1)
    sent = []
    assert _notify(j, sent)["sent"] == 1 and _notify(j, sent)["sent"] == 0
    assert len(sent) == 1 and sent[0]["data"]["thing1"]["value"] == follow["name"]
    assert sent[0]["data"]["number2"]["value"] == "1" and sent[0]["page"] == "pages/follow/index"
    assert _json(j.client.get(f"/api/v1/me/follows/{follow['id']}/new"))["total"] == 1
    history = _json(j.client.get("/api/v1/me/data"))["subscriptions"]
    assert next(row for row in history if row["id"] == first["id"])["status"] == "sent"
    assert (
        next(row for row in history if row["id"] == other_subscription["id"])["status"] == "unused"
    )
    assert _json(j.client.get("/api/v1/me/subscriptions"))["items"] == [other_subscription]
    assert _json(j.client.get("/api/v1/me"))["subscriptions"]["newDocs"] == 1
    assert (
        next(
            row
            for row in _json(j.client.get("/api/v1/me/follows"))["items"]
            if row["id"] == follow["id"]
        )["newCount"]
        == 1
    )
    selected = _json(j.client.get("/api/v1/search", params=follow["conditions"]))["items"]
    selected_docs = [j.session.get(Document, int(row["id"])) for row in selected]
    for format_ in ("xlsx", "docx"):
        _export(j.client, selected_docs, format_)
    for name in (COMPANY, "北京旅程 (科技) 有限公司"):
        found = _json(j.client.get("/api/v1/lists/lookup", params={"q": name}))
        assert found["ambiguous"] and {row["id"] for row in found["items"]} == {
            str(listed.id),
            str(same_name.id),
        }
        assert {row["matched"]["creditCode"] for row in found["items"]} == {
            "91110108123456789A",
            "91110108987654321B",
        }
        assert {row["status"] for row in found["items"]} == {"公示", "公布"}
        assert all(row["url"] for row in found["items"])
        assert found["similar"] == [{"name": "北京旅程股份有限公司", "lists": 1}]
    assert (
        len(
            _json(j.client.get("/api/v1/lists/lookup", params={"q": "91110108123456789A"}))["items"]
        )
        == 1
    )
    missing = _json(j.client.get("/api/v1/lists/lookup", params={"q": "虚构完全无关有限公司"}))
    assert missing["items"] == [] and missing["coverage"]["lists"] == 2
    board = _json(j.client.get("/api/v1/reports/lists"))
    assert board["recent"]["total"] == j.session.scalar(
        select(func.count(Document.id)).where(Document.id.in_([listed.id, same_name.id]))
    )
    for row in board["recent"]["items"]:
        assert row["count"] == j.session.scalar(
            select(func.count(ListEntry.id)).where(ListEntry.document_id == int(row["id"]))
        )
    assert sum(row["entries"] for row in board["stats"]["by_type"]) == 3
    assert _json(j.client.post(f"/api/v1/me/follows/{follow['id']}/seen"))["newCount"] == 0
    assert _json(j.client.get(f"/api/v1/me/follows/{follow['id']}/new")) == {
        "total": 0,
        "items": [],
    }


def test_c_consent_feedback_admin_personal_data_delete_and_account_deletion(journey):
    j = journey
    source = _source(j.session)
    doc, _ = _document(j.session, source, "项目申报通知", "feedback-doc", 7, "approved")
    j.session.commit()
    for path, body in (
        ("/api/v1/me/follows", {"name": "未同意", "conditions": {}}),
        (
            "/api/v1/me/subscriptions",
            {"kind": "deadline", "templateId": "journey-deadline", "docId": doc.id},
        ),
        ("/api/v1/feedback", {"docId": doc.id, "content": "核对日期"}),
    ):
        assert (
            _json(j.client.post(path, json=body), 403)["detail"]
            == "请先阅读并同意用户协议和隐私保护指引"
        )
    assert j.session.scalar(select(func.count(User.id))) == 0
    _agree(j)
    follow = _follow(j, "待更正名称", {"q": "项目"})
    _subscribe(j, "deadline", docId=doc.id)
    response = _json(
        j.client.post(
            "/api/v1/feedback",
            json={"docId": doc.id, "content": "截止日期有误，请核对。联系电话13812345678"},
        ),
        201,
    )
    assert response["masked"]
    feedback_id = response["id"]
    row = next(
        row
        for row in _json(j.client.get("/api/admin/feedback", headers=ADMIN))
        if row["id"] == feedback_id
    )
    assert row["docId"] == doc.id and row["checkResult"] == "pass"
    assert "13812345678" not in row["content"] and "［已隐去］" in row["content"]
    assert _json(
        j.client.post(
            f"/api/admin/feedback/{feedback_id}/handle", headers=ADMIN, json={"outcome": "fixed"}
        )
    ) == {
        "handled": True,
        "changed": True,
        "outcome": "fixed",
        "outcomeLabel": "已更正",
        "reply": None,
    }
    handled = next(
        row
        for row in _json(
            j.client.get("/api/admin/feedback", headers=ADMIN, params={"status": "all"})
        )
        if row["id"] == feedback_id
    )
    assert handled["handledAt"]
    assert (
        _json(j.client.put(f"/api/v1/me/follows/{follow['id']}", json={"name": "更正后的名称"}))[
            "name"
        ]
        == "更正后的名称"
    )
    data = _json(j.client.get("/api/v1/me/data"))
    assert data["user"]["openid"] == IDENTITY["X-WX-OPENID"]
    assert data["follows"][0]["name"] == "更正后的名称"
    feedback = data["feedback"][0]
    assert (
        feedback["id"] == feedback_id
        and feedback["docTitle"] == doc.title
        and feedback["content"] == row["content"]
    )
    assert _json(
        j.client.post("/api/v1/events", json={"name": "feedback", "props": {"count": 1}}), 202
    )["accepted"]
    assert j.client.delete(f"/api/v1/me/feedback/{feedback_id}").status_code == 204
    assert _json(j.client.get("/api/v1/me/data"))["feedbackTotal"] == 0
    # 留另一条纠错供注销级联检查，避免只验证空表。
    _json(j.client.post("/api/v1/feedback", json={"content": "另一条建议"}), 201)
    assert _json(j.client.get("/api/v1/me/data"))["eventsTotal"] == 1
    assert j.client.delete("/api/v1/me").status_code == 204
    j.session.expire_all()
    for model in (User, Follow, Subscription, Feedback, Event, AdvisorTask, AdvisorUsage):
        assert j.session.scalar(select(func.count()).select_from(model)) == 0
    assert _json(j.client.get("/api/v1/me")) == {
        "agreed": None,
        "follows": 0,
        "subscriptions": {"newDocs": 0, "deadline": 0},
        # 注销后：不在售，也没有订单（小程序不显示购买入口和「我的订单」）
        "pay": {"enabled": False, "hasOrders": False},
        # 注销后没有使用期：我的页头没有"批量工具"标识
        "pass": None,
        # 纠错送天数开关默认关：纠错页不提送天数
        "correctionReward": None,
    }
    empty = _json(j.client.get("/api/v1/me/data"))
    assert (
        empty["user"] is None
        and empty["follows"] == empty["feedback"] == empty["subscriptions"] == []
    )
    assert j.client.get("/api/v1/me/follows").status_code == 403


@pytest.mark.parametrize("action", ["approve", "correct", "reject", "undo", "reopen"])
def test_d_review_to_live_detail_calendar_and_uncached_radar(journey, action):
    j = journey
    source = _source(j.session)
    doc, deadline = _document(j.session, source, "日期闭环申报通知", "review", 7)
    j.session.commit()
    _support(j, doc)
    # 带支持方式的雷达不缓存；预热缓存的即时更新另有回归用例。
    radar_params = {"days": 30, "support": "资金支持"}
    assert _json(j.client.get("/api/v1/calendar"))["items"] == []
    assert _json(j.client.get("/api/v1/reports/deadlines", params=radar_params))["items"] == []
    assert _json(j.client.get(f"/api/v1/documents/{doc.id}"))["deadline"] is None
    if action == "reject":
        _review(j, deadline, "reject")
        expected = None
    else:
        due = "2026-10-15" if action == "correct" else "2026-10-10"
        _review(j, deadline, **({"date": due} if action == "correct" else {}))
        assert _json(j.client.get(f"/api/v1/documents/{doc.id}"))["deadline"]["date"] == due
        assert _json(j.client.get("/api/v1/calendar"))["items"][0]["due"] == due
        assert (
            _json(j.client.get("/api/v1/reports/deadlines", params=radar_params))["items"][0][
                "date"
            ]
            == due
        )
        expected = due
        if action == "undo":
            assert _json(
                j.client.post(
                    f"/api/admin/deadlines/{deadline.id}/undo",
                    headers=ADMIN,
                    json={"reviewer": REVIEWER},
                )
            ) == {"status": "pending"}
            expected = None
        if action == "reopen":
            assert _review(j, deadline, "reopen") == {"status": "pending"}
            expected = None
    detail = _json(j.client.get(f"/api/v1/documents/{doc.id}"))
    calendar = _json(j.client.get("/api/v1/calendar"))
    radar = _json(j.client.get("/api/v1/reports/deadlines", params=radar_params))
    if expected:
        assert (
            detail["deadline"]["date"]
            == calendar["items"][0]["due"]
            == radar["items"][0]["date"]
            == expected
        )
    else:
        assert detail["deadline"] is None and calendar["items"] == radar["items"] == []
        assert j.client.get("/api/v1/search", params={"open": 1}).json()["total"] == 0
    j.session.refresh(deadline)
    assert deadline.status == (
        "rejected"
        if action == "reject"
        else "pending"
        if action in ("undo", "reopen")
        else "approved"
    )
    if action in ("undo", "reopen"):
        assert any(
            row["id"] == deadline.id
            for row in _json(j.client.get("/api/admin/deadlines", headers=ADMIN))
        )


def test_d_support_approval_publishes_label_and_original_evidence(journey):
    j = journey
    source = _source(j.session)
    doc, _ = _document(j.session, source, "支持方式申报通知", "support", 7, "approved")
    j.session.commit()
    field = _support(j, doc, approve=False)
    assert "fields" not in _json(j.client.get(f"/api/v1/documents/{doc.id}"))
    queue = _json(j.client.get("/api/admin/fields", headers=ADMIN))
    assert queue[0]["document_id"] == doc.id
    _json(
        j.client.post(
            f"/api/admin/fields/{field.id}/approve", headers=ADMIN, json={"reviewed_by": REVIEWER}
        )
    )
    detail = _json(j.client.get(f"/api/v1/documents/{doc.id}"))
    assert detail["fields"]["support"] == [field.value]
    assert detail["fields"]["support_evidence"][field.value] == field.evidence
    assert field.evidence in detail["body"]["text"]


@pytest.mark.parametrize("kind", ["申报通知", "名单"])
def test_d_takedown_and_restore_all_user_surfaces(journey, kind):
    j = journey
    source = _source(j.session)
    title = "2026年高新技术企业名单公示" if kind == "名单" else "高新技术企业申报通知"
    doc, _ = _document(
        j.session, source, title, "remove", 7, "approved", seen=NOW + timedelta(minutes=1)
    )
    if kind == "名单":
        _entry(j.session, doc)
    j.session.commit()
    _agree(j)
    follow = _follow(j, "高新技术企业", {"q": "高新技术企业"})
    _subscribe(j, "newDocs", followId=follow["id"])
    if kind == "申报通知":
        _subscribe(j, "deadline", docId=doc.id)
    paths = [
        "/api/v1/home",
        "/api/v1/search",
        "/api/v1/calendar",
        "/api/v1/reports/deadlines",
        "/api/v1/reports/lists",
        "/api/v1/lists/recent",
        f"/api/v1/me/follows/{follow['id']}/new",
    ]
    # 先取数据预热缓存，下架后不能靠等 TTL 才消失。
    before = {path: _json(j.client.get(path)) for path in paths}
    lookup = _json(j.client.get("/api/v1/lists/lookup", params={"q": COMPANY}))
    for format_ in ("xlsx", "docx"):
        _export(j.client, [doc], format_)

    assert _json(
        j.client.post(
            f"/api/admin/takedown/{doc.id}/remove",
            headers=ADMIN,
            json={"reason": "侵权投诉", "note": "离线测试下架", "operated_by": REVIEWER},
        )
    )["document_ids"] == [doc.id]
    for path in paths:
        assert doc.title not in j.client.get(path).text
    response = j.client.get(f"/api/v1/documents/{doc.id}")
    assert _json(response, 404)["detail"] == "该文件已下架"
    assert _json(j.client.get("/api/v1/lists/lookup", params={"q": COMPANY}))["items"] == []
    assert _json(j.client.get("/api/v1/me/follows"))["items"][0]["newCount"] == 0
    for format_ in ("xlsx", "docx"):
        assert (
            j.client.post(
                "/api/v1/exports", json={"ids": [str(doc.id)], "format": format_}
            ).status_code
            == 422
        )
    sent = []
    j.clock.stamp = NOW + timedelta(days=4)
    assert _notify(j, sent)["sent"] == 0 and sent == []
    j.clock.stamp = NOW
    _json(
        j.client.post(
            f"/api/admin/takedown/{doc.id}/restore",
            headers=ADMIN,
            json={"note": "离线测试恢复", "operated_by": REVIEWER},
        )
    )
    for path in paths:
        assert _json(j.client.get(path)) == before[path]
    assert _json(j.client.get("/api/v1/lists/lookup", params={"q": COMPANY})) == lookup
    assert _json(j.client.get(f"/api/v1/documents/{doc.id}"))["title"] == doc.title
    for format_ in ("xlsx", "docx"):
        _export(j.client, [doc], format_)
    j.clock.stamp = NOW + timedelta(days=4)
    assert _notify(j, sent)["sent"] == (2 if kind == "申报通知" else 1)
    assert _notify(j, sent)["sent"] == 0


@pytest.mark.parametrize("kind", ["compare", "track"])
def test_e_advisor_submit_process_download_quota_delete(journey, monkeypatch, kind):
    j = journey
    monkeypatch.setenv("YOUJU_ADVISOR_ENABLED", "true")
    j.settings.advisor_enabled = Settings(_env_file=None).advisor_enabled
    assert j.settings.advisor_enabled
    source = _source(j.session)
    listed, _ = _document(j.session, source, "2026年高新技术企业名单公布", "tool-list")
    application, _ = _document(
        j.session, source, "高新技术企业项目申报通知", "tool-apply", 7, "approved"
    )
    _entry(j.session, listed)
    j.session.commit()
    _support(j, application)
    _agree(j)
    _tools_ready(j)
    _json(
        j.client.post(
            "/api/v1/me/agreement",
            json={"scope": "advisor", "version": advisor_tasks.ADVISOR_AGREEMENT_VERSION},
        )
    )
    storage = {}

    def upload(client, env, remote_path, local_path, **kwargs):
        assert env == j.settings.cloud_env
        file_id = f"cloud://journey/{kind}"
        storage[file_id] = local_path.read_bytes()
        return file_id

    def delete_files(client, env, file_ids):
        for file_id in file_ids:
            del storage[file_id]
        return file_ids

    monkeypatch.setattr(advisor_tasks, "upload", upload)
    monkeypatch.setattr(
        advisor_tasks,
        "download_url",
        lambda client, env, file_id: (
            f"https://journey-download.test/{kind}.xlsx" if file_id in storage else None
        ),
    )
    monkeypatch.setattr(advisor_tasks, "delete_result_files", delete_files)
    rows = (
        [COMPANY, "完全无关有限公司"]
        if kind == "compare"
        else [
            {
                "name": COMPANY,
                "province": "北京",
                "keywords": "高新技术企业",
                "supports": "资金支持",
            }
        ]
    )
    body = {"kind": kind, "rows": rows}
    created = _json(
        j.client.post("/api/v1/advisor/tasks", headers={"Idempotency-Key": kind}, json=body), 202
    )
    assert created["status"] == "queued" and created["rowCount"] == len(rows)
    assert _json(j.client.get("/api/v1/advisor/tasks"))["items"][0]["id"] == created["id"]
    assert (
        _json(
            j.client.post("/api/v1/advisor/tasks", headers={"Idempotency-Key": kind}, json=body),
            202,
        )["id"]
        == created["id"]
    )
    assert j.client.get(f"/api/v1/advisor/tasks/{created['id']}/download").status_code == 410
    assert (
        advisor_tasks.run_advisor_tasks(j.session, settings=j.settings, stamp=NOW, max_tasks=1)[
            "done"
        ]
        == 1
    )
    done = _json(j.client.get("/api/v1/advisor/tasks"))["items"][0]
    assert done["status"] == "done" and done["hasFile"] and done["resultRows"] >= 1
    # 结果文件经接口分段取回，拼起来就是 Excel
    stored = storage[f"cloud://journey/{kind}"]
    monkeypatch.setattr(advisor_tasks, "result_size", lambda client, url: len(stored))
    monkeypatch.setattr(
        advisor_tasks,
        "read_result_part",
        lambda client, url, index, total: stored[
            index * advisor_tasks.RESULT_PART_SIZE : (index + 1) * advisor_tasks.RESULT_PART_SIZE
        ],
    )
    info = _json(j.client.get(f"/api/v1/advisor/tasks/{created['id']}/download"))
    assert info["size"] == len(stored) and info["parts"] >= 1 and "url" not in info
    downloaded = b"".join(
        j.client.get(f"/api/v1/advisor/tasks/{created['id']}/file?part={n}").content
        for n in range(info["parts"])
    )
    book = load_workbook(BytesIO(downloaded))
    try:
        text = "\n".join(str(cell.value) for sheet in book for row in sheet for cell in row)
        target = listed if kind == "compare" else application
        assert COMPANY in text and target.title in text and target.url in text
    finally:
        book.close()
    assert _json(j.client.get("/api/v1/me/data"))["advisorTasks"][0]["status"] == "done"
    for n in (2, 3):
        extra = _json(
            j.client.post(
                "/api/v1/advisor/tasks", headers={"Idempotency-Key": f"{kind}-{n}"}, json=body
            ),
            202,
        )
        # 已完成任务删除后仍占额度，任务列表留给原闭环验证。
        row = j.session.get(AdvisorTask, extra["id"])
        row.status = "done"
        j.session.commit()
        assert j.client.delete(f"/api/v1/advisor/tasks/{extra['id']}").status_code == 204
    limited = j.client.post(
        "/api/v1/advisor/tasks", headers={"Idempotency-Key": f"{kind}-blocked"}, json=body
    )
    assert _json(limited, 429)["detail"]["code"] == "FREE_MONTHLY_LIMIT"
    assert j.client.delete(f"/api/v1/advisor/tasks/{created['id']}").status_code == 204
    assert storage == {} and _json(j.client.get("/api/v1/advisor/tasks"))["items"] == []
    assert _json(j.client.get("/api/v1/me/data"))["advisorTasksTotal"] == 0
    # 删已完成任务不能刷出当前周期的新额度。
    assert (
        j.client.post(
            "/api/v1/advisor/tasks", headers={"Idempotency-Key": f"{kind}-deleted"}, json=body
        ).status_code
        == 429
    )


def test_e_advisor_disabled_blocks_all_endpoints_and_worker(journey, monkeypatch):
    j = journey
    _agree(j)
    monkeypatch.setenv("YOUJU_ADVISOR_ENABLED", "false")
    j.settings.advisor_enabled = Settings(_env_file=None).advisor_enabled
    for method, path, kwargs in (
        ("get", "/api/v1/advisor/tasks", {}),
        (
            "post",
            "/api/v1/advisor/tasks",
            {"headers": {"Idempotency-Key": "off"}, "json": {"kind": "compare", "rows": [COMPANY]}},
        ),
        ("get", "/api/v1/advisor/tasks/1", {}),
        ("get", "/api/v1/advisor/tasks/1/download", {}),
        ("get", "/api/v1/advisor/tasks/1/file?part=0", {}),
        ("post", "/api/v1/advisor/tasks/1/cancel", {}),
    ):
        assert _json(getattr(j.client, method)(path, **kwargs), 404)["detail"] == "没有找到接口"
    # 删除自己的记录不看开关（D-093 上线门槛第 8 条）：删除接口一直在，没有这条记录照样"没有找到"
    response = j.client.delete("/api/v1/advisor/tasks/1")
    assert _json(response, 404)["detail"] == "没有找到顾问任务"
    assert advisor_tasks.run_advisor_tasks(j.session, settings=j.settings)["processed"] == 0
    assert _json(j.client.get("/api/v1/me/data"))["advisorTasksTotal"] == 0


@pytest.mark.parametrize("errcode,expected", [(43101, "refused"), (-1, "failed")])
def test_reminder_failure_consumes_once_and_user_can_resubscribe(journey, errcode, expected):
    j = journey
    source = _source(j.session)
    doc, _ = _document(j.session, source, "提醒失败申报通知", "failed", 3, "approved")
    j.session.commit()
    _agree(j)
    old = _subscribe(j, "deadline", docId=doc.id)
    sent = []
    assert _notify(j, sent, errcode)[expected] == 1
    assert _notify(j, sent, errcode)[expected] == 0 and len(sent) == 1
    assert _json(j.client.get("/api/v1/me/data"))["subscriptions"][0]["status"] == expected
    new = _subscribe(j, "deadline", docId=doc.id)
    assert new["id"] != old["id"]
    assert j.client.delete(f"/api/v1/me/subscriptions/{new['id']}").status_code == 204


@pytest.mark.parametrize("security", ["reject", "unavailable"])
def test_feedback_security_rejects_without_saving_and_retry_succeeds(
    journey, monkeypatch, security
):
    j = journey
    _agree(j)

    def check(openid, content):
        assert openid == IDENTITY["X-WX-OPENID"]
        if security == "unavailable":
            raise httpx.ConnectError("offline fake")
        return False

    monkeypatch.setattr(me, "check_feedback", check)
    response = j.client.post("/api/v1/feedback", json={"content": "测试不合规内容"})
    assert _json(response, 503 if security == "unavailable" else 422)["detail"] == (
        "内容安全校验暂不可用" if security == "unavailable" else "内容可能不合规，请修改后再提交"
    )
    assert _json(j.client.get("/api/v1/me/data"))["feedbackTotal"] == 0
    monkeypatch.setattr(me, "check_feedback", lambda openid, content: True)
    saved = _json(j.client.post("/api/v1/feedback", json={"content": "修改后建议"}), 201)
    assert _json(j.client.get("/api/v1/me/data"))["feedback"][0]["id"] == saved["id"]


def test_identity_missing_headers_never_creates_user(journey):
    j = journey
    with TestClient(admin.app) as anonymous:
        assert _json(anonymous.get("/api/v1/me"), 401)["detail"] == "请在微信小程序里使用"
        assert (
            anonymous.post("/api/v1/me/agreement", json={"version": "2026-10"}).status_code == 401
        )
        assert anonymous.get("/api/v1/home").status_code == 200
    assert j.session.scalar(select(func.count(User.id))) == 0


def test_account_deletion_cascades_nonempty_tools_and_keeps_other_user(journey, monkeypatch):
    j = journey
    j.settings.advisor_enabled = True
    _agree(j)
    _tools_ready(j)
    _json(
        j.client.post(
            "/api/v1/me/agreement",
            json={"scope": "advisor", "version": advisor_tasks.ADVISOR_AGREEMENT_VERSION},
        )
    )
    owner_id = j.session.scalar(select(User.id).where(User.openid == IDENTITY["X-WX-OPENID"]))
    follow = _follow(j, "本人关注", {})
    _subscribe(j, "newDocs", followId=follow["id"])
    _json(j.client.post("/api/v1/feedback", json={"content": "本人纠错"}), 201)
    _json(j.client.post("/api/v1/events", json={"name": "follow", "props": {"count": 1}}), 202)
    for kind, rows in (("compare", [COMPANY]), ("track", [{"name": COMPANY, "province": "北京"}])):
        created = _json(
            j.client.post(
                "/api/v1/advisor/tasks",
                headers={"Idempotency-Key": kind},
                json={"kind": kind, "rows": rows},
            ),
            202,
        )
        task = j.session.get(AdvisorTask, created["id"])
        task.status, task.result_file_id = "done", f"cloud://account/{kind}"
        j.session.commit()
    deleted = []
    monkeypatch.setattr(
        advisor_tasks,
        "delete_result_files",
        lambda client, env, file_ids: deleted.extend(file_ids) or file_ids,
    )
    # 另一位用户使用同样的 API 保存自己的数据；本人无权读取或删除它。
    other_headers = {"X-WX-OPENID": "journey-other-user"}
    _json(j.client.post("/api/v1/me/agreement", headers=other_headers, json={"version": "2026-10"}))
    other_follow = _json(
        j.client.post(
            "/api/v1/me/follows", headers=other_headers, json={"name": "他人关注", "conditions": {}}
        ),
        201,
    )
    other_feedback = _json(
        j.client.post("/api/v1/feedback", headers=other_headers, json={"content": "他人纠错"}), 201
    )
    other_data = _json(j.client.get("/api/v1/me/data", headers=other_headers))
    own_data = _json(j.client.get("/api/v1/me/data"))
    assert own_data["advisorTasksTotal"] == 2
    assert [row["id"] for row in own_data["follows"]] == [follow["id"]]
    assert j.client.delete(f"/api/v1/me/follows/{other_follow['id']}").status_code == 404
    assert j.client.delete(f"/api/v1/me/feedback/{other_feedback['id']}").status_code == 404
    assert j.client.delete("/api/v1/me").status_code == 204
    j.session.expire_all()
    assert j.session.get(User, owner_id) is None
    for model in (Follow, Subscription, Feedback, Event, AdvisorTask, AdvisorUsage):
        assert (
            j.session.scalar(
                select(func.count()).select_from(model).where(model.user_id == owner_id)
            )
            == 0
        )
    assert sorted(deleted) == ["cloud://account/compare", "cloud://account/track"]
    assert _json(j.client.get("/api/v1/me"))["agreed"] is None
    remaining = _json(j.client.get("/api/v1/me/data", headers=other_headers))
    for key in ("user", "follows", "feedback", "subscriptions", "advisorTasks"):
        assert remaining[key] == other_data[key]


def test_fixed_sent_subscription_can_be_deleted_individually(journey, record_property):
    j = journey
    source = _source(j.session)
    doc, _ = _document(j.session, source, "历史授权申报通知", "history", 3, "approved")
    j.session.commit()
    _agree(j)
    sub = _subscribe(j, "deadline", docId=doc.id)
    sent = []
    assert _notify(j, sent)["sent"] == 1
    assert j.client.delete(f"/api/v1/me/subscriptions/{sub['id']}").status_code == 204
    assert _json(j.client.get("/api/v1/me/data"))["subscriptions"] == []
    assert _notify(j, sent)["sent"] == 0
    assert len(sent) == 1
    record_property("acceptance", "已修：J-08，已用授权可单条删除且不会重复发送")


def test_fixed_coverage_prefix_and_input_share_maximum(journey, record_property):
    j = journey
    _agree(j)
    # 小程序覆盖输入上限 495，与五字前缀共用服务端 500 字限额。
    content = "申请覆盖：" + "字" * 495
    assert len(content) == 500
    saved = _json(j.client.post("/api/v1/feedback", json={"content": content}), 201)
    row = _json(j.client.get("/api/v1/me/data"))["feedback"][0]
    assert row["id"] == saved["id"] and row["content"] == content
    _json(j.client.post("/api/v1/feedback", json={"content": content + "字"}), 422)
    record_property("acceptance", "已修：J-10，495 字覆盖内容含前缀可完整保存")


def test_fixed_open_search_can_be_saved_as_follow(journey, record_property):
    j = journey
    _agree(j)
    conditions = {
        "q": "高新技术企业",
        "kind": "申报通知",
        "regions": "北京",
        "time": "近 30 天",
        "open": 1,
    }
    saved = _follow(j, "未截止的高企", conditions)
    assert saved["conditions"]["open"] == 1
    source = _source(j.session)
    first_seen = NOW + timedelta(seconds=1)
    live, deadline = _document(
        j.session, source, "高新技术企业申报通知", "open-new", 7, "approved", seen=first_seen
    )
    _document(
        j.session, source, "高新技术企业申报通知", "closed-new", -1, "approved", seen=first_seen
    )
    _document(j.session, source, "高新技术企业申报通知", "pending-new", 7, seen=first_seen)
    j.session.commit()
    search = _json(j.client.get("/api/v1/search", params=saved["conditions"]))
    new = _json(j.client.get(f"/api/v1/me/follows/{saved['id']}/new"))
    assert [row["id"] for row in search["items"]] == [str(live.id)]
    assert new["items"] == search["items"] and new["total"] == 1
    _review(j, deadline, "reopen")
    assert _json(j.client.get(f"/api/v1/me/follows/{saved['id']}/new"))["total"] == 0
    assert _json(j.client.get("/api/v1/me/follows"))["items"][0]["newCount"] == 0
    sent = []
    _subscribe(j, "newDocs", followId=saved["id"])
    assert _notify(j, sent)["sent"] == 0
    record_property("acceptance", "已修：J-02，open 条件保存、搜索、新文件及提醒口径一致")


def test_fixed_feedback_handling_visible_in_personal_data(journey, record_property):
    j = journey
    _agree(j)
    saved = _json(j.client.post("/api/v1/feedback", json={"content": "请修正日期"}), 201)
    before = _json(j.client.get("/api/v1/me/data"))["feedback"]
    _json(
        j.client.post(
            f"/api/admin/feedback/{saved['id']}/handle",
            headers=ADMIN,
            json={"outcome": "fixed", "reply": "已改正"},
        )
    )
    after = _json(j.client.get("/api/v1/me/data"))["feedback"]
    j.session.expire_all()
    assert j.session.get(Feedback, saved["id"]).handled_at is not None
    assert before[0]["handled"] is False and before[0]["handledAt"] is None
    assert after[0]["handled"] is True and after[0]["handledAt"] is not None
    assert (after[0]["outcome"], after[0]["outcomeLabel"], after[0]["reply"]) == (
        "fixed",
        "已更正",
        "已改正",
    )
    assert (before[0]["outcome"], before[0]["outcomeLabel"], before[0]["reply"]) == (None,) * 3
    assert after[0]["checkResult"] == before[0]["checkResult"]
    _json(j.client.post(f"/api/admin/feedback/{saved['id']}/unhandle", headers=ADMIN))
    assert _json(j.client.get("/api/v1/me/data"))["feedback"] == before
    record_property("acceptance", "已修：J-04，处理状态和时间独立于内容安全结果")


@pytest.mark.parametrize("action", ["approve", "correct", "undo", "reopen"])
def test_fixed_warmed_radar_reflects_deadline_review(journey, action, record_property):
    j = journey
    source = _source(j.session)
    doc, deadline = _document(j.session, source, "缓存核对申报通知", "cached-review", 7)
    j.session.commit()
    if action in ("correct", "undo", "reopen"):
        _review(j, deadline)
    endpoint = "/api/v1/reports/deadlines"
    before = _json(j.client.get(endpoint))
    if action == "undo":
        _json(
            j.client.post(
                f"/api/admin/deadlines/{deadline.id}/undo",
                headers=ADMIN,
                json={"reviewer": REVIEWER},
            )
        )
    elif action == "reopen":
        _review(j, deadline, "reopen")
    elif action == "correct":
        _review(j, deadline, "reopen")
        _review(j, deadline, date="2026-10-09")
    else:
        _review(j, deadline)
    live = _json(j.client.get(f"/api/v1/documents/{doc.id}"))["deadline"]
    assert (live is None) == (action in ("undo", "reopen"))
    if live:
        assert live["date"] == ("2026-10-09" if action == "correct" else "2026-10-10")
    current = _json(j.client.get(endpoint))
    assert current != before
    assert current["total"] == (0 if live is None else 1)
    if live:
        assert current["items"][0]["date"] == live["date"]
    calendar = _json(j.client.get("/api/v1/calendar"))
    assert calendar["total"] == current["total"]
    record_property("acceptance", f"已修：J-01，{action} 后预热雷达立即更新")


def test_fixed_calendar_returns_recent_closed_items(journey, record_property):
    j = journey
    source = _source(j.session)
    past, _ = _document(j.session, source, "昨日截止申报通知", "closed", -1, "approved")
    today, _ = _document(j.session, source, "今日截止申报通知", "today", 0, "approved")
    j.session.commit()
    assert _json(j.client.get(f"/api/v1/documents/{past.id}"))["app_status"]["code"] == "closed"
    calendar = _json(j.client.get("/api/v1/calendar"))
    assert [row["docId"] for row in calendar["items"]] == [str(today.id)]
    assert [row["docId"] for row in calendar["past"]] == [str(past.id)]
    assert calendar["pastTotal"] == 1 and calendar["total"] == 1
    record_property("acceptance", "已修：J-03，已核定昨日截止返回独立 past 分区")


@pytest.mark.parametrize("action", ["reject", "change"])
def test_warmed_radar_reject_and_change(journey, action):
    j = journey
    source = _source(j.session)
    doc, deadline = _document(j.session, source, "缓存申报通知", "extra-review", 7)
    newer, _ = _document(j.session, source, "延期申报通知", "extension")
    j.session.commit()
    if action == "change":
        _review(j, deadline)
    before = _json(j.client.get("/api/v1/reports/deadlines"))
    _review(
        j,
        deadline,
        action,
        **({"reason": "延期", "new_document_id": newer.id} if action == "change" else {}),
    )
    after = _json(j.client.get("/api/v1/reports/deadlines"))
    assert after["total"] == 0
    assert _json(j.client.get("/api/v1/calendar"))["items"] == []
    if action == "change":
        assert before["total"] == 1 and before != after
    else:
        # pending -> rejected 的输出都是空，但数据库版本必须变化。
        assert len(v1._REPORT_CACHE) == 2


def test_warmed_list_board_reflects_reparse_and_kind_decisions(journey):
    j = journey
    source = _source(j.session)
    doc, _ = _document(j.session, source, "高新技术企业认定名单公布", "list-revision")
    j.session.commit()
    classification = classify_title(doc.title)
    save_list_parse(
        j.session,
        doc,
        doc.url,
        "html",
        classification,
        [
            ParsedEntry(
                name=COMPANY,
                normalized_name=normalize_name(COMPANY),
                credit_code=None,
                sequence="1",
                region=None,
            )
        ],
    )
    endpoint = "/api/v1/reports/lists"
    before = _json(j.client.get(endpoint))
    assert before["recent"]["items"][0]["count"] == 1
    # 模拟解析工作进程重写名单；独立 session 提交，API 不执行进程内清缓存。
    with j.session.get_bind().connect() as connection:
        from sqlalchemy.orm import Session

        with Session(connection) as worker:
            save_list_parse(
                worker,
                doc,
                doc.url,
                "html",
                classification,
                [
                    ParsedEntry(
                        name=name,
                        normalized_name=normalize_name(name),
                        credit_code=None,
                        sequence=str(i),
                        region=None,
                    )
                    for i, name in enumerate(
                        ("北京新增一号科技有限公司", "北京新增二号科技有限公司"), 1
                    )
                ],
            )
    after = _json(j.client.get(endpoint))
    assert after["recent"]["items"][0]["count"] == 2
    _json(
        j.client.post(
            f"/api/admin/documents/{doc.id}/kind",
            headers=ADMIN,
            json={"kind": "政策", "decided_by": REVIEWER},
        )
    )
    assert _json(j.client.get(endpoint))["recent"]["total"] == 0
    _json(j.client.post(f"/api/admin/documents/{doc.id}/kind/undo", headers=ADMIN))
    assert _json(j.client.get(endpoint)) == after


@pytest.mark.parametrize("status", ["failed", "expired", "refused"])
def test_delete_historical_subscription_preserves_other_user(journey, status):
    j = journey
    _agree(j)
    f = _follow(j, "历史授权", {"q": "历史", "time": "不限"})
    sub = _subscribe(j, "newDocs", followId=f["id"])
    row = j.session.get(Subscription, sub["id"])
    row.status = status
    j.session.commit()
    other = {"X-WX-SOURCE": "miniapp", "X-WX-OPENID": "journey-other-user"}
    _json(
        j.client.post(
            "/api/v1/me/agreement", headers=other, json={"version": j.settings.agreement_version}
        )
    )
    assert j.client.delete(f"/api/v1/me/subscriptions/{row.id}", headers=other).status_code == 404
    assert j.client.delete(f"/api/v1/me/subscriptions/{row.id}").status_code == 204
    assert _json(j.client.get("/api/v1/me/data"))["subscriptionsTotal"] == 0


def test_calendar_closed_boundaries_filters_and_pagination(journey):
    j = journey
    source = _source(j.session)
    due30, _ = _document(j.session, source, "30天前申报通知", "past30", -30, "approved")
    due1, _ = _document(j.session, source, "昨日申报通知", "past1", -1, "approved")
    _document(j.session, source, "31天前申报通知", "past31", -31, "approved")
    _document(j.session, source, "未核对申报通知", "unverified", -1)
    hidden, _ = _document(j.session, source, "下架申报通知", "hidden-past", -1, "approved")
    hidden.taken_down_at = NOW
    j.session.commit()
    first = _json(j.client.get("/api/v1/calendar", params={"limit": 1}))
    second = _json(j.client.get("/api/v1/calendar", params={"limit": 1, "offset": 1}))
    assert first["total"] == 0 and first["items"] == []
    assert first["pastTotal"] == second["pastTotal"] == 2
    assert [x["docId"] for x in first["past"] + second["past"]] == [str(due1.id), str(due30.id)]


@pytest.mark.parametrize("report", ["deadlines", "lists"])
def test_report_detail_selection_refreshes_after_ttl(journey, report, monkeypatch):
    j = journey
    source = _source(j.session)
    title = "合成申报通知" if report == "deadlines" else "合成认定名单公布"
    master, _ = _document(j.session, source, title, "detail-master", 7, "approved")
    twin, _ = _document(j.session, source, title, "detail-twin")
    master_detail = j.session.get(DocumentDetail, master.id)
    twin_detail = j.session.get(DocumentDetail, twin.id)
    master_detail.issuer, master_detail.doc_number = "合成主机关", "主〔2026〕1号"
    twin_detail.issuer, twin_detail.doc_number = "合成副机关", "副〔2026〕2号"
    master_detail.body_text = "主" * 100
    twin_detail.body_text = "副" * 10
    j.session.add(
        DocumentMerge(
            duplicate_id=twin.id,
            canonical_id=master.id,
            active=True,
            match_key="synthetic-detail-selection",
        )
    )
    j.session.commit()
    endpoint = f"/api/v1/reports/{report}"
    before = _json(j.client.get(endpoint))
    items = before["items"] if report == "deadlines" else before["recent"]["items"]
    assert items[0]["issuer"] == "合成主机关"
    # 正文/机关变化不参与每次请求的版本聚合，按 300 秒 TTL 更新展示。
    twin_detail.body_text = "副" * 200
    j.session.commit()
    assert _json(j.client.get(endpoint)) == before
    elapsed = v1._time.monotonic() + v1.REPORT_CACHE_TTL + 1
    monkeypatch.setattr(v1._time, "monotonic", lambda: elapsed)
    after = _json(j.client.get(endpoint))
    items = after["items"] if report == "deadlines" else after["recent"]["items"]
    assert items[0]["issuer"] == "合成副机关"
    if report == "deadlines":
        assert items[0]["doc_number"] == "副〔2026〕2号"
    assert before != after


@pytest.mark.parametrize("report", ["deadlines", "lists"])
@pytest.mark.parametrize("hidden", [False, True])
def test_publication_revision_uses_three_aggregate_queries(journey, report, hidden):
    j = journey
    source = _source(j.session)
    doc, _ = _document(j.session, source, "聚合申报通知", "aggregate", 7, "approved")
    twin, _ = _document(j.session, source, "聚合申报通知", "aggregate-twin")
    j.session.add(
        DocumentMerge(duplicate_id=twin.id, canonical_id=doc.id, active=True, match_key="aggregate")
    )
    j.session.commit()
    if hidden:
        change_state(
            j.session, doc.id, action="takedown", note="合成测试", operator=REVIEWER, reason="其他"
        )

    statements = []

    def capture(connection, cursor, statement, parameters, context, executemany):
        statements.append((statement.lower(), context.compiled.statement))

    engine = j.session.get_bind()
    event.listen(engine, "before_cursor_execute", capture)
    try:
        first = publication_revision(j.session, report)
        assert len(statements) == 3
        assert publication_revision(j.session, report) == first
        assert len(statements) == 6
    finally:
        event.remove(engine, "before_cursor_execute", capture)

    for sql, query in statements:
        assert "list_entries" not in sql and "document_details" not in sql
        # 每条顶层 SELECT 的所有投影都必须是聚合（包括 COALESCE(SUM(...), 0)）。
        for column in query.selected_columns:
            assert re.search(r"\b(count|sum|max|min)\s*\(", str(column).lower()), sql
        # 隐藏归并组子查询可带 WHERE，但不能藏一个无聚合、无限定的整表读取。
        for node in visitors.iterate(query):
            if isinstance(node, Select):
                aggregate = any(
                    re.search(r"\b(count|sum|max|min)\s*\(", str(c).lower())
                    for c in node.selected_columns
                )
                assert (
                    aggregate or node.whereclause is not None or node._limit_clause is not None
                ), sql


def test_revision_detects_withdrawal_of_older_review(journey):
    j = journey
    source = _source(j.session)
    _, older = _document(j.session, source, "较早申报通知", "older", 7, "approved")
    _, latest = _document(j.session, source, "较新申报通知", "latest", 9, "approved")
    older.reviewed_at = NOW - timedelta(minutes=1)
    j.session.commit()
    before = publication_revision(j.session, "deadlines")
    older.status, older.reviewed_at, older.approved_date = "pending", None, None
    j.session.commit()
    assert latest.reviewed_at == NOW
    assert publication_revision(j.session, "deadlines") != before


@pytest.mark.parametrize(
    "field,value",
    [("status", "failed"), ("entry_count", 2), ("attempted_at", NOW - timedelta(hours=1))],
)
def test_list_revision_tracks_older_parse_changes(journey, field, value):
    j = journey
    source = _source(j.session)
    doc, _ = _document(j.session, source, "旧解析名单公布", "parse-old")
    twin, _ = _document(j.session, source, "新解析名单公布", "parse-new")
    older = ListParse(
        document_id=doc.id,
        source_url=doc.url,
        source_kind="html",
        status="ok",
        entry_count=1,
        attempted_at=NOW - timedelta(hours=2),
    )
    j.session.add_all(
        [
            older,
            ListParse(
                document_id=twin.id,
                source_url=twin.url,
                source_kind="html",
                status="ok",
                entry_count=1,
                attempted_at=NOW,
            ),
        ]
    )
    j.session.commit()
    before = publication_revision(j.session, "lists")
    setattr(older, field, value)
    j.session.commit()
    assert j.session.scalar(select(func.max(ListParse.attempted_at))) == NOW
    assert publication_revision(j.session, "lists") != before
