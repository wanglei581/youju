"""转发组对外展示（D-055 第二步）：只有确认过的组才按组折叠、按组计数；截止按版本标，不互相套用。"""

from datetime import UTC, date, datetime, time
from io import BytesIO
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from youju.api import app as admin
from youju.api import v1
from youju.db import get_session
from youju.dedup import reconcile_documents
from youju.models import (
    Document,
    DocumentDeadline,
    DocumentDetail,
    RepostGroup,
    RepostGroupMember,
    Source,
)
from youju.repost import member_facts, rebuild

# 北京时间 2026-09-27（周日）；本周是 9/21–9/27
NOW = datetime(2026, 9, 27, 10, 0, tzinfo=ZoneInfo("Asia/Shanghai"))
HEADERS = {"X-Admin-Token": "repost-public-token"}
UPSTREAM = (
    "根据《工业和信息化部等七部门办公厅（办公室）关于组织开展2026年度视听系统典型案例征集工作的通知》"
    "（工信厅联电子函〔2026〕376号）要求，现就有关工作通知如下。"
)
TITLE = "关于组织开展2026年度视听系统典型案例征集工作的通知"


@pytest.fixture
def client(session_factory, monkeypatch):
    def override():
        with session_factory() as s:
            yield s

    monkeypatch.setattr(v1, "now", lambda: NOW)
    monkeypatch.setattr(
        admin,
        "get_settings",
        lambda: SimpleNamespace(
            role="all", admin_token=HEADERS["X-Admin-Token"], advisor_enabled=False
        ),
    )
    admin.app.dependency_overrides[get_session] = override
    with TestClient(admin.app) as test_client:
        yield test_client
    admin.app.dependency_overrides.clear()


def _source(session, province, dept="工信"):
    src = Source(
        level="国家级" if province == "国家" else "省级",
        province=province,
        dept=dept,
        dept_full_name="工业和信息化部" if province == "国家" else f"{province}省工业和信息化厅",
        column_name="通知公告",
        url=f"https://{province}-{dept}.example.gov.cn/tzgg/",
        enabled=True,
        parser_config={},
        rights_level=1,
    )
    session.add(src)
    session.flush()
    return src


def _doc(session, province, title=TITLE, day=22, body=UPSTREAM):
    src = _source(session, province)
    stamp = datetime.now(UTC)
    doc = Document(
        source_id=src.id,
        url=f"https://{province}.example.gov.cn/{day}.html",
        url_key=f"{province}.example.gov.cn/{day}.html",
        title=title,
        publish_date=date(2026, 9, day),
        first_seen_at=stamp,
        last_seen_at=stamp,
    )
    session.add(doc)
    session.flush()
    session.add(
        DocumentDetail(
            document_id=doc.id,
            body_text=body + province,
            fetched_at=stamp,
            status="ok",
            http_status=200,
        )
    )
    session.flush()
    return doc


def _deadline(session, doc, kind, day, status="approved", clock=None):
    row = DocumentDeadline(
        document_id=doc.id,
        fingerprint=f"{doc.id}-{kind}-{day}",
        kind=kind,
        candidate_date=date(2026, 10, day),
        candidate_time=clock,
        evidence=f"请于10月{day}日前报送",
        confidence="高",
        status=status,
    )
    if status == "approved":
        row.reviewed_by = "创始人"
        row.reviewed_at = datetime(2026, 9, 26, tzinfo=UTC)
        row.approved_date = row.candidate_date
        row.approved_time = clock
    session.add(row)
    session.flush()
    return row


@pytest.fixture
def world(session):
    """视听系统典型案例三个省的版本（同一份部里通知），各自的截止不同；另有一份不相干的通知。"""
    docs = {
        "shaanxi": _doc(session, "陕西", day=22),
        "tianjin": _doc(session, "天津", day=23),
        "fujian": _doc(
            session,
            "福建",
            "福建省工业和信息化厅等六部门关于开展2026年度视听系统典型案例征集工作的通知",
            day=24,
        ),
        "jiangsu": _doc(session, "江苏", "关于组织申报2026年度省级工业设计中心的通知", 25, "江苏"),
    }
    _deadline(session, docs["shaanxi"], "推荐或审核截止", 10)
    _deadline(session, docs["tianjin"], "企业申报截止", 12, clock=time(16, 0))
    _deadline(session, docs["fujian"], "推荐或审核截止", 15, status="pending")
    session.commit()
    rebuild(session)
    session.commit()
    return docs


def _confirm(session):
    rows = session.scalars(select(RepostGroupMember)).all()
    _, fingerprints = member_facts(session, [r.document_id for r in rows])
    for row in rows:
        row.confirmed_by, row.confirmed_at = "创始人", datetime.now(UTC)
        row.confirmed_fingerprint = fingerprints[row.document_id]
    session.commit()


def _search(client, **params):
    response = client.get("/api/v1/search", params={"time": "不限", **params})
    assert response.status_code == 200
    return response.json()


def test_unconfirmed_group_is_not_folded(client, world) -> None:
    """D-055：所有组都要人工确认后才对外合并；没确认的组仍按单条显示、按条计数。"""
    body = _search(client, group=1)
    assert body["total"] == 4 and body["reposts"] == 0
    assert not any("group" in item for item in body["items"])
    stats = client.get("/api/v1/home", params={"group": 1}).json()["stats"]
    assert stats["week_total"] == 4 and stats["week_reposts"] == 0
    detail = client.get(f"/api/v1/documents/{world['tianjin'].id}").json()
    assert "repost" not in detail


def test_confirmed_group_folds_and_counts_by_group(client, session, world) -> None:
    _confirm(session)
    body = _search(client, group=1)
    # 一个组算 1 份：三个省的版本折叠成一条主条目（最早发布的陕西版）
    assert body["total"] == 2 and body["total_all"] == 2 and body["reposts"] == 2
    main = next(item for item in body["items"] if item["id"] == str(world["shaanxi"].id))
    assert main["group"]["versions"] == 2
    assert main["group"]["regions"] == ["天津", "福建"]
    assert [v["label"] for v in main["group"]["items"]] == ["天津版", "福建版"]
    # 每个版本各自的原文链接
    assert main["group"]["items"][0]["url"] == world["tianjin"].url
    # 翻页按主条目：翻到底的条数 = total
    page = _search(client, group=1, limit=1, offset=1)
    assert len(page["items"]) == 1 and page["total"] == 2
    # 首页本周新增按组算：本周有新版本的组计 1 份，另写并进组里的版本数
    home = client.get("/api/v1/home", params={"group": 1}).json()
    assert home["stats"]["week_total"] == 2 and home["stats"]["apply"] == 2
    assert home["stats"]["week_reposts"] == 2
    assert len(home["items"]) == 2
    # 旧版小程序不传 group：返回和以前一样（按条计数、不带 group、reposts）
    old = _search(client)
    assert old["total"] == 4 and "reposts" not in old
    assert not any("group" in item for item in old["items"])
    old_home = client.get("/api/v1/home").json()
    assert old_home["stats"]["week_total"] == 4 and "week_reposts" not in old_home["stats"]
    lists = client.get("/api/v1/lists/recent", params={"group": 1}).json()
    assert lists["total"] == 0


def test_filter_first_then_fold(client, session, world) -> None:
    """先筛选、再折叠：选了福建，主条目是福建版；“另有 N 个”数的是组里其他全部版本。"""
    _confirm(session)
    body = _search(client, group=1, regions="福建")
    assert body["total"] == 1
    item = body["items"][0]
    assert item["id"] == str(world["fujian"].id)
    assert item["group"]["versions"] == 2 and item["group"]["regions"] == ["陕西", "天津"]
    # 关键词只命中某一版时也一样
    body = _search(client, group=1, q="六部门")
    assert body["total"] == 1 and body["items"][0]["group"]["versions"] == 2


def test_deadlines_stay_per_version(client, session, world) -> None:
    """不同版本的截止不互相套用：主条目、详情都只用这一版自己核对过的截止，其他版本按版本标。"""
    _confirm(session)
    main = next(
        item for item in _search(client, group=1)["items"] if item["id"] == str(world["shaanxi"].id)
    )
    assert main["deadline"]["date"] == "2026-10-10"  # 陕西版自己的，不是天津版的 10-12
    assert main["app_status"]["label"] == "推荐或审核截止"

    fujian = client.get(f"/api/v1/documents/{world['fujian'].id}").json()
    # 福建版自己的候选还没核对：不借用陕西、天津已通过的日期
    assert fujian["deadline"] is None and fujian["deadline_state"] == "unverified"
    repost = fujian["repost"]
    assert repost["label"] == "福建版"
    assert repost["upstream"] == {
        "title": "工业和信息化部等七部门办公厅（办公室）关于组织开展2026年度视听系统典型案例征集工作的通知",
        "number": "工信厅联电子函〔2026〕376号",
    }
    versions = {v["label"]: v for v in repost["versions"]}
    assert [v["label"] for v in repost["versions"]] == ["福建版", "陕西版", "天津版"]
    assert versions["福建版"]["current"] is True
    assert versions["福建版"]["deadlines"] == []
    assert versions["福建版"]["deadline_state"] == "unverified"
    assert versions["陕西版"]["deadlines"] == [
        {"label": "推荐或审核截止", "date": "2026-10-10", "time": ""}
    ]
    assert versions["天津版"]["deadlines"] == [
        {"label": "企业申报截止", "date": "2026-10-12", "time": "16:00"}
    ]
    assert versions["天津版"]["url"] == world["tianjin"].url
    # 库里没有部里原件：不出现“部委”
    assert all(v["label"] != "部委" for v in repost["versions"])

    tianjin = client.get(f"/api/v1/documents/{world['tianjin'].id}").json()
    assert tianjin["deadline"]["label"] == "企业申报截止"
    assert tianjin["deadline"]["date"] == "2026-10-12" and tianjin["deadline"]["notes"] == []


def test_national_original_is_main_and_called_ministry(client, session, world) -> None:
    national = _doc(
        session,
        "国家",
        "工业和信息化部等七部门办公厅（办公室）关于组织开展2026年度视听系统典型案例征集工作的通知",
        day=26,
        body="各有关单位：现组织开展2026年度视听系统典型案例征集工作。",
    )
    _deadline(session, national, "企业申报截止", 20)
    session.commit()
    rebuild(session)
    session.commit()
    _confirm(session)
    body = _search(client, group=1)
    # 部里原件在库就用原件做主条目（即使它发布得晚）
    main = next(item for item in body["items"] if item.get("group"))
    assert main["id"] == str(national.id) and main["group"]["versions"] == 3
    detail = client.get(f"/api/v1/documents/{world['shaanxi'].id}").json()
    first = detail["repost"]["versions"][0]
    assert first["label"] == "部委" and first["deadlines"][0]["label"] == "企业申报截止"
    assert detail["repost"]["versions"][1]["label"] == "陕西版"


def _admin(client, path, **body):
    return client.post(
        f"/api/admin/repost-groups{path}",
        headers=HEADERS,
        json={"reviewed_by": "创始人", **body},
    )


def _seen(client, group_id):
    """管理台页面上看到的这一组：修订号和在组版本。"""
    for state in ("pending", "confirmed", "removed"):
        for group in client.get(
            "/api/admin/repost-groups", params={"state": state}, headers=HEADERS
        ).json():
            if group["id"] == group_id:
                ids = [m["document_id"] for m in group["members"] if m["active"]]
                return {"revision": group["revision"], "document_ids": ids}
    raise AssertionError(group_id)


def test_withdraw_and_remove_restore_single_items(client, session, world) -> None:
    group_id = session.scalars(select(RepostGroupMember.group_id)).first()
    assert _admin(client, f"/{group_id}/confirm", **_seen(client, group_id)).status_code == 200
    assert _search(client, group=1)["total"] == 2
    # 撤回确认后马上恢复成单条
    revision = _seen(client, group_id)["revision"]
    assert _admin(client, f"/{group_id}/withdraw", revision=revision).status_code == 200
    body = _search(client, group=1)
    assert body["total"] == 4 and not any("group" in item for item in body["items"])
    # 再确认、移出福建版：福建版单独一条，另两版仍是一组
    assert _admin(client, f"/{group_id}/confirm", **_seen(client, group_id)).status_code == 200
    remove = _admin(client, f"/{group_id}/members/{world['fujian'].id}/remove")
    assert remove.status_code == 200
    body = _search(client, group=1)
    assert body["total"] == 3
    fujian = next(item for item in body["items"] if item["id"] == str(world["fujian"].id))
    assert "group" not in fujian
    assert "repost" not in client.get(f"/api/v1/documents/{world['fujian'].id}").json()


def test_grouped_fold_counts_groups_actually_hidden(client, session) -> None:
    """一个转发组里同时有“资讯与指南”和“其他”的版本：折叠资讯与指南后这一组仍以“其他”那版列出，
    folded 不把它算进去（按实际少掉的条数算）。"""
    title = "关于开展人工智能应用服务商培育专项行动的通知"
    docs = [_doc(session, p, title, day, "正文") for p, day in (("黑龙江", 22), ("四川", 23))]
    guide = _doc(session, "陕西", "关于2026年度工业设计大赛培训班报名的通知", 24, "正文")
    docs[1].kind_override = "资讯与指南"
    guide.kind_override = "资讯与指南"
    session.commit()
    rebuild(session)
    session.commit()
    _confirm(session)
    assert v1.effective_kind(docs[1].title, docs[1].kind_override) == "资讯与指南"
    assert v1.effective_kind(guide.title, guide.kind_override) == "资讯与指南"
    shown = _search(client, group=1, fold=0)
    assert shown["total"] == 2 and shown["guides"] == 1 and shown["folded"] == 0
    folded = _search(client, group=1, fold=1)
    # 只有单独那份资讯与指南被折叠；转发组仍以“其他”那版（黑龙江）列出
    assert folded["folded"] == 1 and folded["guides"] == 1 and folded["total"] == 1
    assert folded["items"][0]["id"] == str(docs[0].id)
    assert shown["total"] - folded["total"] == folded["folded"]


def test_taken_down_version_is_not_listed(client, session, world) -> None:
    _confirm(session)
    world["tianjin"].taken_down_at = datetime.now(UTC)
    session.commit()
    main = next(
        item for item in _search(client, group=1)["items"] if item["id"] == str(world["shaanxi"].id)
    )
    assert main["group"]["regions"] == ["福建"]


def test_same_url_different_issuers_stay_separate(client, session) -> None:
    """网址相同但机关不同仍按 main 各自保留，对外算两份。"""
    src = _source(session, "辽宁")
    stamp = datetime.now(UTC)
    url = "https://gxt.ln.gov.cn/tzgg/2026092511/index.html"
    docs = [
        Document(
            source_id=src.id,
            url=url,
            url_key=key,
            title="辽宁省工业和信息化厅关于2026年度先进制造业企业名单制定工作有关事项的通知",
            publish_date=date(2026, 9, 25),
            first_seen_at=stamp,
            last_seen_at=stamp,
        )
        for key in ("gxt.ln.gov.cn/tzgg/2026092511/index.html", "gxt.ln.gov.cn/tzgg/2026092511/")
    ]
    session.add_all(docs)
    session.flush()
    session.add(
        DocumentDetail(
            document_id=docs[0].id,
            body_text="",
            issuer="辽宁省工业和信息化厅、辽宁省通信管理局",
            fetched_at=stamp,
            status="ok",
        )
    )
    session.commit()
    assert _search(client)["total"] == 2
    reconcile_documents(session)
    session.commit()
    assert _search(client)["total"] == 2 and _search(client, group=1)["total"] == 2


@pytest.mark.parametrize(
    "change", ["date", "missing_date", "agency", "issuer", "title", "body_tail", "span"]
)
def test_changed_confirmed_group_stops_folding_until_reconfirmed(
    client, session, world, change
) -> None:
    """确认快照覆盖完整正文；任一版本变化，全部停止折叠、按单条计数并进入待办。"""
    if change == "body_tail":
        session.get(DocumentDetail, world["fujian"].id).body_text += "原正文内容" * 400
        session.commit()
    group = session.scalars(select(RepostGroup)).one()
    ids = [r.document_id for r in session.scalars(select(RepostGroupMember))]
    response = client.post(
        f"/api/admin/repost-groups/{group.id}/confirm",
        headers=HEADERS,
        json={"reviewed_by": "创始人", "revision": group.revision, "document_ids": ids},
    )
    assert response.status_code == 200
    session.expire_all()
    assert all(r.confirmed_fingerprint for r in session.scalars(select(RepostGroupMember)))
    assert _search(client, group=1)["total"] == 2
    original_date = world["fujian"].publish_date
    original_source = world["fujian"].source_id
    if change == "date":
        world["fujian"].publish_date = date(2026, 9, 25)
    elif change == "missing_date":
        world["fujian"].publish_date = None
    elif change == "agency":
        world["fujian"].source_id = world["tianjin"].source_id
    elif change == "issuer":
        session.get(DocumentDetail, world["fujian"].id).issuer = "福建省联合发文机关"
    elif change == "title":
        world["fujian"].title += "（更正）"
    elif change == "body_tail":
        detail = session.get(DocumentDetail, world["fujian"].id)
        detail.body_text += "正文尾部的更正"
    else:
        world["fujian"].publish_date = date(2027, 1, 1)
    session.commit()
    rebuild(session)
    session.commit()
    session.refresh(group)
    assert group.needs_review is True
    body = _search(client, group=1)
    assert body["total"] == 4 and body["reposts"] == 0
    assert not any("group" in item for item in body["items"])
    home = client.get("/api/v1/home?group=1").json()["stats"]
    expected_week = 3 if change in ("span", "missing_date") else 4
    assert home["week_total"] == expected_week and home["week_reposts"] == 0
    assert "repost" not in client.get(f"/api/v1/documents/{world['shaanxi'].id}").json()
    assert client.get("/api/admin/repost-groups/count", headers=HEADERS).json() == {"pending": 1}
    assert client.get("/api/admin/overview", headers=HEADERS).json()["todos"]["repost"] == 1
    pending = client.get("/api/admin/repost-groups", headers=HEADERS).json()[0]
    assert pending["needs_review"] and not pending["public"] and pending["pending"] == 3
    assert client.get("/api/admin/repost-groups?state=confirmed", headers=HEADERS).json() == []
    # 数据恢复也不能自动恢复确认；人重新确认后才继续折叠。
    world["fujian"].publish_date = original_date
    world["fujian"].source_id = original_source
    session.commit()
    rebuild(session)
    session.commit()
    session.refresh(group)
    assert group.needs_review
    response = client.post(
        f"/api/admin/repost-groups/{group.id}/confirm",
        headers=HEADERS,
        json={"reviewed_by": "创始人复核", "revision": group.revision, "document_ids": ids},
    )
    assert response.status_code == 200
    assert _search(client, group=1)["total"] == 2
    rebuild(session)
    session.commit()
    session.refresh(group)
    assert not group.needs_review


@pytest.mark.parametrize("state", ["pending", "confirmed", "needs_review", "withdrawn"])
def test_group_actions_never_copy_deadline_approval_in_any_output(
    client, session, world, state
) -> None:
    """同类型、日期、证据也不能跨版本复制审批；所有公开出口只读取本版本已核对日期。"""
    group = session.scalars(select(RepostGroup)).one()
    ids = [r.document_id for r in session.scalars(select(RepostGroupMember))]
    a = session.scalars(
        select(DocumentDeadline).where(DocumentDeadline.document_id == world["shaanxi"].id)
    ).one()
    for doc in (world["tianjin"], world["fujian"]):
        row = session.scalars(
            select(DocumentDeadline).where(DocumentDeadline.document_id == doc.id)
        ).one()
        row.kind, row.candidate_date, row.candidate_time, row.evidence = (
            a.kind,
            a.candidate_date,
            a.candidate_time,
            a.evidence,
        )
        row.status, row.approved_date, row.approved_time, row.reviewed_by, row.reviewed_at = (
            "pending",
            None,
            None,
            None,
            None,
        )
    session.commit()
    if state != "pending":
        response = client.post(
            f"/api/admin/repost-groups/{group.id}/confirm",
            headers=HEADERS,
            json={"reviewed_by": "创始人", "revision": group.revision, "document_ids": ids},
        )
        assert response.status_code == 200
        if state == "needs_review":
            world["tianjin"].title += "（更正）"
            session.commit()
        elif state == "withdrawn":
            session.refresh(group)
            response = client.post(
                f"/api/admin/repost-groups/{group.id}/withdraw",
                headers=HEADERS,
                json={"reviewed_by": "创始人", "revision": group.revision},
            )
            assert response.status_code == 200
    rebuild(session)
    reconcile_documents(session)
    from youju.dedup import sync_reviewed_twins

    sync_reviewed_twins(session, ids)
    session.commit()
    for doc in (world["tianjin"], world["fujian"]):
        detail = client.get(f"/api/v1/documents/{doc.id}").json()
        assert detail["deadline"] is None and detail["deadline_state"] == "unverified"
        row = session.scalars(
            select(DocumentDeadline)
            .where(DocumentDeadline.document_id == doc.id)
            .execution_options(populate_existing=True)
        ).one()
        assert row.status == "pending" and row.approved_date is None and row.reviewed_by is None
    calendar = client.get("/api/v1/calendar").json()
    assert [i["docId"] for i in calendar["items"]] == [str(world["shaanxi"].id)]
    v1._REPORT_CACHE.clear()
    radar = client.get("/api/v1/reports/deadlines?days=30").json()
    assert [i["id"] for i in radar["items"]] == [str(world["shaanxi"].id)]
    exported = client.post("/api/v1/exports", json={"ids": [str(i) for i in ids], "format": "xlsx"})
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.content))["文件清单"]
    assert sheet.max_row == 4
    assert "2026-10-10" in sheet["I2"].value
    assert [sheet.cell(i, 9).value for i in (3, 4)] == ["见原文", "见原文"]


def test_confirm_new_version_cannot_overwrite_changed_confirmed_snapshot(
    client, session, world
) -> None:
    """补详情先改了旧版本、归组尚未重跑时，确认新版本也必须先触发复核、刷新后再确认。"""
    group_id = session.scalars(select(RepostGroup.id)).one()
    assert _admin(client, f"/{group_id}/confirm", **_seen(client, group_id)).status_code == 200
    newcomer = _doc(session, "北京", day=25)
    session.commit()
    rebuild(session)
    session.commit()
    seen = _seen(client, group_id)
    old_fingerprint = session.get(RepostGroupMember, world["tianjin"].id).confirmed_fingerprint
    world["tianjin"].title += "（更正）"
    session.commit()
    assert _admin(client, f"/{group_id}/confirm", **seen).status_code == 409
    session.expire_all()
    assert session.get(RepostGroup, group_id).needs_review
    assert session.get(RepostGroupMember, newcomer.id).confirmed_at is None
    assert (
        session.get(RepostGroupMember, world["tianjin"].id).confirmed_fingerprint == old_fingerprint
    )
    assert _search(client, group=1)["total"] == 5
    assert _admin(client, f"/{group_id}/confirm", **_seen(client, group_id)).status_code == 200
    assert _search(client, group=1)["total"] == 2
