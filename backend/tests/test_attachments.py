"""附件副本：定时任务只存允许的附件、校验格式和个人信息、更新和清理；api 分段读取、绝不碰备份。
对象存储、开放接口和政府网站全部用模拟，不联网。"""

import io
import json
import random
import re
import struct
import zipfile
import zlib
from collections import Counter
from datetime import UTC, date, datetime, timedelta
from hashlib import sha256
from pathlib import Path
from types import SimpleNamespace

import httpx
import openpyxl
import pytest
from alembic import command
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from sqlalchemy import inspect, select
from sqlalchemy import text as sql

from youju import attachments as copies
from youju.api import attachments as attachment_api
from youju.api import v1
from youju.api.app import app
from youju.crawler.fetcher import FetchResult
from youju.db import get_session
from youju.migrations import alembic_config
from youju.models import AttachmentCopy, Document, DocumentAttachment, DocumentDetail, Source

FIXTURES = Path(__file__).parent / "fixtures"
NOW = datetime(2026, 10, 6, 2, 0, tzinfo=UTC)
ENV = "prod-1"
PREFIX = "cloud://prod-1.7072-prod-1-1258/"
BACKUP_ID = PREFIX + "backups/youju-20261005.jsonl.gz"


# ---- 测试用的文件 ----


def make_pdf(text: str | None = None, images: tuple = ()) -> bytes:
    """最小的 PDF（一页，300×200）：给 text 时有这行文字，不给就是空白页（像扫描件，取不出文字）；
    images 里每个 (x, y, 宽, 高) 在页面上画一张位图。"""
    content = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode() if text else b""
    for x, y, w, h in images:
        content += f" q {w} 0 0 {h} {x} {y} cm /Im1 Do Q".encode()
    pixels = bytes([0, 128, 128, 255])
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> /XObject << /Im1 6 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /XObject /Subtype /Image /Width 2 /Height 2 /ColorSpace /DeviceGray"
        b" /BitsPerComponent 8 /Length 4 >>\nstream\n" + pixels + b"\nendstream",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, obj in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return out


def make_docx(*lines: str) -> bytes:
    doc = WordDocument()
    for line in lines:
        doc.add_paragraph(line)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def make_xlsx(*rows: tuple) -> bytes:
    book = openpyxl.Workbook()
    for row in rows:
        book.active.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def make_pptx(text: str) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("ppt/presentation.xml", "<p:presentation/>")
        zf.writestr("ppt/slides/slide1.xml", f"<p:sld><a:p><a:t>{text}</a:t></a:p></p:sld>")
    return out.getvalue()


XLS = (FIXTURES / "lists" / "heilongjiang_30.xls").read_bytes()


def make_ole(stream: str) -> bytes:
    """把真实 xls 的主数据流改名，得到旧版 Word（WordDocument）或 PowerPoint 的结构。"""
    data = bytearray(XLS)
    at = data.find("Workbook".encode("utf-16-le") + b"\x00\x00")
    name = (stream + "\x00").encode("utf-16-le")
    data[at : at + 64] = name.ljust(64, b"\x00")
    data[at + 64 : at + 66] = len(name).to_bytes(2, "little")
    return bytes(data)


def make_png(side: int, seed: int = 1) -> bytes:
    """随机像素的 PNG（几乎压缩不了）：side=150 约 67KB，side=16 不到 1KB（像小图标）。"""
    from PIL import Image

    pixels = random.Random(seed).randbytes(side * side * 3)
    out = io.BytesIO()
    Image.frombytes("RGB", (side, side), pixels).save(out, "PNG")
    return out.getvalue()


def make_flat_png(width: int, height: int) -> bytes:
    """纯色 PNG：尺寸可以很大，文件却很小（像压缩得很小的扫描页）。"""
    from PIL import Image

    out = io.BytesIO()
    Image.new("L", (width, height), 255).save(out, "PNG")
    return out.getvalue()


def docx_with_picture(png: bytes, text: str = "申报指南") -> bytes:
    doc = WordDocument()
    doc.add_paragraph(text)
    doc.add_picture(io.BytesIO(png))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def with_parts(data: bytes, extra: dict[str, bytes]) -> bytes:
    """在 Office 文件（zip）里加几个部件。"""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, src.read(info))
        for name, body in extra.items():
            dst.writestr(name, body)
    return out.getvalue()


def make_cfb(streams: dict[str, bytes], storages: dict[str, dict[str, bytes]] | None = None):
    """最小的复合文档（旧版 Office 的容器，v3、512 字节扇区）：根目录下的数据流和一层子目录。
    每个数据流补到至少 4096 字节，都放在普通扇区里。"""
    end, free, fat_mark, none = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF
    nodes = [{"name": "Root Entry", "type": 5, "data": b"", "kids": []}]

    def add(name, kind, data, parent):
        nodes.append({"name": name, "type": kind, "data": data.ljust(4096, b"\0"), "kids": []})
        nodes[parent]["kids"].append(len(nodes) - 1)
        return len(nodes) - 1

    for name, data in streams.items():
        add(name, 2, data, 0)
    for name, inner in (storages or {}).items():
        folder = add(name, 1, b"", 0)
        nodes[folder]["data"] = b""
        for child, data in inner.items():
            add(child, 2, data, folder)
    n_dir = -(-len(nodes) * 128 // 512)
    data_sectors = sum(-(-len(n["data"]) // 512) for n in nodes if n["type"] == 2)
    n_fat = 1
    while n_fat * 128 < n_fat + n_dir + data_sectors:
        n_fat += 1
    fat = [fat_mark] * n_fat + [n_fat + i + 1 for i in range(n_dir - 1)] + [end]
    body = b""
    for node in nodes:
        if node["type"] != 2:
            node["start"] = end if node["type"] == 5 else 0
            continue
        node["start"] = len(fat)
        count = -(-len(node["data"]) // 512)
        fat += [len(fat) + i + 1 for i in range(count - 1)] + [end]
        body += node["data"].ljust(count * 512, b"\0")
    fat += [free] * (n_fat * 128 - len(fat))
    entries = b""
    for index, node in enumerate(nodes):
        kids = node["kids"]
        parent = next((n for n in nodes if index in n["kids"]), None)
        siblings = parent["kids"] if parent else []
        pos = siblings.index(index) if index in siblings else -1
        right = siblings[pos + 1] if 0 <= pos < len(siblings) - 1 else none
        name = (node["name"] + "\0").encode("utf-16-le")
        entries += name.ljust(64, b"\0") + struct.pack("<HBB", len(name), node["type"], 1)
        entries += struct.pack("<III", none, right, kids[0] if kids else none) + bytes(36)
        entries += struct.pack("<IQ", node["start"], len(node["data"]) if node["type"] == 2 else 0)
    entries = entries.ljust(n_dir * 512, b"\0")
    header = bytearray(512)
    header[:8] = bytes.fromhex("d0cf11e0a1b11ae1")
    struct.pack_into("<HHHHH", header, 24, 0x3E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<9I", header, 40, 0, n_fat, n_fat, 0, 4096, end, 0, end, 0)
    struct.pack_into("<109I", header, 76, *(list(range(n_fat)) + [free] * (109 - n_fat)))
    return bytes(header) + struct.pack(f"<{len(fat)}I", *fat) + entries + body


# 旧版 Office 里图片的记录头（[MS-ODRAW]）：PNG、DIB、EMF
PNG_BLIP = bytes.fromhex("006e1ef0")
DIB_BLIP = bytes.fromhex("807a1ff0")
EMF_BLIP = bytes.fromhex("403d1af0")


def bitmap_blip(head: bytes, image: bytes) -> bytes:
    """位图记录：16 字节编号 + 1 字节标记 + 图片文件本身。"""
    return blip(head, bytes(16) + b"\xff" + image)


def dib(width: int, height: int) -> bytes:
    return struct.pack("<IiiHH", 40, width, height, 1, 8) + bytes(24) + bytes(64)


def blip(head: bytes, payload: bytes) -> bytes:
    return head + len(payload).to_bytes(4, "little") + payload


def emf_blip(text: str) -> bytes:
    """一张压缩存放的小矢量图（EMF），里面的文字按 UTF-16LE 存。"""
    raw = b"\x01\x00\x00\x00" + text.encode("utf-16-le")
    packed = zlib.compress(raw)
    header = struct.pack("<I16s8sIBB", len(raw), bytes(16), bytes(8), len(packed), 0, 0xFE)
    return blip(EMF_BLIP, bytes(16) + header + packed)


PDF = make_pdf("Application guide 2026")
DOCX = make_docx("申报指南", "咨询电话：024-12345678")


def test_detect_type_and_content_checks() -> None:
    assert copies.detect_type(PDF) == "pdf"
    assert copies.detect_type(DOCX) == "docx"
    assert copies.detect_type(make_xlsx(("名称",))) == "xlsx"
    assert copies.detect_type(make_pptx("说明")) == "pptx"
    assert copies.detect_type(XLS) == "xls"
    assert copies.detect_type(make_ole("WordDocument")) == "doc"
    assert copies.detect_type(make_ole("PowerPoint Document")) == "ppt"
    for junk in (b"<html>404</html>", b"PK\x03\x04broken", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1x"):
        assert copies.detect_type(junk) is None
    plain_zip = io.BytesIO()
    with zipfile.ZipFile(plain_zip, "w") as zf:
        zf.writestr("a.txt", "x")
    assert copies.detect_type(plain_zip.getvalue()) is None

    # 同一类格式可以（.doc 实际是 docx 照样能打开，用实际格式）；不同类、网页都不行
    assert copies.check_content(DOCX, "doc") == "docx"
    assert copies.check_content(XLS, "xlsx") == "xls"
    assert copies.check_content(make_ole("WordDocument"), "doc") == "doc"
    assert copies.check_content(PDF, "docx").code == "mismatch"
    assert copies.check_content(b"<html>not found</html>", "pdf").code == "mismatch"
    assert copies.check_content(b"", "pdf").code == "empty"
    # 取不出文字的（空白页像扫描件）没法查个人信息，不存
    assert copies.check_content(make_pdf(), "pdf").code == "unchecked"
    # 身份证号、个人手机号、银行账号：附件没法脱敏，有就不存；办公电话不算
    assert copies.check_content(make_pdf("Contact 13812345678"), "pdf").code == "personal"
    assert copies.check_content(make_docx("联系人手机：13912345678"), "docx").code == "personal"
    assert copies.check_content(make_docx("身份证号 110105199001011234"), "docx").code == (
        "personal"
    )
    # 被拆成几段的手机号、存成数字的手机号也查得出
    split = make_docx("x")
    with zipfile.ZipFile(io.BytesIO(split)) as zf:
        parts = {n: zf.read(n) for n in zf.namelist()}
    parts["word/document.xml"] = parts["word/document.xml"].replace(
        b"<w:t>x</w:t>", b"<w:t>138</w:t></w:r><w:r><w:t>12345678</w:t>"
    )
    rebuilt = io.BytesIO()
    with zipfile.ZipFile(rebuilt, "w") as zf:
        for name, body in parts.items():
            zf.writestr(name, body)
    assert copies.check_content(rebuilt.getvalue(), "docx").code == "personal"
    assert copies.check_content(make_xlsx(("联系人", 13712345678)), "xlsx").code == "personal"
    assert copies.check_content(make_pptx("账号：6222020200112233445"), "pptx").code == "personal"


def test_personal_numbers_written_with_spaces_or_dashes_are_found() -> None:
    """手机号、身份证号中间带空格或横线，银行卡号前面没写"账号"，都查得出（附件没法脱敏，宁可多拦）。"""
    for text in (
        "联系人手机：138 1234 5678",
        "联系人手机：138-1234-5678",
        "联系人手机：138    1234    5678",
        "联系人手机：138\t1234\n5678",
        "联系人手机：138\u3000\u30001234 \uff0d  5678",
        "手机 +86 138 1234 5678",
        "电话：0571-87654321 138 1234 5678",
        "身份证号 110105 19900101 1234",
        "身份证号 1101 0519 9001 0112 3X",
        "身份证号 110105    19900101\n1234",
        "4111 1111 1111 1111",  # 没写"账号"、通过 Luhn 校验的 16 位
        "6222-0212-3456-7890-128",
        "4111   1111\n1111    1111",
        "收款 6212261202011584340",
    ):
        assert copies.check_content(make_docx("申报指南", text), "docx").code == "personal", text
    # 办公电话、日期、金额、通不过 Luhn 校验的长数字不算
    for text in (
        "咨询电话：024-12345678、0571-8765 4321",
        "申报时间 2026-10-07 至 2026-11-30，邮编 110001",
        "合计 120 35 18 173 万元",
        "项目编号 4111 1111 1111 1112",
    ):
        assert copies.check_content(make_docx("申报指南", text), "docx") == "docx", text


def test_office_files_with_large_pictures_are_not_copied() -> None:
    big, logo = make_png(150), make_png(16)
    assert len(big) > copies.OBJECT_BYTES and len(logo) < copies.BITMAP_BYTES
    image = copies.check_content(docx_with_picture(big), "docx")
    assert image.code == "image" and "无法核验个人信息" in image.detail
    # 小图标、logo 不算；位图超过 10KB、任一边超过 200 像素（文件再小也算）、认不出尺寸的都不存
    assert copies.check_content(docx_with_picture(logo), "docx") == "docx"
    medium = make_png(64)
    assert copies.BITMAP_BYTES < len(medium) < copies.OBJECT_BYTES
    assert copies.check_content(docx_with_picture(medium), "docx").code == "image"
    wide = make_flat_png(1200, 40)
    assert len(wide) < copies.BITMAP_BYTES
    assert "200 像素" in copies.check_content(docx_with_picture(wide), "docx").detail
    assert copies.check_content(docx_with_picture(make_flat_png(200, 200)), "docx") == "docx"
    unknown = with_parts(DOCX, {"word/media/image9.wdp": b"not a picture"})
    assert copies.check_content(unknown, "docx").code == "image"
    assert (
        copies.check_content(
            with_parts(make_xlsx(("名称",)), {"xl/media/image1.png": big}), "xlsx"
        ).code
        == "image"
    )
    assert (
        copies.check_content(
            with_parts(make_pptx("说明"), {"ppt/media/image1.jpeg": big}), "pptx"
        ).code
        == "image"
    )
    # 嵌入的对象（嵌入的表格、文档等）：大的查不了，小的取出里面的文字一起查
    embedded = with_parts(DOCX, {"word/embeddings/Workbook1.xlsx": big})
    assert copies.check_content(embedded, "docx").code == "image"
    people = make_xlsx(("联系人", "13712345678"))
    assert len(people) < copies.OBJECT_BYTES
    embedded = with_parts(DOCX, {"word/embeddings/Workbook1.xlsx": people})
    assert copies.check_content(embedded, "docx").code == "personal"
    # 嵌入的旧版表格按单元格取文字
    assert "黑龙江" in copies._embedded_text(XLS, nested=False)
    assert copies.check_content(with_parts(DOCX, {"word/embeddings/o1.bin": XLS}), "docx") == "docx"
    # 小的矢量图（如粘贴成图片的表格）：取出里面的文字一起查
    emf = b"\x01\x00\x00\x00" + "联系人 13812345678".encode("utf-16-le")
    assert (
        copies.check_content(with_parts(DOCX, {"word/media/image2.emf": emf}), "docx").code
        == "personal"
    )
    # 首页缩略图是正文首页的样子，不算
    thumbnail = with_parts(DOCX, {"docProps/thumbnail.jpeg": big})
    assert copies.check_content(thumbnail, "docx") == "docx"


def test_pdf_pages_mostly_covered_by_pictures_are_not_copied() -> None:
    text = "Application guide 2026"
    # 页面 300×200：任一张图片超过 1/10 页面就不存；公章、logo 这类小图不算
    assert copies.check_content(make_pdf(text, ((0, 0, 300, 150),)), "pdf").code == "image"
    single = copies.check_content(make_pdf(text, ((0, 0, 100, 70),)), "pdf")
    assert single.code == "image" and "1/10" in single.detail
    assert copies.check_content(make_pdf(text, ((10, 10, 40, 40),)), "pdf") == "pdf"
    assert copies.check_content(make_pdf(text, ((0, 0, 100, 55),)), "pdf") == "pdf"
    # 每张都不到 1/10、合计超过半页（切成条的扫描页）也不存
    tiles = tuple((x, y, 100, 55) for x in (0, 100, 200) for y in (0, 60))
    assert "合计" in copies.check_content(make_pdf(text, tiles), "pdf").detail
    # 超出页面的部分不算
    assert copies.check_content(make_pdf(text, ((290, 0, 300, 200),)), "pdf") == "pdf"
    # 几乎没有文字的页面有图片，仍按扫描件处理
    assert copies.check_content(make_pdf(None, ((10, 10, 40, 40),)), "pdf").code == "unchecked"


def test_old_office_files_with_large_pictures_are_not_copied() -> None:
    def word(**extra):
        return make_cfb({"WordDocument": b"Word 97", **extra.pop("streams", {})}, **extra)

    assert copies.detect_type(word()) == "doc"
    assert copies.check_content(word(), "doc") == "doc"
    big = word(streams={"Data": blip(PNG_BLIP, bytes(60_000))})
    assert copies.check_content(big, "doc").code == "image"
    # 小位图：不超过 10KB、两边都不超过 200 像素才放行；认不出尺寸的不放行
    small = word(streams={"Data": bitmap_blip(PNG_BLIP, make_png(16))})
    assert copies.check_content(small, "doc") == "doc"
    wide = word(streams={"Data": bitmap_blip(PNG_BLIP, make_flat_png(1200, 40))})
    assert copies.check_content(wide, "doc").code == "image"
    assert (
        copies.check_content(
            word(streams={"Data": bitmap_blip(DIB_BLIP, dib(300, 10))}), "doc"
        ).code
        == "image"
    )
    assert (
        copies.check_content(word(streams={"Data": bitmap_blip(DIB_BLIP, dib(16, -16))}), "doc")
        == "doc"
    )
    junk = word(streams={"Data": blip(PNG_BLIP, bytes(2_000))})
    assert copies.check_content(junk, "doc").code == "image"
    # 小的矢量图是压缩存的：解压后取出文字一起查
    phone = word(streams={"Data": emf_blip("联系人 13812345678")})
    assert copies.check_content(phone, "doc").code == "personal"
    assert copies.check_content(word(streams={"Data": emf_blip("申报流程图")}), "doc") == "doc"
    # 嵌入的对象：Word 放在 ObjectPool 下，大的查不了
    assert (
        copies.check_content(word(storages={"ObjectPool": {"_1": bytes(60_000)}}), "doc").code
        == "image"
    )
    assert copies.check_content(word(storages={"ObjectPool": {"_1": bytes(3_000)}}), "doc") == "doc"
    # PowerPoint 的图片在 Pictures 流里
    slides = make_cfb({"PowerPoint Document": b"ppt", "Pictures": blip(PNG_BLIP, bytes(60_000))})
    assert copies.check_content(slides, "ppt").code == "image"
    # Excel 的图片在 MSODRAWINGGROUP 记录里，超过 8224 字节的部分接在 CONTINUE 记录里
    data = blip(PNG_BLIP, bytes(60_000))
    chunks = [data[i : i + 8224] for i in range(0, len(data), 8224)]
    workbook = b"".join(
        struct.pack("<HH", 0x00EB if i == 0 else 0x003C, len(c)) + c for i, c in enumerate(chunks)
    )
    sheet = make_cfb({"Workbook": struct.pack("<HH", 0x0809, 0) + workbook})
    assert copies.detect_type(sheet) == "xls"
    assert copies.check_content(sheet, "xls").code == "image"


# ---- 数据和模拟 ----


def add_doc(session, number, title, day, attachments, *, rights_level=1, taken_down=False):
    source = Source(
        level="省级",
        province="辽宁",
        dept="科技",
        dept_full_name="辽宁省科学技术厅",
        column_name="通知",
        url=f"https://kjt.ln.gov.cn/list{number}/",
        enabled=True,
        parser_config={},
        rights_level=rights_level,
    )
    session.add(source)
    session.flush()
    doc = Document(
        source_id=source.id,
        url=f"https://kjt.ln.gov.cn/{number}.html",
        url_key=f"kjt.ln.gov.cn/{number}.html",
        title=title,
        publish_date=day,
        first_seen_at=NOW,
        last_seen_at=NOW,
        status="fetched",
        taken_down_at=NOW if taken_down else None,
    )
    session.add(doc)
    session.flush()
    session.add(
        DocumentDetail(
            document_id=doc.id,
            body_text="正文",
            fetched_at=NOW,
            status="ok",
            attachments=[
                DocumentAttachment(position=i, name=name, url=url, extension=ext)
                for i, (name, url, ext) in enumerate(attachments)
            ],
        )
    )
    session.commit()
    return doc


APPLY = "关于组织申报2026年度省级科技计划项目的通知"
LIST = "关于公布2026年第一批高新技术企业名单的通知"
RECENT = date(2026, 10, 1)


class FakeFetcher:
    def __init__(self, files):
        self.files = files  # 网址 → 内容（bytes）或 FetchResult
        self.calls = []

    def fetch(self, url):
        self.calls.append(url)
        value = self.files[url]
        if isinstance(value, FetchResult):
            return value
        return FetchResult(url, "ok", final_url=url, http_status=200, content=value)

    def close(self):
        pass


class FakeCloud:
    """模拟开放接口服务（上传、删除、取临时链接）和对象存储（上传、按段读取）。"""

    def __init__(self, *, upload_code=0, link_code=0, honor_range=True):
        self.upload_code, self.link_code, self.honor_range = upload_code, link_code, honor_range
        self.uploads: list[str] = []
        self.deleted: list[str] = []
        self.delete_fails: set[str] = set()
        self.links: list[dict] = []
        self.ranges: list[str | None] = []
        self.objects: dict[str, bytes] = {}
        self.expired_once = False
        self.client = self.new_client()

    def new_client(self) -> httpx.Client:
        return httpx.Client(transport=httpx.MockTransport(self.handle))

    def handle(self, request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if url == "http://api.weixin.qq.com/tcb/uploadfile":
            body = json.loads(request.content)
            assert body["env"] == ENV
            self.uploads.append(body["path"])
            if self.upload_code:
                return httpx.Response(200, json={"errcode": self.upload_code, "errmsg": "no"})
            return httpx.Response(
                200,
                json={
                    "errcode": 0,
                    "url": "https://cos.example.test/upload",
                    "token": "tok",
                    "authorization": "sig",
                    "file_id": PREFIX + body["path"],
                    "cos_file_id": "cosid",
                },
            )
        if url == "https://cos.example.test/upload":
            return httpx.Response(204)
        if url == "http://api.weixin.qq.com/tcb/batchdeletefile":
            ids = json.loads(request.content)["fileid_list"]
            self.deleted += ids
            items = [{"fileid": i, "status": -1 if i in self.delete_fails else 0} for i in ids]
            return httpx.Response(200, json={"errcode": 0, "delete_list": items})
        if url == "http://api.weixin.qq.com/tcb/batchdownloadfile":
            body = json.loads(request.content)
            self.links.append(body)
            if self.link_code:
                return httpx.Response(200, json={"errcode": self.link_code, "errmsg": "no"})
            fileid = body["file_list"][0]["fileid"]
            path = fileid.removeprefix(PREFIX)
            return httpx.Response(
                200,
                json={
                    "errcode": 0,
                    "file_list": [
                        {
                            "fileid": fileid,
                            "download_url": f"https://bucket.example.test/{path}?sign=s",
                            "status": 0,
                        }
                    ],
                },
            )
        if url.startswith("https://bucket.example.test/"):
            if self.expired_once:
                self.expired_once = False
                return httpx.Response(403)
            data = self.objects.get(request.url.path.lstrip("/"))
            if data is None:
                return httpx.Response(404)
            header = request.headers.get("range")
            self.ranges.append(header)
            if header and self.honor_range:
                start, end = map(int, re.match(r"bytes=(\d+)-(\d+)", header).groups())
                return httpx.Response(
                    206,
                    content=data[start : end + 1],
                    headers={"Content-Range": f"bytes {start}-{end}/{len(data)}"},
                )
            return httpx.Response(200, content=data)
        return httpx.Response(404)


def run(session, fetcher, cloud, *, now=NOW, **kwargs):
    return copies.cache_attachments(
        session, ENV, client=cloud.client, fetcher=fetcher, now=now, sleep=lambda _: None, **kwargs
    )


def rows_by_url(session):
    session.expire_all()
    return {r.url: r for r in session.scalars(select(AttachmentCopy))}


def obj(content: bytes, ext: str) -> str:
    return f"attachments/{sha256(content).hexdigest()}.{ext}"


# ---- 定时任务 ----


def test_cache_run_stores_only_allowed_attachments(session) -> None:
    base = "https://kjt.ln.gov.cn/files/"
    add_doc(
        session,
        1,
        APPLY,
        RECENT,
        [
            ("附件1 申报指南", base + "guide.pdf", "pdf"),
            ("附件2 申报书", base + "form.doc", "doc"),  # 实际是 docx
            ("附件3 同一份申报书", base + "form-copy.docx", "docx"),
            ("附件4 汇总表", base + "sheet.wps", "wps"),
            ("附件5 外站文件", "https://example.com/x.pdf", "pdf"),
            ("附件6 大文件", base + "big.pdf", "pdf"),
            ("附件7 其实是网页", base + "page.pdf", "pdf"),
            ("附件8 联系人名单", base + "contacts.docx", "docx"),
            ("附件9 已删除", base + "gone.pdf", "pdf"),
            ("附件10 网络不好", base + "slow.xlsx", "xlsx"),
            ("附件11 扫描件", base + "scan.pdf", "pdf"),
        ],
    )
    add_doc(session, 2, LIST, RECENT, [("名单", base + "list.pdf", "pdf")])
    add_doc(session, 3, APPLY, date(2026, 3, 1), [("旧附件", base + "old.pdf", "pdf")])
    add_doc(session, 4, APPLY, RECENT, [("下架", base + "down.pdf", "pdf")], taken_down=True)
    add_doc(session, 5, APPLY, RECENT, [("二级来源", base + "lvl2.pdf", "pdf")], rights_level=2)
    files = {
        base + "guide.pdf": PDF,
        base + "form.doc": DOCX,
        base + "form-copy.docx": DOCX,
        base + "big.pdf": FetchResult(base + "big.pdf", "too_large", error="文件超过 20MB"),
        base + "page.pdf": b"<html><body>page not found</body></html>",
        base + "contacts.docx": make_docx("联系人：张三 13812345678"),
        base + "gone.pdf": FetchResult(base + "gone.pdf", "http_error", http_status=404),
        base + "slow.xlsx": FetchResult(base + "slow.xlsx", "network_error", error="超时"),
        base + "scan.pdf": make_pdf(),
    }
    fetcher, cloud = FakeFetcher(files), FakeCloud()
    report = run(session, fetcher, cloud)

    # 只下载该下载的：wps、外站、名单类、超过 90 天、下架、二级来源的都不碰
    assert sorted(fetcher.calls) == sorted(files)
    assert report.added == len(files) and report.checked == len(files)
    # 只往 attachments/ 存，内容相同的只存一次
    assert cloud.uploads == [obj(PDF, "pdf"), obj(DOCX, "docx")]
    assert report.stored == 3 and report.removed == 0 and report.storage_error is None
    rows = rows_by_url(session)
    guide = rows[base + "guide.pdf"]
    assert (guide.status, guide.file_type, guide.size) == ("ready", "pdf", len(PDF))
    assert guide.file_id == PREFIX + obj(PDF, "pdf") and guide.next_check_at == NOW + copies.RECHECK
    assert rows[base + "form.doc"].file_type == "docx"
    assert rows[base + "form.doc"].file_id == rows[base + "form-copy.docx"].file_id
    failures = {url.removeprefix(base): (r.status, r.failure) for url, r in rows.items()}
    assert failures["big.pdf"] == ("failed", "too_large")
    assert failures["page.pdf"] == ("failed", "mismatch")
    assert failures["contacts.docx"] == ("failed", "personal")
    assert failures["gone.pdf"] == ("failed", "gone")
    assert failures["slow.xlsx"] == ("failed", "network")
    assert failures["scan.pdf"] == ("failed", "unchecked")
    # 临时问题 6 小时后重试，其他的 7 天后再看
    assert rows[base + "slow.xlsx"].next_check_at == NOW + timedelta(hours=6)
    assert rows[base + "gone.pdf"].next_check_at == NOW + copies.RECHECK
    assert report.failed["network"] == 1 and report.failed["personal"] == 1

    # 再跑一次：都没到期，不下载
    fetcher.calls.clear()
    again = run(session, fetcher, cloud)
    assert fetcher.calls == [] and again.added == 0 and again.checked == 0
    # 立即重试没存成的
    retried = run(session, fetcher, cloud, retry_failed=True)
    assert retried.checked == 6 and len(fetcher.calls) == 6


def test_limit_takes_new_files_first(session) -> None:
    base = "https://kjt.ln.gov.cn/files/"
    add_doc(session, 1, APPLY, date(2026, 9, 1), [("早", base + "early.pdf", "pdf")])
    add_doc(session, 2, APPLY, date(2026, 10, 2), [("新", base + "new.pdf", "pdf")])
    fetcher = FakeFetcher({base + "early.pdf": PDF, base + "new.pdf": make_pdf("New")})
    report = run(session, fetcher, FakeCloud(), limit=1)
    assert fetcher.calls == [base + "new.pdf"]
    assert report.checked == 1 and report.deferred == 1
    with pytest.raises(ValueError):
        run(session, fetcher, FakeCloud(), limit=0)


def test_recheck_updates_changed_files_and_drops_removed_ones(session) -> None:
    base = "https://kjt.ln.gov.cn/files/"
    names = ("same.pdf", "changed.pdf", "removed.pdf", "flaky.pdf")
    add_doc(session, 1, APPLY, RECENT, [(n, base + n, "pdf") for n in names])
    first = {base + n: make_pdf(f"Version one {n}") for n in names}
    fetcher, cloud = FakeFetcher(first), FakeCloud()
    run(session, fetcher, cloud)
    before = rows_by_url(session)
    old_changed = before[base + "changed.pdf"].file_id
    old_removed = before[base + "removed.pdf"].file_id
    old_flaky = before[base + "flaky.pdf"].file_id

    later = NOW + timedelta(days=8)
    new_content = make_pdf("Version two")
    fetcher.files.update(
        {
            base + "changed.pdf": new_content,
            base + "removed.pdf": FetchResult(base + "removed.pdf", "http_error", http_status=410),
            base + "flaky.pdf": FetchResult(base + "flaky.pdf", "http_error", http_status=503),
        }
    )
    cloud.uploads.clear()
    report = run(session, fetcher, cloud, now=later)
    rows = rows_by_url(session)
    assert report.unchanged == 1 and report.stored == 1
    assert cloud.uploads == [obj(new_content, "pdf")]
    assert rows[base + "changed.pdf"].file_id == PREFIX + obj(new_content, "pdf")
    assert rows[base + "changed.pdf"].fetched_at == later
    # 原网站删了：不再提供旧副本；服务器出错：旧副本照常可用
    assert (
        rows[base + "removed.pdf"].status == "failed" and rows[base + "removed.pdf"].file_id is None
    )
    assert (
        rows[base + "flaky.pdf"].status == "ready" and rows[base + "flaky.pdf"].file_id == old_flaky
    )
    assert rows[base + "flaky.pdf"].failure == "http"
    assert rows[base + "flaky.pdf"].next_check_at == later + timedelta(hours=6)
    # 换下来的旧副本当轮删掉
    assert sorted(cloud.deleted) == sorted([old_changed, old_removed])
    assert report.removed == 2
    assert all(r.stale_file_ids is None for r in rows.values())


def test_expired_copies_are_deleted_but_backups_never(session, monkeypatch) -> None:
    retry = copies._retrying
    monkeypatch.setattr(copies, "_retrying", lambda action, sleep: retry(action, lambda _: None))
    base = "https://kjt.ln.gov.cn/files/"
    doc = add_doc(session, 1, APPLY, RECENT, [("a", base + "a.pdf", "pdf")])
    add_doc(session, 2, APPLY, RECENT, [("b", base + "b.pdf", "pdf")])
    fetcher = FakeFetcher({base + "a.pdf": PDF, base + "b.pdf": make_pdf("Other")})
    cloud = FakeCloud()
    run(session, fetcher, cloud)
    rows = rows_by_url(session)
    kept = rows[base + "b.pdf"]
    # 就算表里混进了备份的文件 ID（不该发生），删除时也绝不发出去
    kept.stale_file_ids = [BACKUP_ID, "cloud://other-env.7072-x/attachments/" + "0" * 64 + ".pdf"]
    session.commit()
    a_file = rows[base + "a.pdf"].file_id

    # 下架后：副本删掉、记录删掉；删除失败的下一轮再删
    session.get(Document, doc.id).taken_down_at = NOW
    session.commit()
    cloud.delete_fails = {a_file}
    report = run(session, fetcher, cloud)
    assert a_file in cloud.deleted and BACKUP_ID not in cloud.deleted
    assert all("/attachments/" in i and i.startswith(PREFIX) for i in cloud.deleted)
    assert report.removed == 0 and base + "a.pdf" in rows_by_url(session)
    assert report.storage_error and "删除未确认" in report.storage_error
    cloud.delete_fails = set()
    report = run(session, fetcher, cloud)
    assert report.removed == 1 and base + "a.pdf" not in rows_by_url(session)
    assert BACKUP_ID not in cloud.deleted
    # 删不掉的不是附件副本，留在记录里，不当作删了
    assert set(rows_by_url(session)[base + "b.pdf"].stale_file_ids) == set(kept.stale_file_ids)

    # 超过申报通知 180 天的同样删掉
    report = run(session, fetcher, cloud, now=NOW + timedelta(days=200))
    assert report.removed == 1


def test_storage_failure_stops_the_run(session, monkeypatch, capsys) -> None:
    base = "https://kjt.ln.gov.cn/files/"
    add_doc(session, 1, APPLY, RECENT, [("a", base + "a.pdf", "pdf"), ("b", base + "b.pdf", "pdf")])
    fetcher = FakeFetcher({base + "a.pdf": PDF, base + "b.pdf": make_pdf("B")})
    report = run(session, fetcher, FakeCloud(upload_code=85088))
    assert report.storage_error and "85088" in report.storage_error
    # 第一个上传失败就停下，不再去原网站下载；记录保持待下载
    assert fetcher.calls == [base + "a.pdf"]
    assert {r.status for r in rows_by_url(session).values()} == {"pending"}

    from youju import cli

    settings = SimpleNamespace(cloud_env=None)
    monkeypatch.setattr(cli, "get_settings", lambda: settings)
    assert cli.main(["cache-attachments"]) == 1
    settings.cloud_env = ENV
    calls = []

    def fake_cache(session, env, **kwargs):
        calls.append((env, kwargs))
        return copies.CacheReport(checked=2, stored=1)

    monkeypatch.setattr(copies, "cache_attachments", fake_cache)
    assert cli.main(["cache-attachments", "--limit", "5", "--retry-failed"]) == 0
    assert calls == [(ENV, {"limit": 5, "retry_failed": True})]
    assert "本轮下载 2" in capsys.readouterr().out
    monkeypatch.setattr(
        copies, "cache_attachments", lambda *a, **k: copies.CacheReport(storage_error="上传失败")
    )
    assert cli.main(["cache-attachments"]) == 1


def test_robots_disallowed_attachments_are_not_downloaded(session) -> None:
    """下载附件的抓取器不解码内容，robots.txt 照样要遵守：网站禁止的附件不下载，记下原因。"""
    from youju.crawler.fetcher import PoliteFetcher

    base = "https://kjt.ln.gov.cn/"
    add_doc(
        session,
        1,
        APPLY,
        RECENT,
        [("指南", base + "private/guide.pdf", "pdf"), ("表", base + "files/form.pdf", "pdf")],
    )
    requested = []

    def handler(request: httpx.Request) -> httpx.Response:
        requested.append(request.url.path)
        if request.url.path == "/robots.txt":
            return httpx.Response(200, content=b"User-agent: *\nDisallow: /private/\n")
        return httpx.Response(200, content=PDF)

    fetcher = PoliteFetcher(
        "TestAgent YoujuBot/0.1",
        per_host_interval=0,
        transport=httpx.MockTransport(handler),
        resolver=lambda host: ["1.2.3.4"],
        decode_text=False,
    )
    cloud = FakeCloud()
    report = run(session, fetcher, cloud)
    rows = rows_by_url(session)
    blocked = rows[base + "private/guide.pdf"]
    assert blocked.status == "failed" and blocked.failure == "robots"
    assert blocked.failure_detail == "robots.txt 不允许下载" and blocked.file_id is None
    assert "/private/guide.pdf" not in requested
    assert rows[base + "files/form.pdf"].status == "ready"
    assert report.failed == {"robots": 1} and cloud.uploads == [obj(PDF, "pdf")]


def test_recheck_drops_stored_copies_that_fail_the_new_rules(session, monkeypatch, capsys) -> None:
    base = "https://kjt.ln.gov.cn/files/"
    names = ("guide.pdf", "form.docx", "form-copy.docx", "lost.pdf", "odd.pdf")
    doc = add_doc(session, 1, APPLY, RECENT, [(n, base + n, n.rsplit(".", 1)[1]) for n in names])
    picture = docx_with_picture(make_png(150))  # 规则收紧前存下的：有大图
    lost = make_pdf("Lost")
    cloud = FakeCloud()
    for content, ext in ((PDF, "pdf"), (picture, "docx")):
        cloud.objects[obj(content, ext)] = content
    add_copy(session, base + "guide.pdf", PDF, "pdf")
    add_copy(session, base + "form.docx", picture, "docx")
    add_copy(session, base + "form-copy.docx", picture, "docx")  # 内容相同，共用一个副本
    add_copy(session, base + "lost.pdf", lost, "pdf")  # 对象存储里读不到
    # 表里混进了备份的文件 ID（不该发生）：不是附件副本，不读也不删
    add_copy(session, base + "odd.pdf", PDF, "pdf", file_id=BACKUP_ID)

    def recheck(**kwargs):
        return copies.recheck_copies(session, ENV, client=cloud.client, now=NOW, **kwargs)

    # 只检查：报告不合格的，不改表、不删文件
    report = recheck(dry_run=True)
    assert (report.checked, report.kept, report.unreadable) == (2, 1, 1)
    assert report.dropped == {"image": 1} and report.rows == 2 and report.removed == 0
    assert cloud.deleted == [] and rows_by_url(session)[base + "form.docx"].status == "ready"

    report = recheck()
    rows = rows_by_url(session)
    for name in ("form.docx", "form-copy.docx"):
        row = rows[base + name]
        assert row.status == "failed" and row.failure == "image" and row.file_id is None
        assert "无法核验个人信息" in row.failure_detail and row.stale_file_ids is None
    assert cloud.deleted == [PREFIX + obj(picture, "docx")] and report.removed == 1
    assert rows[base + "guide.pdf"].status == "ready"
    assert rows[base + "lost.pdf"].status == "ready"  # 读不到的这次不动
    assert rows[base + "odd.pdf"].file_id == BACKUP_ID
    assert all("backups/" not in link["file_list"][0]["fileid"] for link in cloud.links)
    # 小程序：说明原因、复制原文链接
    items = {
        item["name"]: item
        for item in copies.attachment_items(
            session, doc.detail.attachments, kind="申报通知", recent=True
        )
    }
    assert items["form.docx"]["preview"] is False and items["form.docx"]["reason"] == "image"
    assert items["guide.pdf"]["preview"] is True

    from youju import cli

    settings = SimpleNamespace(cloud_env=None)
    monkeypatch.setattr(cli, "get_settings", lambda: settings)
    assert cli.main(["recheck-attachments"]) == 1
    settings.cloud_env = ENV
    calls = []

    def fake_recheck(session, env, **kwargs):
        calls.append((env, kwargs))
        return copies.RecheckReport(checked=3, kept=2, dropped=Counter(image=1), rows=2, removed=1)

    monkeypatch.setattr(copies, "recheck_copies", fake_recheck)
    assert cli.main(["recheck-attachments", "--dry-run"]) == 0
    assert calls == [(ENV, {"dry_run": True})]
    assert "不合格 1 个（image 1）" in capsys.readouterr().out
    monkeypatch.setattr(
        copies, "recheck_copies", lambda *a, **k: copies.RecheckReport(unreadable=1)
    )
    assert cli.main(["recheck-attachments"]) == 1


def test_job_and_cron_are_registered() -> None:
    from youju.jobs import JOB_NAMES

    assert "cache-attachments" in JOB_NAMES
    crontab = (Path(__file__).resolve().parents[1] / "deploy" / "crontab").read_text()
    assert crontab.count("/app/backend/deploy/job.sh cache-attachments") == 2


# ---- 小程序接口 ----


@pytest.fixture
def client(session_factory, monkeypatch):
    def override():
        with session_factory() as s:
            yield s

    monkeypatch.setattr(v1, "now", lambda: NOW)
    app.dependency_overrides[get_session] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture(autouse=True)
def _clear_caches():
    copies._links.clear()
    attachment_api._times.clear()
    yield
    copies._links.clear()
    attachment_api._times.clear()


def add_copy(
    session,
    url,
    content,
    ext,
    *,
    status="ready",
    failure=None,
    file_id=None,
    rules_version=copies.RULES_VERSION,
):
    row = AttachmentCopy(
        url_hash=copies.url_hash(url),
        url=url,
        declared_type=ext,
        status=status,
        failure=failure,
        rules_version=rules_version,
    )
    if status == "ready":
        row.file_type, row.size = ext, len(content)
        row.sha256 = sha256(content).hexdigest()
        row.file_id = file_id or PREFIX + obj(content, ext)
    session.add(row)
    session.commit()
    return row


def test_detail_lists_preview_info_and_reasons(client, session) -> None:
    base = "https://kjt.ln.gov.cn/files/"
    doc = add_doc(
        session,
        1,
        APPLY,
        RECENT,
        [
            ("指南", base + "guide.pdf", "pdf"),
            ("申报书", base + "form.docx", "docx"),
            ("名单", base + "people.xlsx", "xlsx"),
            ("扫描件", base + "scan.pdf", "pdf"),
            ("大文件", base + "big.pdf", "pdf"),
            ("打不开", base + "bad.doc", "doc"),
            ("汇总表", base + "sheet.wps", "wps"),
            ("外站", "https://example.com/a.pdf", "pdf"),
        ],
    )
    add_copy(session, base + "guide.pdf", PDF, "pdf")
    add_copy(session, base + "people.xlsx", b"", "xlsx", status="failed", failure="personal")
    add_copy(session, base + "scan.pdf", b"", "pdf", status="failed", failure="unchecked")
    add_copy(session, base + "big.pdf", b"", "pdf", status="failed", failure="too_large")
    add_copy(session, base + "bad.doc", b"", "doc", status="failed", failure="mismatch")
    items = client.get(f"/api/v1/documents/{doc.id}").json()["attachments"]
    first = items[0]
    assert first["preview"] is True and first["fileType"] == "pdf" and first["size"] == len(PDF)
    assert first["version"] == sha256(PDF).hexdigest()[:16] and first["parts"] == 1
    assert first["partSize"] == copies.PART_SIZE and first["id"].isdigit()
    assert first["ext"] == "PDF" and first["url"] == base + "guide.pdf"
    # 不给对象存储的文件 ID 或链接
    assert "cloud://" not in json.dumps(items)
    reasons = [item.get("reason") for item in items]
    assert reasons == [None, "pending", "personal", "scan", "too_large", "failed", "format", "host"]
    assert all(item["preview"] is False for item in items[1:])

    listed = add_doc(session, 2, LIST, RECENT, [("名单", base + "guide.pdf", "pdf")])
    old = add_doc(session, 3, APPLY, date(2026, 3, 1), [("旧", base + "old.pdf", "pdf")])
    [item] = client.get(f"/api/v1/documents/{listed.id}").json()["attachments"]
    assert item["preview"] is False and item["reason"] == "kind"
    [item] = client.get(f"/api/v1/documents/{old.id}").json()["attachments"]
    assert item["reason"] == "old"


def part_url(att_id, version, part=0):
    return f"/api/v1/attachments/{att_id}/content?v={version}&part={part}"


def test_content_is_served_in_parts_from_attachments_only(client, session, monkeypatch) -> None:
    base = "https://kjt.ln.gov.cn/files/"
    content = bytes(range(256)) * (copies.PART_SIZE // 256 + 1000)  # 两段
    doc = add_doc(session, 1, APPLY, RECENT, [("指南", base + "guide.pdf", "pdf")])
    add_copy(session, base + "guide.pdf", content, "pdf")
    cloud = FakeCloud()
    cloud.objects[obj(content, "pdf")] = content
    monkeypatch.setattr(attachment_api, "_client", cloud.new_client)
    item = client.get(f"/api/v1/documents/{doc.id}").json()["attachments"][0]
    assert item["parts"] == 2
    att, version = item["id"], item["version"]

    first = client.get(part_url(att, version, 0))
    assert first.status_code == 200 and first.content == content[: copies.PART_SIZE]
    assert first.headers["content-type"] == "application/octet-stream"
    second = client.get(part_url(att, version, 1))
    assert second.content == content[copies.PART_SIZE :]
    assert cloud.ranges == [
        f"bytes=0-{copies.PART_SIZE - 1}",
        f"bytes={copies.PART_SIZE}-{len(content) - 1}",
    ]
    # 临时链接按文档格式申请，同一进程里复用
    assert cloud.links == [
        {
            "env": ENV,
            "file_list": [{"fileid": PREFIX + obj(content, "pdf"), "max_age": 1800}],
        }
    ]
    # 链接过期（403）时换一个再读
    copies._links[PREFIX + obj(content, "pdf")] = ("https://bucket.example.test/x", 1e18)
    cloud.expired_once = True
    assert client.get(part_url(att, version, 1)).content == content[copies.PART_SIZE :]
    # 不支持分段（返回整份）也能取对
    cloud.honor_range = False
    assert client.get(part_url(att, version, 1)).content == content[copies.PART_SIZE :]

    assert client.get(part_url(att, "0" * 16, 0)).status_code == 409
    assert client.get(part_url(att, version, 2)).status_code == 422
    assert client.get(part_url(att, "xyz", 0)).status_code == 422
    assert client.get(part_url(999, version, 0)).status_code == 404

    # 对象存储出错：503，小程序改为复制原文链接
    copies._links.clear()
    cloud.link_code = -501002
    assert client.get(part_url(att, version, 0)).status_code == 503

    # 下架后不给
    session.get(Document, doc.id).taken_down_at = NOW
    session.commit()
    assert client.get(part_url(att, version, 0)).status_code == 404


def test_content_never_reads_backups_or_unready_copies(client, session, monkeypatch) -> None:
    base = "https://kjt.ln.gov.cn/files/"
    doc = add_doc(
        session,
        1,
        APPLY,
        RECENT,
        [("a", base + "a.pdf", "pdf"), ("b", base + "b.pdf", "pdf"), ("c", base + "c.wps", "wps")],
    )
    tampered = add_copy(session, base + "a.pdf", PDF, "pdf", file_id=BACKUP_ID)
    add_copy(session, base + "b.pdf", b"", "pdf", status="pending")
    cloud = FakeCloud()
    monkeypatch.setattr(attachment_api, "_client", cloud.new_client)
    ids = [str(a.id) for a in session.get(DocumentDetail, doc.id).attachments]
    version = tampered.sha256[:16]
    for att in ids:
        assert client.get(part_url(att, version, 0)).status_code == 404
    # 文件 ID 不是 attachments/ 下的（比如备份），详情里也不给打开
    items = client.get(f"/api/v1/documents/{doc.id}").json()["attachments"]
    assert [i["preview"] for i in items] == [False, False, False]
    assert cloud.links == [] and cloud.ranges == []
    with pytest.raises(copies.BackupError):
        copies.download_link(cloud.client, BACKUP_ID)
    assert copies.parse_file_id(BACKUP_ID) is None
    good = PREFIX + obj(PDF, "pdf")
    assert copies.parse_file_id(good) == (ENV, obj(PDF, "pdf"))
    for bad in (good + "\n", good.replace("attachments/", "attachments/../backups/"), "x" + good):
        assert copies.parse_file_id(bad) is None
    assert copies.delete_copies(cloud.client, ENV, [BACKUP_ID]) == set()
    assert cloud.deleted == []


def test_copies_checked_under_old_rules_are_not_served(client, session, monkeypatch) -> None:
    """按旧规则版本查过的副本立即不提供：详情里当作还在准备（小程序复制原文链接、说明原因），
    取内容返回 404，也不去读对象存储。"""
    base = "https://kjt.ln.gov.cn/files/"
    doc = add_doc(session, 1, APPLY, RECENT, [("指南", base + "guide.pdf", "pdf")])
    row = add_copy(session, base + "guide.pdf", PDF, "pdf", rules_version=copies.RULES_VERSION - 1)
    cloud = FakeCloud()
    cloud.objects[obj(PDF, "pdf")] = PDF
    monkeypatch.setattr(attachment_api, "_client", cloud.new_client)
    [item] = client.get(f"/api/v1/documents/{doc.id}").json()["attachments"]
    assert item["preview"] is False and item["reason"] == "pending"
    att = str(session.get(DocumentDetail, doc.id).attachments[0].id)
    assert client.get(part_url(att, row.sha256[:16], 0)).status_code == 404
    assert cloud.links == [] and cloud.ranges == []
    # 按现在的规则查过以后照常提供
    row.rules_version = copies.RULES_VERSION
    session.commit()
    [item] = client.get(f"/api/v1/documents/{doc.id}").json()["attachments"]
    assert item["preview"] is True
    assert client.get(part_url(att, row.sha256[:16], 0)).status_code == 200


def test_cache_run_rechecks_old_rule_copies_first(session) -> None:
    """定时任务先从对象存储读出旧规则版本的副本按新规则重查（每轮有上限）：合格的升到当前版本、
    恢复提供；不合格的改成不提供、当轮从对象存储删掉；不访问原网站，不碰备份。"""
    base = "https://kjt.ln.gov.cn/files/"
    names = ("guide.pdf", "form.docx", "form-copy.docx", "odd.pdf")
    add_doc(session, 1, APPLY, RECENT, [(n, base + n, n.rsplit(".", 1)[1]) for n in names])
    picture = docx_with_picture(make_png(150))
    cloud = FakeCloud()
    for content, ext in ((PDF, "pdf"), (picture, "docx")):
        cloud.objects[obj(content, ext)] = content
    old = copies.RULES_VERSION - 1
    for name, content, ext in (
        ("guide.pdf", PDF, "pdf"),
        ("form.docx", picture, "docx"),
        ("form-copy.docx", picture, "docx"),
    ):
        add_copy(session, base + name, content, ext, rules_version=old)
    add_copy(session, base + "odd.pdf", PDF, "pdf", file_id=BACKUP_ID, rules_version=old)
    for row in rows_by_url(session).values():
        row.next_check_at = NOW + timedelta(days=3)  # 都还没到重新下载的时候
    session.commit()
    fetcher = FakeFetcher({})

    # 每轮有上限：这一轮只查第一个文件
    report = run(session, fetcher, cloud, recheck_limit=1)
    rows = rows_by_url(session)
    assert report.recheck.checked == 1 and report.recheck.kept == 1
    assert rows[base + "guide.pdf"].rules_version == copies.RULES_VERSION
    assert copies.servable(rows[base + "guide.pdf"])
    assert not copies.servable(rows[base + "form.docx"])  # 还没轮到，仍不提供
    assert fetcher.calls == [] and cloud.deleted == []

    report = run(session, fetcher, cloud, now=NOW + timedelta(hours=12))
    rows = rows_by_url(session)
    assert report.recheck.dropped == {"image": 1} and report.recheck.rows == 2
    for name in ("form.docx", "form-copy.docx"):
        row = rows[base + name]
        assert row.status == "failed" and row.failure == "image" and row.file_id is None
        assert row.stale_file_ids is None
    # 不合格的当轮删掉，只删 attachments/ 下的，备份不读不删
    assert cloud.deleted == [PREFIX + obj(picture, "docx")] and report.removed == 1
    assert all("backups/" not in link["file_list"][0]["fileid"] for link in cloud.links)
    assert rows[base + "odd.pdf"].file_id == BACKUP_ID and not copies.servable(
        rows[base + "odd.pdf"]
    )
    assert fetcher.calls == []
    # 都查过了：下一轮不再读
    links = len(cloud.links)
    report = run(session, fetcher, cloud, now=NOW + timedelta(hours=13))
    assert report.recheck.checked == 0 and len(cloud.links) == links


def test_old_rule_recheck_failures_do_not_block_the_queue(session) -> None:
    """重查按上次重查的时间排：读不到的记下时间排到后面，不会一直堵在队首。"""
    base = "https://kjt.ln.gov.cn/files/"
    add_doc(session, 1, APPLY, RECENT, [("a", base + "a.pdf", "pdf"), ("b", base + "b.pdf", "pdf")])
    lost, good = make_pdf("Lost"), make_pdf("Good")
    cloud = FakeCloud()
    cloud.objects[obj(good, "pdf")] = good  # a 在对象存储里读不到
    old = copies.RULES_VERSION - 1
    add_copy(session, base + "a.pdf", lost, "pdf", rules_version=old)
    add_copy(session, base + "b.pdf", good, "pdf", rules_version=old)
    for row in rows_by_url(session).values():
        row.next_check_at = NOW + timedelta(days=3)
    session.commit()
    fetcher = FakeFetcher({})
    report = run(session, fetcher, cloud, recheck_limit=1)
    assert report.recheck.unreadable == 1 and report.recheck.checked == 0
    assert rows_by_url(session)[base + "a.pdf"].rechecked_at == NOW
    report = run(session, fetcher, cloud, now=NOW + timedelta(hours=12), recheck_limit=1)
    rows = rows_by_url(session)
    assert report.recheck.kept == 1 and copies.servable(rows[base + "b.pdf"])
    assert not copies.servable(rows[base + "a.pdf"])


def test_cache_run_and_manual_recheck_share_one_lock(session, monkeypatch, capsys) -> None:
    """附件定时任务和手动重查共用一把数据库锁：另一个在跑就跳过本轮，什么都不改。"""
    base = "https://kjt.ln.gov.cn/files/"
    add_doc(session, 1, APPLY, RECENT, [("a", base + "a.pdf", "pdf")])
    add_copy(session, base + "a.pdf", PDF, "pdf", rules_version=copies.RULES_VERSION - 1)
    cloud = FakeCloud()
    cloud.objects[obj(PDF, "pdf")] = PDF
    fetcher = FakeFetcher({base + "a.pdf": PDF})
    with session.get_bind().connect() as other:
        assert other.execute(sql("SELECT GET_LOCK('youju_attachment_copies', 0)")).scalar() == 1
        report = run(session, fetcher, cloud)
        assert report.skipped and report.recheck is None
        manual = copies.recheck_copies(session, ENV, client=cloud.client, now=NOW)
        assert manual.skipped and manual.checked == 0
        assert fetcher.calls == [] and cloud.links == [] and cloud.deleted == []
        assert not copies.servable(rows_by_url(session)[base + "a.pdf"])
        other.execute(sql("SELECT RELEASE_LOCK('youju_attachment_copies')"))
    report = run(session, fetcher, cloud)
    assert not report.skipped and report.recheck.kept == 1
    # 跑完会释放：手动重查接着能拿到锁
    assert not copies.recheck_copies(session, ENV, client=cloud.client, now=NOW).skipped

    from youju import cli

    monkeypatch.setattr(cli, "get_settings", lambda: SimpleNamespace(cloud_env=ENV))
    monkeypatch.setattr(
        copies, "cache_attachments", lambda *a, **k: copies.CacheReport(skipped=True)
    )
    monkeypatch.setattr(
        copies, "recheck_copies", lambda *a, **k: copies.RecheckReport(skipped=True)
    )
    assert cli.main(["cache-attachments"]) == 0
    assert cli.main(["recheck-attachments"]) == 1
    out = capsys.readouterr().out
    assert "本轮跳过" in out and "这次没有重查" in out


def test_objects_still_referenced_are_never_deleted(session) -> None:
    """删除前在同一事务里再确认一次：还有要提供的附件引用着的对象不删。"""
    base = "https://kjt.ln.gov.cn/files/"
    add_doc(session, 1, APPLY, RECENT, [("a", base + "a.pdf", "pdf"), ("b", base + "b.pdf", "pdf")])
    shared = PREFIX + obj(PDF, "pdf")
    add_copy(session, base + "a.pdf", PDF, "pdf")
    stale = add_copy(session, base + "b.pdf", make_pdf("B"), "pdf")
    stale.stale_file_ids = [shared]  # 换下来的旧副本恰好是别的附件正在用的对象
    session.commit()
    live = {copies.url_hash(base + "a.pdf"), copies.url_hash(base + "b.pdf")}
    assert copies._referenced(session, {shared}, live) == {shared}
    assert copies._referenced(session, {shared}, {copies.url_hash(base + "b.pdf")}) == set()
    session.rollback()
    cloud = FakeCloud()
    report = copies.CacheReport()
    copies._cleanup(session, cloud.client, ENV, live, report)
    assert cloud.deleted == [] and report.removed == 0
    assert rows_by_url(session)[base + "b.pdf"].stale_file_ids is None
    assert copies._drop_stale(session, cloud.client, ENV) == 0 and cloud.deleted == []


def test_downloaded_copies_get_the_current_rules_version(session) -> None:
    base = "https://kjt.ln.gov.cn/files/"
    add_doc(session, 1, APPLY, RECENT, [("指南", base + "guide.pdf", "pdf")])
    run(session, FakeFetcher({base + "guide.pdf": PDF}), FakeCloud())
    row = rows_by_url(session)[base + "guide.pdf"]
    assert row.rules_version == copies.RULES_VERSION and copies.servable(row)


def test_migration_0021_rules_version(db_engine) -> None:
    assert db_engine.url.database.startswith("youju_test")
    cfg = alembic_config(db_engine.url.render_as_string(hide_password=False))

    def columns():
        return {c["name"] for c in inspect(db_engine).get_columns("attachment_copies")}

    try:
        command.downgrade(cfg, "0020")
        assert {"rules_version", "rechecked_at"}.isdisjoint(columns())
    finally:
        command.upgrade(cfg, "head")
    assert {"rules_version", "rechecked_at"} <= columns()


def test_content_is_rate_limited(client, session, monkeypatch) -> None:
    base = "https://kjt.ln.gov.cn/files/"
    doc = add_doc(session, 1, APPLY, RECENT, [("a", base + "a.pdf", "pdf")])
    row = add_copy(session, base + "a.pdf", PDF, "pdf")
    cloud = FakeCloud()
    cloud.objects[obj(PDF, "pdf")] = PDF
    monkeypatch.setattr(attachment_api, "_client", cloud.new_client)
    monkeypatch.setattr(attachment_api, "PARTS_PER_10MIN", 2)
    att = session.get(DocumentDetail, doc.id).attachments[0].id
    codes = [client.get(part_url(att, row.sha256[:16])).status_code for _ in range(3)]
    assert codes == [200, 200, 429]


def test_migration_0018_downgrade_upgrade(db_engine) -> None:
    assert db_engine.url.database.startswith("youju_test")
    cfg = alembic_config(db_engine.url.render_as_string(hide_password=False))
    try:
        command.downgrade(cfg, "0017")
        assert "attachment_copies" not in inspect(db_engine).get_table_names()
    finally:
        command.upgrade(cfg, "head")
    assert "attachment_copies" in inspect(db_engine).get_table_names()


def test_application_180_day_window_cache_detail_and_cleanup(client, session):
    base = "https://kjt.ln.gov.cn/files/"
    day = NOW.date() - timedelta(days=120)
    application = add_doc(session, 10, APPLY, day, [("申报书", base + "apply.pdf", "pdf")])
    policy = add_doc(
        session, 11, "关于印发科技项目管理办法的通知", day, [("办法", base + "policy.pdf", "pdf")]
    )
    assert copies.candidates(session, NOW) == {base + "apply.pdf": "pdf"}
    assert (
        client.get(f"/api/v1/documents/{application.id}").json()["attachments"][0]["reason"]
        == "pending"
    )
    assert client.get(f"/api/v1/documents/{policy.id}").json()["attachments"][0]["reason"] == "old"
    cloud = FakeCloud()
    fetcher = FakeFetcher({base + "apply.pdf": PDF})
    run(session, fetcher, cloud)
    assert (
        client.get(f"/api/v1/documents/{application.id}").json()["attachments"][0]["preview"]
        is True
    )
    run(session, fetcher, cloud, now=NOW + timedelta(days=61))
    assert base + "apply.pdf" not in rows_by_url(session)
    assert cloud.deleted


@pytest.mark.parametrize("kind,days", [("申报通知", 180), ("政策", 90)])
def test_attachment_window_boundary_and_missing_publish_date(kind, days):
    boundary = NOW.date() - timedelta(days=days)
    assert copies.is_recent(boundary, None, NOW, kind=kind)
    assert not copies.is_recent(boundary - timedelta(days=1), NOW, NOW, kind=kind)
    assert copies.is_recent(None, NOW - timedelta(days=days - 1), NOW, kind=kind)
    assert not copies.is_recent(None, NOW - timedelta(days=days + 1), NOW, kind=kind)


def test_delete_copies_retries_temporary_busy_then_confirms_only_returned_ids():
    file_id = PREFIX + obj(PDF, "pdf")
    requests, waits = [], []

    def handle(request):
        requests.append(json.loads(request.content))
        if len(requests) == 1:
            return httpx.Response(200, json={"errcode": -1})
        return httpx.Response(
            200, json={"errcode": 0, "delete_list": [{"fileid": file_id, "status": 0}]}
        )

    with httpx.Client(transport=httpx.MockTransport(handle)) as client:
        assert copies.delete_copies(client, ENV, [file_id, BACKUP_ID], sleep=waits.append) == {
            file_id
        }
    assert waits == [5]
    assert requests == [{"env": ENV, "fileid_list": [file_id]}] * 2


def test_delete_copies_permanent_error_does_not_retry():
    file_id = PREFIX + obj(PDF, "pdf")
    requests, waits = [], []

    def handle(request):
        requests.append(request)
        return httpx.Response(200, json={"errcode": 40013})

    with (
        httpx.Client(transport=httpx.MockTransport(handle)) as client,
        pytest.raises(copies.BackupError, match="40013"),
    ):
        copies.delete_copies(client, ENV, [file_id], sleep=waits.append)
    assert len(requests) == 1 and waits == []


def test_delete_copies_exhaustion_preserves_stale_reference_for_next_run(session, monkeypatch):
    row = add_copy(session, "https://example.gov.cn/retry-synthetic.pdf", PDF, "pdf")
    stale_id = PREFIX + obj(b"synthetic obsolete content", "pdf")
    current_id = row.file_id
    row.stale_file_ids = [stale_id]
    session.commit()
    requests, waits = [], []
    retry = copies._retrying
    monkeypatch.setattr(copies, "_retrying", lambda action, sleep: retry(action, waits.append))

    def handle(request):
        requests.append(json.loads(request.content))
        return httpx.Response(200, json={"errcode": -1})

    with (
        httpx.Client(transport=httpx.MockTransport(handle)) as client,
        pytest.raises(copies.BackupError, match="errcode -1"),
    ):
        copies._drop_stale(session, client, ENV)
    session.expire_all()
    assert row.stale_file_ids == [stale_id] and row.file_id == current_id
    assert waits == [5, 15] and len(requests) == 3
    assert all(r == {"env": ENV, "fileid_list": [stale_id]} for r in requests)


def test_delete_partial_temporary_failure_retries_only_unconfirmed_items():
    ids = [PREFIX + f"attachments/{i:064x}.pdf" for i in range(2)]
    requests, waits = [], []

    def handle(request):
        pending = json.loads(request.content)["fileid_list"]
        requests.append(pending)
        return httpx.Response(
            200,
            json={
                "errcode": 0,
                "delete_list": [
                    {
                        "fileid": file_id,
                        "status": -1 if len(requests) == 1 and file_id == ids[1] else 0,
                    }
                    for file_id in pending
                ],
            },
        )

    with httpx.Client(transport=httpx.MockTransport(handle)) as client:
        assert copies.delete_copies(client, ENV, ids, sleep=waits.append) == set(ids)
    assert requests == [ids, [ids[1]]] and waits == [5]


@pytest.mark.parametrize("status", [None, 40097])
def test_delete_missing_or_permanent_item_is_not_silent_success(status):
    file_id = PREFIX + obj(PDF, "pdf")
    requests, waits = [], []

    def handle(request):
        requests.append(request)
        return httpx.Response(
            200,
            json={
                "errcode": 0,
                "delete_list": [] if status is None else [{"fileid": file_id, "status": status}],
            },
        )

    with (
        httpx.Client(transport=httpx.MockTransport(handle)) as client,
        pytest.raises(copies.DeleteCopiesError, match="未确认") as error,
    ):
        copies.delete_copies(client, ENV, [file_id], sleep=waits.append)
    assert len(requests) == 1 and waits == [] and error.value.confirmed == set()


def test_partial_cleanup_settles_only_confirmed_and_keeps_failure_for_alert(session):
    row = add_copy(session, "https://example.gov.cn/partial-synthetic.pdf", PDF, "pdf")
    good, bad = [PREFIX + f"attachments/{i:064x}.pdf" for i in range(2)]
    row.stale_file_ids = [good, bad]
    session.commit()

    def handle(request):
        ids = json.loads(request.content)["fileid_list"]
        return httpx.Response(
            200,
            json={
                "errcode": 0,
                "delete_list": [
                    {"fileid": file_id, "status": 0 if file_id == good else 40097}
                    for file_id in ids
                ],
            },
        )

    with (
        httpx.Client(transport=httpx.MockTransport(handle)) as client,
        pytest.raises(copies.BackupError, match="未确认"),
    ):
        copies._drop_stale(session, client, ENV)
    session.expire_all()
    assert row.stale_file_ids == [bad]
    assert row.file_id == PREFIX + obj(PDF, "pdf")


def test_delete_batches_preserve_fifty_item_limit_and_ownership():
    ids = [PREFIX + f"attachments/{i:064x}.pdf" for i in range(101)]
    sizes = []

    def handle(request):
        batch = json.loads(request.content)["fileid_list"]
        sizes.append(len(batch))
        assert all(file_id in ids for file_id in batch)
        return httpx.Response(
            200,
            json={
                "errcode": 0,
                "delete_list": [{"fileid": file_id, "status": 0} for file_id in batch],
            },
        )

    with httpx.Client(transport=httpx.MockTransport(handle)) as client:
        assert copies.delete_copies(client, ENV, ids + [BACKUP_ID], sleep=lambda _: None) == set(
            ids
        )
    assert sizes == [50, 50, 1]


def test_expired_partial_cleanup_clears_confirmed_current_id_but_retains_pending_stale(session):
    row = add_copy(session, "https://example.gov.cn/expired-partial.pdf", PDF, "pdf")
    current = row.file_id
    stale = PREFIX + f"attachments/{2:064x}.pdf"
    row.stale_file_ids = [stale]
    session.commit()
    report = copies.CacheReport()

    def handle(request):
        ids = json.loads(request.content)["fileid_list"]
        return httpx.Response(
            200,
            json={
                "errcode": 0,
                "delete_list": [
                    {"fileid": file_id, "status": 0 if file_id == current else 40097}
                    for file_id in ids
                ],
            },
        )

    with (
        httpx.Client(transport=httpx.MockTransport(handle)) as client,
        pytest.raises(copies.BackupError, match="未确认"),
    ):
        copies._cleanup(session, client, ENV, set(), report)
    session.expire_all()
    assert row.file_id is None and row.stale_file_ids == [stale]
    assert report.removed == 1
