"""同一份真实报告导出工作簿和原生结果，核对汇总统计及明细保留。"""

from datetime import UTC, date, datetime, timedelta

import pytest
from openpyxl import load_workbook

from youju import native_results
from youju.compare import (
    CompareInputError,
    CompareReport,
    CompareRow,
    ListMatch,
    RemovedInput,
    export_compare_excel,
)
from youju.track import CustomerProfile, CustomerTrack, TrackItem, TrackReport, export_track_excel

NOW = datetime(2026, 10, 23, tzinfo=UTC)


def _excel_summary(path):
    book = load_workbook(path)
    sheet = book["汇总"]
    headers = [cell.value for cell in sheet[2]]
    rows = [
        {
            label: str(value) if value is not None else "—"
            for label, value in zip(headers, values, strict=True)
        }
        for values in sheet.iter_rows(min_row=3, values_only=True)
    ]
    note = sheet["A1"].value
    book.close()
    return rows, note


def _native_summary(content):
    summaries = [item for item in content["items"] if item["title"] == "汇总"]
    assert all(item["documentId"] is None for item in summaries)
    assert content["items"][: len(summaries)] == summaries
    assert [item["n"] for item in content["items"]] == list(range(len(content["items"])))
    return [{f["label"]: f["value"] for f in item["fields"]} for item in summaries]


def _match(doc_id, year=2026):
    return ListMatch(
        document_id=doc_id,
        title=f"第{doc_id}批高新技术企业名单",
        publish_date=date(2026, 10, 1),
        status="公示",
        issuer="示例科技厅",
        province="山东",
        doc_number="示科字1号",
        url=f"https://example.gov.cn/{doc_id}",
        list_type="高新技术企业",
        year=year,
    )


def test_compare_summary_matches_excel_unique_customers_year_and_removed(tmp_path):
    report = CompareReport(
        [
            CompareRow("甲公司", "甲公司", "name", "命中", [_match(1), _match(2), _match(3, None)]),
            CompareRow("乙公司", "乙公司", "name", "多个同名", [_match(1)]),
            CompareRow("丙公司", "丙公司", "name", "名称相似", [_match(2)]),
            CompareRow("丁公司", "丁公司", "name", "未查到"),
        ],
        [RemovedInput("", "空行"), RemovedInput("", "空行"), RemovedInput("甲公司", "重复行")],
        NOW,
    )
    output = tmp_path / "compare.xlsx"
    assert export_compare_excel(report, output) == 6
    native = native_results.build("compare", report, generated_at=NOW)
    summaries, note = _excel_summary(output)
    assert (
        _native_summary(native)
        == summaries
        == [
            {"名单类型": "高新技术企业", "年度": "2026", "命中家数": "2"},
            {"名单类型": "高新技术企业", "年度": "—", "命中家数": "1"},
        ]
    )
    for removed in ("空行2行", "重复行1行"):
        assert removed in native["notice"] and removed in note
    details = [item for item in native["items"] if item["title"] != "汇总"]
    assert len(details) == report.detail_row_count == 6
    assert [item["documentId"] for item in details] == ["1", "2", "3", "1", "2", None]
    assert any(f == {"label": "客户输入", "value": "丁公司"} for f in details[-1]["fields"])
    # 用词 D-090：命中名单的公示/公布叫"发布阶段"，不再叫"名单状态"
    assert {"label": "发布阶段", "value": "公示"} in details[0]["fields"]
    assert "发布阶段按标题用词标注" in native["notice"]
    assert "名单状态" not in str(native)


def test_compare_no_matches_keeps_no_hit_detail_without_inventing_groups(tmp_path):
    report = CompareReport([CompareRow("甲公司", "甲公司", "name", "未查到")], [], None)
    output = tmp_path / "empty.xlsx"
    assert export_compare_excel(report, output) == 1
    content = native_results.build("compare", report)
    assert _native_summary(content) == _excel_summary(output)[0] == []
    assert len(content["items"]) == 1
    assert content["items"][0]["title"] == "未查到"


def _track_item(doc_id, source, days):
    return TrackItem(
        document_id=doc_id,
        customer_name="甲公司",
        title=f"第{doc_id}份申报通知",
        issuer="示例科技厅",
        province="山东",
        deadline_type="企业申报截止",
        deadline_date=NOW.date() + timedelta(days=days) if days is not None else None,
        deadline_time=None,
        remaining_days=days,
        support="",
        url=f"https://example.gov.cn/{doc_id}",
        deadline_source=source,
    )


def test_track_summary_matches_excel_verified_excerpt_zero_unknown_and_empty(tmp_path):
    report = TrackReport(
        [
            CustomerTrack(
                CustomerProfile("甲公司", "山东"),
                [
                    _track_item(i, source, days)
                    for i, (source, days) in enumerate(
                        [
                            ("verified", 0),
                            ("verified", 7),
                            ("verified", 8),
                            ("verified", 31),
                            ("excerpt", 0),
                            ("excerpt", 14),
                            ("excerpt", 30),
                            ("excerpt", 31),
                            ("unverified", None),
                        ],
                        1,
                    )
                ],
            ),
            CustomerTrack(CustomerProfile("只有摘录", "山东"), [_track_item(10, "excerpt", 0)]),
            CustomerTrack(
                CustomerProfile("只有未知", "山东"), [_track_item(11, "unverified", None)]
            ),
            CustomerTrack(CustomerProfile("没有匹配", "山东")),
            CustomerTrack(CustomerProfile("日期缺失", "山东"), [_track_item(12, "verified", None)]),
        ],
        NOW,
    )
    output = tmp_path / "track.xlsx"
    assert export_track_excel(report, output) == 13
    native = native_results.build("track", report)
    rows = _native_summary(native)
    assert rows == _excel_summary(output)[0]
    fields = (
        "通知数",
        "人工核对过截止日期的",
        "7 天内截止",
        "14 天内截止",
        "30 天内截止",
        "截止未核对的",
    )
    assert [[row[field] for field in fields] for row in rows] == [
        ["9", "4", "2", "5", "6", "5"],
        ["1", "0", "未核对", "1", "1", "1"],
        ["1", "0", "未核对", "未核对", "未核对", "1"],
        ["0", "0", "—", "—", "—", "0"],
        ["1", "1", "未核对", "未核对", "未核对", "0"],
    ]
    assert len(native["items"]) == 5 + report.detail_row_count == 18
    today = next(item for item in native["items"] if item["documentId"] == "1")
    assert {"label": "生成时剩余天数", "value": "0"} in today["fields"]


def test_summary_uses_same_item_and_total_size_budgets(monkeypatch):
    with pytest.raises(CompareInputError, match="单项结果"):
        native_results.build(
            "track", TrackReport([CustomerTrack(CustomerProfile("甲" * 80000, "山东"))], NOW)
        )
    monkeypatch.setattr(native_results, "MAX_BYTES", 1024)
    with pytest.raises(CompareInputError, match="结果内容过多"):
        native_results.build(
            "track", TrackReport([CustomerTrack(CustomerProfile("甲公司", "山东"))], NOW)
        )
