"""给 AI 用的附件文字（youju/attachment_text.py）：只取看得见的字、按阅读顺序分段、表格按行按格记表头、
遮号码、三层存法、撤下旧文字、和附件副本定时任务的配合、实测统计。政府网站和对象存储全部用模拟。"""

import io
import json
import struct
import zipfile
from datetime import date, datetime, timedelta

import openpyxl
import pytest
from sqlalchemy import func, select
from sqlalchemy import text as sql

from youju import attachment_text as at
from youju import attachment_text_report as atr
from youju import attachments as copies
from youju import cli
from youju.crawler import detail as body_rules
from youju.crawler.fetcher import FetchResult
from youju.models import (
    AttachmentCopy,
    AttachmentTextSegment,
    AttachmentTextUrl,
    AttachmentTextVersion,
    Document,
    DocumentAttachment,
    DocumentDetail,
    DocumentMerge,
)

from .test_attachments import (
    APPLY,
    NOW,
    PREFIX,
    RECENT,
    XLS,
    FakeCloud,
    FakeFetcher,
    add_doc,
    client,  # noqa: F401 - pytest 夹具
    make_cfb,
    make_docx,
    make_ole,
    make_pdf,
    make_png,
    obj,
    run,
)

BASE = "https://kjt.ln.gov.cn/files/"
SCAN = make_pdf(images=((0, 0, 300, 200),))  # 整页是图片：扫描件


# ---- 造测试文件 ----

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"


def docx_xml(body: str) -> bytes:
    """用手写的 document.xml 造 docx（测隐藏文字、修订、域、文本框、合并单元格这些细节）。"""
    xml = (
        f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{W}" xmlns:mc="{MC}" '
        f'xmlns:wps="{WPS}"><w:body>{body}</w:body></w:document>'
    )
    rels = (
        '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/'
        '2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("_rels/.rels", rels)
        zf.writestr("word/document.xml", xml)
    return out.getvalue()


def p(*runs: str) -> str:
    return "<w:p>" + "".join(runs) + "</w:p>"


def r(text: str, *, props: str = "") -> str:
    return f'<w:r>{f"<w:rPr>{props}</w:rPr>" if props else ""}<w:t xml:space="preserve">{text}</w:t></w:r>'


def tc(text: str, props: str = "") -> str:
    return f"<w:tc>{f'<w:tcPr>{props}</w:tcPr>' if props else ''}{p(r(text)) if text else '<w:p/>'}</w:tc>"


def tr(*cells: str, props: str = "") -> str:
    return f"<w:tr>{f'<w:trPr>{props}</w:trPr>' if props else ''}{''.join(cells)}</w:tr>"


def make_word97(paras, *, encrypted=False, nfib=0x00C1, deleted=()) -> bytes:
    """最小的 Word 97 二进制文件（WordDocument + 1Table 两个数据流，[MS-DOC]）。paras 每项是
    (文字, 结尾标记 0x0D 或 0x07, 段落属性 {in_table, ttp, itap, inner_cell, inner_ttp})：
    文字按 UTF-16LE 放在 0x400（位置、字数都按 UTF-16 码元算，扩展区的字占两个），段落属性写进
    一页 FKP；deleted 里的几段字在字符属性里标成"修订删除"。Table 流里放 PlcBteChpx、PlcBtePapx
    和 Clx（一个分块）。"""
    text = "".join(t + m for t, m, _ in paras)
    text_fc = 0x400
    raw = text.encode("utf-16-le")
    papx_at = ((text_fc + len(raw)) // 512 + 1) * 512
    chpx_at = papx_at + 512
    fcs = [text_fc]
    for t, m, _ in paras:
        fcs.append(fcs[-1] + len((t + m).encode("utf-16-le")))
    crun = len(paras)
    page = bytearray(512)
    struct.pack_into(f"<{crun + 1}I", page, 0, *fcs)
    free = 511
    for i, (_t, _m, props) in enumerate(paras):
        grpprl = b""
        if props.get("in_table"):
            grpprl += struct.pack("<HB", 0x2416, 1)
        if props.get("ttp"):
            grpprl += struct.pack("<HB", 0x2417, 1)
        if props.get("itap"):
            grpprl += struct.pack("<Hi", 0x6649, props["itap"])
        if props.get("inner_cell"):
            grpprl += struct.pack("<HB", 0x244B, 1)
        if props.get("inner_ttp"):
            grpprl += struct.pack("<HB", 0x244C, 1)
        body = struct.pack("<H", 0) + grpprl  # istd + 属性
        papx = (
            bytes([(len(body) + 1) // 2]) + body
            if len(body) % 2
            else bytes([0, len(body) // 2]) + body
        )
        start = free - len(papx)
        start -= start % 2
        page[start : start + len(papx)] = papx
        free = start
        page[4 * (crun + 1) + 13 * i] = start // 2
    page[511] = crun
    # 字符属性：删掉的几段一段一条（标 sprmCFRMarkDel），其余的一条不带属性
    marks = sorted(
        (text_fc + len(text[: text.index(d)].encode("utf-16-le")), len(d.encode("utf-16-le")))
        for d in deleted
    )
    bounds, flags, at = [text_fc], [], text_fc
    for begin, size in marks:
        if begin > at:
            bounds.append(begin)
            flags.append(False)
        bounds.append(begin + size)
        flags.append(True)
        at = begin + size
    if at < fcs[-1]:
        bounds.append(fcs[-1])
        flags.append(False)
    chpx = bytearray(512)
    struct.pack_into(f"<{len(flags) + 1}I", chpx, 0, *bounds)
    chpx_body = bytes([3]) + struct.pack("<HB", 0x0800, 1)
    chpx[500 : 500 + len(chpx_body)] = chpx_body
    for i, flag in enumerate(flags):
        chpx[4 * (len(flags) + 1) + i] = 250 if flag else 0
    chpx[511] = len(flags)
    word = bytearray(chpx_at + 512)
    struct.pack_into("<HH", word, 0, 0xA5EC, nfib)
    struct.pack_into("<H", word, 0x0A, 0x0200 | (0x0100 if encrypted else 0))
    struct.pack_into("<H", word, 0x20, 14)  # csw
    struct.pack_into("<H", word, 0x3E, 22)  # cslw
    struct.pack_into("<i", word, 0x40 + 12, len(raw) // 2)  # ccpText（码元个数）
    struct.pack_into("<H", word, 0x98, 0x5D)  # cbRgFcLcb
    word[text_fc : text_fc + len(raw)] = raw
    word[papx_at : papx_at + 512] = page
    word[chpx_at : chpx_at + 512] = chpx
    plc_papx = struct.pack("<3I", text_fc, fcs[-1], papx_at // 512)
    plc_chpx = struct.pack("<3I", text_fc, fcs[-1], chpx_at // 512)
    plc_pcd = struct.pack("<2i", 0, len(raw) // 2) + struct.pack("<HIH", 0, text_fc, 0)
    clx = b"\x02" + struct.pack("<I", len(plc_pcd)) + plc_pcd
    table = plc_papx + plc_chpx + clx
    struct.pack_into("<II", word, 0x9A + 12 * 8, len(plc_papx), len(plc_chpx))
    struct.pack_into("<II", word, 0x9A + 13 * 8, 0, len(plc_papx))
    struct.pack_into("<II", word, 0x9A + 33 * 8, len(plc_papx) + len(plc_chpx), len(clx))
    return make_cfb({"WordDocument": bytes(word), "1Table": table})


def make_ppt(slides: list[list[str]], boxes: list[list[str]] | None = None) -> bytes:
    """旧版 PowerPoint：文档容器里一个幻灯片文字列表（实例 0）和一个母版文字列表（实例 1）；
    boxes 是各张幻灯片（Slide 容器）绘图里文本框的字；最后放一个母版容器，里面的字不该取。"""

    def rec(ver_inst: int, kind: int, body: bytes) -> bytes:
        return struct.pack("<HHI", ver_inst, kind, len(body)) + body

    def chars(text: str) -> bytes:
        return rec(0, 0x0FA0, text.encode("utf-16-le"))

    atoms = b""
    for texts in slides:
        atoms += rec(0, 0x03F3, bytes(20)) + chars("\r".join(texts))
    slides_list = rec(0x000F, 0x0FF0, atoms)
    masters = rec(0x001F, 0x0FF0, rec(0, 0x03F3, bytes(20)) + chars("母版标题"))
    drawings = b"".join(
        rec(0x000F, 0x03EE, rec(0x000F, 0xF00D, chars("\r".join(t)))) for t in boxes or []
    )
    main_master = rec(0x000F, 0x03F8, chars("单击此处编辑母版"))
    stream = rec(0x000F, 0x03E8, slides_list + masters) + drawings + main_master
    return make_cfb({"PowerPoint Document": stream})


def pdf_pages(pages: list[dict]) -> bytes:
    """多页 PDF（300×400）：每页 {"text": [(x, y, 文字)], "lines": [(x1, y1, x2, y2)],
    "images": [(x, y, 宽, 高)]}。ASCII 文字用 Helvetica；带中文的用宋体（STSong-Light，
    UniGB-UCS2-H 编码，不嵌字体，每个字 10 宽），pdfplumber 照样取得到字。"""
    objs: list[bytes] = [b"", b""]  # 1 目录、2 页面树，最后补
    kids = []
    for page in pages:
        content = b""
        mode = "3 Tr " if page.get("ocr") else ""  # 识别层：文字设成看不见
        for x, y, text in page.get("text", []):
            shown = f"({text})" if text.isascii() else f"<{text.encode('utf-16-be').hex()}>"
            font = "F1" if text.isascii() else "F2"
            content += f"BT {mode}/{font} 10 Tf {x} {y} Td {shown} Tj ET\n".encode()
        for x1, y1, x2, y2 in page.get("lines", []):
            content += f"{x1} {y1} m {x2} {y2} l S\n".encode()
        for i in range(page.get("curves", 0)):  # 文字转成的曲线
            x, y = 20 + (i % 20) * 12, 380 - (i // 20) * 12
            content += f"{x} {y} m {x + 3} {y + 8} {x + 6} {y + 8} {x + 9} {y} c f\n".encode()
        for x, y, w, h in page.get("images", []):
            content += f"q {w} 0 0 {h} {x} {y} cm /Im1 Do Q\n".encode()
        objs.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
        content_id = len(objs)
        objs.append(b"")  # 页面，等字体和图片编号定了再写
        kids.append((len(objs), content_id))
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    font_id = len(objs)
    objs.append(
        b"<< /Type /FontDescriptor /FontName /STSong-Light /Flags 4 /FontBBox [0 -200 1000 900]"
        b" /ItalicAngle 0 /Ascent 880 /Descent -120 /CapHeight 880 /StemV 80 >>"
    )
    objs.append(
        b"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /STSong-Light /FontDescriptor %d 0 R"
        b" /CIDSystemInfo << /Registry (Adobe) /Ordering (GB1) /Supplement 4 >> >>" % (len(objs))
    )
    objs.append(
        b"<< /Type /Font /Subtype /Type0 /BaseFont /STSong-Light /Encoding /UniGB-UCS2-H"
        b" /DescendantFonts [%d 0 R] >>" % len(objs)
    )
    cjk_id = len(objs)
    pixels = bytes([0, 128, 128, 255])
    objs.append(
        b"<< /Type /XObject /Subtype /Image /Width 2 /Height 2 /ColorSpace /DeviceGray"
        b" /BitsPerComponent 8 /Length 4 >>\nstream\n" + pixels + b"\nendstream"
    )
    image_id = len(objs)
    for page_id, content_id in kids:
        objs[page_id - 1] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 400] /Contents %d 0 R"
            b" /Resources << /Font << /F1 %d 0 R /F2 %d 0 R >> /XObject << /Im1 %d 0 R >> >> >>"
            % (content_id, font_id, cjk_id, image_id)
        )
    objs[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objs[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        b" ".join(b"%d 0 R" % page_id for page_id, _ in kids),
        len(kids),
    )
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return out


def texts(result: at.Extraction) -> list[str]:
    return [s.text for s in result.segments]


# ---- 认格式 ----


def test_sniff_recognizes_formats_and_reasons() -> None:
    assert at.sniff(make_pdf("x")) == "pdf"
    assert at.sniff(make_docx("申报指南")) == "docx"
    assert at.sniff(make_ole("WordDocument")) == "doc"  # wps 多半就是这样
    assert at.sniff(XLS) == "xls"  # et 多半就是这样
    assert at.sniff(make_ppt([["标题"]])) == "ppt"
    cases = {
        b"": "broken",
        b"<!DOCTYPE html><html><body>404</body></html>": "webpage",
        b"\r\n  <html><head></head></html>": "webpage",
        b"Rar!\x1a\x07\x00xxxx": "format",
        b"{\\rtf1\\ansi hello}": "format",
        b"\x00\x01\x02 random": "format",
        b"PK\x03\x04garbage": "broken",
    }
    for data, reason in cases.items():
        with pytest.raises(at.Unreadable) as exc:
            at.sniff(data)
        assert exc.value.reason == reason, data
    plain_zip, ofd = io.BytesIO(), io.BytesIO()
    with zipfile.ZipFile(plain_zip, "w") as zf:
        zf.writestr("a.txt", "x")
    with zipfile.ZipFile(ofd, "w") as zf:
        zf.writestr("OFD.xml", "<ofd/>")
    encrypted = make_cfb({"EncryptionInfo": b"x", "EncryptedPackage": b"y"})
    for data, reason, detail in (
        (plain_zip.getvalue(), "format", "压缩包"),
        (ofd.getvalue(), "format", "OFD"),
        (encrypted, "broken", "加密"),
    ):
        with pytest.raises(at.Unreadable) as exc:
            at.sniff(data)
        assert exc.value.reason == reason and detail in exc.value.detail
    # extract 不抛异常，把原因写在结果里
    result = at.extract(b"<html><body>page</body></html>")
    assert (result.ok, result.reason, result.file_type, result.segments) == (
        False,
        "webpage",
        None,
        [],
    )


# ---- 各种格式 ----


def test_docx_takes_visible_text_in_reading_order() -> None:
    textbox = (
        f'<w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wps:txbx>'
        f"<w:txbxContent>{p(r('文本框里的字'))}</w:txbxContent></wps:txbx></w:drawing>"
        f"</mc:Choice><mc:Fallback><w:pict><w:txbxContent>{p(r('文本框里的字'))}"
        f"</w:txbxContent></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
    )
    body = "".join(
        [
            p(r("　　一、申报条件")),
            p(r("看得见"), r("隐藏的字", props="<w:vanish/>"), r("的字")),
            p(
                r("改后"),
                "<w:del><w:r><w:delText>删掉的</w:delText></w:r></w:del>",
                "<w:moveFrom><w:r><w:t>移走的</w:t></w:r></w:moveFrom>",
                r("留下"),
            ),
            p(
                '<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
                '<w:r><w:instrText> HYPERLINK "http://x" </w:instrText></w:r>',
                '<w:r><w:fldChar w:fldCharType="separate"/></w:r>',
                r("链接文字"),
                '<w:r><w:fldChar w:fldCharType="end"/></w:r>',
            ),
            p(
                r("勾选"),
                '<w:r><w:sym w:font="Wingdings" w:char="F0A8"/></w:r>',
                r("企业"),
                textbox,
            ),
            p(),
            "<w:tbl>"
            + tr(tc("2026年申报汇总表", '<w:gridSpan w:val="3"/>'))
            + tr(tc("序号"), tc("项目名称"), tc("截止时间"))
            + tr(tc("1"), tc("智能制造", '<w:vMerge w:val="restart"/>'), tc("2026年11月5日"))
            + tr(tc("2"), tc("", "<w:vMerge/>"), f"<w:tc>{p(r('第一行'))}{p(r('第二行'))}</w:tc>")
            + "</w:tbl>",
            "<w:tbl>" + tr(tc("单位名称"), tc(""), tc("成立时间"), tc("")) + "</w:tbl>",
            p(r("表后段落")),
        ]
    )
    result = at.extract(docx_xml(body))
    assert result.ok and result.file_type == "docx"
    assert texts(result) == [
        "一、申报条件",
        "看得见的字",
        "改后留下",
        "链接文字",
        "勾选□企业",
        "文本框里的字",
        "2026年申报汇总表",
        "序号 | 项目名称 | 截止时间",
        "1 | 智能制造 | 2026年11月5日",
        "2 | 第一行\n第二行",
        "单位名称 | 成立时间",
        "表后段落",
    ]
    segs = {s.text: s for s in result.segments}
    title, head = segs["2026年申报汇总表"], segs["序号 | 项目名称 | 截止时间"]
    assert (title.kind, title.loc, title.cells) == (
        "row",
        {"table": 1, "row": 1},
        [{"c": 0, "t": "2026年申报汇总表"}],
    )
    assert (head.kind, head.loc) == ("head", {"table": 1, "row": 2})
    row = segs["1 | 智能制造 | 2026年11月5日"]
    assert row.cells == [
        {"c": 0, "t": "1", "h": "序号"},
        {"c": 1, "t": "智能制造", "h": "项目名称"},
        {"c": 2, "t": "2026年11月5日", "h": "截止时间"},
    ]
    # 纵向合并的下一格不重复取字；填写表（中间有空格）不认表头
    assert segs["2 | 第一行\n第二行"].cells[1] == {"c": 2, "t": "第一行\n第二行", "h": "截止时间"}
    form = segs["单位名称 | 成立时间"]
    assert (form.kind, form.loc, form.cells) == (
        "row",
        {"table": 2, "row": 1},
        [{"c": 0, "t": "单位名称"}, {"c": 2, "t": "成立时间"}],
    )
    assert [s.n for s in result.segments] == list(range(1, 13))


def test_xlsx_rows_with_headers_dates_and_hidden_parts() -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "汇总"
    sheet.append(["2026年项目申报汇总表"])
    sheet.merge_cells("A1:D1")
    sheet.append(["序号", "单位", "截止日期", "占比"])
    sheet.append([1, "甲公司", datetime(2026, 11, 5), 0.25])
    sheet.append([2, "隐藏行", datetime(2026, 11, 6), 0.5])
    sheet.append([3, "乙公司", datetime(2026, 11, 7, 17, 30), 1])
    sheet.row_dimensions[4].hidden = True
    sheet["C3"].number_format = "yyyy年m月d日"
    sheet["C5"].number_format = "yyyy-mm-dd hh:mm"
    sheet["D3"].number_format = "0%"
    sheet.column_dimensions["E"].hidden = True
    sheet["E3"] = "隐藏列"
    hidden = book.create_sheet("说明")
    hidden.append(["隐藏工作表"])
    hidden.sheet_state = "hidden"
    book.create_sheet("附表").append(["备注：以原文为准"])
    out = io.BytesIO()
    book.save(out)
    result = at.extract(out.getvalue())
    assert result.ok and result.file_type == "xlsx"
    assert texts(result) == [
        "2026年项目申报汇总表",
        "序号 | 单位 | 截止日期 | 占比",
        "1 | 甲公司 | 2026-11-05 | 25%",
        "3 | 乙公司 | 2026-11-07 17:30 | 1",
        "备注：以原文为准",
    ]
    first = result.segments[2]
    assert first.loc == {"sheet": "汇总", "row": 3} and first.kind == "row"
    assert first.cells[1] == {"c": 1, "t": "甲公司", "h": "单位"}
    assert result.segments[0].cells == [{"c": 0, "t": "2026年项目申报汇总表"}]
    assert result.segments[3].loc == {"sheet": "汇总", "row": 5}
    assert result.segments[4].loc == {"sheet": "附表", "row": 1}


def test_pptx_slides_in_order_and_hidden_slides_skipped() -> None:
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    pns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    rns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    def slide(body: str, show: str = "") -> str:
        return (
            f'<p:sld xmlns:p="{pns}" xmlns:a="{a}"{show}><p:cSld><p:spTree>{body}'
            "</p:spTree></p:cSld></p:sld>"
        )

    shape = (
        "<p:sp><p:txBody><a:p><a:r><a:t>申报</a:t></a:r><a:r><a:t>指南</a:t></a:r></a:p>"
        "<a:p><a:r><a:t>第二段</a:t></a:r></a:p></p:txBody></p:sp>"
    )

    def cell(text: str) -> str:
        return f"<a:tc><a:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></a:txBody></a:tc>"

    table = (
        "<p:graphicFrame><a:graphic><a:graphicData><a:tbl>"
        f"<a:tr>{cell('指标')}{cell('分值')}{cell('说明')}</a:tr>"
        f"<a:tr>{cell('创新性')}{cell('30')}{cell('看核心技术')}</a:tr>"
        "</a:tbl></a:graphicData></a:graphic></p:graphicFrame>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        zf.writestr(
            "ppt/presentation.xml",
            f'<p:presentation xmlns:p="{pns}" xmlns:r="{rns}"><p:sldIdLst>'
            '<p:sldId id="257" r:id="rId3"/><p:sldId id="256" r:id="rId2"/>'
            '<p:sldId id="258" r:id="rId4"/></p:sldIdLst></p:presentation>',
        )
        zf.writestr(
            "ppt/_rels/presentation.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId2" Target="slides/slide1.xml"/>'
            '<Relationship Id="rId3" Target="slides/slide2.xml"/>'
            '<Relationship Id="rId4" Target="slides/slide3.xml"/></Relationships>',
        )
        zf.writestr("ppt/slides/slide1.xml", slide(table))
        zf.writestr("ppt/slides/slide2.xml", slide(shape))
        zf.writestr("ppt/slides/slide3.xml", slide(shape, ' show="0"'))
    result = at.extract(out.getvalue())
    assert texts(result) == ["申报指南", "第二段", "指标 | 分值 | 说明", "创新性 | 30 | 看核心技术"]
    assert [s.loc for s in result.segments] == [
        {"slide": 1},
        {"slide": 1},
        {"slide": 2, "table": 1, "row": 1},
        {"slide": 2, "table": 1, "row": 2},
    ]
    assert result.segments[3].cells[1] == {"c": 1, "t": "30", "h": "分值"}


def test_pdf_lines_tables_page_numbers_and_scanned_pages() -> None:
    long1 = "Applicants must be registered in the province and must hav"
    long2 = "e at least three years of operating history before the dea"
    table_lines = [(40, 300, 260, 300), (40, 280, 260, 280), (40, 260, 260, 260)]
    table_lines += [(40, 300, 40, 260), (120, 300, 120, 260), (190, 300, 190, 260)]
    table_lines += [(260, 300, 260, 260)]
    data = pdf_pages(
        [
            {
                "text": [
                    (40, 370, long1),
                    (40, 355, long2),
                    (40, 340, "dline."),
                    (40, 325, "Materials: application form and licence copy, all signed."),
                    (45, 285, "Item"),
                    (125, 285, "Score"),
                    (195, 285, "Note"),
                    (45, 265, "Innovation"),
                    (125, 265, "30"),
                    (195, 265, "Core"),
                    (140, 20, "- 1 -"),
                ],
                "lines": table_lines,
            },
            {"images": [(0, 0, 300, 400)], "text": [(40, 300, "HIDDEN OCR TEXT LAYER")], "ocr": 1},
            {"images": [(50, 50, 200, 200)]},
            {"text": [(40, 370, "Page four keeps going.")]},
            # 整页底图上的正常文字（演示文稿导出的 PDF）照样取
            {"images": [(0, 0, 300, 400)], "text": [(40, 300, "Slide text over a background")]},
            # 只有几个字加一枚公章：字照样取，不算扫描页
            {"images": [(200, 20, 60, 60)], "text": [(40, 370, "Seal.")]},
        ]
    )
    result = at.extract(data)
    assert result.ok and (result.pages, result.scanned_pages) == (6, 2)
    assert texts(result) == [
        long1 + long2 + "dline.",
        "Materials: application form and licence copy, all signed.",
        "Item | Score | Note",
        "Innovation | 30 | Core",
        "Page four keeps going.",
        "Slide text over a background",
        "Seal.",
    ]
    assert result.segments[0].loc == {"page": 1}
    assert result.segments[2].loc == {"page": 1, "table": 1, "row": 1}
    assert result.segments[3].cells == [
        {"c": 0, "t": "Innovation", "h": "Item"},
        {"c": 1, "t": "30", "h": "Score"},
        {"c": 2, "t": "Core", "h": "Note"},
    ]
    assert result.segments[4].loc == {"page": 4}
    # 每页都是扫描页：取不出
    scan = at.extract(pdf_pages([{"images": [(0, 0, 300, 400)]}, {"images": [(10, 10, 200, 200)]}]))
    assert (scan.reason, scan.scanned_pages, scan.segments) == ("scan", 2, [])
    assert at.extract(pdf_pages([{}])).reason == "empty"
    assert at.extract(b"%PDF-1.4 broken").reason == "broken"


def test_word97_and_wps_paragraphs_tables_and_fields() -> None:
    cell = {"in_table": 1}
    end = {"in_table": 1, "ttp": 1}
    astral = chr(0x2B404)  # 扩展区的字（UTF-16 两个码元），后面的位置不能错开
    data = make_word97(
        [
            ("附件1", "\r", {}),
            ('申报条件：注册满一年。详见\x13 HYPERLINK "http://x" \x14申报系统\x15', "\r", {}),
            ("修订：已删掉的字留下的字", "\r", {}),
            ("序号", "\x07", cell),
            ("姓名", "\x07", cell),
            ("要求", "\x07", cell),
            ("", "\x07", end),
            ("1", "\x07", cell),
            ("王" + astral, "\x07", cell),
            ("第一行\r第二行", "\x07", cell),
            ("", "\x07", end),
            ("2", "\x07", cell),
            ("李四", "\x07", cell),
            ("无", "\x07", cell),
            ("", "\x07", end),
            ("表后\x0b换行\x01", "\r", {}),
        ],
        deleted=("已删掉的字",),
    )
    result = at.extract(data)
    assert result.ok and result.file_type == "doc"
    assert texts(result) == [
        "附件1",
        "申报条件：注册满一年。详见申报系统",
        "修订：留下的字",
        "序号 | 姓名 | 要求",
        f"1 | 王{astral} | 第一行\n第二行",
        "2 | 李四 | 无",
        "表后\n换行",
    ]
    assert result.segments[4].loc == {"table": 1, "row": 2}
    assert result.segments[4].cells == [
        {"c": 0, "t": "1", "h": "序号"},
        {"c": 1, "t": "王" + astral, "h": "姓名"},
        {"c": 2, "t": "第一行\n第二行", "h": "要求"},
    ]
    # 表格里套的表格：内层的格、行接成一句放进外面那一格
    outer = {"in_table": 1, "itap": 1}
    inner = {"in_table": 1, "itap": 2}
    nested = make_word97(
        [
            ("外层", "\x07", outer),
            ("甲", "\r", {**inner, "inner_cell": 1}),
            ("乙", "\r", {**inner, "inner_cell": 1}),
            ("", "\r", {**inner, "inner_ttp": 1}),
            ("", "\x07", outer),
            ("", "\x07", {**outer, "ttp": 1}),
            ("表后", "\r", {}),
        ]
    )
    assert texts(at.extract(nested)) == ["外层 | 甲 | 乙", "表后"]
    assert at.extract(make_word97([("x", "\r", {})], encrypted=True)).reason == "broken"
    # 只有图片（0x01 占位）没有字：扫描件
    assert at.extract(make_word97([("\x01", "\r", {}), ("\x01", "\r", {})])).reason == "scan"
    old = at.extract(make_word97([("x", "\r", {})], nfib=0x0068))
    assert (old.reason, old.detail) == ("format", "Word 95 及更早的格式")


def test_xls_and_ppt() -> None:
    result = at.extract(XLS)
    assert result.ok and result.file_type == "xls" and result.segments
    rows = [s for s in result.segments if s.kind in ("row", "head")]
    assert rows and all(s.loc["sheet"] and s.loc["row"] >= 1 for s in rows)
    assert all(s.cells for s in rows)
    ppt = at.extract(make_ppt([["第一页标题", "第一页正文"], ["第二页"]], [["文本框"], []]))
    assert texts(ppt) == ["第一页标题", "第一页正文", "文本框", "第二页"]
    assert [s.loc for s in ppt.segments] == [{"slide": 1}] * 3 + [{"slide": 2}]


def test_pictures_only_is_scan_and_no_text_is_empty() -> None:
    from docx import Document as WordDocument

    doc = WordDocument()
    doc.add_picture(io.BytesIO(make_png(150)))
    out = io.BytesIO()
    doc.save(out)
    assert at.extract(out.getvalue()).reason == "scan"
    assert at.extract(docx_xml(p())).reason == "empty"
    # 带公章等大图片的文件照样取字（附件副本因为图片不存，文字照取）
    seal = io.BytesIO()
    doc = WordDocument()
    doc.add_paragraph("盖章页：申报单位（公章）")
    doc.add_picture(io.BytesIO(make_png(150)))
    doc.save(seal)
    assert copies.check_content(seal.getvalue(), "docx").code == "image"
    assert texts(at.extract(seal.getvalue())) == ["盖章页：申报单位（公章）"]


def test_pdf_continued_tables_dedupe_body_numbers_and_outlined_text() -> None:
    def ruled(top: float, rows: int) -> list:
        """三列的表格线：从 top 往下 rows 行，每行 20 高。"""
        bottom = top - 20 * rows
        lines = [(40, top - 20 * i, 260, top - 20 * i) for i in range(rows + 1)]
        return lines + [(x, top, x, bottom) for x in (40, 120, 190, 260)]

    def row(y: float, *cells: str) -> list:
        return [(x, y, text) for x, text in zip((45, 125, 195), cells, strict=True)]

    data = pdf_pages(
        [
            {
                "text": [(40, 370, "Annex 1: application guide for provincial key projects 2026.")]
                + row(85, "Item", "Amount", "Deadline")
                + row(65, "AI", "500", "Nov 5")
                + [(150, 200, "2025"), (150, 20, "1")],  # 正文里单独一行的数字留着，页脚页码删
                "lines": ruled(100, 2),
            },
            # 续表：这一页一开始就是表、竖线位置一样；第一行是没有数字的数据，也不当表头
            {
                "text": row(385, "Biotech", "up to two", "Dec first")
                + row(365, "Robot", "300", "Nov 9")
                + row(345, "Chip", "200", "Nov 9")
                + [(40, 300, "Bold title"), (40.3, 300, "Bold title")],  # 为了加粗重画一遍
                "lines": ruled(400, 3),
            },
            {"curves": 150},  # 字被转成线条：算扫描页
        ]
    )
    result = at.extract(data)
    assert result.ok and (result.pages, result.scanned_pages) == (3, 1)
    assert texts(result) == [
        "Annex 1: application guide for provincial key projects 2026.",
        "2025",
        "Item | Amount | Deadline",
        "AI | 500 | Nov 5",
        "Biotech | up to two | Dec first",
        "Robot | 300 | Nov 9",
        "Chip | 200 | Nov 9",
        "Bold title",
    ]
    biotech = result.segments[4]
    assert (biotech.kind, biotech.loc) == ("row", {"page": 2, "table": 1, "row": 3})
    assert biotech.cells[0] == {"c": 0, "t": "Biotech", "h": "Item"}


def test_pdf_repeated_header_rows_and_a_new_table_at_the_top_of_a_page() -> None:
    def ruled(top: float, rows: int, xs=(40, 120, 190, 260)) -> list:
        bottom = top - 20 * rows
        lines = [(xs[0], top - 20 * i, xs[-1], top - 20 * i) for i in range(rows + 1)]
        return lines + [(x, top, x, bottom) for x in xs]

    def row(y: float, *cells: str, xs=(45, 125, 195)) -> list:
        return [(x, y, text) for x, text in zip(xs, cells, strict=True)]

    data = pdf_pages(
        [
            {
                "text": row(85, "Item", "Amount", "Deadline") + row(65, "AI", "500", "Nov 5"),
                "lines": ruled(100, 2),
            },
            # Word"标题行重复"导出：每页开头重复表头
            {
                "text": row(385, "Item", "Amount", "Deadline") + row(365, "Robot", "300", "Nov 9"),
                "lines": ruled(400, 2),
            },
            # 开头是另一张表（竖线位置不一样）：不当续表
            {
                "text": row(385, "Name", "Score", "Note", xs=(45, 105, 225))
                + row(365, "Ann", "90", "fine", xs=(45, 105, 225)),
                "lines": ruled(400, 2, xs=(40, 100, 220, 260)),
            },
        ]
    )
    result = at.extract(data)
    kinds = [(x.kind, x.text, x.loc["table"], x.loc["row"]) for x in result.segments]
    assert kinds == [
        ("head", "Item | Amount | Deadline", 1, 1),
        ("row", "AI | 500 | Nov 5", 1, 2),
        ("head", "Item | Amount | Deadline", 1, 3),
        ("row", "Robot | 300 | Nov 9", 1, 4),
        ("head", "Name | Score | Note", 2, 1),
        ("row", "Ann | 90 | fine", 2, 2),
    ]
    assert result.segments[3].cells[0] == {"c": 0, "t": "Robot", "h": "Item"}


def test_pdf_continued_table_keeps_account_columns() -> None:
    xs = (20, 50, 110, 290)

    def ruled(top: float, rows: int) -> list:
        bottom = top - 20 * rows
        lines = [(xs[0], top - 20 * i, xs[-1], top - 20 * i) for i in range(rows + 1)]
        return lines + [(x, top, x, bottom) for x in xs]

    def row(y: float, *cells: str) -> list:
        return [(x + 5, y, text) for x, text in zip(xs[:-1], cells, strict=True)]

    data = pdf_pages(
        [
            {
                "text": row(85, "序号", "单位名称", "银行账号")
                + row(65, "1", "甲公司", "110060149018001234567"),
                "lines": ruled(100, 2),
            },
            # 续表：没有重复表头，这一页也没有"账号"两个字，15 位的对公账号照样遮
            {
                "text": row(385, "2", "乙公司", "123456789012345")
                + row(365, "3", "丙公司", "待补"),
                "lines": ruled(400, 2),
            },
        ]
    )
    result = at.extract(data)
    assert texts(result) == [
        "序号 | 单位名称 | 银行账号",
        "1 | 甲公司 | [银行账号已脱敏]",
        "2 | 乙公司 | [银行账号已脱敏]",
        "3 | 丙公司 | 待补",
    ]
    assert result.segments[2].loc == {"page": 2, "table": 1, "row": 3}
    assert result.segments[2].cells[2] == {"c": 2, "t": "[银行账号已脱敏]", "h": "银行账号"}
    assert result.masked == 2


def test_few_characters_with_big_pictures_count_as_scans(monkeypatch) -> None:
    from docx import Document as WordDocument

    doc = WordDocument()
    doc.add_paragraph("附件2")
    doc.add_picture(io.BytesIO(make_png(300)))  # 约 270KB：一整页扫描图
    out = io.BytesIO()
    doc.save(out)
    result = at.extract(out.getvalue())
    assert (result.reason, result.segments) == ("scan", [])
    assert "只有 3 个字" in result.detail
    # 字多的照样能读（公章、徽标再大也不算扫描件）：几十个字的延期说明加一枚大公章
    doc = WordDocument()
    doc.add_paragraph("关于延长申报时间的说明：受理截止时间延长至2026年11月20日17:00。")
    doc.add_picture(io.BytesIO(make_png(300)))
    out = io.BytesIO()
    doc.save(out)
    assert at.extract(out.getvalue()).ok


def test_xlsx_escapes_time_limit_and_part_size_limit(monkeypatch) -> None:
    book = openpyxl.Workbook()
    book.active.append(["材料清单", "备注", "份数"])
    book.active.append(["营业执照_x000D_\n申报书", "原件", "1"])
    out = io.BytesIO()
    book.save(out)
    raw = out.getvalue()
    assert texts(at.extract(raw))[1] == "营业执照\n申报书 | 原件 | 1"
    # 单个部件太大：不整份载入内存，按太大记
    monkeypatch.setattr(at, "_TREE_MAX", 100)
    assert at.extract(make_docx("申报指南")).reason == "too_large"
    monkeypatch.undo()
    # 一份文件取太久：按太大记，附件任务不卡住
    import time

    def slow(_data):
        time.sleep(5)

    monkeypatch.setattr(at, "EXTRACT_SECONDS", 1)
    monkeypatch.setitem(at._EXTRACTORS, "docx", slow)
    started = time.monotonic()
    result = at.extract(make_docx("申报指南"))
    assert (result.reason, result.detail) == ("too_large", "取文字超过 1 秒")
    assert time.monotonic() - started < 3


# ---- 分段、表头、上限、遮号码 ----


def test_long_paragraphs_split_at_sentence_ends_and_char_limit(monkeypatch) -> None:
    sentence = "申报单位须在省内注册并具有独立法人资格。"
    result = at.extract(make_docx(sentence * 120))
    assert all(len(s.text) <= at.SEGMENT_MAX for s in result.segments)
    assert all(s.text.endswith("。") for s in result.segments)
    assert "".join(texts(result)) == sentence * 120
    monkeypatch.setattr(at, "MAX_CHARS", 100)
    big = at.extract(make_docx(*[sentence] * 10))
    assert (big.reason, big.segments) == ("too_large", [])


def _row(*texts: str, header: bool = False) -> at._Row:
    return at._Row([at._Cell(i, t) for i, t in enumerate(texts)], {}, header)


def test_header_rules() -> None:
    rows = [_row("汇总表"), _row("序号", "名称", "金额"), _row("1", "甲", "10")]
    assert at._header_rows(rows) == {1}
    # 填写行跳过，下一行是表头
    form = [_row("名称", "", "日期", ""), _row("类别", "说明", "备注"), _row("a", "b", "c")]
    assert at._header_rows(form) == {1}
    assert at._header_rows([_row("1", "2", "3"), _row("4", "5", "6")]) == set()  # 全是数字
    assert at._header_rows([_row("序号", "名称", "金额")]) == set()  # 只有一行
    assert at._header_rows([_row("x" * 31, "名称", "金额"), _row("1", "2", "3")]) == set()
    marked = [_row("a", "b", header=True), _row("c", "d", header=True), _row("1", "2")]
    assert at._header_rows(marked) == {0, 1}
    # 长的说明行、只占一部分列的"填报单位 | 填报日期"跳过，下一行才是表头
    note = _row(
        "请各单位认真填写本表，逐项核对无误后加盖公章，于规定时间内报送省科技厅",
        "填表人",
        "编号：001",
    )
    head, data = _row("序号", "项目名称", "截止日期"), _row("1", "甲", "2026-11-05")
    assert at._header_rows([note, head, data]) == {1}
    meta = _row("填报单位", "填报日期")
    wide = _row("序号", "单位", "项目", "金额", "备注")
    assert at._header_rows([meta, wide, _row("1", "甲", "乙", "10", "无")]) == {1}
    # 数据行（有一格只是数字）不当表头；左上角空着的表头照样算
    assert at._header_rows([_row("1", "甲公司", "乙"), _row("2", "乙公司", "丙")]) == set()
    corner = at._Row([at._Cell(1, "2024年"), at._Cell(2, "2025年")], {})
    assert at._header_rows([corner, _row("收入", "100", "200")]) == {0}
    # 两列的表（名目 | 内容）不认列表头：第一行也是内容
    pairs = [
        _row("申报主体", "省内注册企业"),
        _row("支持方式", "无偿资助"),
        _row("截止", "11月5日"),
    ]
    assert at._header_rows(pairs) == set()
    # 三列表的备注大多空着，照样认表头；表头前有好几行说明（前 6 行里找）；年份格可以是表头
    remarks = [_row("序号", "材料名称", "备注"), _row("1", "申报书", ""), _row("2", "营业执照", "")]
    assert at._header_rows(remarks) == {0}
    lead = [_row("附件2"), _row("申报材料清单"), _row("填报单位：", "", ""), _row("单位：份")]
    assert at._header_rows([*lead, *remarks]) == {4}
    years = [_row("项目", "2023", "2024", "2025"), _row("收入", "100", "200", "300")]
    assert at._header_rows(years) == {0}


def test_mask_uses_body_rules_then_spaced_forms() -> None:
    cases = {
        "联系人：张三，电话13812345678": "联系人：[姓名已脱敏]，电话[手机号已脱敏]",
        "联系电话：138 1234 5678": "联系电话：[手机号已脱敏]",
        "手机：１３８－１２３４－５６７８": "手机：[手机号已脱敏]",
        "身份证 110101 19900307 1234": "身份证 [身份证号已脱敏]",
        "身份证号：110101199003071234": "身份证号：[身份证号已脱敏]",
        "账号：6222 0212 3456 7890 123": "账号：[银行账号已脱敏]",
        "卡号 4111 1111 1111 1111": "卡号 [银行账号已脱敏]",
        "联 系 人：张蕾": "联 系 人：[姓名已脱敏]",
        # 只遮邮箱地址：正文规则会把"于7月2日前发送至"一起遮掉
        "PDF电子版于7月2日前发送至ln_cx@163.com邮箱。": "PDF电子版于7月2日前发送至[邮箱已脱敏]邮箱。",
    }
    for raw, masked in cases.items():
        assert at.mask(raw) == masked, raw
    # 正文规则本身会吞字（见 attachment_text._EMAIL）：附件不能照搬
    assert body_rules.redact_sensitive("于7月2日前发送至a@b.com") == "[邮箱已脱敏]"
    # 写了"账号"的：对公账号过不了 Luhn、可能超过 19 位，也不留尾巴；后面跟着的年份不算进去
    assert at.mask("开户账号：33050161613500000123") == "开户账号：[银行账号已脱敏]"
    assert at.mask("账号：3305 0161 6135 0000 0123") == "账号：[银行账号已脱敏]"
    assert at.mask("账号：1234567890 2025年11月5日前") == "账号：[银行账号已脱敏] 2025年11月5日前"
    # 写了"身份证""手机"的怎么分组都认；没写的只认整串、常见分法（6-8-4、6-4-4-4、4-4-4-4-2、
    # 3-4-4、连写），跨行写的也一样
    for raw, masked in {
        "身份证号：110101 1990 0307 1234": "身份证号：[身份证号已脱敏]",
        "身份证：11010 11990 03071 234": "身份证：[身份证号已脱敏]",
        "1101 0119 9003 0712 34": "[身份证号已脱敏]",
        "手机：1381 234 5678": "手机：[手机号已脱敏]",
        "手机：13 8123 45678": "手机：[手机号已脱敏]",
        "138 1234 5678": "[手机号已脱敏]",
        "+86 138 1234 5678": "+[手机号已脱敏]",
        "4111111111111111": "[银行账号已脱敏]",
        "4111 1111 1111 1111": "[银行账号已脱敏]",
        "138\n1234\n5678": "[手机号已脱敏]",
    }.items():
        assert at.mask(raw) == masked, raw
    for keep in (
        "请于2026年10月20日前提交，支持金额 1,000,000 元",
        "统一社会信用代码 91110113071669887D",
        "项目编号 2026-1234-5678",
        "第 1 | 13812345 | 678 行",  # 不同格子的数字不连起来
        # 日期区间、表格里空格隔开的一串数字：不拼成卡号（以前约一成会被当成卡号遮掉）
        "申报时间：2025-10-10—2025-10-24",
        "2025-10-13-2025-11-12",
        "合计 1099 800 450 300 200",
        "2025年 1 2 3 4 5 6 7 8 9 10 11 12 月",
        # 四组年份、整数：不当卡号；"13 800 000 000 元"是金额，不当手机号
        "年度 2023 2024 2025 2026",
        "补助标准（元） 1000 1500 2000 8000",
        "5000 6000 7000 8000",
        "13 800 000 000 元",
        "13 8123 45678",  # 没写"手机"、不是常见分法：不当手机号
        "13812 345678",
        "11010 11990 03071 234",  # 没写"身份证"、分法少见：不当身份证号（写了标签的才任意分组）
        "2026110100000002",  # 过得了 Luhn，但 2 开头不像卡号（编号、年份多半这样）
        "4100 4100 4100 4500",  # 过得了 Luhn，但每组都是整百：金额
        # 一串金额里任取几组拼起来像号码的，不遮（上一版有 4%–21% 的金额行被误遮）
        "合计 1350 820 4500",
        "经费 1600 300 5000 万元",
        "预算 159 7000 7400",
        "138 1234 5678 元",
        "1 2 3 4 5 6 7 8 9 10 11 12 13 14 15",
        "1350\n820\n4500",
    ):
        assert at.mask(keep) == keep, keep
    builder = at._Builder()
    builder.para("联系电话：138 1234 5678")
    builder.table([_row("姓名", "手机", "备注"), _row("张三", "13912345678", "无")])
    assert builder.masked == 2
    assert builder.segments[0].text == "联系电话：[手机号已脱敏]"
    assert builder.segments[2].cells[1] == {"c": 1, "t": "[手机号已脱敏]", "h": "手机"}
    # 表格里"账号"写在左边那格或列表头：这一格的长数字也遮（对公账号过不了 Luhn）
    builder = at._Builder()
    builder.table(
        [_row("开户银行", "工商银行", "银行账号", "1202021309001234568"), _row("备注", "无")]
    )
    builder.table([_row("单位", "账号", "金额"), _row("甲", "33050161613500000123", "100")])
    assert builder.segments[0].text == "开户银行 | 工商银行 | 银行账号 | [银行账号已脱敏]"
    assert builder.segments[3].text == "甲 | [银行账号已脱敏] | 100"
    # 分散对齐的"账  号"、多级表头的最后一级是"银行账号"
    builder = at._Builder()
    builder.table([_row("单位", "账  号", "金额"), _row("甲", "33050161613500000123", "10")])
    top = at._Row([at._Cell(0, "基本信息", 3)], {}, True)
    labels = at._Row([at._Cell(0, "单位"), at._Cell(1, "银行账号"), at._Cell(2, "金额")], {}, True)
    builder.table([top, labels, _row("甲", "33050161613500000123", "10")])
    assert builder.segments[1].text == "甲 | [银行账号已脱敏] | 10"
    assert builder.segments[4].text == "甲 | [银行账号已脱敏] | 10"
    # 认不出表头也遮：表头前有三行说明、三列表的备注大多空着
    builder = at._Builder()
    builder.table(
        [
            _row("2026年度补助资金拨付表"),
            _row("填报单位：", "", "", "", ""),
            _row("单位：万元"),
            _row("序号", "单位", "开户行", "银行账号", "金额"),
            _row("1", "某某科技有限公司", "工行杭州分行", "1202021309001234568", "100"),
        ]
    )
    builder.table(
        [
            _row("单位名称", "银行账号", "备注"),
            _row("某某公司", "1202021309001234568", ""),
            _row("另一公司", "1202021309001234569", ""),
        ]
    )
    texts_now = [x.text for x in builder.segments]
    assert "1 | 某某科技有限公司 | 工行杭州分行 | [银行账号已脱敏] | 100" in texts_now
    assert "某某公司 | [银行账号已脱敏]" in texts_now and "另一公司 | [银行账号已脱敏]" in texts_now
    assert [x.kind for x in builder.segments[-3:]] == ["head", "row", "row"]
    # 只认像标签的格：跨整行的表名、"基本账户"这类数据不把别的列当账号列；账号列里的日期不遮
    builder = at._Builder()
    builder.table(
        [
            at._Row([at._Cell(0, "各单位银行账户信息表", 4)], {}),
            _row("单位", "账户类型", "项目编号", "有效期"),
            _row("甲", "基本账户", "2025330106000123", "2018-05-20—2026-10-09"),
        ]
    )
    builder.table([_row("单位", "银行账号", "备注"), _row("甲", "2018-05-20—2026-10-09", "无")])
    assert builder.segments[2].text == "甲 | 基本账户 | 2025330106000123 | 2018-05-20—2026-10-09"
    assert builder.segments[4].text == "甲 | 2018-05-20—2026-10-09 | 无"
    assert at.mask("账 号：3305 0161 6135 0000 0123") == "账 号：[银行账号已脱敏]"
    # 存成数字的长号码不写成科学计数法（写成科学计数法就遮不到了）
    epoch = datetime(1899, 12, 30)
    assert at._sheet_number("6222021234567890", False, False, epoch) == "6222021234567890"
    assert at._sheet_number("4.111111111111111e+15", False, False, epoch) == "4111111111111111"
    assert at.mask(at._sheet_number("4111111111111111", False, False, epoch)) == "[银行账号已脱敏]"


def test_mask_numbers_next_to_other_numbers_and_wider_account_labels() -> None:
    # 号码挨着别的数字（不带表格线的 PDF 里几列连成一行、一格里写了两个号码）：整串能分成号码和
    # 序号、日期、座机的，分别遮
    for raw, masked in {
        "1 138 1234 5678": "1 [手机号已脱敏]",
        "张三 110101 19900307 1234 138 1234 5678": "张三 [身份证号已脱敏] [手机号已脱敏]",
        "138 1234 5678\n139 8765 4321": "[手机号已脱敏]\n[手机号已脱敏]",
        "110101 19900307 1234\n110101 19900307 1235": "[身份证号已脱敏]\n[身份证号已脱敏]",
        "2025-10-10 138 1234 5678": "2025-10-10 [手机号已脱敏]",
        "20251010 138 1234 5678": "20251010 [手机号已脱敏]",
        "电话：0571-87654321 138 1234 5678": "电话：[电话已脱敏] [手机号已脱敏]",
        # 金额词后面隔着别的字的不算金额；身份证号前面紧挨着金额词也遮
        "补助对象：张三 138 1234 5678": "补助对象：张三 [手机号已脱敏]",
        "资金联系人李四 138 1234 5678": "资金联系人李四 [手机号已脱敏]",
        "奖励对象张三 110101 19900307 1234": "奖励对象张三 [身份证号已脱敏]",
        "合计 110101 19900307 1234": "合计 [身份证号已脱敏]",
    }.items():
        assert at.mask(raw) == masked, raw
    for keep in (
        "合计金额为 159 7000 7400",
        "资金（万元）：159 7000 7400",
        "12 1350 820 4500",  # 有一组既不是号码、也不是序号日期座机：整串不遮
        "2025 138 1234 5678",  # 单独的年份不算可以夹着的一段
    ):
        assert at.mask(keep) == keep, keep
    # 表格里的账号标签：前面最多 8 个字、后面可以带"/卡号"、括号说明、星号
    for label in (
        "开户银行及账号",
        "银行账号（基本户）",
        "账号（必填）",
        "收款单位开户账号",
        "账号/卡号",
        "*银行账号",
        "收款账户",
        "账 户：",
    ):
        assert at._account_cell(label), label
    for other in ("账户名称", "账户类型", "基本账户", "零余额账户", "账号信息表", "开户银行"):
        assert not at._account_cell(other), other
    builder = at._Builder()
    builder.table(
        [
            _row("单位", "开户银行及账号", "金额"),
            _row("甲", "工行杭州分行 1202021309001234568", "100"),
        ]
    )
    assert builder.segments[1].text == "甲 | 工行杭州分行 [银行账号已脱敏] | 100"


def test_mask_label_wording_bank_names_decimals_and_account_cells() -> None:
    # 标签后面跟"为""是"、"……号码"这样的标签、标签和号码中间隔着开户行
    for raw, masked in {
        "收款账号为1202021309001234568": "收款账号为[银行账号已脱敏]",
        "账号是 3305 0161 6135 0000 0123": "账号是 [银行账号已脱敏]",
        "账户号码：1202021309001234568": "账户号码：[银行账号已脱敏]",
        "银行账户号码 1202021309001234568": "银行账户号码 [银行账号已脱敏]",
        "开户行及账号：中国工商银行股份有限公司杭州西湖支行 1202021309001234568": (
            "开户行及账号：中国工商银行股份有限公司杭州西湖支行 [银行账号已脱敏]"
        ),
        "开户银行及账号：工行西湖支行 33050161613500000123": (
            "开户银行及账号：工行西湖支行 [银行账号已脱敏]"
        ),
        "联系电话为 1381 234 5678": "联系电话为 [手机号已脱敏]",  # 分法少见：靠标签认
        "身份证号是 11010 11990 03071 234": "身份证号是 [身份证号已脱敏]",
        # 小数不当号码，也不和前面的数连成一串；句末的英文句点不算小数点
        "1 138 1234 5678 9.5": "1 [手机号已脱敏] 9.5",
        "联系电话：138 1234 5678.": "联系电话：[手机号已脱敏].",
    }.items():
        assert at.mask(raw) == masked, raw
    for keep in (
        "项目产值 13812345678.50元",  # 正文规则会把小数点前面的 11 位当手机号：小数不交给它
        "1350 820 4500.5",
        "账户名称：某某科技有限公司 1350 820",
    ):
        assert at.mask(keep) == keep, keep
    # 表格账号标签：字间有空格（PDF 分散对齐）、"……号码"
    for label in ("银 行 账 号", "收 款 账 户", "账户号码", "银行卡号码"):
        assert at._account_cell(label), label
    assert not at._account_cell("基 本 账 户")
    assert not at._account_cell("基本账户（人民币）")
    # 账号列：一格写了两个账号都遮；中间夹着认不出的组也遮；8 位连写的日期区间不遮
    builder = at._Builder()
    builder.table(
        [
            _row("单位", "银行账号", "备注"),
            _row("甲", "1202021309001234568\n33050161613500000123", "无"),
            _row("乙", "1202021309001234568 2025年起", "无"),
            _row("丙", "20180520-20261009", "无"),
        ]
    )
    assert [x.text for x in builder.segments[1:]] == [
        "甲 | [银行账号已脱敏]\n[银行账号已脱敏] | 无",
        "乙 | [银行账号已脱敏] 2025年起 | 无",
        "丙 | 20180520-20261009 | 无",
    ]


def test_mask_bank_names_and_labels_do_not_swallow_neighbours() -> None:
    account = "1202021309001234568"
    # 开户行：农商行、信用联社、合作社、带英文字母、结尾带括号；"账号（开户行）："；只写"开户行："
    for raw in (
        f"开户行及账号：杭州联合农商行 {account}",
        f"账号：某某县农村信用合作社 {account}",
        f"账号：某某县信用联社 {account}",
        f"账号：ICBC杭州分行 {account}",
        f"账号：中国工商银行杭州分行（营业部） {account}",
        f"账号（工商银行杭州分行）：{account}",
        f"开户行：中国工商银行杭州分行 {account}",
        f"开户行及账号：中国工商银行杭州分行：{account}",  # 行名后面跟冒号、"为"
        f"账号（工商银行杭州分行）为 {account}",
    ):
        assert at.mask(raw) == raw.replace(account, "[银行账号已脱敏]"), raw
    # 没写"账号""开户行"、只有行名的不遮（设计分歧，写进 PR 请创始人定）
    assert at.mask(f"中国工商银行杭州分行 {account}") == f"中国工商银行杭州分行 {account}"
    # 写了标签的不把后面挨着的日期、金额一起遮：手机、身份证标签只管第一段；账号第一组连写
    # 12 位以上的不往后接，也不接日期
    for raw, masked in {
        f"账号：{account} 2025-10-10": "账号：[银行账号已脱敏] 2025-10-10",
        "账号：3305 0161 6135 0000 0123 2025-10-10": "账号：[银行账号已脱敏] 2025-10-10",
        "联系电话：0571-87654321 1350 820 4500": "联系电话：[电话已脱敏] 1350 820 4500",
        "手机：138 1234 5678 139 8765 4321": "手机：[手机号已脱敏] [手机号已脱敏]",
    }.items():
        assert at.mask(raw) == masked, raw
    for raw, masked in {
        f"{account}\n2018-05-20": "[银行账号已脱敏]\n2018-05-20",
        f"{account} 1500 3000": "[银行账号已脱敏] 1500 3000",
        "3305 0161 6135 0000 0123": "[银行账号已脱敏]",  # 四位一组的照样接起来
    }.items():
        assert at.mask(raw, column="account") == masked, raw


def test_mask_wrapped_numbers_in_labeled_columns_and_bank_short_names() -> None:
    # PDF 格子里折成两行的账号（两行用空格接起来）照样接上；写了"身份证号"的怎么分组都认
    for raw in ("1202021309001 234568", "120202130900 1234568"):
        assert at.mask(raw, column="account") == "[银行账号已脱敏]", raw
    for raw in (
        "身份证号：110101 1990 03 07 1234",
        "身份证号：110101-1990-03-07-1234",
        "身份证号：1101 01 19900307 1234",
        "身份证号：11010119900307 1234",
    ):
        assert at.mask(raw) == "身份证号：[身份证号已脱敏]", raw
    # 开户行写简称
    for raw in (
        "开户行：建行 33050161613500000123",
        "开户行及账号：工行 1202021309001234568",
        "账号：中行 1202021309001234568",
    ):
        assert at.mask(raw).endswith(" [银行账号已脱敏]"), raw
    # 表格里身份证号、手机号的列（窄列折行写开）和左边那格写着"身份证号"的
    for label, kind in (("法人身份证号", "id"), ("证件号码", "id"), ("联系人电话", "phone")):
        assert at._cell_label(label) == kind, label
    assert at._cell_label("证件类型") is None
    builder = at._Builder()
    builder.table(
        [
            _row("姓名", "身份证号", "手机号码"),
            _row("张三", "1101011990030 71234", "1381234 5678"),
        ]
    )
    builder.table(  # 格子里自己写了标签的按格子里的认
        [
            _row("姓名", "手机号码", "身份证号"),
            _row("甲", "开户行：中国工商银行杭州分行 1202021309001234568", "手机：1381 234 5678"),
            _row("乙", "身份证号：110101 1990 03 07 1234", "无"),
        ]
    )
    builder.table(  # 名目加内容的两列表：标签在左边那格
        [
            _row("姓名", "张三"),
            _row("身份证号", "1101011990030 71234"),
            _row("手机", "1381234 5678"),
        ]
    )
    assert [x.text for x in builder.segments[1:]] == [
        "张三 | [身份证号已脱敏] | [手机号已脱敏]",
        "姓名 | 手机号码 | 身份证号",
        "甲 | 开户行：中国工商银行杭州分行 [银行账号已脱敏] | 手机：[手机号已脱敏]",
        "乙 | 身份证号：[身份证号已脱敏] | 无",
        "姓名 | 张三",
        "身份证号 | [身份证号已脱敏]",
        "手机 | [手机号已脱敏]",
    ]


def test_clean_is_linear_on_long_runs_of_spaces() -> None:
    import time

    started = time.monotonic()
    assert at.clean("甲" + chr(0x3000) * 40000 + "乙") == "甲 乙"  # 以前要 5 秒
    assert time.monotonic() - started < 0.5


def test_clean_joins_vertical_text_and_maps_symbol_fonts() -> None:
    assert at.clean("资\n格\n审\n核") == "资格审核"
    assert at.clean("　　第一行\n\n第二行  \t多空格 ") == "第一行\n第二行 多空格"
    assert at.clean(chr(0xF0A3) + "企业 " + chr(0xF06C) + "要点" + chr(0xE123)) == "□企业 ●要点"


# ---- 存法：三层、共用、撤下 ----


def _url_rows(session) -> dict[str, AttachmentTextUrl]:
    session.expire_all()
    return {r.url: r for r in session.scalars(select(AttachmentTextUrl))}


def _segment_texts(session, version_id) -> list[str]:
    return [s.text for s in at.segments(session, version_id)]


def test_record_three_layers_shares_content_and_masks_before_storing(session) -> None:
    report = at.TextReport()
    content = make_docx("申报指南", "联系人：张三 13812345678")
    at.record(session, BASE + "a.docx", "docx", NOW, report, content=content)
    at.record(session, BASE + "b.doc", "doc", NOW, report, content=content)
    rows = _url_rows(session)
    a, b = rows[BASE + "a.docx"], rows[BASE + "b.doc"]
    assert (a.status, a.reason, a.extractor, a.checked_at, a.changed_at) == (
        "ok",
        None,
        at.EXTRACTOR_VERSION,
        NOW,
        NOW,
    )
    assert a.version_id == b.version_id and b.declared_type == "doc"
    version = session.get(AttachmentTextVersion, a.version_id)
    assert (version.file_type, version.status, version.segments, version.masked) == (
        "docx",
        "ok",
        2,
        1,
    )
    assert _segment_texts(session, version.id) == [
        "申报指南",
        "联系人：[姓名已脱敏] [手机号已脱敏]",
    ]
    # 数据库里没有原号码
    stored = session.scalar(
        select(func.count()).where(AttachmentTextSegment.text.like("%13812345678%"))
    )
    assert stored == 0
    assert (report.extracted, report.reused, report.unchanged) == (1, 1, 0)
    # 同一链接再下载到同样的内容：不重取
    at.record(session, BASE + "a.docx", "docx", NOW + timedelta(days=7), report, content=content)
    assert report.unchanged == 1 and report.extracted == 1
    assert _url_rows(session)[BASE + "a.docx"].changed_at == NOW


def test_changed_gone_webpage_and_temporary_failures(session) -> None:
    report = at.TextReport()
    url = BASE + "guide.pdf"
    at.record(session, url, "pdf", NOW, report, content=make_docx("第一版"))
    first = _url_rows(session)[url].version_id
    later = NOW + timedelta(days=7)
    # 内容变了：改指新版本，旧版本没人指向，清理时连同文字删掉
    at.record(session, url, "pdf", later, report, content=make_docx("第二版"))
    row = _url_rows(session)[url]
    assert row.version_id != first and row.changed_at == later and report.retracted == 1
    at.cleanup(session, {at.url_key(url)}, report)
    assert session.get(AttachmentTextVersion, first) is None
    assert (
        session.scalar(select(func.count()).where(AttachmentTextSegment.version_id == first)) == 0
    )
    assert report.versions_removed == 1
    # 网络等临时问题：保留文字，规则版本不改，隔一天再试
    temporary = copies.Failure("network", "超时", True)
    at.record(session, url, "pdf", later + timedelta(hours=1), report, failure=temporary)
    row = _url_rows(session)[url]
    assert (row.status, row.changed_at, report.kept) == ("ok", later, 1)
    assert row.retry_at == later + timedelta(hours=1) + at.TEXT_RETRY
    # 原网站删了：撤下文字
    gone = copies.Failure("gone", "原网站已没有这个附件（HTTP 404）")
    at.record(session, url, "pdf", later + timedelta(days=1), report, failure=gone)
    row = _url_rows(session)[url]
    assert (row.status, row.reason, row.version_id) == ("failed", "gone", None)
    assert report.retracted == 2 and report.failed["gone"] == 1
    # 下载到的是网页：也撤下（这里本来就没有文字了），原因是网页
    page = b"<html><body>not found</body></html>"
    at.record(session, url, "pdf", later + timedelta(days=2), report, content=page)
    row = _url_rows(session)[url]
    assert (row.status, row.reason) == ("failed", "webpage")
    version = session.get(AttachmentTextVersion, row.version_id)
    assert (version.status, version.reason, version.file_type) == ("failed", "webpage", None)
    # 第一次就遇到临时问题：还没有文字，仍是待取；隔一天（或者手动不等）再算到期
    at.record(session, BASE + "slow.pdf", "pdf", NOW, report, failure=temporary)
    slow = _url_rows(session)[BASE + "slow.pdf"]
    assert (slow.status, slow.extractor) == ("pending", 0)
    key = {at.url_key(BASE + "slow.pdf")}
    assert at.due(session, key, NOW + timedelta(hours=1)) == set()
    assert at.due(session, key, NOW + timedelta(hours=1), retry_now=True) == key
    assert at.due(session, key, NOW + at.TEXT_RETRY) == key


def test_new_extractor_version_reextracts_and_marks_every_link(session, monkeypatch) -> None:
    report = at.TextReport()
    content = make_docx("旧规则取的字")
    at.record(session, BASE + "a.docx", "docx", NOW, report, content=content)
    at.record(session, BASE + "b.docx", "docx", NOW, report, content=content)
    keys = {at.url_key(BASE + "a.docx"), at.url_key(BASE + "b.docx"), at.url_key(BASE + "new")}
    assert at.due(session, keys, NOW) == {at.url_key(BASE + "new")}
    monkeypatch.setattr(at, "EXTRACTOR_VERSION", at.EXTRACTOR_VERSION + 1)
    assert at.due(session, keys, NOW) == keys
    # 规则升级后遇到临时问题：不算已按新规则处理，过了重试时间照样要重取
    failure = copies.Failure("network", "超时", True)
    at.record(session, BASE + "b.docx", "docx", NOW, report, failure=failure)
    assert at.due(session, keys, NOW + at.TEXT_RETRY) == keys
    later = NOW + timedelta(days=1)
    at.record(session, BASE + "a.docx", "docx", later, report, content=content)
    rows = _url_rows(session)
    assert rows[BASE + "a.docx"].changed_at == later and rows[BASE + "b.docx"].changed_at == later
    version = session.get(AttachmentTextVersion, rows[BASE + "a.docx"].version_id)
    assert version.extractor == at.EXTRACTOR_VERSION and version.extracted_at == later
    assert _segment_texts(session, version.id) == ["旧规则取的字"]
    assert report.extracted == 2


def test_interrupted_extraction_is_not_retried_and_shared_links_follow(
    session, monkeypatch
) -> None:
    report = at.TextReport()
    content = make_docx("会让进程中断的文件")

    def crash(_data):
        raise RuntimeError("进程被杀")  # 代替内存不足被杀、超时被停

    monkeypatch.setattr(at, "extract", crash)
    with pytest.raises(RuntimeError):
        at.record(session, BASE + "a.docx", "docx", NOW, report, content=content)
    session.rollback()
    row = _url_rows(session)[BASE + "a.docx"]
    version = session.get(AttachmentTextVersion, row.version_id)
    assert row.status == "pending" and version.detail == at._IN_PROGRESS

    # 中断不一定是文件的问题（部署重启也会）：再取一次；第二次也没取完，才按损坏记
    with pytest.raises(RuntimeError):
        at.record(session, BASE + "a.docx", "docx", NOW, report, content=content)
    session.rollback()
    version = session.get(AttachmentTextVersion, _url_rows(session)[BASE + "a.docx"].version_id)
    assert version.detail == at._IN_PROGRESS_AGAIN

    def never(_data):
        raise AssertionError("不该再取")

    monkeypatch.setattr(at, "extract", never)
    later = NOW + at.TEXT_RETRY
    at.record(session, BASE + "a.docx", "docx", later, report, content=content)
    at.record(session, BASE + "b.docx", "docx", later, report, content=content)  # 同样的内容
    rows = _url_rows(session)
    for url in (BASE + "a.docx", BASE + "b.docx"):
        assert (rows[url].status, rows[url].reason) == ("failed", "broken")
    # 按新规则重取共用的内容：其他链接的状态跟着改，不会还写"能读、0 字"
    monkeypatch.undo()
    shared = make_docx("共用的内容")
    at.record(session, BASE + "c.docx", "docx", NOW, report, content=shared)
    at.record(session, BASE + "d.docx", "docx", NOW, report, content=shared)
    monkeypatch.setattr(at, "EXTRACTOR_VERSION", at.EXTRACTOR_VERSION + 1)
    monkeypatch.setattr(
        at, "extract", lambda _data: at.Extraction("docx", reason="empty", detail="空")
    )
    at.record(session, BASE + "c.docx", "docx", later, report, content=shared)
    rows = _url_rows(session)
    assert (rows[BASE + "d.docx"].status, rows[BASE + "d.docx"].reason) == ("failed", "empty")
    assert rows[BASE + "d.docx"].changed_at == later
    # 共用的内容已按新规则取过：另一个链接不用再为它重新下载
    assert at.due(session, {at.url_key(BASE + "d.docx")}, later) == set()


def test_cleanup_drops_out_of_scope_links_and_orphan_texts(session) -> None:
    report = at.TextReport()
    at.record(session, BASE + "keep.docx", "docx", NOW, report, content=make_docx("留下"))
    at.record(session, BASE + "old.docx", "docx", NOW, report, content=make_docx("超出范围"))
    at.cleanup(session, {at.url_key(BASE + "keep.docx")}, report)
    assert list(_url_rows(session)) == [BASE + "keep.docx"]
    assert session.scalar(select(func.count()).select_from(AttachmentTextVersion)) == 1
    assert (report.removed, report.versions_removed) == (1, 1)


# ---- 和附件副本定时任务一起跑 ----


def test_cache_run_downloads_wps_and_et_only_for_text(session, client) -> None:  # noqa: F811
    doc = add_doc(
        session,
        1,
        APPLY,
        RECENT,
        [
            ("附件1 申报书", BASE + "form.wps", "wps"),
            ("附件2 汇总表", BASE + "sheet.et", "et"),
            ("附件3 指南", BASE + "guide.docx", "docx"),
            ("附件4 打包", BASE + "all.zip", "zip"),
        ],
    )
    wps = make_word97([("申报书：项目名称", "\r", {})])
    docx = make_docx("申报指南：截止时间为2026年11月5日。")
    fetcher = FakeFetcher(
        {BASE + "form.wps": wps, BASE + "sheet.et": XLS, BASE + "guide.docx": docx}
    )
    cloud = FakeCloud()
    report = run(session, fetcher, cloud)
    # zip 不下载；wps、et 下载了但不存副本
    assert sorted(fetcher.calls) == sorted(
        [BASE + "form.wps", BASE + "sheet.et", BASE + "guide.docx"]
    )
    assert cloud.uploads == [obj(docx, "docx")]
    assert report.text_only == 2 and not report.failed and report.stored == 1
    copies_by_url = {r.url: r for r in session.scalars(select(AttachmentCopy))}
    assert (copies_by_url[BASE + "form.wps"].status, copies_by_url[BASE + "form.wps"].failure) == (
        "failed",
        "format",
    )
    # 能不能在小程序里打开：规则不变，wps 仍是"格式不支持"
    detail = client.get(f"/api/v1/documents/{doc.id}").json()
    reasons = {a["name"]: (a["preview"], a.get("reason")) for a in detail["attachments"]}
    assert reasons["附件1 申报书"] == (False, "format")
    assert reasons["附件2 汇总表"] == (False, "format")
    assert reasons["附件3 指南"] == (True, None)
    # 三个都取到了文字（wps 实际是 doc、et 实际是 xls）
    rows = _url_rows(session)
    assert {u: r.status for u, r in rows.items()} == {
        BASE + "form.wps": "ok",
        BASE + "sheet.et": "ok",
        BASE + "guide.docx": "ok",
    }
    types = {u: session.get(AttachmentTextVersion, r.version_id).file_type for u, r in rows.items()}
    assert types == {
        BASE + "form.wps": "doc",
        BASE + "sheet.et": "xls",
        BASE + "guide.docx": "docx",
    }
    assert (report.texts.extracted, report.texts.due) == (3, 0)
    # 读的时候：zip 写格式不支持，其余能读
    states = [(x.ext, x.state) for x in at.readings(session, doc.id, NOW)]
    assert states == [("wps", "ok"), ("et", "ok"), ("docx", "ok"), ("zip", "format")]


def test_existing_copies_get_their_text_on_the_next_run(session) -> None:
    add_doc(session, 1, APPLY, RECENT, [("指南", BASE + "guide.pdf", "pdf")])
    pdf = make_pdf("Application guide 2026")
    session.add(
        AttachmentCopy(
            url_hash=copies.url_hash(BASE + "guide.pdf"),
            url=BASE + "guide.pdf",
            declared_type="pdf",
            status="ready",
            file_type="pdf",
            size=len(pdf),
            sha256=copies.sha256(pdf).hexdigest(),
            file_id=PREFIX + obj(pdf, "pdf"),
            rules_version=copies.RULES_VERSION,
            next_check_at=NOW + timedelta(days=5),
        )
    )
    session.commit()
    fetcher, cloud = FakeFetcher({BASE + "guide.pdf": pdf}), FakeCloud()
    report = run(session, fetcher, cloud)
    # 副本没到期，为了补文字下载了一次；内容没变，不重新上传
    assert fetcher.calls == [BASE + "guide.pdf"] and cloud.uploads == []
    assert report.unchanged == 1 and report.texts.extracted == 1
    assert _url_rows(session)[BASE + "guide.pdf"].status == "ok"
    # 再跑：文字有了，副本也没到期，不下载
    fetcher.calls.clear()
    run(session, fetcher, cloud)
    assert fetcher.calls == []


def test_text_only_rows_wait_behind_copy_rechecks_and_copyable_extensions_win(session) -> None:
    # 新通知：副本没到期，只为补文字；旧通知：副本到期了。每轮只下 1 个时先复查副本
    add_doc(session, 1, APPLY, date(2026, 10, 2), [("新", BASE + "new.pdf", "pdf")])
    add_doc(session, 2, APPLY, date(2026, 9, 1), [("旧", BASE + "old.pdf", "pdf")])
    pdf = make_pdf("Guide")
    for url, next_check in ((BASE + "new.pdf", NOW + timedelta(days=5)), (BASE + "old.pdf", NOW)):
        session.add(
            AttachmentCopy(
                url_hash=copies.url_hash(url),
                url=url,
                declared_type="pdf",
                status="ready",
                file_type="pdf",
                size=len(pdf),
                sha256=copies.sha256(pdf).hexdigest(),
                file_id=PREFIX + obj(pdf, "pdf"),
                rules_version=copies.RULES_VERSION,
                next_check_at=next_check,
            )
        )
    session.commit()
    fetcher = FakeFetcher({BASE + "new.pdf": pdf, BASE + "old.pdf": pdf})
    run(session, fetcher, FakeCloud(), limit=1)
    assert fetcher.calls == [BASE + "old.pdf"]
    # 同一链接在新通知里写成 wps、老通知里是 doc：按 doc 算（能存副本），老通知那个附件照样能打开
    add_doc(session, 3, APPLY, date(2026, 10, 3), [("新写法", BASE + "same", "wps")])
    add_doc(session, 4, APPLY, date(2026, 9, 3), [("老写法", BASE + "same", "doc")])
    assert copies.candidates(session, NOW)[BASE + "same"] == "doc"


def test_text_step_errors_never_stop_copies(session, monkeypatch, caplog) -> None:
    add_doc(session, 1, APPLY, RECENT, [("指南", BASE + "guide.pdf", "pdf")])
    pdf = make_pdf("Application guide 2026")

    def broken(*_args, **_kwargs):
        raise RuntimeError("数据库出错")

    monkeypatch.setattr(at, "record", broken)
    monkeypatch.setattr(at, "cleanup", broken)
    report = run(session, FakeFetcher({BASE + "guide.pdf": pdf}), FakeCloud())
    assert report.stored == 1 and report.storage_error is None
    assert report.texts.errors == 2
    assert session.scalar(select(AttachmentCopy.status)) == "ready"


def test_text_only_run_never_touches_copies_or_storage(session) -> None:
    add_doc(
        session,
        1,
        APPLY,
        RECENT,
        [
            ("附件1", BASE + "a.docx", "docx"),
            ("附件2", BASE + "b.pdf", "pdf"),
            ("附件3", BASE + "slow.doc", "doc"),
        ],
    )
    add_doc(session, 2, APPLY, date(2026, 3, 1), [("旧附件", BASE + "old.pdf", "pdf")])
    slow = FetchResult(BASE + "slow.doc", "network_error", error="超时")
    files = {
        BASE + "a.docx": make_docx("甲"),
        BASE + "b.pdf": make_pdf("B"),
        BASE + "slow.doc": slow,
    }
    fetcher = FakeFetcher(files)
    report = copies.extract_texts(session, fetcher=fetcher, now=NOW, limit=2)
    assert len(fetcher.calls) == 2 and report.deferred == 1 and report.texts.due == 1
    assert session.scalar(select(func.count()).select_from(AttachmentCopy)) == 0
    report = copies.extract_texts(session, fetcher=fetcher, now=NOW)
    assert fetcher.calls[-1] == BASE + "slow.doc" and report.texts.kept == 1
    rows = _url_rows(session)
    assert rows[BASE + "slow.doc"].status == "pending" and BASE + "old.pdf" not in rows
    # 临时问题的不自动重试，加 --retry-pending 才重试
    fetcher.calls.clear()
    copies.extract_texts(session, fetcher=fetcher, now=NOW)
    assert fetcher.calls == []
    files[BASE + "slow.doc"] = make_word97([("乙", "\r", {})])
    copies.extract_texts(session, fetcher=fetcher, now=NOW, retry_pending=True)
    assert (
        fetcher.calls == [BASE + "slow.doc"]
        and _url_rows(session)[BASE + "slow.doc"].status == "ok"
    )
    with pytest.raises(ValueError):
        copies.extract_texts(session, fetcher=fetcher, now=NOW, limit=0)


def test_readings_states(session) -> None:
    doc = add_doc(
        session,
        1,
        APPLY,
        RECENT,
        [
            ("能读", BASE + "a.docx", "docx"),
            ("扫描件", BASE + "scan.pdf", "pdf"),
            ("没了", BASE + "gone.pdf", "pdf"),
            ("外站", "https://example.com/x.docx", "docx"),
            ("还没取", BASE + "later.docx", "docx"),
            ("压缩包", BASE + "x.rar", "rar"),
        ],
    )
    old = add_doc(session, 2, APPLY, date(2026, 3, 1), [("旧", BASE + "a.docx", "docx")])
    report = at.TextReport()
    at.record(session, BASE + "a.docx", "docx", NOW, report, content=make_docx("申报指南"))
    at.record(session, BASE + "scan.pdf", "pdf", NOW, report, content=SCAN)
    at.record(
        session, BASE + "gone.pdf", "pdf", NOW, report, failure=copies.Failure("gone", "HTTP 404")
    )
    found = at.readings(session, doc.id, NOW)
    assert [(x.name, x.state) for x in found] == [
        ("能读", "ok"),
        ("扫描件", "scan"),
        ("没了", "gone"),
        ("外站", "unavailable"),
        ("还没取", "pending"),
        ("压缩包", "format"),
    ]
    assert found[0].chars == 4 and found[0].segments == 1 and found[0].file_type == "docx"
    assert [x.state for x in at.readings(session, old.id, NOW)] == ["old"]
    # 合并转载：详情页列的是组里正文最长那份的附件；那份是二级来源，不在下载范围里
    master = add_doc(session, 3, APPLY, RECENT, [])
    session.get(DocumentDetail, master.id).body_text = ""
    copy = add_doc(
        session, 4, APPLY, RECENT, [("转载附件", BASE + "c.docx", "docx")], rights_level=2
    )
    session.get(DocumentDetail, copy.id).body_text = "转载的正文更长"
    session.add(
        DocumentMerge(canonical_id=master.id, duplicate_id=copy.id, active=True, match_key="test")
    )
    session.commit()
    assert [x.state for x in at.readings(session, master.id, NOW)] == ["source"]


# ---- 实测统计 ----


def test_report_counts_formats_reasons_limits_and_sample(session, monkeypatch) -> None:
    long_body = "申报条件：注册满一年。" * 2000  # 2.2 万字
    a = add_doc(
        session,
        1,
        APPLY,
        RECENT,
        [("附件1 指南", BASE + "a.docx", "docx"), ("附件2 扫描", BASE + "s.pdf", "pdf")],
    )
    session.get(DocumentDetail, a.id).body_text = long_body
    b = add_doc(session, 2, APPLY, RECENT, [("附件1 申报书", BASE + "b.wps", "wps")])
    session.get(DocumentDetail, b.id).body_text = ""
    add_doc(session, 3, APPLY, RECENT, [("附件1 打包", BASE + "c.zip", "zip")])
    # 库里存的类型是旧规则算的申报通知，按现在的标题规则是"资讯与指南"：在下载范围里，不进抽样
    stale = add_doc(session, 5, "关于举办2026年度政策宣讲培训班的通知", RECENT, [])
    session.get(Document, stale.id).rule_kind = "申报通知"
    stale_detail = session.get(DocumentDetail, stale.id)
    stale_detail.attachments.append(
        DocumentAttachment(position=0, name="回执", url=BASE + "d.docx", extension="docx")
    )
    session.commit()
    report = at.TextReport()
    guide = make_docx(*["评分标准：创新性30分。截止时间2026年11月5日前。"] * 100)
    at.record(session, BASE + "a.docx", "docx", NOW, report, content=guide)
    at.record(session, BASE + "s.pdf", "pdf", NOW, report, content=SCAN)
    at.record(
        session,
        BASE + "b.wps",
        "wps",
        NOW,
        report,
        content=make_word97([("申报材料清单", "\r", {})]),
    )
    at.record(session, BASE + "d.docx", "docx", NOW, report, content=make_docx("回执"))
    result = atr.build(session, NOW, sample=5, seed=1)
    notice = result.kinds[0]
    assert notice.kind == "申报通知"
    assert notice.attachments == {
        "docx": {"ok": 2},
        "pdf": {"scan": 1},
        "wps": {"ok": 1},
        "zip": {"format": 1},
    }
    assert notice.actual_types == {"wps→doc": 1}
    assert (notice.notices, notice.with_attachments, notice.with_readable) == (4, 4, 3)
    assert notice.stale_kind == 1
    assert (notice.with_body, notice.no_body_but_readable) == (3, 1)
    assert (notice.over_limit, notice.over_only_with_attachments, notice.body_too_long) == (1, 1, 0)
    assert len(result.sample) == 2 and result.sample_from == 2  # 类型过期的那份不进抽样
    assert stale.id not in {x.doc_id for x in result.sample}
    item = next(x for x in result.sample if x.doc_id == a.id)
    assert item.verdicts["评分"] == "only_attachments" and item.verdicts["条件"] == "body_only"
    text = atr.render(result)
    assert "能读 3 个（60%）" in text and "加上附件才超的 1 份" in text
    assert "其中 1 份库里存的类型是申报通知" in text
    assert "评分：只在附件里 1" in text
    sheet = atr.worksheet(result)
    assert "附件1 第1段" in sheet and "粗判 只在附件里" in sheet


def test_cli_commands(session, monkeypatch, capsys, tmp_path) -> None:
    add_doc(session, 1, APPLY, RECENT, [("附件1", BASE + "a.docx", "docx")])
    fetcher = FakeFetcher({BASE + "a.docx": make_docx("申报指南")})
    real = copies.extract_texts
    monkeypatch.setattr(
        copies, "extract_texts", lambda s, **kw: real(s, **{**kw, "fetcher": fetcher, "now": NOW})
    )
    assert cli.main(["attachment-texts", "--limit", "5"]) == 0
    out = capsys.readouterr().out
    assert "新取 1" in out and "还没取 0" in out
    sheet = tmp_path / "sheet.md"
    assert (
        cli.main(["attachment-text-report", "--now", "2026-10-06", "--worksheet", str(sheet)]) == 0
    )
    out = capsys.readouterr().out
    assert "【申报通知（近 180 天）】" in out and "docx：1 个，能读 1（100%）" in out
    assert sheet.read_text(encoding="utf-8").startswith("# 附件读入实测：抽样核对表")
    assert cli.main(["attachment-text-report", "--now", "2026-10-06", "--json"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert data["kinds"][0]["attachments"] == {"docx": {"ok": 1}}


def test_text_tables_are_not_backed_up() -> None:
    from youju.backup import backed_up_tables

    names = {t.name for t in backed_up_tables()}
    assert not names & {
        "attachment_text_urls",
        "attachment_text_versions",
        "attachment_text_segments",
    }


def test_migration_0040_downgrade_upgrade(db_engine) -> None:
    from alembic import command

    from youju.migrations import alembic_config

    from .conftest import TEST_DB_URL

    config = alembic_config(TEST_DB_URL)
    command.downgrade(config, "0039")
    with db_engine.connect() as conn:
        tables = set(conn.scalars(sql("SHOW TABLES LIKE 'attachment_text%'")))
    assert tables == set()
    command.upgrade(config, "head")
    with db_engine.connect() as conn:
        tables = set(conn.scalars(sql("SHOW TABLES LIKE 'attachment_text%'")))
    assert tables == {
        "attachment_text_urls",
        "attachment_text_versions",
        "attachment_text_segments",
    }
