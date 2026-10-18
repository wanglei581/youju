"""检查实际导出文件中的说明和可选小程序码，不依赖生产资源。"""

from datetime import UTC, datetime
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document as WordDocument
from openpyxl import load_workbook
from PIL import Image as PillowImage

from youju import export_notes
from youju.api import exports
from youju.compare import CompareReport, CompareRow, export_compare_excel
from youju.track import CustomerProfile, CustomerTrack, TrackReport, export_track_excel

STAMP = datetime(2026, 10, 10, tzinfo=UTC)
DISCLAIMER = "第三方按公开发布的信息整理，仅供参考，以发文机关原文为准。"
SEARCH_NOTICE = "在微信里搜索小程序「有据申报与行业助手」可以查看原文和更新。"
SCAN_NOTICE = "微信扫码查看原文和更新"


def item():
    return exports.ExportItem(
        number=1,
        doc_id=12,
        title="关于组织申报项目的通知",
        kind="申报通知",
        status="",
        issuer="示例科技厅",
        doc_number="",
        published="2026-10-10",
        region="北京",
        deadline="见原文",
        date_role="截止时间见原文",
        support="见原文",
        unchecked="截止日期、支持方式",
        urls=["https://example.gov.cn/12.html"],
        sites="示例科技厅",
        source_checked="2026-10-10",
        deadline_checked="未记录",
        change_checked="未记录",
    )


def excel(kind, tmp_path):
    if kind == "public":
        return exports._xlsx([item()], "2026-10-10 08:00")
    output = tmp_path / f"{kind}.xlsx"
    if kind == "compare":
        report = CompareReport(
            [CompareRow("示例企业", "示例企业", "企业名称", "未命中")], [], STAMP
        )
        assert export_compare_excel(report, output) == 1
    else:
        report = TrackReport([CustomerTrack(CustomerProfile("示例企业", "北京"))], STAMP)
        assert export_track_excel(report, output) == 1
    return output.read_bytes()


@pytest.fixture
def mini_code(tmp_path, monkeypatch):
    code = tmp_path / "assets" / "mini-code.png"
    monkeypatch.setattr(export_notes, "MINI_CODE_PATH", code)
    return code


@pytest.mark.parametrize("kind", ["public", "compare", "track"])
@pytest.mark.parametrize("has_image", [False, True])
def test_excel_notes_with_optional_image(kind, has_image, mini_code, tmp_path):
    if has_image:
        mini_code.parent.mkdir()
        PillowImage.new("RGB", (32, 32), "white").save(mini_code)
    content = excel(kind, tmp_path)
    book = load_workbook(BytesIO(content))
    note = book["数据说明"]
    values = [cell.value for row in note for cell in row]
    assert values.count(DISCLAIMER) == 1
    assert values.count(SEARCH_NOTICE) == 1
    assert len(note._images) == int(has_image)
    assert (SCAN_NOTICE in values) is has_image
    with ZipFile(BytesIO(content)) as archive:
        images = [name for name in archive.namelist() if name.startswith("xl/media/")]
        assert len(images) == int(has_image)
        if has_image:
            assert archive.read(images[0]) == mini_code.read_bytes()
    if has_image:
        anchor = note._images[0].anchor
        assert anchor._from.col == 1
        assert note.cell(anchor._from.row + 1, 3).value == SCAN_NOTICE
        assert anchor.ext.cx == anchor.ext.cy == 160 * 9525
    # 新说明不挪动已有业务行，用户上传的名称也不会混入说明页。
    if kind == "public":
        assert book["文件清单"]["B2"].value == item().title
        assert note["D4"].value == "2026-10-10"
    else:
        assert book.worksheets[0]["A3"].value == "示例企业"
        assert "示例企业" not in values


@pytest.mark.parametrize("has_image", [False, True])
def test_word_notes_always_include_text(has_image, mini_code):
    if has_image:
        mini_code.parent.mkdir()
        PillowImage.new("RGB", (32, 32), "white").save(mini_code)
    book = WordDocument(BytesIO(exports._docx([item()], "2026-10-10 08:00")))
    paragraphs = [paragraph.text for paragraph in book.paragraphs]
    assert paragraphs.count(DISCLAIMER) == 1
    assert paragraphs.count(SEARCH_NOTICE) == 1
    assert paragraphs.index(DISCLAIMER) < paragraphs.index("1. " + item().title)
    assert paragraphs.index(SEARCH_NOTICE) < paragraphs.index("1. " + item().title)


@pytest.mark.parametrize("kind", ["public", "compare", "track"])
def test_broken_optional_image_keeps_text_export(kind, mini_code, tmp_path):
    mini_code.parent.mkdir()
    mini_code.write_bytes(b"not a png")
    book = load_workbook(BytesIO(excel(kind, tmp_path)))
    note = book["数据说明"]
    assert not note._images
    assert SEARCH_NOTICE in [cell.value for row in note for cell in row]
