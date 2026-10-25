#!/usr/bin/env python3
"""人工代做一次（D-058 验证愿不愿意付费用）：按顾问给的客户清单，用线上数据出一张 Excel。

做什么：
  1. 逐家查客户有没有上已收录的公示、认定名单（企业全称精确相同才算，和小程序"查名单"同一个接口）；
  2. 按客户所在省份（加国家部委）和产业方向，列出近 90 天发布的申报通知，只去掉已人工核对且已截止的
     （和小程序"按客户筛通知"同一个口径）。截止日期按首发版规则写（D-061，和小程序列表、导出一致）：
     人工核对过的写日期；8–30 天后到期、原句在原文里的写日期并标"机器摘录·未经人工核对"、附原句；
     其余不写日期。
不判断"能不能报"，不打分，不推荐。产业方向只认小程序里那 8 项的原文，写别的一律按"不限产业方向"列，
不替客户猜行业。

D-086 加的三样（"直接给结果"第 1 阶段：先人工交付，记下客户的反应）：
  - 每行有行号（通知 N1、N2……，名单 L1、L2……），每家客户有代号（A、B……）。顾问只说行号和代号，
    就能告诉我们客户对哪几行有回应，不用说企业名称。
  - 另出一份"反馈记录"表：只有行号、代号、省份、产业方向和公开通知的标题，不含企业名称、名单标题；
    交付当天删掉结果表以后，这份可以留下来记反馈。
  - 每家客户另出一个"给客户的短表"文件，只有这一家的通知和上榜情况，可以直接转给这家客户。
  客户清单可以多一列"在办项目"（如高企、专精特新）：通知标题里带这几个字的排在这家的前面并标出来，
  只按字找，没标的不代表无关。

园区版（--park，推广与获客方案第五节的"园区企业上榜汇总"）：园区给入驻企业名单（只要企业全称），
只查上榜、不查通知，出三页：说明（写给园区的）、按名单看（每份名单上有园区哪几家，年底统计时可以拿来
查漏）、按企业看。

怎么取数：api 服务不对公网开放，只能从小程序里调用。这里借微信开发者工具的自动化接口，在模拟器里
用 wx.cloud.callContainer 调线上接口（只读的查询接口）。所以运行前：开发者工具要开着"有据"项目、
模拟器已经加载出首页。查名单接口每个来源每分钟最多 30 次，这里每 2.2 秒查一家，200 家约 8 分钟。

用法（用后端虚拟环境里的 Python，里面有 openpyxl）：
  /Users/wanglei/youju-worktrees/ai/backend/.venv/bin/python tools/team/concierge.py [--park] <清单.xlsx|.csv> [输出.xlsx]

客户清单：第一行是表头，列名里带"企业""公司"或"名称"的是企业全称（必填）；带"地区"或"省"的是所在省份
（可空）；带"产业"或"方向"的是产业方向（可空，多个用顿号或逗号隔开）；带"在办"的是在办项目（可空，
多个用顿号或逗号隔开）。只有一列、没有表头也可以。
园区版只用企业全称这一列。
**清单里不要有联系人、电话**；清单放在仓库以外的文件夹，不上传到服务器，交付当天由创始人删除本机文件
（见 docs/商业/首批试用外联包_v2.md 第二节）。
"""

from __future__ import annotations

import csv
import json
import re
import subprocess
import sys
import time
from datetime import date, datetime
from pathlib import Path
from typing import TYPE_CHECKING
from urllib.parse import quote
from zoneinfo import ZoneInfo

# openpyxl 只在读写 Excel 时才导入：CI 跑 tools/tests 用的 Python 里没有它，测试只用到不碰 Excel 的部分
if TYPE_CHECKING:
    from openpyxl import Workbook

WECHATIDE = "/Applications/wechatwebdevtools.app/Contents/MacOS/wechatide"
PROJECT = "/Users/wanglei/youju-worktrees/miniapp-release/apps/miniapp"
CLOUD_ENV = "prod-d5getqftefb61fc92"
TZ = ZoneInfo("Asia/Shanghai")
MAX_CLIENTS = 500
LOOKUP_GAP_MS = 2200  # 查名单限流每分钟 30 次
SEARCH_GAP_MS = 400
NOTICE_LIMIT = 50  # 每家客户最多列多少条通知
EXCERPT_MIN_DAYS, EXCERPT_MAX_DAYS = 8, 30  # 机器摘录只在这个区间写日期（D-061）
INDUSTRIES = (
    "数字经济",
    "先进制造",
    "新材料与化工",
    "生物医药与健康",
    "新能源与绿色低碳",
    "电子信息与集成电路",
    "现代农业",
    "文旅与现代服务业",
)
PROVINCES = ["北京", "天津", "河北", "山西", "内蒙古", "辽宁", "吉林", "黑龙江", "上海", "江苏", "浙江", "安徽", "福建", "江西", "山东", "河南", "湖北", "湖南", "广东", "广西", "海南", "重庆", "四川", "贵州", "云南", "西藏", "陕西", "甘肃", "青海", "宁夏", "新疆"]
# 推广材料的固定写法（推广与获客方案第八节，长版）
STATEMENT = (
    "有据是第三方整理政府网站公开信息的工具，不是政府平台，不代办申报；"
    "每条附政府原文链接，截止日期写明人工核对或机器摘录，以原文为准。"
)
# 名单怎么处理（和外联包 v2 第二节的保密说明一致：名称发去比对，但服务器不保存、不记录）
HANDLING = (
    "查的时候企业名称逐个发到查询接口比对，服务器不保存名单，也不记录查了哪些企业；"
    "我们在交付当天删除自己电脑上的名单和生成的文件，以及自己微信里收发的副本；您微信里的文件请您自己删。"
)
# 10/18 试跑发现：名单解析目前不收"某某公司（某某工厂）"这类按工厂写的条目（后端 valid_org_name 不认），
# 查不到要提醒对方看原文；后端改好以后删掉这句
PLANT_NOTE = "名单里按工厂写的条目（如“某某公司（某某工厂）”）目前可能没有收录，查不到时请打开原文再看一眼。"
_SPLIT = re.compile(r"[、,，;；/|\s]+")
# 在办项目的简称：除了顾问写的字，标题里出现这个全称也算。只用来标出、排到前面，不推断有没有关系。
# 不收"小巨人"一类的别名：各地还有"科技小巨人"等别的项目，会标错（D-086 代码审阅）
PROJECT_WORDS = {
    "高企": ("高新技术企业",),
}
# 给客户的短表开头的说明（D-086 合规审阅：不代表符合条件、不是代办）
DISCLAIMER = (
    "这份表只是按地区和产业方向，从政府公开发布的申报通知里整理的线索，不代表符合申报条件，"
    "也不是代办申报服务；请以政府原文为准，自主核实、办理。"
)
DEADLINE_NOTE = (
    "【截止日期】写“已人工核对”的，是对照原文核对过的；写“机器摘录·未经人工核对”的，是程序从原文里"
    "摘出的日期（只给 8–30 天后到期的，附原句），可能有误；写“截止未核对”或“截止时间见原文”的，"
    "请打开原文看。"
)

# 装进模拟器的小工具：排队逐个调接口（间隔由任务自己带），调完马上按任务的 slim 只留要用的字段，
# 结果留在 getApp().__cg.done 里等 Python 来取（命令行一次传回的内容有限，正文和多余字段都不要）。
# slim：不写是查名单、查通知；report 是截止雷达；detail 是文件详情；coverage 是收录范围。
_INSTALL_JS = (
    "function(){ const app=getApp(); app.__cg={queue:[],done:{},running:false};"
    " app.__cgRun=function(){ const s=app.__cg; if(s.running) return; s.running=true;"
    " const next=function(){ const job=s.queue.shift(); if(!job){ s.running=false; return; }"
    " const finish=function(v){ s.done[job.id]=app.__cgSlim(v, job.slim); setTimeout(next, job.gap); };"
    ' wx.cloud.callContainer({config:{env:"%s"},path:job.path,method:"GET",header:{"X-WX-SERVICE":"api"},'
    " success:function(r){ finish({status:r.statusCode,data:r.data}); },"
    " fail:function(e){ finish({status:0,error:String(e&&e.errMsg)}); }}); };"
    " next(); };"
    " app.__cgSlim=function(v, kind){ const d=(v&&v.data)||{};"
    " const out={status:v.status,error:v.error||(typeof d.detail==='string'?d.detail:'')};"
    " const pick=function(x, keys){ const y={}; keys.forEach(function(k){ y[k]=x[k]; }); return y; };"
    " if(kind==='detail'){ out.title=d.title; out.issuer=d.issuer; out.url=d.url; return out; }"
    " if(kind==='coverage'){ out.cities=d.cities; out.districts=d.districts; out.notCovered=d.notCovered;"
    " return out; }"
    " if(kind==='report'){ const soon=d.unverifiedSoon||{}; out.total=d.total;"
    " out.items=(d.items||[]).map(function(x){ return pick(x,['id','title','issuer','province','city',"
    "'district','level','label','date','time','remaining_days','source','quote']); });"
    " out.soon=(soon.items||[]).map(function(x){ return pick(x,['id','title','issuer','province','city',"
    "'district','level','kind']); }); out.soonTotal=soon.total; return out; }"
    " out.total=d.total; out.items=(d.items||[]).map(function(x){ const y=pick(x,['title','issuer','date',"
    "'url','status','region','category','deadline','deadline_state','due','matched']);"
    " if(x.group){ y.group={label:x.group.label,versions:x.group.versions,regions:x.group.regions}; }"
    " return y; });"
    " out.candidates=(d.candidates||[]).slice(0,5); out.similar=(d.similar||[]).slice(0,5);"
    " out.ambiguous=d.ambiguous; out.truncated=d.truncated; out.coverage=d.coverage; return out; };"
    ' return "ok"; }'
) % CLOUD_ENV


def evaluate(js: str, timeout: int = 60):
    """在模拟器里执行一段 JS，返回它的返回值。"""
    proc = subprocess.run(
        [
            WECHATIDE,
            "-c",
            "claude-code",
            "automation_evaluate",
            "--project",
            PROJECT,
            "--fn-source",
            js,
        ],
        capture_output=True,
        text=True,
        timeout=timeout,
        check=False,
    )
    found = re.search(r"\{.*", proc.stdout, re.DOTALL)
    if not found:
        raise RuntimeError(
            f"开发者工具没有返回结果：{(proc.stdout or proc.stderr)[-200:]}"
        )
    result = json.loads(found.group(0))
    result = result.get("result", result)
    while isinstance(result, dict) and "result" in result and len(result) <= 3:
        result = result["result"]
    if isinstance(result, dict) and result.get("ok") is False:
        raise RuntimeError(f"开发者工具执行出错：{result.get('message')}")
    return result


def run_jobs(jobs: list[dict], label: str) -> dict[str, dict]:
    """把任务排进模拟器里的队列，等它逐个调完，取回精简后的结果。失败的最多重试 3 次。"""
    results: dict[str, dict] = {}
    pending = {job["id"]: job for job in jobs}
    for attempt in range(4):
        if not pending:
            break
        if attempt:
            print(f"  {label}：{len(pending)} 个没成功，第 {attempt} 次重试")
            time.sleep(20)
        batch = list(pending.values())
        for start in range(0, len(batch), 30):
            chunk = json.dumps(batch[start : start + 30], ensure_ascii=False)
            evaluate(
                "function(){ const app=getApp(); const jobs=%s; jobs.forEach(function(j){ app.__cg.queue.push(j); });"
                " app.__cgRun(); return app.__cg.queue.length; }" % chunk
            )
        fetched: set[str] = set()
        while len(fetched) < len(batch):
            time.sleep(4)
            done = json.loads(
                evaluate(
                    "function(){ return JSON.stringify(Object.keys(getApp().__cg.done)); }"
                )
            )
            ready = [i for i in done if i in pending and i not in fetched]
            for start in range(0, len(ready), 5):
                ids = ready[start : start + 5]
                raw = evaluate(
                    "function(){ const app=getApp(); const ids=%s; const out={};"
                    " ids.forEach(function(i){ out[i]=app.__cg.done[i]; delete app.__cg.done[i]; });"
                    " return JSON.stringify(out); }" % json.dumps(ids)
                )
                for job_id, value in json.loads(raw).items():
                    results[job_id] = value
                    fetched.add(job_id)
            print(f"  {label}：{len(results)}/{len(jobs)}", end="\r", flush=True)
        # 限流、服务端临时出错、网关错误、调用失败的再试；422（像个人姓名或太短）不重试
        pending = {
            i: job
            for i, job in pending.items()
            if results.get(i, {}).get("status") in (0, 429, 500, 502, 503, 504)
        }
    print()
    return results


def read_clients(path: Path) -> list[dict]:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = [[cell.strip() for cell in row] for row in csv.reader(handle)]
    else:
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).active
        rows = [
            ["" if cell is None else str(cell).strip() for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    rows = [row for row in rows if any(row)]
    if not rows:
        raise SystemExit("清单是空的")
    header = rows[0]

    def column(*words: str) -> int | None:
        return next((i for i, name in enumerate(header) if any(w in name for w in words)), None)

    name_col, region_col, industry_col, project_col = (
        column("企业", "公司", "名称"),
        column("地区", "省"),
        column("产业", "方向"),
        column("在办"),
    )
    if name_col is None or not looks_like_header(header):
        # 没有表头：第一列就是企业名称（第一行"某某有限公司"也带"公司"二字，不能当表头）
        name_col, region_col, industry_col, project_col, body = 0, None, None, None, rows
        print(f"没有表头，第一行就是第一家：{rows[0][0]}")
    else:
        body = rows[1:]
        print(f"第一行当表头（不算企业）：{'、'.join(c for c in header if c)}")
    clients, by_name = [], {}
    for row in body:
        name = row[name_col] if name_col < len(row) else ""
        if not name:
            continue
        region_text = row[region_col] if region_col is not None and region_col < len(row) else ""
        industry_text = (
            row[industry_col] if industry_col is not None and industry_col < len(row) else ""
        )
        project_text = (
            row[project_col] if project_col is not None and project_col < len(row) else ""
        )
        provinces = [p for p in PROVINCES if p in region_text]
        wanted = [v for v in _SPLIT.split(industry_text) if v]
        client = by_name.get(name)
        if client is None:
            client = by_name[name] = {
                "name": name,
                "provinces": [],
                "industries": [],
                "ignored": [],
                "projects": [],
                "region_text": region_text,
            }
            clients.append(client)
        # 同一家写了几行（如两个产业方向）：合并成一家，省份和方向都留下
        for key, values in (
            ("provinces", provinces),
            ("industries", [v for v in wanted if v in INDUSTRIES]),
            ("ignored", [v for v in wanted if v not in INDUSTRIES]),
            ("projects", [v for v in _SPLIT.split(project_text) if v]),
        ):
            client[key] += [v for v in values if v not in client[key]]
    if len(clients) > MAX_CLIENTS:
        raise SystemExit(f"清单有 {len(clients)} 家，一次最多 {MAX_CLIENTS} 家")
    for index, client in enumerate(clients):
        client["code"] = client_code(index)
    return clients


def client_code(index: int) -> str:
    """客户代号：A、B……Z、AA、AB……。反馈只说代号和行号，不说企业名称（D-072、D-086）。"""
    code, index = "", index + 1
    while index:
        index, rest = divmod(index - 1, 26)
        code = chr(ord("A") + rest) + code
    return code


def project_hits(title: str | None, projects: list[str]) -> list[str]:
    """通知标题里实际出现的在办项目用词（顾问写的字，或 PROJECT_WORDS 里的全称），按清单顺序去重。
    只按字找，不推断；表里写的就是标题里出现的字。"""
    title, found = title or "", []
    for project in projects:
        for word in (project, *PROJECT_WORDS.get(project, ())):
            if word in title and word not in found:
                found.append(word)
                break
    return found


# 第一行当不当表头：有一格是短列名（带这些词，12 个字以内），而且没有一格带"有限""股份"
# （企业全称几乎都带）。这是猜的，所以读完会把第一行怎么处理的打印出来，让人看一眼。
_HEADER_WORDS = ("企业", "公司", "名称", "地区", "省", "产业", "方向")
_ORG_MARKS = ("有限", "股份")


def looks_like_header(row: list[str]) -> bool:
    cells = [cell for cell in row if cell]

    def org_like(cell: str) -> bool:
        return any(mark in cell for mark in _ORG_MARKS)

    def label(cell: str) -> bool:
        return len(cell) <= 12 and any(w in cell for w in _HEADER_WORDS) and not org_like(cell)

    # 有一格超过 12 个字的不当表头：没表头时第一家的名称常常很长，第二格又可能是"高新技术企业"这类带表头词的短词
    return (
        any(label(c) for c in cells)
        and not any(org_like(c) for c in cells)
        and all(len(c) <= 12 for c in cells)
    )


def notice_key(client: dict) -> str:
    return "|".join([",".join(client["provinces"]), ",".join(client["industries"])])


def notice_path(client: dict) -> str:
    params = {
        "kind": "申报通知",
        "time": "近 90 天",
        "exclude_closed": "true",
        "group": "1",
        "excerpts": "1",  # 每条带主截止 due（D-061），和小程序列表、导出一致
    }
    params["limit"] = str(NOTICE_LIMIT)
    # 和小程序"按客户筛通知"一样：有省份时查"国家部委 + 这些省"，没有时全国
    params["regions"] = (
        ",".join(["国家部委", *client["provinces"]])
        if client["provinces"]
        else "不限地区"
    )
    if client["industries"]:
        params["industries"] = ",".join(client["industries"])
    return "/api/v1/search?" + "&".join(f"{k}={quote(v)}" for k, v in params.items())


def now_cn() -> datetime:
    return datetime.now(TZ)


def today_cn() -> date:
    return now_cn().date()


def deadline_text(item: dict, now: datetime | None = None) -> str:
    """截止日期一格，和小程序 utils/format.js 的 dueView、deadlineNote 同一个口径（D-061）。

    机器摘录这里再拦一次：没带原句、按今天算不在 8–30 天里的，都不写日期。
    已核对的过了日期，或者就是今天、已经过了写明的时刻，写"已截止"（和小程序 daysUntil 一样）。"""
    now = now or now_cn()
    today = now.date()
    due = item.get("due") or {}
    if due.get("date"):
        at = due["date"] + (f" {due['time']}" if due.get("time") else "")
        kind = f"{due['kind']}，" if due.get("kind") else ""
        days = (date.fromisoformat(due["date"]) - today).days
        if due.get("source") == "verified":
            closed = days < 0 or (
                days == 0 and bool(due.get("time")) and now.strftime("%H:%M") >= due["time"]
            )
            return f"{at}{'已截止' if closed else ''}（{kind}已人工核对）"
        if (
            due.get("source") == "excerpt"
            and due.get("quote")
            and EXCERPT_MIN_DAYS <= days <= EXCERPT_MAX_DAYS
        ):
            return f"{at}（{kind}机器摘录·未经人工核对；原文：{due['quote']}）"
    elif (item.get("deadline") or {}).get("date"):  # 不带 excerpts 的旧接口：deadline 只有核对过的
        return f"{item['deadline']['date']}（已人工核对）"
    if item.get("deadline_state") == "none":
        return "截止时间见原文"
    return "截止未核对，请以原文为准"


def version_note(item: dict) -> str:
    """转发组（D-055）：确认过的组折成一行时，写明这一行是哪一版、另有几个转发版本。
    截止一格只写这一版自己的截止（接口就是这样给的），不合成一个日期。"""
    group = item.get("group") or {}
    if not group.get("versions"):
        return ""
    regions = list(group.get("regions") or [])
    where = "、".join(regions[:3]) + ("等" if len(regions) > 3 else "")
    label = group.get("label") or "其中一版"
    return f"这一行是{label}；另有 {group['versions']} 个转发版本（{where}），截止以各自原文为准"


def lookup_rows(
    clients: list[dict], lookups: dict[str, dict]
) -> tuple[list[list], int, dict[str, dict]]:
    """逐家整理查名单的结果：表格行、没查成的个数、按名单归拢的结果（园区版"按名单看"用）。"""
    rows, by_list, failed = [], {}, 0
    for i, client in enumerate(clients):
        found = lookups.get(f"l{i}") or {}
        status = found.get("status")
        if status == 200 and found.get("items"):
            for item in found["items"]:
                matched = item.get("matched") or {}
                note = []
                if matched.get("via") == "creditCode":
                    note.append("按统一社会信用代码相同找到（名单里的名称不同）")
                if found.get("ambiguous"):
                    note.append("名单里有同名但信用代码不同的企业，请核对")
                if found.get("truncated"):  # 和小程序查名单一样的说法
                    note.append(f"只列出最近的 {len(found['items'])} 份，可能还有更早的")
                rows.append(
                    [
                        client["name"],
                        f"上榜 {len(found['items'])} 份",
                        item.get("title"),
                        item.get("issuer"),
                        item.get("date"),
                        item.get("status") or "",
                        matched.get("name") or "",
                        " ".join(
                            str(v)
                            for v in (matched.get("batch"), matched.get("sequence"))
                            if v
                        ),
                        item.get("url"),
                        "；".join(note),
                    ]
                )
                entry = by_list.setdefault(
                    item.get("url") or item.get("title") or "",
                    {"item": item, "clients": []},
                )
                entry["clients"].append((client["name"], matched.get("name") or ""))
            continue
        if status == 200:
            hints = []
            if found.get("similar"):
                hints.append(
                    "名称相近的（没有算作同一家，是不是同一家请核对原文）："
                    + "、".join(
                        f"{s['name']}（{s['lists']} 份）" for s in found["similar"]
                    )
                )
            if found.get("candidates"):
                hints.append("名称里带这几个字的有：" + "、".join(found["candidates"]))
            result = (
                "没查完（只查了最近的名单），请在小程序里再查"
                if found.get("truncated")
                else "已收录的名单里没查到"
            )
            note = "；".join(hints)
        elif status == 422:
            result, note = "没查：名称太短或像个人姓名，请给企业全称", ""
        else:
            failed += 1
            result, note = f"没查成（{found.get('error') or status}），请重跑", ""
        rows.append([client["name"], result, "", "", "", "", "", "", "", note])
    return rows, failed, by_list


def truncated_count(clients: list[dict], lookups: dict[str, dict]) -> int:
    return sum(
        1 for i in range(len(clients)) if (lookups.get(f"l{i}") or {}).get("truncated")
    )


def coverage_text(lookups: dict[str, dict]) -> str:
    coverage = (
        next((r["coverage"] for r in lookups.values() if r and r.get("coverage")), None)
        or {}
    )
    if not coverage:
        return "收录范围有限。"
    return (
        f"目前只收录了 {coverage.get('lists')} 份名单（{coverage.get('since')} 以来各省科技、工信部门"
        "发布、能读出企业名称的；市、区发布的名单大多没有收录）。"
    )


def write_intro(book: Workbook, lines: list[str]) -> None:
    from openpyxl.styles import Alignment, Font

    intro = book.active
    intro.title = "说明"
    for line in lines:
        intro.append([line])
    intro.column_dimensions["A"].width = 110
    intro["A1"].font = Font(bold=True, size=14)
    for row in intro.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")


def write_sheet(
    book: Workbook, title: str, header: list[str], rows: list[list], widths: list[int]
):
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    sheet = book.create_sheet(title)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
                cell.font = Font(color="2A5CE0", underline="single")
    sheet.freeze_panes = "A2"
    return sheet


LIST_HEADER = [
    "企业",
    "结果",
    "名单标题",
    "发文机关",
    "发布日期",
    "公示/公布",
    "名单里的写法",
    "批次/序号",
    "原文链接",
    "备注",
]
LIST_WIDTHS = [30, 22, 50, 22, 12, 10, 30, 12, 40, 40]


def park_book(
    clients: list[dict], lookups: dict[str, dict], stamp: str
) -> tuple[Workbook, int, int]:
    """园区版：说明、按名单看、按企业看。返回工作簿、查到的企业数、没查成的个数。"""
    from openpyxl import Workbook

    rows, failed, by_list = lookup_rows(clients, lookups)
    hit = len({row[0] for row in rows if str(row[1]).startswith("上榜")})
    cut = truncated_count(clients, lookups)
    book = Workbook()
    write_intro(
        book,
        [
            "园区企业上榜汇总",
            (
                f"生成时间：{stamp}；园区给的企业 {len(clients)} 家，其中 {hit} 家在已收录的名单里查到，"
                f"共涉及 {len(by_list)} 份名单。"
            ),
            "",
            STATEMENT,
            "这张表只写名单上有什么，不评价企业，也不判断能不能申报。",
            "",
            "【怎么查的】企业全称和名单原文完全相同才算查到；名单里名称不同、但统一社会信用代码相同的，"
            "也列出来并在备注里写明。没查到不等于没上榜：" + coverage_text(lookups) + PLANT_NOTE,
            (
                "【按名单看】每份名单上有园区的哪几家企业，按发布日期从新到旧排。年底统计时可以拿来查漏，"
                "但不能当完整的统计：只含已收录的名单（见上一条）。"
            ),
            "【按企业看】每家企业上了哪些名单；没查到的也列在这里，名称相近的写在备注里（没有算作同一家，请核对原文）。",
            (
                "【公示和公布】标“公示”的是公示文件，不是最终结果；有没有正式公布，请看原文和后续文件。"
                "统计时请看这一列。"
            ),
            f"【名单怎么处理】园区给的企业名单只用来做这一张表；{HANDLING}",
            *(
                [f"【注意】有 {cut} 家企业的查询结果没取全（只列了最近的名单），“结果”或“备注”列里有标注，请在小程序里再查。"]
                if cut
                else []
            ),
        ],
    )
    list_rows = []
    for url, entry in sorted(
        by_list.items(),
        key=lambda kv: (kv[1]["item"].get("date") or "", kv[1]["item"].get("title") or ""),
        reverse=True,
    ):
        item = entry["item"]
        names = [name for name, _ in entry["clients"]]
        renamed = [
            f"{name}（名单里写作：{written}）"
            for name, written in entry["clients"]
            if written and written != name
        ]
        list_rows.append(
            [
                item.get("title"),
                item.get("issuer"),
                item.get("date"),
                item.get("status") or "",
                len(names),
                "、".join(names),
                "；".join(renamed),
                item.get("url") or url,
            ]
        )
    write_sheet(
        book,
        "按名单看",
        [
            "名单标题",
            "发文机关",
            "发布日期",
            "公示/公布",
            "园区企业数",
            "园区企业",
            "名单里写法不同的",
            "原文链接",
        ],
        list_rows,
        [56, 22, 12, 10, 10, 60, 40, 40],
    )
    write_sheet(book, "按企业看", LIST_HEADER, rows, LIST_WIDTHS)
    return book, hit, failed


def region_of(client: dict) -> str:
    return "、".join(client["provinces"]) + "（含国家部委）" if client["provinces"] else "全国"


def industry_of(client: dict) -> str:
    text = "、".join(client["industries"]) or "不限产业方向"
    if client["ignored"]:
        ignored = "、".join(client["ignored"])
        text += f"（清单里写的“{ignored}”不在可筛选的 8 项里，没有用）"
    return text


def prepare(clients: list[dict]) -> None:
    """测试和旧清单里的客户可能没有代号、在办项目：补上，后面几张表按同一套代号和行号。"""
    for index, client in enumerate(clients):
        client.setdefault("code", client_code(index))
        client.setdefault("projects", [])


def notice_entries(clients: list[dict], notices: dict[str, dict]) -> tuple[list[dict], int, int]:
    """每家客户的相关通知，编好行号（N1、N2……）：和在办项目有关的排在这家的前面，其余按接口给的顺序。
    没查成、没有通知、只列了一部分的，各占一行说明（不编行号）。返回行、没查成的家数、只列了一部分的家数。"""
    prepare(clients)
    entries, failed, cut, number = [], 0, 0, 0
    for client in clients:
        result = notices.get(notice_key(client)) or {}
        items = result.get("items") or []
        if result.get("status") != 200:
            failed += 1
            entries.append({"client": client, "note": "没查成，请重跑"})
            continue
        if not items:
            entries.append({"client": client, "note": "近 90 天没有列出的通知"})
        hits = [project_hits(item.get("title"), client["projects"]) for item in items]
        for hit, item in sorted(zip(hits, items), key=lambda pair: not pair[0]):  # 稳定排序
            number += 1
            entries.append({"client": client, "no": f"N{number}", "item": item, "hits": hit})
        if (result.get("total") or 0) > len(items) and items:
            cut += 1
            entries.append(
                {
                    "client": client,
                    "note": f"（共 {result['total']} 条，这里只列最新的 {len(items)} 条）",
                }
            )
    return entries, failed, cut


def numbered_list_rows(clients: list[dict], lookups: dict[str, dict]) -> tuple[list[list], int]:
    """上榜情况加行号（L1、L2……）和代号：[行号, 客户, 代号, 结果, 名单标题……备注]。"""
    prepare(clients)
    rows, failed, _ = lookup_rows(clients, lookups)
    codes = {client["name"]: client["code"] for client in clients}
    return [
        [f"L{n}", row[0], codes.get(row[0], ""), *row[1:]] for n, row in enumerate(rows, start=1)
    ], failed


def advisor_book(
    clients: list[dict],
    lookups: dict[str, dict],
    notices: dict[str, dict],
    stamp: str,
    now: datetime | None = None,
) -> tuple[Workbook, int, int, int]:
    """顾问版四页：说明、上榜情况、相关通知、按通知看。返回工作簿、查到的客户数、通知条数、没查成的个数。"""
    from openpyxl import Workbook

    now = now or now_cn()
    list_rows, failed = numbered_list_rows(clients, lookups)
    entries, notice_failed, cut_notices = notice_entries(clients, notices)
    failed += notice_failed
    notice_rows, by_notice = [], {}
    for entry in entries:
        client, item = entry["client"], entry.get("item")
        head = [client["name"], client["code"], region_of(client), industry_of(client)]
        if item is None:
            notice_rows.append(["", *head, entry["note"], "", "", "", "", "", "", ""])
            continue
        notice_rows.append(
            [
                entry["no"],
                *head,
                item.get("title"),
                item.get("issuer"),
                item.get("date"),
                item.get("category") or "",
                "、".join(entry["hits"]),
                deadline_text(item, now),
                version_note(item),
                item.get("url"),
            ]
        )
        by_notice.setdefault(item.get("url"), {"item": item, "clients": []})["clients"].append(
            client["name"]
        )

    book = Workbook()
    hit = len({row[1] for row in list_rows if str(row[3]).startswith("上榜")})
    cut = truncated_count(clients, lookups)
    with_projects = any(client["projects"] for client in clients)
    write_intro(
        book,
        [
            "客户名单与申报通知对照表",
            f"生成时间：{stamp}；客户 {len(clients)} 家，其中 {hit} 家在已收录的名单里查到。",
            "",
            STATEMENT,
            "这张表是按政府网站公开发布的信息整理的，不是政府平台出具的。",
            "",
            "【上榜情况】企业全称和名单原文完全相同才算查到；没查到不等于没上榜："
            + coverage_text(lookups)
            + PLANT_NOTE,
            (
                "【相关通知】从我们目前已收录的国家部委和省级部门栏目里，按客户所在省份（加国家部委）和产业方向，"
                "列出近 90 天发布的申报通知，只去掉截止日期已人工核对、确认已截止的；市、区县发的大多没有收录，"
                "没列出的不代表没有。产业方向按通知标题归类，仅供检索参考，不是官方分类。"
            ),
            "列出来不代表客户符合申报条件，这里不作任何资格判断。",
            *(
                [
                    "【在办项目】清单里写了在办项目的（如高企、专精特新），通知标题里带这些字的排在这家客户的前面，"
                    "“标题里有在办项目的字”一列写的是标题里实际出现的词；写了“高企”的，标题里的“高新技术企业”也算。"
                    "只按标题里的字找，不是资格判断，没标的不代表无关。"
                ]
                if with_projects
                else []
            ),
            DEADLINE_NOTE,
            "【按通知看】同一条通知和哪些客户有关，方便一次通知多家。",
            "【行号和代号】每行有行号（通知 N1、N2……，名单 L1、L2……），每家客户有代号（A、B……）。",
            (
                "【请帮我们记反馈】客户对哪几行有回应、问了什么、你改了哪里，告诉我们行号和代号就行，"
                "不用说企业名称、联系人。"
            ),
            (
                "【给客户的短表】每家客户另有一个文件，里面只有这一家的通知和上榜情况；截止日期和名单请看过再转，"
                "只发这家对应的那一个文件，不要发整个文件夹。这张总表里有你所有客户，请不要整张转给客户。"
            ),
            f"【客户清单怎么处理】只用来做这一张表；{HANDLING}"
            "我们只留一份不带企业名称的反馈记录（客户用代号 A、B 代替），用来记客户对哪几行有回应。",
            *(
                [
                    (
                        f"【注意】有 {cut_notices} 家客户的相关通知超过 {NOTICE_LIMIT} 条，只列了最新的 {NOTICE_LIMIT} 条"
                        "（“相关通知”页有标注），“按通知看”也只含列出来的这些。"
                    )
                ]
                if cut_notices
                else []
            ),
            *(
                [
                    f"【注意】有 {cut} 家客户的上榜结果没取全（只列了最近的名单），“结果”或“备注”列里有标注，请在小程序里再查。"
                ]
                if cut
                else []
            ),
        ],
    )
    write_sheet(
        book,
        "上榜情况",
        ["行号", "客户", "代号", *LIST_HEADER[1:]],
        list_rows,
        [7, LIST_WIDTHS[0], 6, *LIST_WIDTHS[1:]],
    )
    write_sheet(
        book,
        "相关通知",
        [
            "行号",
            "客户",
            "代号",
            "地区",
            "产业方向",
            "通知标题",
            "发文机关",
            "发布日期",
            "项目类别",
            "标题里有在办项目的字",
            "截止日期",
            "转发版本",
            "原文链接",
        ],
        notice_rows,
        [7, 30, 6, 20, 22, 56, 22, 12, 16, 14, 40, 30, 40],
    )
    write_sheet(
        book,
        "按通知看",
        [
            "通知标题",
            "发文机关",
            "发布日期",
            "截止日期",
            "有关的客户数",
            "有关的客户",
            "原文链接",
        ],
        [
            [
                e["item"].get("title"),
                e["item"].get("issuer"),
                e["item"].get("date"),
                deadline_text(e["item"], now),
                len(e["clients"]),
                "、".join(e["clients"]),
                url,
            ]
            for url, e in sorted(
                by_notice.items(),
                key=lambda kv: kv[1]["item"].get("date") or "",
                reverse=True,
            )
        ],
        [56, 22, 12, 40, 12, 60, 40],
    )
    return book, hit, len(by_notice), failed


FEEDBACK_HEADER = [
    "行号",
    "客户代号",
    "省份",
    "产业方向",
    "类型",
    "内容",
    "标题是否带在办用词",
    "截止日期",
    "原文链接",
    "客户有没有回应（回了 / 没回）",
    "问的是什么（截止 / 材料 / 附件 / 条件 / 别的）",
    "顾问改了哪里",
    "备注",
]


def feedback_book(
    clients: list[dict],
    lookups: dict[str, dict],
    notices: dict[str, dict],
    stamp: str,
    now: datetime | None = None,
) -> Workbook:
    """反馈记录（D-086 ②）：只有行号、代号、省份、产业方向、公开通知的标题和截止，不含企业名称。
    名单行只写结果（如"上榜 2 份（公示）"），不写名单标题和链接：名单小、再加上省份，可能猜出是哪家。
    在办项目只记标题里有没有（"是"），不记是哪一项：省份、产业再加在办项目，细分行业里可能认出是哪家。"""
    from openpyxl import Workbook

    now = now or now_cn()
    entries, _, _ = notice_entries(clients, notices)
    list_rows, _ = numbered_list_rows(clients, lookups)
    by_code = {client["code"]: client for client in clients}

    def where(client: dict) -> list[str]:
        return ["、".join(client["provinces"]) or "全国", "、".join(client["industries"]) or "不限"]

    rows = []
    for entry in entries:
        client = entry["client"]
        if not entry.get("no"):  # 没查成、近 90 天没有、只列了一部分：照记，分得清"没查成"和"确实没有"
            rows.append(["", client["code"], *where(client), "说明", entry["note"], *[""] * 7])
            continue
        item = entry["item"]
        rows.append(
            [
                entry["no"],
                client["code"],
                *where(client),
                "通知",
                item.get("title") or "",
                "是" if entry["hits"] else "",
                deadline_text(item, now),
                item.get("url"),
                "",
                "",
                "",
                "",
            ]
        )
    for row in list_rows:
        # [行号, 客户, 代号, 结果, 名单标题, 发文机关, 发布日期, 公示/公布, 名单里的写法……]：只取结果和公示/公布
        client = by_code[row[2]]
        status = f"（{row[7]}）" if row[7] else ""
        rows.append(
            [row[0], row[2], *where(client), "名单", f"{row[3]}{status}", "", "", "", "", "", "", ""]
        )

    book = Workbook()
    write_intro(
        book,
        [
            "样品表反馈记录（不含企业名称）",
            f"生成时间：{stamp}；对应的结果表里有 {len(clients)} 家客户，这里用代号 A、B……代替企业名称。",
            "",
            "怎么用：顾问说客户对哪几行有回应、问了什么、自己改了哪里，按行号填后四列。",
            "只记行为：不写企业名称（也不写简称）、联系人、电话，不写企业的财务、技术数据（D-072、D-086）。",
            "“标题是否带在办用词”一列只写“是”或空着，不要补项目名称或企业简称。",
            "名单行只写结果，不写名单标题：名单小、再加上省份，可能猜出是哪家。",
            "交付当天，我们删除自己电脑上的名单、结果表、给客户的短表文件夹和自己微信里收发的副本；顾问微信里的文件由顾问自己删；这份记录表只给我们自己用，可以留着。",
        ],
    )
    write_sheet(
        book,
        "反馈记录",
        FEEDBACK_HEADER,
        rows,
        [7, 8, 14, 18, 6, 56, 10, 40, 40, 14, 22, 30, 30],
    )
    return book


def client_books(
    clients: list[dict],
    lookups: dict[str, dict],
    notices: dict[str, dict],
    stamp: str,
    now: datetime | None = None,
) -> list[tuple[dict, Workbook]]:
    """给客户的短表：一家一个工作簿，只有这一家的通知和上榜情况，顾问可以直接转给这家客户。
    客户再多也一家一个文件：总表里有所有客户，不能转给客户（D-086 合规审阅）。"""
    from openpyxl import Workbook

    now = now or now_cn()
    entries, _, _ = notice_entries(clients, notices)
    list_rows, _ = numbered_list_rows(clients, lookups)
    books = []
    for client in clients:
        mine = [e for e in entries if e["client"] is client]
        # 没收这家的在办项目，就不出这一列和那句说明：整列空着，客户会以为是"判定无关"
        projects = bool(client["projects"])
        rows = [
            [
                e["item"].get("title"),
                e["item"].get("issuer"),
                e["item"].get("date"),
                deadline_text(e["item"], now),
                version_note(e["item"]),
                *(["、".join(e["hits"])] if projects else []),
                e["item"].get("url"),
            ]
            for e in mine
            if e.get("item")
        ]
        rows += [
            [CLIENT_NOTES.get(e["note"], e["note"]), *[""] * (6 if projects else 5)]
            for e in mine
            if not e.get("item")
        ]
        book = Workbook()
        write_intro(
            book,
            [
                f"申报线索 · {client['name']}",
                f"生成时间：{stamp}；地区：{region_of(client)}；产业方向：{industry_of(client)}。",
                "",
                STATEMENT,
                DISCLAIMER,
                "",
                (
                    "【覆盖范围】只含我们目前已收录的国家部委和省级部门栏目里近 90 天发布的申报通知；"
                    "市、区县发的大多没有收录，没列出的不代表没有。"
                ),
                *(
                    ["【标题里有在办项目的字】只按通知标题里的字找，不是资格判断；没标的不代表无关。"]
                    if projects
                    else []
                ),
                DEADLINE_NOTE,
                "【上榜情况】企业全称和名单原文完全相同才算查到；没查到不等于没上榜："
                + coverage_text(lookups)
                + PLANT_NOTE,
            ],
        )
        write_sheet(
            book,
            "申报线索",
            [
                "通知标题",
                "发文机关",
                "发布日期",
                "截止日期",
                "转发版本",
                *(["标题里有在办项目的字"] if projects else []),
                "原文链接",
            ],
            rows,
            [56, 22, 12, 40, 30, *([14] if projects else []), 40],
        )
        write_sheet(
            book,
            "上榜情况",
            LIST_HEADER[1:],
            client_list_rows([row for row in list_rows if row[2] == client["code"]]),
            LIST_WIDTHS[1:],
        )
        books.append((client, book))
    return books


# 给客户看的说法：总表里的"请重跑"是写给我们和顾问的
CLIENT_NOTES = {"没查成，请重跑": "本次没查出通知，请联系顾问再查"}
AMBIGUOUS = "名单里有同名但信用代码不同的企业"


def client_list_rows(rows: list[list]) -> list[list]:
    """给客户的短表里的上榜情况（D-086 ③：名单同名没核对的不能给客户）：
    同名没核对的整行换成一句"顾问核对后再告诉您"，没查成的换成对客户的说法；其余照总表。"""
    out = []
    for row in rows:
        result, note = str(row[3]), str(row[11] or "")
        if AMBIGUOUS in note:
            row = ["名单里有同名但统一社会信用代码不同的企业，请对照信用代码核实，或问顾问", *[""] * 8]
        elif result.startswith("没查成"):
            row = ["本次没查出，请联系顾问再查", *[""] * 8]
        else:
            row = row[3:]
        if row not in out:  # 同名的企业上了几份名单：只说一次
            out.append(row)
    return out


def save_client_files(folder: Path, books: list[tuple[dict, Workbook]]) -> None:
    """一家一个文件。先清掉文件夹里上一次的短表：重跑或去掉一家以后，不留别的客户的文件。"""
    folder.mkdir(exist_ok=True)
    for old in folder.glob("*.xlsx"):
        old.unlink()
    for client, book in books:
        book.save(folder / f"{client['code']}_{safe_file_name(client['name'])}.xlsx")


def safe_file_name(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", name).strip("_")[:40] or "客户"


def main() -> int:
    args = sys.argv[1:]
    park = "--park" in args
    args = [a for a in args if a != "--park"]
    if not args:
        print(__doc__)
        return 1
    source = Path(args[0]).expanduser()
    suffix = "_园区汇总.xlsx" if park else "_结果.xlsx"
    target = (
        Path(args[1]).expanduser()
        if len(args) > 1
        else source.with_name(source.stem + suffix)
    )
    clients = read_clients(source)
    print(
        f"读到 {len(clients)} 家{'企业' if park else '客户'}；"
        f"查名单约需 {len(clients) * LOOKUP_GAP_MS // 60000 + 1} 分钟"
    )
    evaluate(_INSTALL_JS)

    lookups = run_jobs(
        [
            {
                "id": f"l{i}",
                "path": "/api/v1/lists/lookup?q=" + quote(c["name"]),
                "gap": LOOKUP_GAP_MS,
            }
            for i, c in enumerate(clients)
        ],
        "查名单",
    )
    stamp = datetime.now(TZ).strftime("%Y-%m-%d %H:%M")
    if park:
        book, hit, failed = park_book(clients, lookups, stamp)
        book.save(target)
        print(f"已生成：{target}")
        print(
            f"查到 {hit} 家；没查成 {failed} 项" + ("（请重跑）" if failed else "")
        )
        return 1 if failed else 0

    combos = {notice_key(c): c for c in clients}
    searches = run_jobs(
        [
            {"id": f"s{n}", "path": notice_path(c), "gap": SEARCH_GAP_MS}
            for n, c in enumerate(combos.values())
        ],
        "查通知",
    )
    notices = {key: searches.get(f"s{n}") or {} for n, key in enumerate(combos)}
    book, hit, notice_count, failed = advisor_book(clients, lookups, notices, stamp)
    book.save(target)
    print(f"已生成：{target}")
    feedback = target.with_name(target.stem + "_反馈记录.xlsx")
    feedback_book(clients, lookups, notices, stamp).save(feedback)
    print(f"反馈记录（不含企业名称，可以留着）：{feedback}")
    folder = target.with_name(target.stem + "_给客户的短表")
    books = client_books(clients, lookups, notices, stamp)
    save_client_files(folder, books)
    print(f"给客户的短表 {len(books)} 个：{folder}（只发对应的那一个文件，不要发整个文件夹）")
    print(
        f"上榜 {hit} 家；通知 {notice_count} 条；没查成 {failed} 项"
        + ("（请重跑；重跑之前不要发给客户的短表）" if failed else "")
    )
    print("交付当天删除：顾问给的名单、结果表、给客户的短表文件夹、微信里收发的副本；反馈记录表可以留着。")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
