import importlib.util
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

TEAM = Path(__file__).parents[1] / "team"
sys.path.insert(0, str(TEAM))
SPEC = importlib.util.spec_from_file_location("concierge", TEAM / "concierge.py")
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)
try:  # CI 的工具测试环境没有 openpyxl：出 Excel 的测试只在本机跑
    import openpyxl  # noqa: F401

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
NOW = datetime(2026, 10, 19, 9, 0, tzinfo=MODULE.TZ)
QUOTE = "请于2026年10月30日17:00前完成网上申报"


def item(**due):
    return {"due": due or None, "deadline_state": "unverified"}


class DeadlineTextTests(unittest.TestCase):
    """截止一格和小程序 dueView 同一个口径（D-061）。"""

    def test_verified(self):
        text = MODULE.deadline_text(
            item(kind="企业申报截止", source="verified", date="2026-10-21", time="17:00"), NOW
        )
        self.assertEqual(text, "2026-10-21 17:00（企业申报截止，已人工核对）")

    def test_verified_past_says_closed(self):
        text = MODULE.deadline_text(item(source="verified", date="2026-10-18"), NOW)
        self.assertEqual(text, "2026-10-18已截止（已人工核对）")

    def test_verified_today_after_its_time_says_closed(self):
        due = {"source": "verified", "date": "2026-10-19", "time": "08:30"}
        self.assertEqual(MODULE.deadline_text(item(**due), NOW), "2026-10-19 08:30已截止（已人工核对）")
        due["time"] = "17:00"
        self.assertEqual(MODULE.deadline_text(item(**due), NOW), "2026-10-19 17:00（已人工核对）")

    def test_excerpt_in_window_has_label_and_quote(self):
        text = MODULE.deadline_text(
            item(kind="企业申报截止", source="excerpt", date="2026-10-30", quote=QUOTE), NOW
        )
        self.assertIn("机器摘录·未经人工核对", text)
        self.assertIn(QUOTE, text)
        self.assertTrue(text.startswith("2026-10-30"))

    def test_excerpt_without_quote_or_outside_8_to_30_days_has_no_date(self):
        for due in (
            {"source": "excerpt", "date": "2026-10-30", "quote": ""},
            {"source": "excerpt", "date": "2026-10-26", "quote": QUOTE},  # 7 天
            {"source": "excerpt", "date": "2026-11-19", "quote": QUOTE},  # 31 天
        ):
            with self.subTest(due=due):
                self.assertEqual(MODULE.deadline_text(item(**due), NOW), "截止未核对，请以原文为准")

    def test_old_api_and_no_candidate(self):
        self.assertEqual(
            MODULE.deadline_text({"deadline": {"date": "2026-10-25"}}, NOW), "2026-10-25（已人工核对）"
        )
        self.assertEqual(
            MODULE.deadline_text({"due": None, "deadline_state": "none"}, NOW), "截止时间见原文"
        )

    def test_notice_search_asks_for_excerpts(self):
        path = MODULE.notice_path({"provinces": ["重庆"], "industries": ["先进制造"]})
        self.assertIn("excerpts=1", path)
        self.assertIn("regions=%E5%9B%BD%E5%AE%B6%E9%83%A8%E5%A7%94%2C%E9%87%8D%E5%BA%86", path)


LIST_A = {"title": "绿色工厂名单的公示", "issuer": "某市经济和信息化委员会", "date": "2026-10-16",
          "status": "公示", "url": "https://example.gov.cn/a.html"}
LIST_B = {"title": "企业技术中心名单", "issuer": "某省工业和信息化厅", "date": "2025-12-30",
          "status": "公布", "url": "https://example.gov.cn/b.html"}


def found(*lists, via="name", written=None):
    return {"status": 200, "items": [
        {**x, "matched": {"name": written, "via": via, "batch": None, "sequence": "3"}} for x in lists
    ], "coverage": {"lists": 1200, "since": "2025-01-02"}}


CLIENTS = [{"name": n, "provinces": [], "industries": [], "ignored": [], "region_text": ""}
           for n in ("甲公司", "乙公司", "丙公司", "丁", "戊公司")]
LOOKUPS = {
    "l0": found(LIST_A, LIST_B, written="甲公司"),
    "l1": found(LIST_A, via="creditCode", written="乙公司（更名前）"),
    "l2": {"status": 200, "items": [], "similar": [{"name": "丙公司（一厂）", "lists": 2}],
           "candidates": [], "coverage": {"lists": 1200, "since": "2025-01-02"}},
    "l3": {"status": 422, "error": "只查企业和机构"},
    "l4": {"status": 0, "error": "timeout"},
}


class ReadClientsTests(unittest.TestCase):
    def read(self, text):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "x.csv"
            path.write_text(text, encoding="utf-8")
            return MODULE.read_clients(path)

    def test_single_column_without_header_keeps_first_company(self):
        clients = self.read("甲科技有限公司\n乙制药股份有限公司\n")
        self.assertEqual([c["name"] for c in clients], ["甲科技有限公司", "乙制药股份有限公司"])

    def test_header_is_skipped_and_columns_found(self):
        clients = self.read("企业名称,所在地区,产业方向\n甲科技有限公司,山东省青岛市,数字经济\n")
        self.assertEqual(clients[0]["name"], "甲科技有限公司")
        self.assertEqual(clients[0]["provinces"], ["山东"])
        self.assertEqual(clients[0]["industries"], ["数字经济"])

    def test_branch_name_in_brackets_is_still_a_company(self):
        clients = self.read("北京甲有限公司（北京分公司）\n乙公司\n")
        self.assertEqual(clients[0]["name"], "北京甲有限公司（北京分公司）")

    def test_short_header_words_are_a_header(self):
        clients = self.read("公司,省,产业方向\n甲科技有限公司,山东,数字经济\n")
        self.assertEqual([c["name"] for c in clients], ["甲科技有限公司"])
        self.assertEqual((clients[0]["provinces"], clients[0]["industries"]), (["山东"], ["数字经济"]))

    def test_same_company_on_two_rows_is_merged(self):
        clients = self.read("企业名称,地区,产业方向\n甲公司,山东,数字经济\n甲公司,江苏,先进制造、食品\n")
        self.assertEqual(len(clients), 1)
        self.assertEqual(clients[0]["provinces"], ["山东", "江苏"])
        self.assertEqual(clients[0]["industries"], ["数字经济", "先进制造"])
        self.assertEqual(clients[0]["ignored"], ["食品"])


class LookupRowsTests(unittest.TestCase):
    def test_rows_failures_and_by_list(self):
        rows, failed, by_list = MODULE.lookup_rows(CLIENTS, LOOKUPS)
        self.assertEqual(failed, 1)
        self.assertEqual([r[1] for r in rows if r[0] == "甲公司"], ["上榜 2 份", "上榜 2 份"])
        self.assertIn("按统一社会信用代码相同找到", next(r[9] for r in rows if r[0] == "乙公司"))
        self.assertIn("丙公司（一厂）（2 份）", next(r[9] for r in rows if r[0] == "丙公司"))
        self.assertIn("名称太短或像个人姓名", next(r[1] for r in rows if r[0] == "丁"))
        self.assertIn("请重跑", next(r[1] for r in rows if r[0] == "戊公司"))
        self.assertEqual(
            by_list[LIST_A["url"]]["clients"], [("甲公司", "甲公司"), ("乙公司", "乙公司（更名前）")]
        )
        self.assertEqual(len(by_list), 2)

    def test_truncated_results_are_marked(self):
        lookups = {"l0": {**found(LIST_A, written="甲公司"), "truncated": True},
                   "l1": {"status": 200, "items": [], "truncated": True}}
        rows, failed, _ = MODULE.lookup_rows(CLIENTS[:2], lookups)
        self.assertIn("只列出最近的 1 份", rows[0][9])
        self.assertIn("没查完", rows[1][1])
        self.assertEqual(failed, 0)
        self.assertEqual(MODULE.truncated_count(CLIENTS[:2], lookups), 2)

    def test_coverage_text_mentions_count_and_start_date(self):
        text = MODULE.coverage_text(LOOKUPS)
        self.assertIn("1200 份名单", text)
        self.assertIn("2025-01-02", text)


@unittest.skipUnless(HAS_OPENPYXL, "没有 openpyxl（CI 的工具测试环境），只在本机跑")
class ParkBookTests(unittest.TestCase):
    def test_three_sheets_counts_and_order(self):
        book, hit, failed = MODULE.park_book(CLIENTS, LOOKUPS, "2026-10-19 09:00")
        self.assertEqual(book.sheetnames, ["说明", "按名单看", "按企业看"])
        self.assertEqual((hit, failed), (2, 1))
        intro = "\n".join(str(c.value) for c in book["说明"]["A"] if c.value)
        self.assertIn("园区给的企业 5 家，其中 2 家", intro)
        self.assertIn(MODULE.STATEMENT, intro)
        self.assertIn("交付当天删除本机上的名单和生成的文件", intro)
        self.assertIn("我们这边微信里收发的副本", intro)
        lists = list(book["按名单看"].iter_rows(min_row=2, values_only=True))
        self.assertEqual([r[0] for r in lists], [LIST_A["title"], LIST_B["title"]])  # 新的在前
        self.assertEqual(lists[0][4], 2)
        self.assertEqual(lists[0][5], "甲公司、乙公司")
        self.assertEqual(lists[0][6], "乙公司（名单里写作：乙公司（更名前））")
        with tempfile.TemporaryDirectory() as tmp:
            book.save(Path(tmp) / "x.xlsx")

    def test_advisor_book_keeps_four_sheets(self):
        clients = [{"name": "甲公司", "provinces": ["重庆"], "industries": [], "ignored": ["汽车"],
                    "region_text": "重庆"}]
        notices = {MODULE.notice_key(clients[0]): {"status": 200, "total": 1, "items": [
            {"title": "关于申报的通知", "issuer": "某厅", "date": "2026-10-10", "url": "https://example.gov.cn/n",
             "category": "", "due": {"kind": "企业申报截止", "source": "excerpt", "date": "2026-11-05",
                                      "quote": "2026年11月5日前"}}]}}
        book, hit, count, failed = MODULE.advisor_book(clients, {"l0": LOOKUPS["l0"]}, notices, "2026-10-19 09:00")
        self.assertEqual(book.sheetnames, ["说明", "上榜情况", "相关通知", "按通知看"])
        self.assertEqual((hit, count, failed), (1, 1, 0))
        row = next(book["相关通知"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(row[:3], ("N1", "甲公司", "A"))
        self.assertIn("“汽车”不在可筛选的 8 项里", row[4])


def notice(title, url, date="2026-10-10"):
    return {
        "title": title,
        "issuer": "某厅",
        "date": date,
        "url": url,
        "category": "",
        "due": {
            "kind": "企业申报截止",
            "source": "excerpt",
            "date": "2026-11-05",
            "quote": "2026年11月5日前",
        },
    }


TWO = [
    {
        "name": "甲科技有限公司",
        "provinces": ["山东"],
        "industries": ["数字经济"],
        "ignored": [],
        "projects": ["高企"],
        "region_text": "山东",
    },
    {
        "name": "乙制药有限公司",
        "provinces": ["江苏"],
        "industries": [],
        "ignored": [],
        "projects": [],
        "region_text": "江苏",
    },
]
TWO_NOTICES = {
    MODULE.notice_key(TWO[0]): {
        "status": 200,
        "total": 2,
        "items": [
            notice("关于申报省级科技计划的通知", "https://example.gov.cn/1", "2026-10-12"),
            notice(
                "关于开展2026年高新技术企业认定工作的通知", "https://example.gov.cn/2", "2026-10-01"
            ),
        ],
    },
    MODULE.notice_key(TWO[1]): {
        "status": 200,
        "total": 1,
        "items": [
            notice("关于申报省级科技计划的通知", "https://example.gov.cn/1", "2026-10-12"),
        ],
    },
}
TWO_LOOKUPS = {
    "l0": found(LIST_A, written="甲科技有限公司"),
    "l1": {
        "status": 200,
        "items": [],
        "similar": [],
        "candidates": [],
        "coverage": {"lists": 1200, "since": "2025-01-02"},
    },
}


def cells(book):
    return [
        str(c) for ws in book.worksheets for row in ws.iter_rows(values_only=True) for c in row if c
    ]


class D086Tests(unittest.TestCase):
    """D-086 第 1 阶段：行号和代号、在办项目、反馈记录不含企业名称、给客户的短表只有这一家。"""

    def clients(self):
        return [dict(c, provinces=list(c["provinces"]), projects=list(c["projects"])) for c in TWO]

    def test_client_codes(self):
        self.assertEqual(
            [MODULE.client_code(i) for i in (0, 25, 26, 27, 51, 52)],
            ["A", "Z", "AA", "AB", "AZ", "BA"],
        )

    def test_project_hits_show_the_words_found_in_the_title(self):
        self.assertEqual(MODULE.project_hits("关于开展高新技术企业认定的通知", ["高企"]), ["高新技术企业"])
        self.assertEqual(MODULE.project_hits("关于高企培育入库的通知", ["高企"]), ["高企"])
        self.assertEqual(MODULE.project_hits("第八批专精特新“小巨人”企业", ["专精特新"]), ["专精特新"])
        # 不收"小巨人"别名：各地的"科技小巨人"是别的项目
        self.assertEqual(MODULE.project_hits("关于开展科技小巨人企业认定的通知", ["专精特新"]), [])
        self.assertEqual(MODULE.project_hits("关于申报科技计划的通知", ["高企", "科技计划"]), ["科技计划"])
        self.assertEqual(MODULE.project_hits(None, ["高企"]), [])

    def test_read_clients_reads_projects_and_assigns_codes(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "x.csv"
            path.write_text(
                "企业名称,省,在办项目\n甲科技有限公司,山东,高企、专精特新\n乙公司,江苏,\n",
                encoding="utf-8",
            )
            clients = MODULE.read_clients(path)
        self.assertEqual(
            [(c["code"], c["projects"]) for c in clients], [("A", ["高企", "专精特新"]), ("B", [])]
        )

    def test_entries_put_project_hits_first_and_number_rows(self):
        entries, failed, cut = MODULE.notice_entries(self.clients(), TWO_NOTICES)
        self.assertEqual((failed, cut), (0, 0))
        numbered = [(e["no"], e["client"]["code"], e["item"]["url"], e["hits"]) for e in entries]
        self.assertEqual(
            numbered,
            [
                ("N1", "A", "https://example.gov.cn/2", ["高新技术企业"]),  # 标题里有在办项目的字的排前面
                ("N2", "A", "https://example.gov.cn/1", []),
                ("N3", "B", "https://example.gov.cn/1", []),
            ],
        )

    @unittest.skipUnless(HAS_OPENPYXL, "没有 openpyxl（CI 的工具测试环境），只在本机跑")
    def test_feedback_book_has_no_company_names_or_list_titles(self):
        book = MODULE.feedback_book(
            self.clients(), TWO_LOOKUPS, TWO_NOTICES, "2026-10-19 09:00", NOW
        )
        self.assertEqual(book.sheetnames, ["说明", "反馈记录"])
        text = "\n".join(cells(book))
        for secret in ("甲科技有限公司", "乙制药有限公司", LIST_A["title"], LIST_A["url"]):
            self.assertNotIn(secret, text)
        rows = list(book["反馈记录"].iter_rows(min_row=2, values_only=True))
        self.assertEqual([r[0] for r in rows], ["N1", "N2", "N3", "L1", "L2"])
        self.assertEqual(rows[0][5:7], ("关于开展2026年高新技术企业认定工作的通知", "是"))
        self.assertNotIn("高企", text)  # 在办项目只记"是"，不记是哪一项
        self.assertIn("机器摘录", rows[0][7])
        self.assertEqual(rows[3][4:6], ("名单", "上榜 1 份（公示）"))

    @unittest.skipUnless(HAS_OPENPYXL, "没有 openpyxl（CI 的工具测试环境），只在本机跑")
    def test_client_books_only_carry_their_own_client(self):
        books = MODULE.client_books(
            self.clients(), TWO_LOOKUPS, TWO_NOTICES, "2026-10-19 09:00", NOW
        )
        self.assertEqual([c["code"] for c, _ in books], ["A", "B"])
        first = "\n".join(cells(books[0][1]))
        self.assertIn("甲科技有限公司", first)
        self.assertNotIn("乙制药有限公司", first)
        self.assertIn(MODULE.DISCLAIMER, first)
        second = books[1][1]
        self.assertNotIn("甲科技有限公司", "\n".join(cells(second)))
        self.assertEqual(
            list(second["上榜情况"].iter_rows(min_row=2, values_only=True))[0][0],
            "已收录的名单里没查到",
        )
        with tempfile.TemporaryDirectory() as tmp:
            books[0][1].save(Path(tmp) / "a.xlsx")

    @unittest.skipUnless(HAS_OPENPYXL, "没有 openpyxl（CI 的工具测试环境），只在本机跑")
    def test_every_client_gets_a_file_however_many(self):
        many = [dict(TWO[1], name=f"客户{i}有限公司") for i in range(31)]
        books = MODULE.client_books(many, {}, {}, "2026-10-19 09:00", NOW)
        self.assertEqual(len(books), 31)
        self.assertNotIn("客户1有限公司", "\n".join(cells(books[0][1])))

    @unittest.skipUnless(HAS_OPENPYXL, "没有 openpyxl（CI 的工具测试环境），只在本机跑")
    def test_advisor_book_numbers_lists_and_says_not_to_forward_whole_table(self):
        book, hit, count, failed = MODULE.advisor_book(
            self.clients(), TWO_LOOKUPS, TWO_NOTICES, "2026-10-19 09:00", NOW
        )
        self.assertEqual((hit, count, failed), (1, 2, 0))
        lists = list(book["上榜情况"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(
            [r[:3] for r in lists], [("L1", "甲科技有限公司", "A"), ("L2", "乙制药有限公司", "B")]
        )
        intro = "\n".join(str(c.value) for c in book["说明"]["A"] if c.value)
        self.assertIn("请不要整张转给客户", intro)
        self.assertIn("【在办项目】", intro)

    def test_no_header_row_with_a_zaiban_cell_is_still_a_company(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "x.csv"
            path.write_text("甲集团控股公司华东区域运营中心,在办高企\n乙公司,\n", encoding="utf-8")
            clients = MODULE.read_clients(path)
        self.assertEqual([c["name"] for c in clients], ["甲集团控股公司华东区域运营中心", "乙公司"])

    @unittest.skipUnless(HAS_OPENPYXL, "没有 openpyxl（CI 的工具测试环境），只在本机跑")
    def test_failures_cuts_and_same_names_in_feedback_and_client_files(self):
        clients = self.clients()
        notices = {
            MODULE.notice_key(clients[0]): {"status": 200, "total": 60, "items": TWO_NOTICES[
                MODULE.notice_key(TWO[0])]["items"]},
            MODULE.notice_key(clients[1]): {"status": 0, "error": "timeout"},
        }
        ambiguous = {**found(LIST_A, LIST_B, written="甲科技有限公司"), "ambiguous": True}
        lookups = {"l0": ambiguous, "l1": {"status": 0, "error": "timeout"}}
        fb = MODULE.feedback_book(clients, lookups, notices, "2026-10-19 09:00", NOW)
        rows = list(fb["反馈记录"].iter_rows(min_row=2, values_only=True))
        notes = [(r[0], r[1], r[4], r[5]) for r in rows if r[4] == "说明"]
        self.assertEqual(notes, [("", "A", "说明", "（共 60 条，这里只列最新的 2 条）"),
                                 ("", "B", "说明", "没查成，请重跑")])
        text = "\n".join(cells(fb))
        for secret in ("甲科技有限公司", "乙制药有限公司", LIST_A["title"], LIST_B["title"]):
            self.assertNotIn(secret, text)
        # 行号和总表是同一套
        book, *_ = MODULE.advisor_book(clients, lookups, notices, "2026-10-19 09:00", NOW)
        master = [r[0] for r in book["相关通知"].iter_rows(min_row=2, values_only=True) if r[0]]
        self.assertEqual(master, [r[0] for r in rows if r[4] == "通知"])
        books = MODULE.client_books(clients, lookups, notices, "2026-10-19 09:00", NOW)
        a_lists = list(books[0][1]["上榜情况"].iter_rows(min_row=2, values_only=True))
        self.assertEqual([r[0] for r in a_lists], ["名单里有同名但统一社会信用代码不同的企业，请对照信用代码核实，或问顾问"])
        self.assertNotIn(LIST_A["title"], "\n".join(cells(books[0][1])))
        b_text = "\n".join(cells(books[1][1]))
        self.assertIn("本次没查出通知，请联系顾问再查", b_text)
        self.assertIn("本次没查出，请联系顾问再查", b_text)
        self.assertNotIn("请重跑", b_text)

    @unittest.skipUnless(HAS_OPENPYXL, "没有 openpyxl（CI 的工具测试环境），只在本机跑")
    def test_client_files_folder_keeps_no_old_files(self):
        books = MODULE.client_books(self.clients()[:1], TWO_LOOKUPS, TWO_NOTICES, "2026-10-19 09:00", NOW)
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / "x_给客户的短表"
            folder.mkdir()
            (folder / "B_乙制药有限公司.xlsx").write_bytes(b"old")
            MODULE.save_client_files(folder, books)
            self.assertEqual(sorted(f.name for f in folder.iterdir()), ["A_甲科技有限公司.xlsx"])

    @unittest.skipUnless(HAS_OPENPYXL, "没有 openpyxl（CI 的工具测试环境），只在本机跑")
    def test_repost_group_row_says_which_version_its_deadline_belongs_to(self):
        grouped = dict(notice("关于申报2026年度省级科技计划项目的通知", "https://example.gov.cn/1", "2026-10-12"),
                       group={"label": "山东版", "versions": 4, "regions": ["江苏", "浙江", "安徽", "河南"]})
        self.assertEqual(MODULE.version_note(grouped),
                         "这一行是山东版；另有 4 个转发版本（江苏、浙江、安徽等），截止以各自原文为准")
        self.assertEqual(MODULE.version_note(notice("x", "u")), "")
        clients = self.clients()[:1]
        notices = {MODULE.notice_key(clients[0]): {"status": 200, "total": 1, "items": [grouped]}}
        book, *_ = MODULE.advisor_book(clients, {}, notices, "2026-10-19 09:00", NOW)
        header = [c.value for c in book["相关通知"][1]]
        row = next(book["相关通知"].iter_rows(min_row=2, values_only=True))
        self.assertTrue(row[header.index("截止日期")].startswith("2026-11-05"))  # 只写这一版自己的截止
        self.assertIn("山东版", row[header.index("转发版本")])
        books = MODULE.client_books(clients, {}, notices, "2026-10-19 09:00", NOW)
        sheet = books[0][1]["申报线索"]
        header = [c.value for c in sheet[1]]
        self.assertIn("另有 4 个转发版本", next(sheet.iter_rows(min_row=2, values_only=True))[header.index("转发版本")])

    @unittest.skipUnless(HAS_OPENPYXL, "没有 openpyxl（CI 的工具测试环境），只在本机跑")
    def test_client_file_has_no_project_column_when_none_given(self):
        books = MODULE.client_books(self.clients(), TWO_LOOKUPS, TWO_NOTICES, "2026-10-19 09:00", NOW)
        with_projects = [c.value for c in books[0][1]["申报线索"][1]]
        without = [c.value for c in books[1][1]["申报线索"][1]]
        self.assertIn("标题里有在办项目的字", with_projects)
        self.assertNotIn("标题里有在办项目的字", without)
        self.assertNotIn("在办项目", "\n".join(cells(books[1][1])))

    def test_safe_file_name(self):
        self.assertEqual(MODULE.safe_file_name('甲/乙:丙 "丁"'), "甲_乙_丙_丁")



if __name__ == "__main__":
    unittest.main()
